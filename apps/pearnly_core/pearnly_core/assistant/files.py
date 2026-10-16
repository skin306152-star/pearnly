"""对话里的文件（ADR-019）：附件进会话、结果文件列出与下载。

- 附件：浏览器传到本站 → 校验类型与大小（HEIC 照片先转 JPEG，与三个产品的上传口一致）→ 交给引擎放进该会话，
  智能体按返回的位置读取。
- 结果：智能体产出的文件按会话列出，可预览、下载。Excel 与 Word（.docx）由这里读成纯数据交给结果面板（`preview`），
  页面按文字渲染，文档里的脚本、外链到不了页面；读不出就给下载。
- 文件不脱离会话存在：下载前先核对「会话属于当前用户」且「文件属于这个会话」，猜编号读不到别人的文件。
- 附件原件在本站也留一份：引擎的文件库不让下载用户上传的文件，产品工具要用原件（例如 Cowork 把对话里发来的
  票据交给录入工作台识别）只能从这里取，见 `attachment`。存在站点私有目录 `private/assistant-attachments/<会话>/`
  （不进 File 表、没有下载地址）：同一张票识别过以后，它在 File 表里是受保护的原件，再存一份同内容的 File 会被
  Cowork 的原件守卫拒绝（`pearnly_cowork/files/guard.py`，2026-10-10 测试：第二次发同一张票上传失败）。删对话时一并删。
- 工具放进对话的文件也在本站留一份（2026-10 CC 闭环 L3）：工具产出（复核工作簿、回读的总账 / 税报……）能再交给别的
  工具当输入（例如对账），产品从文件夹取来的文件同样。来源分三种（`origin`）：用户上传（会话目录本身）、工具生成
  （`tool/`）、从文件夹取来（`folder/`）；产品据此拒绝把自己生成的文件再拿去识别。
- 对话接续到新上下文后，沙箱里的旧文件没了，本站这份还在（按对话留，不分段）：`kept` 列给智能体，工具照样按文件名取。
- 用户发的附件也能在右侧面板预览（2026-10-16 CC 外壳，同 Claude 点消息里的附件卡片）：引擎不让下载用户上传的文件，
  `download` / `preview` 找不到可下载的结果时取本站留存的那份；列表里的附件带上大小与 PDF 页数（附件卡片的说明行）。
- PDF 的一页画成图（`page_image`，第 5 轮测试 1、2）：附件卡片的缩略图（第一页，`thumb`）、手机上的 PDF 预览（安卓 Chrome
  不能在 iframe 里显示 PDF，逐页 `page`）。用镜像里已有的 PyMuPDF（识别流水线同一个，frappe.Dockerfile 锁定），不给前端
  加 PDF 渲染库；只有两种宽度，回应带私有缓存头（文件编号对应的内容不变），浏览器同一张只取一次。
"""
import io
import mimetypes
import os
import re
import shutil

import frappe
from frappe import _

from pearnly_core.assistant import contracts, docx_preview, engine, sessions
from pearnly_core.desk import heic

# 只有这些类型可以在本站内联打开。网页、SVG 一律按下载给出：它们在本站域名下打开会执行脚本；
# 对话页的预览把网页放进无同源权限的沙箱 iframe 里显示（前端读文本后用 srcdoc）。
INLINE_TYPES = frozenset({"image/png", "image/jpeg", "image/gif", "image/webp", "application/pdf"})


def upload(session_id) -> dict:
    row = sessions.owned(session_id)
    sent = (frappe.request.files or {}).get("file") if frappe.request else None
    if not sent:
        frappe.throw(_("No file was received."))
    content = sent.stream.read(contracts.UPLOAD_MAX_BYTES + 1)
    name = sessions._guard(contracts.upload_name, sent.filename, len(content))
    mime = sent.mimetype or "application/octet-stream"
    if heic.is_heic(content):
        content, name, mime = heic.to_jpeg(content), heic.jpeg_name(name), "image/jpeg"
    api = engine.adapter()
    stored = api.attach(sessions.segment_for_turn(row, api), name, content, mime)
    _keep(row, stored["file_id"], name, content)
    meta = _remember_meta(stored["file_id"], name, content)
    return {"fileId": stored["file_id"], "filename": name, "mountPath": stored["path"], **meta}


KEPT = "assistant-attachments"
# 用户上传以外的两种来源，各占会话目录下的一个子目录：工具生成的、产品从文件夹取来的。
ORIGINS = ("tool", "folder")
USER = "user"


def _folder(session_name: str, origin: str = USER) -> str:
    base = frappe.get_site_path("private", KEPT, re.sub(r"[^A-Za-z0-9_-]", "_", session_name))
    return base if origin == USER else os.path.join(base, origin)


def _keep(row, file_id, name, content, origin=USER) -> None:
    folder = _folder(row.name, origin)
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"{file_id}--{os.path.basename(name)}")
    with open(path + ".part", "wb") as handle:
        handle.write(content)
    os.replace(path + ".part", path)


def forget(session_name: str) -> None:
    """删对话时把留存的附件一起删（提交以后再删，回滚时不丢）。"""
    folder = _folder(session_name)
    frappe.db.after_commit.add(lambda: shutil.rmtree(folder, ignore_errors=True))


def _read(path: str) -> bytes:
    with open(path, "rb") as handle:
        return handle.read()


def _kept(file_id=None, row=None) -> list[tuple[str, callable, str]]:
    """[(「<编号>--<文件名>」, 取内容, 来源 user / tool / folder)]，旧 → 新：先是 2026-10-10 晚以前存在 File 表里的旧留存
    （都是用户上传），再是私有目录（按放进来的先后）。不给 `row` 就是当前工具调用所在的对话。"""
    if row is None:
        context = getattr(frappe.local, "pearnly_assistant", None) or {}
        row = sessions.owned(context.get("session_id"), allow_archived=True)
    found = []
    filters = {"attached_to_doctype": sessions.SESSION, "attached_to_name": row.name}
    if file_id:
        filters["file_name"] = ["like", file_id + "--%"]
    for old in frappe.get_all("File", filters=filters, fields=["name", "file_name"], order_by="creation asc"):
        found.append((old.file_name, lambda name=old.name: frappe.get_doc("File", name).get_content(), USER))
    local = []
    for origin in (USER,) + ORIGINS:
        folder = _folder(row.name, origin)
        if not os.path.isdir(folder):
            continue
        for entry in os.listdir(folder):
            path = os.path.join(folder, entry)
            if ("--" in entry and not entry.endswith(".part") and (not file_id or entry.startswith(file_id + "--"))
                    and os.path.isfile(path)):
                local.append((os.path.getmtime(path), entry, path, origin))
    return found + [(entry, lambda path=path: _read(path), origin) for _at, entry, path, origin in sorted(local)]


def attachments() -> list[dict]:
    """当前工具调用所在对话里留存的文件：[{file_id, filename}]（工具拿错编号时据此告诉智能体有哪些）；不是用户上传的
    另带 `origin`（tool / folder）。"""
    out = []
    for entry, _read, origin in _kept():
        file_id, sep, filename = entry.partition("--")
        if sep and not any(item["file_id"] == file_id for item in out):
            out.append({"file_id": file_id, "filename": filename, **({} if origin == USER else {"origin": origin})})
    return out


def kept(row) -> list[dict]:
    """这段对话（`row`：会话记录）留存的全部文件，新 → 旧：[{file_id, filename, origin}]。接续新上下文时告诉智能体
    还有哪些文件能交给工具（`sessions.segment_for_turn`）：沙箱里的那份随旧上下文没了，本站这份还在、按文件名照样取得到。"""
    out = []
    for entry, _read, origin in reversed(_kept(row=row)):
        file_id, sep, filename = entry.partition("--")
        if sep:
            out.append({"file_id": file_id, "filename": filename, "origin": origin})
    return out


def add_to_conversation(filename: str, content: bytes, mime: str, *, origin: str = "tool") -> str:
    """工具生成的文件（例如 Cowork 导出的复核工作簿）放进当前对话，返回智能体能读到的位置；它再放进结果目录给用户。

    同时在本站留一份（`origin`：tool = 工具生成，folder = 产品从已授权文件夹取来），之后的工具能像用户附件一样经
    `attachment` 取用（对账、盘点……）。类型与大小同用户上传（`contracts.upload_name`），HEIC 照片同样先转 JPEG。"""
    if origin not in ORIGINS:
        raise ValueError("unknown origin")
    context = getattr(frappe.local, "pearnly_assistant", None) or {}
    row = sessions.owned(context.get("session_id"))
    name = sessions._guard(contracts.upload_name, filename, len(content))
    if heic.is_heic(content):
        content, name, mime = heic.to_jpeg(content), heic.jpeg_name(name), "image/jpeg"
    api = engine.adapter()
    stored = api.attach(sessions.current(row), name, content, mime)
    _keep(row, stored["file_id"], name, content, origin)
    return stored["path"]


def _resolve(ref) -> str:
    """模型给的可能是编号，也可能是沙箱里的挂载路径（`/mnt/session/uploads/<编号>/<文件名>`）或文件名
    （2026-10-10 测试：先传路径被拒、再传编号，多绕一轮）：按本对话的附件认成编号；认不出原样返回。"""
    text = str(ref or "").strip()
    known = attachments()
    ids = {item["file_id"] for item in known}
    if text in ids:
        return text
    parts = [part for part in text.split("/") if part]
    for part in parts:
        if part in ids:
            return part
    name = parts[-1] if parts else text
    named = [item["file_id"] for item in known if item["filename"] == name]
    return named[-1] if named else text


def resolve(ref) -> str:
    """模型给的编号、挂载路径或文件名 → 本对话里那个文件的编号（同名取最新的）。工具按它定幂等键：同一个文件不论
    怎么称呼都是同一件请求；接续新上下文后智能体按文件名称呼旧文件，也不会被当成另一份、也不会与别的对话里的同名文件撞。
    不在这个对话里报与 `attachment` 相同的错。"""
    fid = sessions._guard(contracts.file_id, _resolve(ref))
    if not _kept(fid):
        frappe.throw(_("This attachment is not in this conversation."))
    return fid


def attachment(file_id) -> tuple[str, bytes]:
    """当前工具调用所在对话里的一个文件（文件名, 内容）：用户附件、工具生成的、从文件夹取来的都算。只认这个对话、这个人的，
    猜编号拿不到别人的。编号、挂载路径、文件名都认（`_resolve`）。"""
    fid = sessions._guard(contracts.file_id, _resolve(file_id))
    for entry, read, _origin in _kept(fid):
        if entry.startswith(fid + "--"):
            return entry[len(fid) + 2:], read()
    frappe.throw(_("This attachment is not in this conversation."))


def origin(file_id) -> str:
    """这个文件从哪来：user（用户上传）/ tool（工具生成）/ folder（从已授权文件夹取来）。认法与 `attachment` 相同，
    不在这个对话里就报同样的错。"""
    fid = sessions._guard(contracts.file_id, _resolve(file_id))
    for entry, _read, source in _kept(fid):
        if entry.startswith(fid + "--"):
            return source
    frappe.throw(_("This attachment is not in this conversation."))


def listing(session_id) -> list[dict]:
    """本会话的全部文件：挂进去的附件（只作输入）+ 智能体产出的结果（可下载）。按生成时间从早到晚排：引擎列文件是
    新的在前，结果面板默认打开最后一个，新导出的文件出来时面板却停在最早那份（2026-10-10 测试）。
    附件（kind = upload）按本站留存的那份补上大小与 PDF 页数（`pages`），对话里的附件卡片显示「PDF · N 页」/ 大小。"""
    row = sessions.owned(session_id, allow_archived=True)
    api = engine.adapter()
    found = sorted((f for remote in sessions.segments(row) for f in api.files(remote)),
                   key=lambda f: str(f.get("created_at") or ""))
    for f in found:
        if f.get("kind") == "upload":
            f.update(_upload_meta(row, f["id"]))
    return found


# 附件的大小与页数按文件编号记一份（引擎的文件编号全局唯一、内容不变）：列表每次打开对话、每轮跑完都要取，不每次读原件。
# 本站没留存的附件（很早以前的）也记一次「没有」，不每次重扫目录；只记一小时，上传时正好撞上列表的那一下不会久留。
META_CACHE = "pearnly_assistant:upload-meta:"
META_TTL = 7 * 24 * 3600
META_MISSING_TTL = 3600


def _pdf_pages(content: bytes) -> int | None:
    try:
        import pypdf

        return len(pypdf.PdfReader(io.BytesIO(content)).pages) or None
    except Exception:
        return None


def _remember_meta(fid: str, name: str, content: bytes) -> dict:
    """附件的说明数据：{"size_bytes"}，PDF 另带 {"pages"}（读不出页数就不带）。记进缓存，列表时不再读原件。"""
    meta = {"size_bytes": len(content)}
    if str(name).lower().endswith(".pdf"):
        pages = _pdf_pages(content)
        if pages:
            meta["pages"] = pages
    try:
        frappe.cache.set_value(META_CACHE + fid, meta, expires_in_sec=META_TTL)
    except Exception:
        pass
    return meta


def _upload_meta(row, fid: str) -> dict:
    """列表里一个附件的大小与页数：先看缓存，没有就读本站留存的那份；没留存（很早以前的附件）就不补。"""
    try:
        cached = frappe.cache.get_value(META_CACHE + fid)
    except Exception:
        cached = None
    if isinstance(cached, dict):
        return cached
    for entry, read, _origin in _kept(fid, row=row):
        if entry.startswith(fid + "--"):
            try:
                return _remember_meta(fid, entry[len(fid) + 2:], read())
            except OSError:
                return {}
    try:
        frappe.cache.set_value(META_CACHE + fid, {}, expires_in_sec=META_MISSING_TTL)
    except Exception:
        pass
    return {}


def _stored(row, api, fid) -> tuple[dict, callable]:
    """本对话里能打开的文件：(说明, 取内容)。先找本站留存的那份（用户上传的附件，以及工具放进对话的文件；只读本地、
    只认这个对话）；没有再去引擎里找可下载的结果（任一段，每段两次引擎接口调用）——引擎不让下载用户上传的文件，
    附件卡片的缩略图、预览先走引擎会白白多付这几次调用（审查第 4 轮：一条消息几十张照片就是几十轮）。
    两处都没有就当作不存在，猜编号拿不到别的对话、别人的文件。"""
    for entry, read, _origin in _kept(fid, row=row):
        if entry.startswith(fid + "--"):
            name = entry[len(fid) + 2:]
            mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
            return {"id": fid, "filename": name, "mime_type": mime, "size_bytes": None}, read
    for remote in sessions.segments(row):
        for f in api.files(remote):
            if f["id"] == fid and f["downloadable"]:
                return f, lambda remote=remote: api.read_file(remote, fid)
    frappe.throw(_("File not found"), frappe.DoesNotExistError)


def download(session_id, file_id, inline=False) -> None:
    row = sessions.owned(session_id, allow_archived=True)
    fid = sessions._guard(contracts.file_id, file_id)
    api = engine.adapter()
    meta, read = _stored(row, api, fid)
    frappe.local.response.filename = meta["filename"]
    frappe.local.response.filecontent = read()
    # Frappe 的 binary 回应是「附件 + octet-stream」：图片用 <img> 照样显示，PDF 放进结果面板却成了下载、一片空白
    # （2026-10-10 测试，桌面端每打开一次还多存一份）。PDF 用 Frappe 的 pdf 回应（inline + application/pdf）。
    if inline and meta["mime_type"] == "application/pdf" or inline and str(meta["filename"]).lower().endswith(".pdf"):
        frappe.local.response.type = "pdf"
    else:
        frappe.local.response.type = "binary" if inline and meta["mime_type"] in INLINE_TYPES else "download"


# 结果面板的结构化预览：表格只读前几张表的前若干行列；Word 只读正文（上限见 docx_preview）。
# 单个文件超过上限就不预览（下载照常）。
PREVIEW_MAX_BYTES = 5 * 1024 * 1024
PREVIEW_SHEETS = 10
PREVIEW_ROWS = 500
PREVIEW_COLS = 40
SHEET_TYPES = (".xlsx", ".xlsm")
DOCUMENT_TYPES = (".docx",)


def preview_kind(filename) -> str | None:
    """按文件名定结构化预览的类型：sheets（Excel）/ document（Word .docx）；别的类型没有结构化预览。
    老格式 .doc / .xls 是二进制，不读。"""
    name = str(filename or "").strip().lower()
    if name.endswith(SHEET_TYPES):
        return "sheets"
    if name.endswith(DOCUMENT_TYPES):
        return "document"
    return None


def preview(session_id, file_id) -> dict:
    """结果文件的结构化预览：Excel `{"type": "sheets", "sheets": [{"name", "rows"}]}`；Word
    `{"type": "document", "blocks": [...]}`（块的写法见 `docx_preview.read`）；PDF `{"type": "pdf", "pages": N}`
    （页面按页取 `page_image`，手机上用）；其余类型、超过上限或读不出 `{"type": "none"}`，页面给下载。"""
    row = sessions.owned(session_id, allow_archived=True)
    fid = sessions._guard(contracts.file_id, file_id)
    api = engine.adapter()
    meta, read = _stored(row, api, fid)
    if _is_pdf(meta):
        if int(meta.get("size_bytes") or 0) > PDF_MAX_BYTES:
            return {"type": "none"}
        content = read()
        pages = _pdf_pages(content) if len(content) <= PDF_MAX_BYTES else None
        return {"type": "pdf", "pages": pages} if pages else {"type": "none"}
    kind = preview_kind(meta["filename"])
    if not kind or int(meta.get("size_bytes") or 0) > PREVIEW_MAX_BYTES:
        return {"type": "none"}
    content = read()
    if len(content) > PREVIEW_MAX_BYTES:
        return {"type": "none"}
    if kind == "sheets":
        return {"type": "sheets", "sheets": sheets(content)}
    try:
        blocks = docx_preview.read(content)
    except docx_preview.Unreadable:
        return {"type": "none"}
    return {"type": "document", "blocks": blocks} if blocks else {"type": "none"}


def _is_pdf(meta) -> bool:
    return meta.get("mime_type") == "application/pdf" or str(meta.get("filename") or "").lower().endswith(".pdf")


# PDF 一页画成图：附件卡片缩略图（卡片宽 120 的两倍）、整页（手机屏宽的三倍上下）。只收这两种，缓存不碎、也不能要一张超大图。
PAGE_WIDTHS = {"thumb": 240, "page": 1240}
# 页面太长（长条收据）时高最多是宽的这么多倍，免得一张图几十兆像素
PAGE_TALLEST = 6
PAGE_MAX = 500
# 页数、画页只读这么大以内的 PDF（上传上限 30MB；工具产出的结果可能更大，太大的给下载）
PDF_MAX_BYTES = 60 * 1024 * 1024
# 同一个文件编号的内容不变：浏览器留一周，同一张缩略图 / 同一页只取一次（只给本人的浏览器，不进共享缓存）
PAGE_CACHE = "private, max-age=604800"


def render_page(content: bytes, number: int, width: int) -> bytes | None:
    """PDF 第 ``number`` 页画成宽 ``width`` 的 PNG；页码超出、加密、读不出回 None。"""
    try:
        import pymupdf
    except ImportError:
        return None
    try:
        with pymupdf.open(stream=content, filetype="pdf") as document:
            if document.needs_pass or not 1 <= number <= document.page_count:
                return None
            page = document[number - 1]
            box = page.rect
            if box.width <= 0 or box.height <= 0:
                return None
            scale = min(width / box.width, width * PAGE_TALLEST / box.height)
            return page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).tobytes("png")
    except Exception:  # 坏文件、不支持的写法：PyMuPDF 抛的类型不一（FileDataError / RuntimeError / ValueError…），一律当画不出
        return None


def page_image(session_id, file_id, page=1, size="thumb"):
    """本对话里一个 PDF 的一页（PNG）：附件卡片的缩略图（第一页、thumb）、手机上的 PDF 预览（逐页、page）。
    文件的取法、认法同 `download`（只认这个对话的，先读本站留存的那份）；不是 PDF、页码不对、画不出都当作不存在。"""
    from werkzeug.wrappers import Response

    row = sessions.owned(session_id, allow_archived=True)
    fid = sessions._guard(contracts.file_id, file_id)
    width = PAGE_WIDTHS.get(str(size or ""))
    try:
        number = int(page)
    except (TypeError, ValueError):
        number = 0
    if not width or not 1 <= number <= PAGE_MAX:
        frappe.throw(_("File not found"), frappe.DoesNotExistError)
    meta, read = _stored(row, engine.adapter(), fid)
    if not _is_pdf(meta) or int(meta.get("size_bytes") or 0) > PDF_MAX_BYTES:
        frappe.throw(_("File not found"), frappe.DoesNotExistError)
    content = read()
    png = render_page(content, number, width) if len(content) <= PDF_MAX_BYTES else None
    if not png:
        frappe.throw(_("File not found"), frappe.DoesNotExistError)
    return Response(png, mimetype="image/png", headers={"Cache-Control": PAGE_CACHE, "X-Content-Type-Options": "nosniff"})


def sheets(content: bytes) -> list[dict]:
    import io

    from openpyxl import load_workbook

    values = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    out = []
    for sheet in values.worksheets[:PREVIEW_SHEETS]:
        raw = formulas[sheet.title].iter_rows(max_row=PREVIEW_ROWS, max_col=PREVIEW_COLS)
        rows = [[_cell(v, f) for v, f in zip(vrow, frow)]
                for vrow, frow in zip(sheet.iter_rows(max_row=PREVIEW_ROWS, max_col=PREVIEW_COLS), raw)]
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
        out.append({"name": sheet.title, "rows": [r[:width] for r in rows]})
    return out


def _cell(value_cell, formula_cell) -> str:
    """单元格显示文字：按数字格式粗略排版；没有缓存值的公式显示公式本身。"""
    import datetime

    value = getattr(value_cell, "value", None)
    if value is None:
        formula = getattr(formula_cell, "value", None)
        return formula if isinstance(formula, str) and formula.startswith("=") else ""
    fmt = str(getattr(value_cell, "number_format", "") or "")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if "0.00" in fmt:
            return f"{value:,.2f}" if "," in fmt else f"{value:.2f}"
        if "#,##0" in fmt:
            return f"{value:,.0f}"
        return f"{value:g}" if isinstance(value, float) else str(value)
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)
