"""Port of ec2e5c957e6b table_path plus its separate billing character rules.

Pages deliberately retain legacy column collisions and primary-table selection.
Provenance exposes those omissions so admission can refuse incomplete coverage.
"""
import csv
import hashlib
import io
import zipfile
from xml.etree import ElementTree

from pearnly_core.originals.rules import OcrInputError, TABLE_EXTENSIONS, file_extension, validate_input_metadata
from pearnly_cowork.ocr.schema import Page

from .archive import check_archive
from .model import DecodeBudget, MAX_ROWS, MAX_TEXT_CHARACTERS, TableDocument, TableInputError


def _decode(content):
    for encoding in ("utf-8-sig", "utf-8", "cp874", "tis-620", "cp1252", "latin-1"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace"), "utf-8-replace"


def _headers(values):
    return ["" if value is None else str(value).strip() for value in values]


def _roundtrip_header(value):
    # openpyxl represents array/data-table formulae as objects whose default
    # repr includes a process address. Neither can equal a literal roundtrip
    # header; use a stable non-header marker in the durable manifest.
    if type(value).__module__ == "openpyxl.worksheet.formula":
        return "[" + type(value).__name__ + "]"
    return "" if value is None else str(value)


def _page(number, text, headers=None, rows=None, confidence=1.0):
    return Page(page_number=number, width=0, height=0, full_text=text,
                avg_confidence=confidence, blocks=[], table_headers=headers, table_rows=rows)


def _render(label, headers, rows):
    lines = [f"[{label}]"]
    if headers:
        lines.append(" | ".join(headers))
    lines.extend(" | ".join(cells) for cells in rows)
    return "\n".join(lines)


def _row(number, cells, **extra):
    return {"source_row": number, "cells": cells, "table_row_index": None, **extra}


def _issue(issues, code, page, **details):
    value = {"code": code, "page_number": page, **details}
    if value not in issues:
        issues.append(value)


def _coverage(headers, rows, issues, page):
    if len(set(headers)) != len(headers):
        _issue(issues, "duplicate_headers", page)
    for row in rows:
        cells = row["cells"]
        if any(cell.strip() for cell in cells[len(headers):]):
            _issue(issues, "columns_outside_headers", page, source_row=row["source_row"])
        if any(not header and index < len(cells) and cells[index].strip()
               for index, header in enumerate(headers)):
            _issue(issues, "unnamed_column_value", page, source_row=row["source_row"])


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError as error:
        raise TableInputError("ocr.table_dependency_unavailable", dependency="openpyxl") from error


def _docx():
    try:
        import docx
        return docx
    except ImportError as error:
        raise TableInputError("ocr.table_dependency_unavailable", dependency="python-docx") from error


def _excel(content, budget, issues):
    worksheets = check_archive(content)
    openpyxl = _openpyxl()
    pages, units, characters = [], [], 0
    # Roundtrip fingerprints use formula expressions, whereas the legacy
    # recognition/billing read uses cached values. Read only row 1 of that view.
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    formula_workbook = None
    try:
        formula_workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
        for index, name in enumerate(workbook.sheetnames, 1):
            sheet = workbook[name]
            headers, body, evidence = [], [], []
            header_row = None
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_number > MAX_ROWS:
                    raise TableInputError("ocr.table_row_limit")
                cells = ["" if value is None else str(value) for value in values]
                budget.row(cells)
                characters += sum(len(str(value)) for value in values if value is not None)
                raw = _row(row_number, cells)
                evidence.append(raw)
                if not any(cell.strip() for cell in cells):
                    continue
                if header_row is None:
                    headers, header_row = _headers(cells), row_number
                else:
                    raw["table_row_index"] = len(body) + 1
                    body.append(raw)
            _coverage(headers, body, issues, index)
            table_rows = [dict(zip(headers, (row["cells"] + [""] * len(headers))[:len(headers)]))
                          for row in body]
            pages.append(_page(index, _render(name, headers, [row["cells"] for row in body]),
                               headers, table_rows))
            raw_first = next(formula_workbook[name].iter_rows(min_row=1, max_row=1, values_only=True), ())
            missing = worksheets.get(sheet._worksheet_path, {}).get("missing_formula_cache", [])
            for cell in missing:
                _issue(issues, "formula_cache_missing", index, **cell)
            for feature in worksheets.get(sheet._worksheet_path, {}).get("unread_features", []):
                _issue(issues, "worksheet_unread_structure", index, feature=feature)
            units.append({"page_number": index, "source_kind": "worksheet", "sheet_name": name,
                          "sheet_index": index, "header_source_row": header_row, "rows": evidence,
                          "roundtrip_headers": [_roundtrip_header(value) for value in raw_first]})
    finally:
        workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
    if not pages:
        pages = [_page(1, "(empty workbook)", [], [])]
        units = [{"page_number": 1, "source_kind": "empty_workbook", "rows": []}]
    return pages, units, characters, "table_path_excel", "worksheet"


def _word_extra_coverage(document, issues):
    unread_tags = {"drawing", "pict", "object", "txbxContent", "altChunk", "sdt", "ins", "del",
                   "footnoteReference", "endnoteReference"}
    if any(element.tag.rsplit("}", 1)[-1] in unread_tags for element in document.element.iter()):
        _issue(issues, "docx_unread_body_structure", 1)
    for part in document.part.package.parts:
        name = str(part.partname).lower()
        if name.startswith(("/word/media/", "/word/embeddings/", "/word/charts/", "/word/diagrams/")):
            _issue(issues, "docx_embedded_content", 1, part=name)
        elif name.startswith(("/word/header", "/word/footer", "/word/footnotes", "/word/endnotes")):
            root = ElementTree.fromstring(part.blob)
            if any((element.text or "").strip() for element in root.iter()
                   if element.tag.rsplit("}", 1)[-1] == "t"):
                _issue(issues, "docx_non_body_text", 1, part=name)


def _csv(content, delimiter, budget, issues):
    text, encoding = _decode(content)
    if len(text) > MAX_TEXT_CHARACTERS:
        raise TableInputError("ocr.table_text_limit")
    # The billing source counts raw UTF-8 with errors ignored, NOT this decoded text.
    characters = len(content.decode("utf-8", errors="ignore"))
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    evidence, previous_line = [], 0
    for number, cells in enumerate(reader, 1):
        if number > MAX_ROWS:
            raise TableInputError("ocr.table_row_limit")
        budget.row(cells)
        evidence.append(_row(number, cells, line_start=previous_line + 1, line_end=reader.line_num))
        previous_line = reader.line_num
    if evidence:
        headers = _headers(evidence[0]["cells"])
        body, table_rows = evidence[1:], []
        _coverage(headers, body, issues, 1)
        for raw in body:
            cells = raw["cells"]
            if not any(cell.strip() for cell in cells):
                continue
            raw["table_row_index"] = len(table_rows) + 1
            padded = (cells + [""] * len(headers))[:len(headers)]
            table_rows.append(dict(zip(headers, padded)))
        page = _page(1, _render("csv", headers, [row["cells"] for row in body]), headers, table_rows)
    else:
        page = _page(1, "(empty csv)", [], [])
    unit = {"page_number": 1, "source_kind": "delimited_records", "encoding": encoding,
            "delimiter": delimiter, "header_source_row": 1 if evidence else None, "rows": evidence}
    return [page], [unit], characters, "table_path_tsv" if delimiter == "\t" else "table_path_csv", "delimited_text"


def _word(content, budget, issues):
    check_archive(content)
    document = _docx().Document(io.BytesIO(content))
    _word_extra_coverage(document, issues)
    paragraphs = []
    for number, paragraph in enumerate(document.paragraphs, 1):
        budget.text(paragraph.text)
        paragraphs.append({"source_paragraph": number, "text": paragraph.text})
        if number > MAX_ROWS:
            raise TableInputError("ocr.table_row_limit")
    characters = sum(len(paragraph["text"]) for paragraph in paragraphs)
    text_parts = ["\n".join(paragraph["text"] for paragraph in paragraphs if paragraph["text"].strip())]
    tables, headers, table_rows = [], [], []
    for table_index, table in enumerate(document.tables, 1):
        evidence = []
        text_parts.append(f"\n[table {table_index}]")
        for row_number, row in enumerate(table.rows, 1):
            if row_number > MAX_ROWS:
                raise TableInputError("ocr.table_row_limit")
            cells = [cell.text for cell in row.cells]
            budget.row(cells)
            evidence.append(_row(row_number, cells))
            text_parts.append(" | ".join(cells))
        tables.append({"table_index": table_index, "rows": evidence})
        if table_index == 1 and evidence:
            headers = _headers(evidence[0]["cells"])
            _coverage(headers, evidence[1:], issues, 1)
            for raw in evidence[1:]:
                cells = (raw["cells"] + [""] * len(headers))[:len(headers)]
                row_dict = dict(zip(headers, cells))
                if any(value.strip() for value in row_dict.values()):
                    raw["table_row_index"] = len(table_rows) + 1
                    table_rows.append(row_dict)
    page = _page(1, "\n".join(text_parts).strip() or "(empty docx)", headers or None, table_rows or None)
    if page.table_headers is not None and page.table_rows is not None:
        if any(paragraph["text"].strip() for paragraph in paragraphs):
            _issue(issues, "docx_paragraphs_outside_primary_grid", 1)
        if any(cell.strip() for table in tables[1:] for row in table["rows"] for cell in row["cells"]):
            _issue(issues, "docx_tables_outside_primary_grid", 1)
    unit = {"page_number": 1, "source_kind": "word_document", "paragraphs": paragraphs,
            "tables": tables, "header_source_row": 1 if headers else None,
            "rows": tables[0]["rows"] if tables else []}
    return [page], [unit], characters, "table_path_docx", "word_document"


def _text(content, legacy_doc, budget, issues):
    text, encoding = _decode(content)
    if legacy_doc:
        _issue(issues, "legacy_doc_best_effort", 1)
        text = "".join(char for char in text if char.isprintable() or char in "\n\t ")
    budget.text(text)
    # .doc retains the old python-docx paragraph attempt for billing only. A
    # real binary Word file takes byte_length//2; recognition remains best effort.
    if legacy_doc:
        if zipfile.is_zipfile(io.BytesIO(content)):
            check_archive(content)
            try:
                document = _docx().Document(io.BytesIO(content))
                characters = sum(len(paragraph.text) for paragraph in document.paragraphs)
            except TableInputError:
                raise
            except Exception:
                characters = len(content) // 2
        else:
            characters = len(content) // 2
    else:
        characters = len(content.decode("utf-8", errors="ignore"))
    lines = text.splitlines(keepends=True)
    if len(lines) > MAX_ROWS:
        raise TableInputError("ocr.table_row_limit")
    page = _page(1, (text or "(legacy .doc — limited extraction)") if legacy_doc else text,
                 confidence=.7 if legacy_doc else 1.0)
    unit = {"page_number": 1, "source_kind": "decoded_text", "encoding": encoding,
            "line_basis": "decoded_printable_text" if legacy_doc else "decoded_text",
            "rows": [_row(number, [line]) for number, line in enumerate(lines, 1)]}
    return [page], [unit], characters, "table_path_doc" if legacy_doc else "table_path_txt", "decoded_text"


def extract_table(file_bytes: bytes, filename: str) -> TableDocument:
    """Decode in memory. Logical sections are never PDF physical page numbers.

    Character units can be zero. Errors return no document and must never be
    charged. The caller verifies source_sha256 against its authoritative original.
    """
    if not isinstance(file_bytes, bytes):
        raise TableInputError("ocr.table_invalid_bytes")
    try:
        meta = validate_input_metadata(filename, len(file_bytes))
    except OcrInputError as error:
        raise TableInputError(error.code) from error
    extension = file_extension(filename)
    if meta["family"] != "table" or extension not in TABLE_EXTENSIONS:
        raise TableInputError("ocr.table_format_not_supported")
    budget, issues = DecodeBudget(), []
    try:
        if extension in {".xlsx", ".xlsm", ".xls"}:
            decoded = _excel(file_bytes, budget, issues)
        elif extension in {".csv", ".tsv"}:
            decoded = _csv(file_bytes, "\t" if extension == ".tsv" else ",", budget, issues)
        elif extension == ".docx":
            decoded = _word(file_bytes, budget, issues)
        else:
            decoded = _text(file_bytes, extension == ".doc", budget, issues)
    except TableInputError:
        raise
    except Exception as error:
        raise TableInputError("ocr.table_parse_failed") from error
    pages, units, characters, engine, kind = decoded
    provenance = {"version": 1, "reference_kind": "logical_section", "unit_kind": kind,
                  "coverage_complete": not issues, "coverage_issues": issues, "units": units}
    return TableDocument(tuple(pages), len(pages), engine, characters, provenance,
                         hashlib.sha256(file_bytes).hexdigest())
