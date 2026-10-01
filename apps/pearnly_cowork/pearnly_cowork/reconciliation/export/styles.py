"""Shared palette and style helpers of the bank-reconciliation workbook.

Source: ec2e5c957e6b ``services/recon/bank_recon_excel_styles.py``. Sheets 2/5/6
share these header, border and date helpers; sheet 1 keeps its own palette. Pure
presentation: no database, clock, file, network or Frappe access.
"""

from copy import copy
from datetime import date
from typing import Optional

from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ── Color palette ──────────────────────────────────────────────────
COLOR_HEADER = "2D6A4F"  # dark green
COLOR_SUBHEAD = "52B788"  # medium green
COLOR_MATCHED = "D8F3DC"  # light green
COLOR_L2 = "FFF3CD"  # amber (date tolerance)
COLOR_L3 = "FFE0CC"  # orange (amount only)
COLOR_L4 = "DBEAFE"  # light blue (preset group total)
COLOR_GL_ONLY = "E8D5F5"  # purple
COLOR_ST_ONLY = "D4E6F1"  # blue
COLOR_DIFF = "FFDAD6"  # red for non-zero diff
COLOR_OK = "D8F3DC"  # green for zero diff
COLOR_ROW_ALT = "F8F9FA"  # alternating row


def hdr_style(ws, row, col, text, color=COLOR_HEADER, bold=True, size=10):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(bold=bold, color="FFFFFF", size=size)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def border_range(ws, min_row, max_row, min_col, max_col):
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    # 加边框后的样式下标只取决于格子原来的样式：同一种只算一次，其余直接用算好的下标（结果与逐格赋值相同；
    # 逐格赋值每次都按值比对一遍样式表，两千多行的银行对账光这一步就几秒）。
    done = {}
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            key = tuple(cell._style) if cell._style is not None else None  # 合并区里的格子可能还没有样式
            if key is not None and key in done:
                cell._style = copy(done[key])
                continue
            cell.border = border
            if key is not None:
                done[key] = copy(cell._style)


def fmt_date(d: Optional[date]) -> str:
    if d is None:
        return ""
    return d.strftime("%d/%m/%Y")
