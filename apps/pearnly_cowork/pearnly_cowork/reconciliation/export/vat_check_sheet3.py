# 销项税报告核查（Cowork，用户 2026-09-26 定迁）：核查工作簿第 3 页：逐张对照（services/vat/vat_excel_sheet3.py）
# 来源：旧线上 ec2e5c957e6b4be7dee93ffdd2dae9fa83f69188 的同名逻辑，逐字移植，只改导入路径。
# 有意改动：作废行（配对 kind=void、没有发票的作废行 void_rep）显示「作废一致」/「作废（无需发票）」、绿底，不计入差异与缺失，
# 汇总另记 n_void（用户 2026-09-30 定；判定见 vat_check_core `_void_pair`，词条在 i18n/messages.tsv）。
"""Sheet 3 · 一对一对账结果(配对引擎 + 状态判定)· vat_excel_build 拆分；Cowork 去掉了标题与 KPI 卡。"""

from typing import List, Dict

from frappe import _
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from pearnly_cowork.reconciliation.normalization import parse_date
from pearnly_cowork.reconciliation.vat_check_core import (
    _derive_period,
    _build_recon_pairs,
    _eq_amount,
    _get_inv_total,
    _get_rep_total,
    _diff_dims,
    diff_labels,
)
from pearnly_cowork.reconciliation.export.vat_check_styles import (
    FONT_NAME,
    F_HEAD,
    F_NORM,
    F_DIFF_RED,
    FILL_HEAD,
    FILL_OK,
    FILL_DIFF,
    FILL_MISS,
    FILL_FUZZY,
    FILL_OCRMSG,
    BORDER_TH,
    AL_C,
    AL_R,
    AL_L,
    HEAD_HEIGHT,
    ROW_HEIGHT,
)


def _void_status(key, lang) -> str:
    """作废两种状态的文字：词条在术语表（``vat_check.status.*``），前面照其它「一致」状态加「✓ 」。"""
    return "✓ " + _(key, lang=lang)


def _build_sheet3(wb, invoices, report_rows, client_name, period_year, period_month, L, lang="th"):
    # ════════════ Sheet 3 · 对账结果(一对一 15 列 · 后端预算结果)════════════
    _lab = diff_labels(lang)  # 逐行差异/金额备注本地化标签(跟用户语言)
    ws3 = wb.create_sheet(L["sh3"])
    ws3.sheet_properties.tabColor = "D97706"  # Tab 橙

    # 跑配对(Bug 2/3/4/5 全在这里)
    match_result = _build_recon_pairs(invoices, report_rows, lang)
    pairs = match_result["pairs"]
    unmatched_inv = match_result["unmatched_inv"]
    unmatched_rep = match_result["unmatched_rep"]
    void_rep = match_result["void_rep"]

    # 计算 KPI（作废的单独计数，不算差异也不算缺失）
    n_total = len(pairs) + len(unmatched_inv) + len(unmatched_rep) + len(void_rep)
    n_void = sum(1 for _p in pairs if _p["kind"] == "void") + len(void_rep)
    n_ok = 0
    for _p in pairs:
        _inv = invoices[_p["inv_idx"]]
        _rep = report_rows[_p["rep_idx"]]
        if not _eq_amount(_get_inv_total(_inv), _get_rep_total(_rep)):
            continue
        if _p["kind"] == "matched_cash":
            n_ok += 1
        elif _p["kind"] == "matched":
            _d = _diff_dims(_inv, _rep, lang)
            if not any(_d.values()):
                n_ok += 1
    n_diff = n_total - n_ok - n_void
    diff_amount_total = 0.0
    for p in pairs:
        if p["kind"] in ("matched", "matched_cash"):
            a = _get_inv_total(invoices[p["inv_idx"]])
            b = _get_rep_total(report_rows[p["rep_idx"]])
            if not _eq_amount(a, b):
                diff_amount_total += abs(a - b)
    for ii in unmatched_inv:
        diff_amount_total += _get_inv_total(invoices[ii])
    for ri in unmatched_rep:
        diff_amount_total += _get_rep_total(report_rows[ri])

    # Cowork 统一复核表（用户 2026-09-26）：不写标题、期间说明与 KPI 卡，表头就在第 1 行（期间见各行“期间”列）。
    # 表头(15 列)
    HEADER_ROW = 1
    DATA_START = 2
    headers3 = L["h_recon"]
    for c, h in enumerate(headers3, 1):
        cell = ws3.cell(row=HEADER_ROW, column=c, value=h)
        cell.font = F_HEAD
        cell.fill = FILL_HEAD
        cell.alignment = AL_C
        cell.border = BORDER_TH
    ws3.row_dimensions[HEADER_ROW].height = HEAD_HEIGHT

    # ── 写数据行 · 配对 + 孤儿
    def _status_for(pair, dims) -> tuple:
        """v118.32.4.10.0 · 维度感知状态判定 · 返回 (status_key, text, fill)"""
        kind = pair["kind"]
        if kind == "void":
            return "st_void", _void_status("vat_check.status.void_match", lang), FILL_OK
        if kind == "fuzzy":
            return "st_fuzzy", L["st_fuzzy"], FILL_FUZZY
        if kind == "ocr_missing":
            return "st_ocr_missing", L["st_ocr_missing"], FILL_OCRMSG
        if dims.get("tax_id", "").startswith("~"):
            return "st_fuzzy", L["st_fuzzy"], FILL_FUZZY
        if kind == "matched_cash":
            a = _get_inv_total(invoices[pair["inv_idx"]])
            b = _get_rep_total(report_rows[pair["rep_idx"]])
            if _eq_amount(a, b):
                return "st_cash", L["st_cash"], FILL_OK
            return "st_diff", L["st_diff"], FILL_DIFF
        diff_keys = [k for k, v in dims.items() if v and not (k == "tax_id" and v.startswith("~"))]
        a = _get_inv_total(invoices[pair["inv_idx"]])
        b = _get_rep_total(report_rows[pair["rep_idx"]])
        if not _eq_amount(a, b):
            diff_keys.append("_amt")
        n = len(diff_keys)
        if n == 0:
            return "st_ok", L["st_ok"], FILL_OK
        if n >= 2:
            return "st_multi_diff", L["st_multi_diff"], FILL_DIFF
        key = diff_keys[0]
        if key == "_amt":
            return "st_diff", L["st_diff"], FILL_DIFF
        if key in ("date", "period"):
            return "st_date_diff", L["st_date_diff"], FILL_DIFF
        if key == "branch":
            return "st_branch_diff", L["st_branch_diff"], FILL_DIFF
        if key == "name":
            return "st_name_diff", L["st_name_diff"], FILL_DIFF
        if key == "inv_no":
            return "st_inv_no_diff", L["st_inv_no_diff"], FILL_DIFF
        if key == "tax_id":
            return "st_tax_id_diff", L["st_tax_id_diff"], FILL_DIFF
        return "st_diff", L["st_diff"], FILL_DIFF

    row_cursor = DATA_START
    seq_no = 0
    diff_col_indices = (9, 10, 11, 12, 13, 14)  # 6 维度差异列(含税号差)
    F_FUZZY_BLUE = Font(name=FONT_NAME, size=10, color="2563EB")
    task_rows: List[Dict] = []  # v118.32.4.10.0 · 任务摘要行

    for pair in pairs:
        seq_no += 1
        inv = invoices[pair["inv_idx"]]
        rep = report_rows[pair["rep_idx"]]
        dims = _diff_dims(inv, rep, lang)
        # Bug 2 · 散客:不比较/不显示分公司和税号差
        if pair["kind"] == "matched_cash":
            dims["branch"] = ""
            dims["tax_id"] = ""
        # 作废一致：单号掉字母、金额不同都是作废造成的，不列差异
        if pair["kind"] == "void":
            dims = {k: "" for k in dims}
        # tax_id fuzzy ~ 前缀处理(显示时去掉 · 用蓝色替换红色)
        tax_id_raw = dims.get("tax_id", "")
        tax_id_is_fuzzy = tax_id_raw.startswith("~")
        tax_id_display = tax_id_raw[1:] if tax_id_is_fuzzy else tax_id_raw
        status_key, status_text, row_fill = _status_for(pair, dims)
        amt_inv = _get_inv_total(inv)
        amt_rep = _get_rep_total(rep)
        amt_diff = 0 if pair["kind"] == "void" else round(amt_inv - amt_rep, 2)
        period_inv = _derive_period(inv.get("invoice_date") or "", inv.get("period") or "")

        # v4.10.13 · 净额/VAT 分字段差异备注
        amt_pre_inv = float(inv.get("amount_pre_vat") or 0)
        amt_vat_inv = float(inv.get("vat_amount") or 0)
        amt_pre_rep = float((rep.get("report_amount_pre_vat") or rep.get("report_amount")) or 0)
        amt_vat_rep = float(rep.get("report_vat_amount") or 0)
        pre_diff = round(amt_pre_inv - amt_pre_rep, 2)
        vat_diff_ = round(amt_vat_inv - amt_vat_rep, 2)
        if not _eq_amount(amt_inv, amt_rep) and pair["kind"] != "void":
            if not _eq_amount(pre_diff, 0) and not _eq_amount(vat_diff_, 0):
                _amt_note = (
                    f"{_lab['amt_pre_diff']} {pre_diff:+,.2f} · "
                    f"{_lab['amt_vat_diff']} {vat_diff_:+,.2f}"
                )
            elif not _eq_amount(pre_diff, 0):
                _amt_note = f"{_lab['amt_pre_diff']} {pre_diff:+,.2f} · {_lab['amt_vat_ok']}"
            else:
                _amt_note = f"{_lab['amt_vat_diff']} {vat_diff_:+,.2f} · {_lab['amt_pre_ok']}"
        else:
            _amt_note = ""

        # v4.10.13 · ocr_missing · 备注追加提醒(dims 原样保留)
        _base_note = pair.get("note") or ""
        if pair["kind"] == "ocr_missing":
            _base_note = (_base_note + " · " + _lab["ocr_incomplete"]).lstrip(" · ")
        _note_val = " · ".join(filter(None, [_base_note, _amt_note]))

        values = [
            seq_no,  # 1 #
            status_text,  # 2 status
            inv.get("buyer_name") or rep.get("report_buyer_name") or "",  # 3 客户
            inv.get("invoice_no") or "",  # 4 发票号
            period_inv,  # 5 期间
            amt_inv,  # 6 金额(发)
            amt_rep,  # 7 金额(报)
            amt_diff,  # 8 差异金额
            dims["inv_no"],  # 9 发票号差
            dims["date"],  # 10 日期差
            dims["period"],  # 11 期间差
            tax_id_display,  # 12 税号差(新)
            dims["branch"],  # 13 分公司差
            dims["name"],  # 14 客户名差
            _note_val,  # 15 备注
        ]
        for c, v in enumerate(values, 1):
            cell = ws3.cell(row=row_cursor, column=c, value=v)
            cell.font = F_NORM
            cell.border = BORDER_TH
            cell.fill = row_fill
            if c in (6, 7, 8):
                cell.alignment = AL_R
                cell.number_format = "#,##0.00"
            elif c == 1:
                cell.alignment = AL_C
            else:
                cell.alignment = AL_L
            # 差异维度列若有内容 · 税号 fuzzy 蓝 · 其余红
            if c in diff_col_indices and v:
                if c == 12 and tax_id_is_fuzzy:
                    cell.font = F_FUZZY_BLUE
                else:
                    cell.font = F_DIFF_RED
        # v118.32.4.10.0 · 收集摘要行
        task_rows.append(
            {
                "status_key": status_key,
                "status": status_text,
                "customer": inv.get("buyer_name") or rep.get("report_buyer_name") or "",
                "invoice_no": inv.get("invoice_no") or "",
                "amount_inv": amt_inv,
                "amount_rep": amt_rep,
                "dims": {
                    k: (v[1:] if k == "tax_id" and v.startswith("~") else v)
                    for k, v in dims.items()
                },
                "kind": pair["kind"],
            }
        )
        ws3.row_dimensions[row_cursor].height = ROW_HEIGHT
        row_cursor += 1

    # 发票孤儿(报告无)
    for ii in unmatched_inv:
        seq_no += 1
        inv = invoices[ii]
        amt_inv = _get_inv_total(inv)
        period_inv = _derive_period(inv.get("invoice_date") or "", inv.get("period") or "")
        values = [
            seq_no,
            L["st_no_rep"],
            inv.get("buyer_name") or "",
            inv.get("invoice_no") or "",
            period_inv,
            amt_inv,
            0,
            amt_inv,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
        for c, v in enumerate(values, 1):
            cell = ws3.cell(row=row_cursor, column=c, value=v)
            cell.font = F_NORM
            cell.border = BORDER_TH
            cell.fill = FILL_MISS
            if c in (6, 7, 8):
                cell.alignment = AL_R
                cell.number_format = "#,##0.00"
            elif c == 1:
                cell.alignment = AL_C
            else:
                cell.alignment = AL_L
        task_rows.append(
            {
                "status_key": "st_no_rep",
                "status": L["st_no_rep"],
                "customer": inv.get("buyer_name") or "",
                "invoice_no": inv.get("invoice_no") or "",
                "amount_inv": amt_inv,
                "amount_rep": 0,
                "dims": {},
                "kind": "invoice_orphan",
            }
        )
        ws3.row_dimensions[row_cursor].height = ROW_HEIGHT
        row_cursor += 1

    # 报告孤儿(发票无)
    for ri in unmatched_rep:
        seq_no += 1
        rep = report_rows[ri]
        amt_rep = _get_rep_total(rep)
        rep_date = rep.get("report_date") or ""
        d = parse_date(rep_date)
        period_rep = f"{d.month:02d}/{d.year}" if d else ""
        values = [
            seq_no,
            L["st_no_inv"],
            rep.get("report_buyer_name") or "",
            rep.get("report_invoice_no") or "",
            period_rep,
            0,
            amt_rep,
            -amt_rep,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
        for c, v in enumerate(values, 1):
            cell = ws3.cell(row=row_cursor, column=c, value=v)
            cell.font = F_NORM
            cell.border = BORDER_TH
            cell.fill = FILL_MISS
            if c in (6, 7, 8):
                cell.alignment = AL_R
                cell.number_format = "#,##0.00"
            elif c == 1:
                cell.alignment = AL_C
            else:
                cell.alignment = AL_L
        task_rows.append(
            {
                "status_key": "st_no_inv",
                "status": L["st_no_inv"],
                "customer": rep.get("report_buyer_name") or "",
                "invoice_no": rep.get("report_invoice_no") or "",
                "amount_inv": 0,
                "amount_rep": amt_rep,
                "dims": {},
                "kind": "report_orphan",
            }
        )
        ws3.row_dimensions[row_cursor].height = ROW_HEIGHT
        row_cursor += 1

    # 作废行没有发票（Cowork 有意改动）：作废本来就不需要发票，绿底、不算缺失
    for ri in void_rep:
        seq_no += 1
        rep = report_rows[ri]
        amt_rep = _get_rep_total(rep)
        d = parse_date(rep.get("report_date") or "")
        period_rep = f"{d.month:02d}/{d.year}" if d else ""
        status_text = _void_status("vat_check.status.void_no_invoice", lang)
        values = [seq_no, status_text, rep.get("report_buyer_name") or "", rep.get("report_invoice_no") or "",
                  period_rep, 0, amt_rep, 0, "", "", "", "", "", "", ""]
        for c, v in enumerate(values, 1):
            cell = ws3.cell(row=row_cursor, column=c, value=v)
            cell.font = F_NORM
            cell.border = BORDER_TH
            cell.fill = FILL_OK
            if c in (6, 7, 8):
                cell.alignment = AL_R
                cell.number_format = "#,##0.00"
            elif c == 1:
                cell.alignment = AL_C
            else:
                cell.alignment = AL_L
        task_rows.append(
            {
                "status_key": "st_void_no_inv",
                "status": status_text,
                "customer": rep.get("report_buyer_name") or "",
                "invoice_no": rep.get("report_invoice_no") or "",
                "amount_inv": 0,
                "amount_rep": amt_rep,
                "dims": {},
                "kind": "report_void",
            }
        )
        ws3.row_dimensions[row_cursor].height = ROW_HEIGHT
        row_cursor += 1

    # 列宽
    widths3 = [5, 22, 26, 16, 10, 14, 14, 13, 22, 14, 16, 22, 18, 26, 22]
    for i, w in enumerate(widths3, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = f"A{DATA_START}"
    return n_total, n_ok, n_diff, diff_amount_total, task_rows, n_void
