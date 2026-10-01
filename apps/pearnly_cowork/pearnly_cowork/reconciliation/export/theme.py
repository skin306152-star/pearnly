"""三种对账复核表的统一外观（用户 2026-09-26：“统一框架，统一视觉效果”，不动逻辑）。

银行对账、收入对账、销项税报告核查的工作簿各自从旧线上移植，配色、字体、字号、行高、边框、数字格式各不相同。
这里在每本工作簿生成之后、保存之前统一走一遍，只改外观，不改任何单元格的值（日期只换显示写法）：

- 字体一种（Arial），正文 10 号，表头 10 号加粗，标题 12 号加粗；
- 表头统一深色底白字、居中；合计 / 期末这类强调行统一浅灰底加粗；分区标题同一浅灰底；
- 状态只有三种颜色，三张表同一口径：绿 = 一致，黄 = 有差异 / 待确认，红 = 缺一边 / 没找到；
  只染状态所在的格，整行染色的表改成只染状态列；不用斜体；
  各表原来的颜色按各自含义归到这三种（同一个颜色在不同表里意思不同，所以按表分别对照）；
- 表格区域一律细灰边框，不用斑马纹、不显示网格线、不设标签页颜色；
- 行高统一（标题 24、表头 22、其余 18），数据格不换行；列宽不超过 48 个字符；
- 金额统一 `#,##0.00`，负数红色带负号；`年-月-日` 的日期文字统一成 `日/月/年`。

页面上的复核表由同一本工作簿逐格画出（`recon_center/sheets.py`），所以页面与下载的 Excel 同步变。
"""
import re
from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "Arial"
TEXT, WHITE = "111827", "FFFFFF"
HEADER_BG, EMPHASIS_BG, LINE = "1F2937", "F3F4F6", "D1D5DB"
STATUS = {"ok": ("DCFCE7", "166534"), "warn": ("FEF3C7", "92400E"), "miss": ("FEE2E2", "991B1B")}
MONEY = "#,##0.00;[Red]-#,##0.00"
TITLE_HEIGHT, HEADER_HEIGHT, ROW_HEIGHT = 24, 22, 18
MAX_WIDTH = 48  # 列宽上限（字符）：汇总表的“项目说明”原来 78 / 90 宽，金额列被挤出屏幕
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")

# 各表原来的底色 → 角色。dark = 表头 / 标题 / 强调行（按位置区分）；None = 去掉底色。
_COMMON = {"FFFFFF": None, "F8F9FA": None, "F9FAFB": None, "FAFBFC": None}
ROLES = {
    "bank": {**_COMMON, "1F2937": "dark", "374151": "dark", "2D6A4F": "dark", "52B788": "dark",
             "D8F3DC": "ok", "D1FAE5": "ok",
             "FFF3CD": "warn", "FFE0CC": "warn", "FFE082": "warn", "FEF3C7": "warn",
             "FFDAD6": "miss", "FEE2E2": "miss", "E8D5F5": "miss", "F3E8FF": "miss", "D4E6F1": "miss",
             "EBF5FB": "miss", "FEF2F2": None,
             "EEF2F6": "emphasis", "E5E7EB": "emphasis", "DBEAFE": "emphasis"},
    "income": {**_COMMON, "1A3C5E": "dark", "D6F5E3": "ok", "DCFCE7": "ok", "FDE8E8": "warn",
               "FFF3CD": "miss", "FFEDD5": "miss", "EFF6FF": "emphasis"},
    "vat": {**_COMMON, "2563EB": "dark", "DCFCE7": "ok", "FEF3C7": "warn", "DBEAFE": "warn",
            "FED7AA": "warn", "FEE2E2": "miss", "EFF6FF": "emphasis"},
}
# 整行染色的页：状态落在哪一列（按工作簿里的页序，列从 1 数）。只染这一列，与只染状态格的页一致。
STATUS_COLUMN = {"bank": {1: 1}, "vat": {0: 11, 2: 2}}
# 原来的字色 → 角色（状态字色跟着三种颜色走，其余一律正文色）。
_FONT_ROLES = {"16A34A": "ok", "065F46": "ok", "166534": "ok",
               "D97706": "warn", "B45309": "warn", "92400E": "warn", "EA580C": "warn",
               "DC2626": "miss", "991B1B": "miss", "B91C1C": "miss"}


def _rgb(color):
    value = getattr(color, "rgb", None) if color is not None else None
    return value[-6:].upper() if isinstance(value, str) and len(value) >= 6 else None


def _fill_role(cell, roles):
    if cell.fill is None or cell.fill.fill_type != "solid":
        return None
    return roles.get(_rgb(cell.fill.fgColor))


def _solid(color):
    return PatternFill("solid", fgColor=color)


def _row_has_value(sheet, row, width):
    cells = (sheet._cells.get((row, column)) for column in range(1, width + 1))
    return any(cell is not None and cell.value not in (None, "") for cell in cells)


def _is_title(sheet, row, width):
    """标题行：第一行、整行合并成一格（银行汇总表的“银行对账 · 银行名”）。"""
    return row == 1 and any(m.min_row == m.max_row == 1 and m.min_col == 1 and m.max_col == width
                            for m in sheet.merged_cells.ranges)


def _style_sheet(sheet, roles, status_column=1, done=None):
    done = {} if done is None else done
    width, height = sheet.max_column, sheet.max_row
    line = Side(style="thin", color=LINE)
    border = Border(left=line, right=line, top=line, bottom=line)
    header_seen = False
    for row in range(1, height + 1):
        has_value = _row_has_value(sheet, row, width)
        if not has_value:
            # 空行（间隔行、冻结窗格顺手建出的空格）：只清掉已有格子的底色与边框，不新建格子。
            for column in range(1, width + 1):
                existing = sheet._cells.get((row, column))
                if existing is not None and existing.has_style:
                    existing.fill, existing.border = PatternFill(fill_type=None), Border()
            continue
        title = _is_title(sheet, row, width)
        dark = has_value and not title and any(
            _fill_role(sheet.cell(row, c), roles) == "dark" for c in range(1, width + 1)
            if sheet.cell(row, c).value not in (None, ""))
        # 第一条深色行是表头；表头之后的深色行（期末余额、合计）是强调行。
        header = dark and not header_seen
        header_seen = header_seen or header
        sheet.row_dimensions[row].height = TITLE_HEIGHT if title else HEADER_HEIGHT if header else ROW_HEIGHT
        # 整行染色（银行对账结果、销项税对账结果、发票校验不过的行）改成只染状态列。
        row_roles = [_fill_role(sheet.cell(row, c), roles) for c in range(1, width + 1)]
        tinted = next((r for r in STATUS if row_roles.count(r) >= max(2, 0.6 * width)), None)
        for column in range(1, width + 1):
            cell = sheet.cell(row, column)
            role = row_roles[column - 1]
            if tinted and role == tinted and column != status_column:
                role = None
            numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            if isinstance(cell.value, str):
                match = _ISO_DATE.match(cell.value.strip())
                if match:
                    cell.value = f"{match.group(3)}/{match.group(2)}/{match.group(1)}"
            # 下面每一项外观只取决于这格原来的样式和这几个判断，所以同一组合第一次算出的样式下标原样复用，
            # 结果与逐格新建样式对象完全相同（openpyxl 按值去重，同样的字体 / 底色 / 边框拿到的就是同一个下标）。
            # 原来每格都新建对象、逐个按值比对，一本几千行的收入对账要几十秒，页面打开就卡在这里（2026-10-01）。
            key = (tuple(cell._style), title, header, dark, role, numeric) if cell._style is not None else None
            if key is not None and key in done:
                cell._style = copy(done[key])
                continue
            font, align = cell.font, cell.alignment
            if numeric and "0.00" in (cell.number_format or ""):
                cell.number_format = MONEY
            if title:
                cell.fill = PatternFill(fill_type=None)
                cell.font = Font(name=FONT, size=12, bold=True, color=TEXT)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = Border()
                if key is not None:
                    done[key] = copy(cell._style)
                continue
            if header:
                cell.fill = _solid(HEADER_BG)
                cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif dark or role == "emphasis":
                cell.fill = _solid(EMPHASIS_BG)
                cell.font = Font(name=FONT, size=10, bold=True, color=TEXT)
            elif role in STATUS:
                background, foreground = STATUS[role]
                cell.fill = _solid(background)
                cell.font = Font(name=FONT, size=10, bold=bool(font and font.b), color=foreground)
            else:
                cell.fill = PatternFill(fill_type=None)
                status = _FONT_ROLES.get(_rgb(font.color) if font else None)
                cell.font = Font(name=FONT, size=10, bold=bool(font and font.b),
                                 color=STATUS[status][1] if status else TEXT)
            if not header:
                horizontal = align.horizontal if align is not None else None
                cell.alignment = Alignment(horizontal=horizontal or ("right" if numeric else "left"),
                                           vertical="center", indent=align.indent if align is not None else 0)
            cell.border = border
            if key is not None:
                done[key] = copy(cell._style)
    for dimension in sheet.column_dimensions.values():
        if dimension.width and dimension.width > MAX_WIDTH:
            dimension.width = MAX_WIDTH
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = None


def apply(book, kind):
    """把一本对账工作簿统一成同一套外观。kind：bank / income / vat。"""
    roles, done = ROLES[kind], {}  # 样式下标是整本工作簿共用的，复用表跟着工作簿走
    for index, sheet in enumerate(book.worksheets):
        _style_sheet(sheet, roles, STATUS_COLUMN.get(kind, {}).get(index, 1), done)
    return book
