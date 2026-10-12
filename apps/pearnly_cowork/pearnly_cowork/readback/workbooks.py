"""回读出的四份工作簿（2026-10 CC 闭环 L1）：总账、销项税报、进项税报（Express 报表的版式，见 `layout`）与差异表。

- 每份一个文件、报表就是第一张表：税报解析器只读第一张表，收入对账读全部表，混放会互相串。
- 只写值不写公式（解析器按缓存值读，openpyxl 写的公式没有缓存值）：文字一律按文字存，以「=」开头的（对方名、说明来自
  第三方发票的识别结果）也不会变成公式；Excel 不收的控制字符去掉。日期存公历日期、按佛历显示；金额两位小数。
- 总账照 Express 的八列（วันที่ / สมุด / ใบสำคัญ / คำอธิบาย / เดบิท / เครดิต / สถานะ / ยอดคงเหลือ），每个科目一行段头
  （「科目号 科目名」+ 期初在余额列）、本月的行（同一天只印第一行的日期）、一行「รวม」；另在最右加一列科目号，收入对账
  据此挑收入科目的行。对账三个解析器（银行 `gl_sheets`、收入 `income_parse`、税报 `vat_parse`）原样接受，见
  `tests/test_cowork_readback_workbooks.py`。
- 差异表的表头按请求人的语言（调用方传入译好的词）。
"""
from __future__ import annotations

import datetime as dt
import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from pearnly_cowork.readback import contracts, layout

BOLD = Font(bold=True)


def _append(sheet, values) -> None:
    """写一行：文字去掉 Excel 不收的控制字符；以「=」开头的文字按文字存（openpyxl 会当成公式）。"""
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _save(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _be(day: dt.date) -> int:
    return day.year + 543


def general_ledger(company: str, period: str, accounts: List[Dict[str, Any]], printed: dt.date) -> bytes:
    """总账（`contracts.ledger` 的结果）。"""
    first, last = contracts.month_range(period)
    month = layout.MONTHS[first.month]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = layout.GL_SHEET
    _append(sheet, [company, None, None, None, None, None, None, layout.GL_PAGE])
    _append(sheet, [layout.GL_TITLE])
    _append(sheet, [layout.GL_FROM, f"1 {month} {_be(first)}      {layout.GL_TO}  {last.day} {month} {_be(last)}",
                    None, None, None, None, None,
                    f"{layout.GL_PRINTED} {printed.day:02d}/{printed.month:02d}/{_be(printed) % 100:02d}"])
    codes = [account["code"] for account in accounts]
    _append(sheet, [layout.GL_ACCOUNTS, f"{codes[0]}         {layout.GL_TO}  {codes[-1]}" if codes else "",
                    None, None, None, None, None, layout.GL_DEPARTMENT])
    _append(sheet, list(layout.GL_HEADER))
    for cell in sheet[sheet.max_row]:
        cell.font = BOLD
    for account in accounts:
        code = account["code"]
        _append(sheet, [f"{code} {account['name'] or code}", None, None, None, None, None, None, account["opening"]])
        shown = None
        for line in account["lines"]:
            day = line["date"] if line["date"] != shown else None
            shown = line["date"]
            _append(sheet, [day, layout.BOOKS.get(line["book"], line["book"]), line["voucher"], line["descrp"],
                            line["debit"], line["credit"], None, line["balance"], code])
            if day:
                sheet.cell(sheet.max_row, 1).number_format = layout.DATE_FORMAT
        _append(sheet, [None, None, None, layout.GL_TOTAL, account["debit"], account["credit"], None,
                        account["closing"]])
    for row in sheet.iter_rows(min_row=6):
        for cell in row[4:8]:
            if cell.value is not None:
                cell.number_format = layout.MONEY_FORMAT
    for column, width in zip("ABCDEFGHI", (12, 8, 16, 52, 16, 16, 8, 18, 14)):
        sheet.column_dimensions[column].width = width
    return _save(workbook)


def vat_report(side: str, company: Dict[str, str], period: str, rows: List[Dict[str, Any]]) -> bytes:
    """进项（P）/ 销项（S）税报（`contracts.vat_rows` 的结果）。"""
    first, _last = contracts.month_range(period)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = layout.VAT_SHEET
    branch = str(company.get("branch") or "")
    _append(sheet, [layout.VAT_TITLE[side]])
    _append(sheet, [f"{layout.VAT_MONTH} {layout.MONTHS_FULL[first.month]} {layout.VAT_YEAR} {_be(first)}"])
    _append(sheet, [layout.VAT_COMPANY, company.get("name") or ""])
    _append(sheet, [layout.VAT_TAX_ID, company.get("tax_id") or "",
                    layout.HEAD_OFFICE if branch in ("", "00000") else f"{layout.BRANCH} {branch}"])
    sheet.cell(4, 2).number_format = "@"
    _append(sheet, list(layout.VAT_HEADER[side]))
    for cell in sheet[sheet.max_row]:
        cell.font = BOLD
    for number, row in enumerate(rows, 1):
        name = f"{row['name']} {layout.VOID}".strip() if row["void"] else row["name"]
        _append(sheet, [number, row["date"], row["invoice"], row["ref"], name, row["tax_id"], row["branch"] or "00000",
                        row["base"], row["vat"], row["total"]])
        here = sheet.max_row
        sheet.cell(here, 2).number_format = layout.DATE_FORMAT
        for column in (3, 4, 6, 7):
            sheet.cell(here, column).number_format = "@"
        for column in (8, 9, 10):
            sheet.cell(here, column).number_format = layout.MONEY_FORMAT
    sums = contracts.totals(rows)
    _append(sheet, [layout.VAT_TOTAL, None, None, None, None, None, None, contracts.amount(sums["before_vat"]),
                    contracts.amount(sums["vat"]), contracts.amount(sums["total"])])
    for column in (8, 9, 10):
        sheet.cell(sheet.max_row, column).number_format = layout.MONEY_FORMAT
    for column, width in zip("ABCDEFGHIJ", (7, 12, 18, 16, 40, 16, 8, 16, 16, 16)):
        sheet.column_dimensions[column].width = width
    return _save(workbook)


def changes(found: Dict[str, Any], labels: Dict[str, Any]) -> bytes:
    """差异表：第一张「改过的」（推送过的单现在在 Express 里怎样了），第二张「直接录入」。``labels`` 是按请求人语言
    译好的表名、列名与 change / field 的值（键见 `service._labels`）。"""
    workbook = Workbook()
    first = workbook.active
    first.title = labels["sheet_changes"][:31]
    _append(first, [labels[key] for key in ("docnum", "direction", "change", "field", "pushed", "now")])
    for item in found["changes"]:
        _append(first, [item["docnum"], labels["directions"].get(item["direction"], item["direction"]),
                        labels["changes"].get(item["change"], item["change"]),
                        labels["fields"].get(item["field"], item["field"]), item["pushed"], item["now"]])
    second = workbook.create_sheet(labels["sheet_direct"][:31])
    _append(second, [labels[key] for key in ("docnum", "date", "kind", "party", "before_vat", "vat", "total",
                                             "entered_by")])
    for item in found["direct"]:
        _append(second, [item["docnum"], item["date"], item["kind"] or layout.BOOKS.get(item["book"], item["book"]),
                         item["party"], item["before_vat"], item["vat"], item["total"], item["entered_by"]])
    for sheet in (first, second):
        for cell in sheet[1]:
            cell.font = BOLD
        for column in "ABCDEFGH":
            sheet.column_dimensions[column].width = 18
    return _save(workbook)
