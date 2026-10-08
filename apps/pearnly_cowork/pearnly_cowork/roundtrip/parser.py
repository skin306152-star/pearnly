"""Bounded byte adapter for the deployed deterministic ERP workbook reader."""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
import hashlib
import io
import math
from typing import Any, Dict, List, Optional, Tuple

from pearnly_cowork.ocr.inputs import OcrInputError, validate_input_metadata
from pearnly_cowork.ocr.table_inputs.archive import check_archive
from pearnly_cowork.ocr.table_inputs.model import DecodeBudget, MAX_ROWS, TableInputError
from . import money, spec as rt
from .contracts import RoundtripNotMatched, RoundtripParseError

_BLANK_STREAK_STOP = 50

def _num(raw):
    """Legacy float parsing is a compatibility boundary, not money arithmetic."""
    if raw is None or isinstance(raw, bool):
        return None
    if not isinstance(raw, (int, float)):
        raw = str(raw).strip().replace(",", "").replace("฿", "")
        if not raw or raw.startswith("="):
            return None
    try:
        value = float(raw)
    except ValueError:
        return None
    if not math.isfinite(value):
        raise RoundtripParseError("ocr.roundtrip_invalid_number")
    return Decimal(str(value))

def _text(raw: Any) -> str:
    if raw is None:
        return ""
    if hasattr(raw, "strftime"):
        return raw.strftime("%Y-%m-%d")
    s = str(raw).strip()
    return "" if s.lower() in ("none", "null", "nan") else s

def _header_map(ws) -> Dict[str, int]:
    """表头行 → {列名: 列号}。重名列取最左一列(会计复制列时不至于串位)。
    列名按规格名记:任一导出语言的表头都认回同一列(见 spec.canonical)。"""
    out: Dict[str, int] = {}
    for col in range(1, (ws.max_column or 0) + 1):
        name = rt.canonical(_text(ws.cell(row=1, column=col).value))
        if name and name not in out:
            out[name] = col
    return out

def _cell(ws, row: int, hmap: Dict[str, int], header: str) -> Any:
    col = hmap.get(header)
    return ws.cell(row=row, column=col).value if col else None

def _row_is_blank(ws, row: int, hmap: Dict[str, int]) -> bool:
    return not any(_text(_cell(ws, row, hmap, h)) for h in hmap)

def _data_rows(ws, hmap: Dict[str, int]):
    """逐条数据行 -> (行号, 解出的回导键)。空行跳过,连续空行够多就停。"""
    streak = 0
    for row in range(2, (ws.max_row or 1) + 1):
        if _row_is_blank(ws, row, hmap):
            streak += 1
            if streak >= _BLANK_STREAK_STOP:
                if any(not _row_is_blank(ws, later, hmap) for later in range(row + 1, ws.max_row + 1)):
                    raise RoundtripParseError("ocr.roundtrip_trailing_rows", sheet=ws.title, stop_row=row)
                return
            continue
        streak = 0
        yield row, rt.decode_row_key(_text(_cell(ws, row, hmap, rt.COL_ROW_KEY)))

def _parse_sales_line(ws, row: int, hmap: Dict[str, int]) -> Dict[str, Any]:
    """销项一行 = 一个商品行。金额缺(公式无缓存)时按合同 数量×单价 补。"""
    qty = _num(_cell(ws, row, hmap, rt.SALES_COL_QTY))
    price = _num(_cell(ws, row, hmap, rt.SALES_COL_PRICE))
    # 会计改成死值就用他的;是公式(无缓存值)就按合同派生 —— 与写侧共用 erp_money
    amount = money.to_money(_cell(ws, row, hmap, rt.SALES_COL_AMOUNT))
    if amount is None:
        amount = money.line_amount(qty, price)
    return {
        "description": _text(_cell(ws, row, hmap, rt.SALES_COL_ITEM)),
        "qty": qty,
        "unit_price": price,
        "amount": amount,
        "erp_item_code": _text(_cell(ws, row, hmap, rt.COL_ERP_ITEM)),
    }

def _sales_doc_totals(items, pre_vat_cells, printed_vat) -> Tuple:
    """单据税基/税额/合计 —— 全走 erp_money,与写侧同一套舍入,不再各算各的。

    税额优先用票面印的那个(法定数字,要进 ภ.พ.30);只有会计改动使税基真的变了才重算。
    """
    explicit = [v for v in pre_vat_cells if v is not None]
    return money.doc_totals(
        [i["amount"] for i in items],
        explicit_base=money.sum_money(explicit) if explicit else None,
        printed_vat=printed_vat,
    )

def _read_sales_sheet(ws, hmap: Dict[str, int]) -> List[Dict[str, Any]]:
    """销项表按单据聚合(每行一个商品行)。聚合键优先用回导键里的 history_id ——
    会计改票号是这个流程的主要目的之一,按票号聚合会把改过号的行拆散。"""
    groups: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []
    for row, decoded in _data_rows(ws, hmap):
        invoice_no = _text(_cell(ws, row, hmap, rt.SALES_COL_INVOICE))
        # 键 > 票号 > 行号:三级兜底,任何一行都不会被静默并进别人的单据
        gkey = decoded[0] if decoded else (f"inv:{invoice_no}" if invoice_no else f"row:{row}")
        if gkey not in groups:
            groups[gkey] = {
                "history_id": decoded[0] if decoded else "",
                "invoice_number": invoice_no,
                "date": _text(_cell(ws, row, hmap, rt.SALES_COL_DATE)),
                "buyer_name": _text(_cell(ws, row, hmap, rt.SALES_COL_PARTY)),
                "buyer_tax": _text(_cell(ws, row, hmap, rt.COL_PARTY_TAX)),
                "buyer_branch": _text(_cell(ws, row, hmap, rt.COL_PARTY_BRANCH)),
                "erp_docnum": _text(_cell(ws, row, hmap, rt.COL_ERP_DOCNUM)),
                "erp_party_code": _text(_cell(ws, row, hmap, rt.COL_ERP_PARTY)),
                "items": [],
                "_pre_vat": [],
                "_doc_vat": money.to_money(_cell(ws, row, hmap, rt.COL_DOC_VAT)),
                "_rows": [],
            }
            order.append(gkey)
        g = groups[gkey]
        g["items"].append(_parse_sales_line(ws, row, hmap))
        g["_pre_vat"].append(_num(_cell(ws, row, hmap, rt.SALES_COL_PRE_VAT)))
        g["_rows"].append(row)

    out: List[Dict[str, Any]] = []
    for gkey in order:
        g = groups.pop(gkey)
        base, vat, total = _sales_doc_totals(g["items"], g.pop("_pre_vat"), g.pop("_doc_vat"))
        rows = g.pop("_rows")
        g.update({"amount_before_vat": base, "vat_amount": vat, "total_amount": total})
        out.append({"fields": g, "source_rows": rows})
    return out

def _read_purchase_sheet(ws, hmap: Dict[str, int]) -> List[Dict[str, Any]]:
    """进项表每行一张票(费用票整张进一个科目 · 不拆商品行)。"""
    out: List[Dict[str, Any]] = []
    for row, decoded in _data_rows(ws, hmap):
        base = money.to_money(_cell(ws, row, hmap, rt.PURCHASE_COL_PRE_VAT))
        printed = money.to_money(_cell(ws, row, hmap, rt.PURCHASE_COL_VAT)) or money.to_money(
            _cell(ws, row, hmap, rt.COL_DOC_VAT)
        )
        base, vat, total = money.doc_totals([], explicit_base=base, printed_vat=printed)
        if base == 0 and printed is None:  # 整行没有金额 · 不编造 0/0/0
            base = vat = total = None
        out.append(
            {
                "fields": {
                    "history_id": decoded[0] if decoded else "",
                    "invoice_number": _text(_cell(ws, row, hmap, rt.PURCHASE_COL_INVOICE)),
                    "date": _text(_cell(ws, row, hmap, rt.PURCHASE_COL_DATE)),
                    "seller_name": _text(_cell(ws, row, hmap, rt.PURCHASE_COL_PARTY)),
                    "seller_tax": (
                        _text(_cell(ws, row, hmap, rt.PURCHASE_COL_PARTY_TAX))
                        or _text(_cell(ws, row, hmap, rt.COL_PARTY_TAX))
                    ),
                    # 法定列 สาขา 优先(会计习惯填那里 · 写侧一直在写),回落回导列。
                    "seller_branch": (
                        _text(_cell(ws, row, hmap, rt.PURCHASE_COL_BRANCH))
                        or _text(_cell(ws, row, hmap, rt.COL_PARTY_BRANCH))
                    ),
                    "category": _text(_cell(ws, row, hmap, rt.PURCHASE_COL_CATEGORY)),
                    "amount_before_vat": base,
                    "vat_amount": vat,
                    "total_amount": total,
                    "erp_docnum": _text(_cell(ws, row, hmap, rt.COL_ERP_DOCNUM)),
                    "erp_party_code": _text(_cell(ws, row, hmap, rt.COL_ERP_PARTY)),
                    "items": [],
                },
                "source_rows": [row],
            }
        )
    return out

def _read_pending_sheet(ws, hmap: Dict[str, int]) -> List[Dict[str, Any]]:
    """待判表:会计没裁决的票。原样带出来给调用方如实汇报,不推。"""
    out: List[Dict[str, Any]] = []
    for row, decoded in _data_rows(ws, hmap):
        out.append(
            {
                "fields": {
                    "history_id": decoded[0] if decoded else "",
                    "invoice_number": _text(_cell(ws, row, hmap, rt.PENDING_COL_INVOICE)),
                    "date": _text(_cell(ws, row, hmap, rt.PENDING_COL_DATE)),
                    "items": [],
                },
                "reason": _text(_cell(ws, row, hmap, rt.PENDING_COL_REASON)),
                "source_rows": [row],
            }
        )
    return out


@dataclass(frozen=True)
class _Cell:
    value: Any


@dataclass(frozen=True)
class _Sheet:
    title: str
    rows: tuple[tuple, ...]
    max_column: int

    @property
    def max_row(self):
        return len(self.rows)

    def cell(self, *, row, column):
        values = self.rows[row - 1] if row <= len(self.rows) else ()
        return _Cell(values[column - 1] if column <= len(values) else None)


def _evidence_value(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if type(value).__module__ == "openpyxl.worksheet.formula":
        return "[" + type(value).__name__ + "]"
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _finite(values):
    if isinstance(values, Decimal) and not values.is_finite():
        raise RoundtripParseError("ocr.roundtrip_invalid_number")
    if isinstance(values, dict):
        for value in values.values():
            _finite(value)
    elif isinstance(values, (tuple, list)):
        for value in values:
            _finite(value)


def _read_sheets(file_bytes):
    try:
        import openpyxl
    except ImportError as error:
        raise RoundtripParseError("ocr.table_dependency_unavailable", dependency="openpyxl") from error
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False, read_only=True)
    hits, ignored, budget = [], [], DecodeBudget()
    try:
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            if not hasattr(worksheet, "iter_rows"):
                if rt.is_data_sheet(rt.canonical(name)):
                    raise RoundtripParseError("ocr.roundtrip_parse_failed", sheet=name)
                ignored.append({"sheet_name": name, "reason": "non_worksheet", "nonempty_rows": []})
                continue
            rows = []
            for number, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                if number > MAX_ROWS:
                    raise RoundtripParseError("ocr.table_row_limit")
                budget.row([_evidence_value(value) for value in values])
                rows.append(tuple(values))
            sheet = _Sheet(name, tuple(rows), max((len(row) for row in rows), default=0))
            headers = _header_map(sheet)
            # 按导出人语言起的 Sheet 名认回规格名(方向仍然只看 Sheet;旧版泰文表原样)
            spec_name = rt.canonical(name) if name == name.strip() else name
            if rt.is_data_sheet(spec_name) and rt.is_roundtrip_sheet(list(headers)):
                if name != name.strip():
                    raise RoundtripParseError("ocr.roundtrip_sheet_name_invalid", sheet=name)
                hits.append((spec_name, sheet, headers))
            else:
                ignored.append({"sheet_name": name, "reason": "summary" if spec_name == rt.SHEET_SUMMARY else "not_matched",
                                "nonempty_rows": [i for i, row in enumerate(rows, 1)
                                                  if any(_text(value) for value in row)]})
    finally:
        workbook.close()
    return hits, ignored


_SHEET_READERS = {rt.SHEET_SALES: _read_sales_sheet, rt.SHEET_PURCHASE: _read_purchase_sheet,
                  rt.SHEET_PENDING: _read_pending_sheet}


def parse_roundtrip_workbook(file_bytes: bytes) -> Dict[str, Any]:
    """Decode the recognized workbook without model calls or source-ID lookup.

    NotMatched alone means generic-table fallback is appropriate. Recognized
    empty/error/trailing data never silently falls through to a paid model.
    Existing archive/cell/character limits apply even to direct callers.
    """
    if not isinstance(file_bytes, bytes):
        raise RoundtripParseError("ocr.table_invalid_bytes")
    try:
        validate_input_metadata("roundtrip.xlsx", len(file_bytes))
        check_archive(file_bytes)
        hits, ignored = _read_sheets(file_bytes)
        if not hits:
            raise RoundtripNotMatched()
        documents, pending, sources = [], [], {}
        for name, sheet, headers in hits:
            sources[name] = (sheet, headers)
            parsed = _SHEET_READERS[name](sheet, headers)
            direction = rt.sheet_direction(name)
            for entry in parsed:
                entry["sheet"] = name
                entry["source_row_keys"] = []
                for row in entry["source_rows"]:
                    raw = _text(_cell(sheet, row, headers, rt.COL_ROW_KEY))
                    decoded = rt.decode_row_key(raw)
                    entry["source_row_keys"].append({"source_row": row, "row_key": raw,
                        "history_id": decoded[0] if decoded else "", "line_index": decoded[1] if decoded else None})
                if direction is None:
                    pending.append(entry)
                else:
                    entry["direction"] = entry["fields"]["direction"] = direction
                    documents.append(entry)
        if not documents and not pending:
            raise RoundtripParseError("ocr.roundtrip_empty")
        _finite((documents, pending))
        units = []
        for number, entry in enumerate([*documents, *pending], 1):
            sheet, headers = sources[entry["sheet"]]
            units.append({"page_number": number, "source_kind": "roundtrip_rows", "sheet_name": sheet.title,
                "source_rows": list(entry["source_rows"]), "header_source_row": 1,
                "header_cells": [_evidence_value(value) for value in (sheet.rows[0] if sheet.rows else ())],
                "header_columns": dict(headers), "pending": "direction" not in entry,
                "rows": [{**reference, "cells": [_evidence_value(value) for value in sheet.rows[reference["source_row"] - 1]]}
                         for reference in entry["source_row_keys"]]})
        provenance = {"version": 1, "reference_kind": "logical_section", "unit_kind": "roundtrip_document",
                      "coverage_complete": False,
                      "coverage_issues": [{"code": "roundtrip_original_not_fully_attributed"}],
                      "units": units, "ignored_sheets": ignored}
        return {"documents": documents, "pending": pending, "sheets": [name for name, _, _ in hits],
                "source_sha256": hashlib.sha256(file_bytes).hexdigest(), "provenance": provenance}
    except RoundtripParseError:
        raise
    except (OcrInputError, TableInputError) as error:
        raise RoundtripParseError(error.code) from error
    except Exception as error:
        raise RoundtripParseError("ocr.roundtrip_parse_failed") from error
