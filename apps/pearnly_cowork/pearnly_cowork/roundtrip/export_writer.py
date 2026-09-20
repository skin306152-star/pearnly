"""In-memory XLSX adapter for the fixed Thai review-workbook protocol."""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pearnly_cowork.roundtrip import money, spec as rt
from pearnly_cowork.roundtrip.export_mapping import PreparedWorkbook, number, sales_date, text

HEADERS_TH = [
    rt.SALES_COL_DATE, rt.SALES_COL_INVOICE, rt.SALES_COL_PARTY, rt.SALES_COL_ITEM,
    rt.SALES_COL_QTY, rt.SALES_COL_PRICE, rt.SALES_COL_AMOUNT, "รวมจำนวนเงิน",
    rt.SALES_COL_PRE_VAT, "VAT", "อ้างอิง", "Status", "สถานะสินค้า", "สถานะลูกค้า",
    *rt.ROUNDTRIP_HEADERS, "กลุ่มบัญชีสินค้าคงคลัง",
]
COLUMN_WIDTHS = [12, 14, 30, 24, 8, 12, 14, 14, 18, 10, 12, 10, 13, 13, *rt.ROUNDTRIP_WIDTHS, 34]
SUMMARY_HEADERS = ("ประเภท", "รหัส", "ชื่อ", "เอกสารที่สร้าง", "กลุ่มบัญชีสินค้าคงคลัง")
SUMMARY_WIDTHS = (16, 18, 40, 20, 34)
KIND_LABELS = {"item": "สินค้า/บริการ", "customer": "ลูกค้า", "supplier": "ผู้ขาย"}
MONEY_FORMAT, QTY_FORMAT, DATE_FORMAT = "#,##0.00", "#,##0.###", "dd/mm/yyyy"
HEADER_FILL = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NEW_FILL = PatternFill(start_color="D6F5DE", end_color="D6F5DE", fill_type="solid")
REUSED_FILL = PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid")
PENDING_FILL = PatternFill(start_color="FDF3D8", end_color="FDF3D8", fill_type="solid")
_THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def _cell(ws, row, column, value=None, *, formula=False):
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and not formula:
        # number_format='@' alone does not prevent openpyxl treating '=...' as
        # a formula. Invoice text is always a string; only fixed formulas below
        # can be emitted as executable spreadsheet cells.
        cell.data_type = "s"
    return cell


def _header(ws, headers, widths):
    for index, label in enumerate(headers, 1):
        cell = _cell(ws, 1, index, label)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, CENTER, BORDER
        ws.column_dimensions[get_column_letter(index)].width = widths[index - 1]
        if label == rt.COL_ROW_KEY:
            ws.column_dimensions[get_column_letter(index)].hidden = True
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _action(value):
    if isinstance(value, bool):
        return "new" if value else "reused"
    raw = text(value).lower()
    if raw in ("new", "created", "auto_created", "erp_auto_created"):
        return "new"
    if raw in ("reused", "reuse", "existing", "matched", "match", "cache_hit", "db_mapping"):
        return "reused"
    return ""


def _fill(action):
    return NEW_FILL if action == "new" else REUSED_FILL if action == "reused" else None


def _status(ws, row, column, action):
    cell = _cell(ws, row, column, {"new": "ใหม่", "reused": "เดิม", "": "-"}[action])
    cell.alignment = CENTER
    fill = _fill(action)
    if fill is not None:
        cell.font, cell.fill = Font(bold=True), fill


def _roundtrip_values(**values):
    result = rt.roundtrip_values(**values)
    # The shared source helper accepts float/int as numeric; keep Decimal as a
    # numeric XLSX cell too, without converting monetary arithmetic to float.
    if isinstance(values.get("doc_vat"), Decimal):
        result[2] = values["doc_vat"]
    return result


def _sales(ws, records):
    _header(ws, HEADERS_TH, COLUMN_WIDTHS)
    row = 2
    for record in records:
        fields = record["merged_fields"]
        date_value = sales_date(fields.get("date"))
        customer_action = _action(fields.get("customer_erp_action"))
        for line_index, item in enumerate(fields["items"]):
            qty = number(item.get("qty") or item.get("quantity"))
            price = number(item.get("unit_price") or item.get("price"))
            amount = number(item.get("amount") or item.get("total"))
            item_action = _action(item.get("erp_action"))
            cell = _cell(ws, row, 1, date_value)
            if date_value:
                cell.number_format = DATE_FORMAT
            _cell(ws, row, 2, text(fields.get("invoice_number"))).number_format = "@"
            cell = _cell(ws, row, 3, text(fields.get("buyer_name")) or text(fields.get("customer_name")))
            if _fill(customer_action):
                cell.fill = _fill(customer_action)
            cell = _cell(ws, row, 4, text(item.get("description")) or text(item.get("product_name")) or text(item.get("name")))
            if _fill(item_action):
                cell.fill = _fill(item_action)
            _cell(ws, row, 5, qty)
            _cell(ws, row, 6, price)
            if qty is not None and price is not None:
                for column in (7, 8, 9):
                    _cell(ws, row, column, f"=E{row}*F{row}", formula=True)
                _cell(ws, row, 10, f"=I{row}*0.07", formula=True)
            elif amount is not None:
                for column in (7, 8, 9):
                    _cell(ws, row, column, amount)
                _cell(ws, row, 10, money.vat_of(amount))
            for column in range(1, len(HEADERS_TH) + 1):
                cell = ws.cell(row, column)
                cell.border = BORDER
                if column == 1 or column in (13, 14):
                    cell.alignment = CENTER
                elif column == 5:
                    cell.alignment, cell.number_format = RIGHT, QTY_FORMAT
                elif column in (6, 7, 8, 9, 10):
                    cell.alignment, cell.number_format = RIGHT, MONEY_FORMAT
                else:
                    cell.alignment = LEFT
            _status(ws, row, 13, item_action)
            _status(ws, row, 14, customer_action)
            values = _roundtrip_values(
                party_tax=text(fields.get("buyer_tax")) or text(fields.get("buyer_tax_id")),
                party_branch=text(fields.get("buyer_branch")),
                doc_vat=number(fields.get("vat_amount")) or number(fields.get("vat")),
                docnum=text(fields.get("erp_docnum")), item_code=item.get("erp_item_code"),
                party_code=text(fields.get("erp_party_code")), push_status=fields.get("push_status"),
                push_reason=fields.get("push_reason"), row_key=rt.encode_row_key(fields["history_id"], line_index))
            for offset, value in enumerate(values, 15):
                _cell(ws, row, offset, value)
            if item_action == "new" and text(fields.get("stock_acc_group")):
                _cell(ws, row, len(HEADERS_TH), text(fields["stock_acc_group"])).fill = NEW_FILL
            row += 1


def _first(fields, *keys):
    for key in keys:
        value = number(fields.get(key), allow_thb=False)
        if value is not None:
            return value
    return None


def _purchase_rows(records):
    rows = []
    for index, record in enumerate(records, 1):
        fields = record["merged_fields"]
        vat = _first(fields, "vat_amount", "vat")
        rows.append([
            index, text(fields.get("date")), text(fields.get("invoice_number")),
            text(fields.get("seller_name")), text(fields.get("seller_tax")), text(fields.get("seller_branch")),
            _first(fields, "amount_before_vat", "subtotal"), vat,
            number(fields.get("total_amount"), allow_thb=False), text(fields.get("category")),
            *_roundtrip_values(party_tax=fields.get("seller_tax") or fields.get("seller_tax_id"),
                party_branch=fields.get("seller_branch"), doc_vat=vat, docnum=fields.get("erp_docnum"),
                party_code=fields.get("erp_party_code"), push_status=fields.get("push_status"),
                push_reason=fields.get("push_reason"), row_key=rt.encode_row_key(fields["history_id"], 0))])
    return rows


def _pending_rows(records):
    rows = []
    for record in records:
        fields = record["merged_fields"]
        rows.append([
            text(fields.get("date")), text(fields.get("invoice_number")),
            text(fields.get("seller_name")) or text(fields.get("buyer_name")),
            text(fields.get("seller_tax")) or text(fields.get("buyer_tax")),
            number(fields.get("total_amount"), allow_thb=False), text(record.get("reason") or fields.get("pending_reason")),
            *_roundtrip_values(party_tax=fields.get("seller_tax") or fields.get("buyer_tax"),
                party_branch=fields.get("seller_branch") or fields.get("buyer_branch"),
                push_status=fields.get("push_status"), push_reason=fields.get("push_reason"),
                row_key=rt.encode_row_key(fields["history_id"], 0))])
    return rows


def _grid(ws, headers, widths, rows, *, money_columns=(), fill=None):
    _header(ws, headers, widths)
    for row, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            cell = _cell(ws, row, column, value)
            cell.border = BORDER
            cell.alignment = RIGHT if column in money_columns else CENTER if column == 1 else LEFT
            if column in money_columns:
                cell.number_format = MONEY_FORMAT
            if fill is not None:
                cell.fill = fill


def write_workbook(prepared: PreparedWorkbook) -> bytes:
    workbook = Workbook()
    sales = workbook.active
    sales.title = rt.SHEET_SALES
    _sales(sales, prepared.sales)
    _grid(workbook.create_sheet(rt.SHEET_PURCHASE), [*rt.PURCHASE_HEADERS, *rt.ROUNDTRIP_HEADERS],
          [*rt.PURCHASE_WIDTHS, *rt.ROUNDTRIP_WIDTHS], _purchase_rows(prepared.purchase), money_columns=(7, 8, 9))
    _grid(workbook.create_sheet(rt.SHEET_PENDING), [*rt.PENDING_HEADERS, *rt.ROUNDTRIP_HEADERS],
          [*rt.PENDING_WIDTHS, *rt.ROUNDTRIP_WIDTHS], _pending_rows(prepared.pending), money_columns=(5,), fill=PENDING_FILL)
    rows = [[KIND_LABELS.get(text(value.get("kind")), text(value.get("kind")) or "-"),
             text(value.get("code")), text(value.get("name")), text(value.get("docnum")), text(value.get("acc_group"))]
            for value in prepared.created_masters]
    _grid(workbook.create_sheet(rt.SHEET_SUMMARY), SUMMARY_HEADERS, SUMMARY_WIDTHS, rows)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
