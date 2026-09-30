# -*- coding: utf-8 -*-
"""表格与 PDF 文字层读取、表头命中（银行单、总账与账套导入共用）。

来源：ec2e5c957e6b ``services/recon/bank_table_io.py``。旧文件里还有汇总行识别
（``_is_summary_row``）——现在由同域 ``summary`` 维护，``bank_statement`` 兼容导出，
所以这里只剩下两件事：把文件读成单元格、判断表头命中。

PDF 文字层（旧 ``_pdf_extract_text_safe``）在第三十九批随 GL/文本读取器一起搬进来：
先 pdfminer（没有 pdfplumber 那个 ``KeyError('date')`` 崩溃），再 pypdf 兜底。

只读传入的 bytes，不打数据库、不走网络、不调模型。读不出来一律返回空列表（旧行为），
由调用方决定是报"读不了"还是走别的路。
"""

from __future__ import annotations

import io
import re
from typing import List, Tuple


def header_hit(header: str, hints: set) -> bool:
	"""表头是否命中某组关键词；**短 ASCII 词（in/out/cr/dr）必须整词匹配**。

	为什么：KTB 真实事故里分行号列 "Init Br." 被当成存款列 —— `in` 是它的子串。
	长词与泰文/中文仍按子串匹配：那些语言没有分词，整词匹配反而漏得更多。
	"""
	h = str(header or "").strip().lower()
	if not h:
		return False
	tokens = None
	for hint in hints:
		hl = hint.lower()
		if hl.isascii() and len(hl) <= 3:
			if tokens is None:
				tokens = set(re.split(r"[\s/().,_\-]+", h))
			if hl in tokens:
				return True
		elif hl in h:
			return True
	return False


def load_workbook_sheets(file_bytes: bytes, *, max_cells=None) -> List[Tuple[str, list]]:
	"""读出所有工作表：``[(sheet_name, [row, ...]), ...]``；读不出来返回 ``[]``。

	先 openpyxl（.xlsx），再退回 xlrd（旧 .xls）；两条都走不通就是空列表。一个文件
	可能有多个账户（每 sheet 一个），所以返回的是全部 sheet，切分交给业务读取器。
	"""
	try:
		import openpyxl

		workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
		try:
			remaining = [max_cells]
			out = []
			for sheet in workbook.worksheets:
				# 只读模式按文件里记的范围读；有的网银导出把范围写成 A1（SCB / Krungsri 实测只读出第一行，
				# 2026-09-30 压测 05-银行导出样本）。openpyxl 的做法：先把范围清掉，按实际内容读到底。
				if hasattr(sheet, "reset_dimensions"):
					sheet.reset_dimensions()
				if max_cells is not None and ((sheet.max_column or 0) > 256 or (sheet.max_row or 0) > 100_000):
					raise ValueError("table dimensions exceed import limit")
				out.append((sheet.title, _limited_rows(sheet.iter_rows(values_only=True), remaining)))
		finally:
			try:
				workbook.close()
			except Exception:
				pass  # Preserve deployed successful reads even if close fails.
		return out
	except Exception:
		pass
	try:
		import xlrd

		workbook = xlrd.open_workbook(file_contents=file_bytes)
		out = []
		remaining = [max_cells]
		for index in range(workbook.nsheets):
			sheet = workbook.sheet_by_index(index)
			if max_cells is not None and (sheet.ncols > 256 or sheet.nrows > 100_000):
				raise ValueError("table dimensions exceed import limit")
			out.append((sheet.name, _limited_rows((sheet.row_values(number) for number in range(sheet.nrows)), remaining)))
		return out
	except Exception:
		return []


def _limited_rows(rows, remaining):
	result = []
	for row in rows:
		if remaining[0] is not None:
			remaining[0] -= len(row)
			if remaining[0] < 0 or len(row) > 256 or len(result) >= 100_000:
				raise ValueError("table cells exceed import limit")
		result.append(list(row))
	return result


def load_csv_sheets(file_bytes: bytes, *, max_cells=None) -> List[Tuple[str, list]]:
	"""CSV/TSV → ``[("csv", rows)]``；读不出来返回 ``[]``。

	编码按 UTF-8 BOM / UTF-8 / 泰文 cp874 / 中文 gbk / latin-1 依次尝试，分隔符先嗅探
	（逗号/分号/制表/竖线），嗅探失败再按出现次数挑。数字里的千分位逗号靠 CSV 的引号
	字段处理：正规导出的 ``"10,620.53"`` 是带引号的。
	"""
	import csv as _csv

	text = None
	# 带 BOM 的 UTF-16（ICBC 等网银导出实测 FF FE 开头）：按 UTF-16 读，否则 latin-1 会把它读成一堆乱码单格。
	encodings = ("utf-16",) if file_bytes[:2] in (b"\xff\xfe", b"\xfe\xff") else ()
	for encoding in encodings + ("utf-8-sig", "utf-8", "cp874", "gbk", "latin-1"):
		try:
			text = file_bytes.decode(encoding)
			break
		except Exception:
			continue
	if text is None:
		return []
	sample = text[:4096]
	delimiter = ","
	try:
		delimiter = _csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
	except Exception:
		for candidate in (";", "\t", "|"):
			if sample.count(candidate) > sample.count(","):
				delimiter = candidate
				break
	try:
		rows = _limited_rows(_csv.reader(io.StringIO(text), delimiter=delimiter), [max_cells])
		return [("csv", rows)] if rows else []
	except Exception:
		return []


def extract_pdf_text_pages(file_bytes: bytes) -> List[str]:
	"""PDF 的文字层 → 每页一段文本；读不出来返回 ``[]``（调用方再决定是否走模型）。

	先 pdfminer：它没有 pdfplumber 那个畸形元数据上的 ``KeyError('date')`` 崩溃，而且
	一次给全文（旧实现就把它当"一页"返回）。pdfminer 拿不到可用文字时退回 pypdf 逐页
	抽；两条都失败就是空列表。
	"""
	try:
		from pdfminer.high_level import extract_text

		text = extract_text(io.BytesIO(file_bytes)) or ""
		if text.strip():
			return [text]
	except Exception:
		pass
	try:
		import pypdf

		reader = pypdf.PdfReader(io.BytesIO(file_bytes))
		return [page.extract_text() or "" for page in reader.pages]
	except Exception:
		pass
	return []
