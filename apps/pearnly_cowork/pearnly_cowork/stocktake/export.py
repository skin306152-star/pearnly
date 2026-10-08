"""Thai scan-count workbook from one authorized, complete report snapshot.

Source: ec2e5c957e6b services/stocktake/reports.py. No Frappe, database,
filesystem, ERP connection or photo fetch. Authorization belongs to report().
"""
from collections import defaultdict
from collections.abc import Mapping
from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from pearnly_cowork.stocktake.calculations import compare_location, summarize_item
from pearnly_cowork.stocktake.contracts import StocktakeValidationError, quantity

SUMMARY = [
    "รหัสสินค้า", "ชื่อสินค้า", "บาร์โค้ด / รหัส QR", "หน่วย", "จำนวนตามบัญชี",
    "จำนวนที่นับได้รวม", "ผลต่าง", "ผลการตรวจนับ", "คลังสินค้าตามบัญชี",
    "คลังสินค้าที่ตรวจนับ", "ผลเทียบคลังสินค้า", "ตำแหน่งตามบัญชี",
    "ตำแหน่งที่ตรวจนับ", "ผลเทียบตำแหน่ง", "ดูรูปภาพ",
]
DETAIL = [
    "รหัสสินค้า", "ชื่อสินค้า", "บาร์โค้ด / รหัส QR", "คลังสินค้า", "ตำแหน่งจัดเก็บ",
    "หน่วย", "จำนวนที่นับครั้งนี้", "ผู้ตรวจนับ", "เวลาที่นับ (UTC)", "สถานะรายการ",
    "ผู้แก้ไขล่าสุด", "เวลาแก้ไข (UTC)", "คลังสินค้าตามบัญชี", "ผลเทียบคลังสินค้า",
    "ตำแหน่งตามบัญชี", "ผลเทียบตำแหน่ง", "ดูรูปภาพ",
]
_LABELS = {
    "not_counted": "ยังไม่ได้นับ", "not_specified": "ไม่ได้ระบุในบัญชี",
    "matched": "ตรงกัน", "mismatched": "ไม่ตรงกัน", "over": "เกินบัญชี", "short": "ขาดบัญชี",
}
# 表头、Sheet 名、结果标签按语义键取（2026-10-08：导出文件也走语言包，按导出人的界面语言）。
# 这里只留泰文默认值（与旧线上逐字一致，泰文导出不变）；其它语言由调用方用翻译包取好后传 `texts` 进来。
SUMMARY_KEYS = (
    "product_code", "product_name", "barcode", "unit", "book_qty", "actual_total", "difference", "count_result",
    "book_warehouse", "counted_warehouse", "warehouse_comparison", "book_location", "counted_location",
    "location_comparison", "photo",
)
DETAIL_KEYS = (
    "product_code", "product_name", "barcode", "warehouse", "location", "unit", "counted_qty", "counted_by",
    "counted_at", "entry_status", "updated_by", "updated_at", "book_warehouse", "warehouse_comparison",
    "book_location", "location_comparison", "photo",
)
TEXT_KEYS = (
    *dict.fromkeys(SUMMARY_KEYS + DETAIL_KEYS), "sheet_summary", "sheet_detail", "location_blank",
    "entry_voided", "entry_counted", *(f"result_{code}" for code in _LABELS),
)
THAI = {
    **dict(zip(SUMMARY_KEYS, SUMMARY)), **dict(zip(DETAIL_KEYS, DETAIL)),
    "sheet_summary": "สรุปผลตรวจนับ", "sheet_detail": "รายละเอียดการนับ", "location_blank": "(ไม่ระบุ)",
    "entry_voided": "ยกเลิก (ไม่รวมยอด)", "entry_counted": "รวมยอด",
    **{f"result_{code}": text for code, text in _LABELS.items()},
}


def _append(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        # Match the source: Excel cannot preserve more than 15 significant
        # digits in a numeric cell. Text remains text, including formula input.
        if isinstance(cell.value, Decimal) and len(cell.value.as_tuple().digits) > 15:
            cell.value = str(cell.value)
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _highlight(sheet, columns):
    for col in columns:
        sheet.cell(sheet.max_row, col).fill = PatternFill("solid", fgColor="FFF2CC")


def _utc_iso(value):
    # Frappe stores naive UTC values. Never let the host timezone reinterpret
    # them; source PostgreSQL timestamps already carried their timezone.
    if isinstance(value, str):
        try:
            value = datetime.fromisoformat(value)
        except ValueError:
            raise StocktakeValidationError("stocktake.request_invalid") from None
    if not isinstance(value, datetime):
        raise StocktakeValidationError("stocktake.request_invalid")
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC).isoformat()


def workbook(snapshot: Mapping, texts: Mapping | None = None) -> bytes:
    """Render service.report() output, preserving its item and entry order.

    `texts` maps every key in TEXT_KEYS to the exporter's language; without it the
    workbook is the Thai one the source produced.

    entry_total is required so the 100-entry detail preview cannot accidentally
    become a complete export. Quantities accept the native exact strings or
    source Decimal values. No caller-owned value is changed.
    """
    if snapshot.get("count_mode") != "scan":
        raise StocktakeValidationError("stocktake.legacy_task")
    if snapshot.get("photos"):
        raise StocktakeValidationError("stocktake.photos_not_supported")
    rows, entries = snapshot["items"], snapshot["entries"]
    if snapshot.get("entry_total") != len(entries):
        raise StocktakeValidationError("stocktake.export_incomplete")
    items = {row["id"]: row for row in rows}
    grouped = defaultdict(list)
    for entry in entries:
        if entry["item_id"] not in items:
            raise StocktakeValidationError("stocktake.request_invalid")
        grouped[entry["item_id"]].append(entry)

    text = {**THAI, **{key: value for key, value in (texts or {}).items() if key in THAI and value}}
    labels = {code: text[f"result_{code}"] for code in _LABELS}
    wb = Workbook()
    summary = wb.active
    summary.title = text["sheet_summary"]
    _append(summary, [text[key] for key in SUMMARY_KEYS])
    for row in rows:
        values = summarize_item(row, grouped[row["id"]])
        _append(summary, [
            row["product_code"], row["product_name"], row["barcode"], row["unit"],
            values["book_qty"], values["actual_qty"], values["difference"], labels[values["count_result"]],
            row["warehouse"], "\n".join(values["warehouses"]), labels[values["warehouse_comparison"]],
            row["location"], "\n".join(value or text["location_blank"] for value in values["locations"]),
            labels[values["location_comparison"]], "",
        ])
        if values["difference"] is not None and values["difference"] != 0:
            _highlight(summary, (5, 6, 7, 8))
        if values["warehouse_comparison"] == "mismatched":
            _highlight(summary, (9, 10, 11))
        if values["location_comparison"] == "mismatched":
            _highlight(summary, (12, 13, 14))

    detail = wb.create_sheet(text["sheet_detail"])
    _append(detail, [text[key] for key in DETAIL_KEYS])
    for entry in entries:
        book = items[entry["item_id"]]
        warehouse = compare_location(book["warehouse"], [entry["warehouse"]])
        location = compare_location(book["location"], [entry["location"]])
        _append(detail, [
            entry["product_code"], entry["product_name"], entry["barcode"], entry["warehouse"],
            entry["location"], entry["unit"], quantity(entry["quantity"], nonnegative=True),
            entry["counted_by_name"], _utc_iso(entry["counted_at"]),
            text["entry_voided"] if entry["voided"] else text["entry_counted"],
            entry["updated_by_name"], _utc_iso(entry["updated_at"]), book["warehouse"],
            labels[warehouse], book["location"], labels[location], "",
        ])
        if not entry["voided"]:
            if warehouse == "mismatched":
                _highlight(detail, (4, 13, 14))
            if location == "mismatched":
                _highlight(detail, (5, 15, 16))

    edge = Side(style="thin", color="A6ACB3")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="202033")
            cell.fill = PatternFill("solid", fgColor="EAF0FF")
            sheet.column_dimensions[cell.column_letter].width = 25
        sheet.column_dimensions["B"].width = 36
        sheet.row_dimensions[1].height = 26
        for cells in sheet:
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
