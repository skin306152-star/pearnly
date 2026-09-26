"""上传盘点（用户 2026-09-26：Cowork 不做在线盘点，账面与实盘一张表上传）。

- 模板：按站点语言出表头，编码 / 条码两列预设文本格式（防止前导 0 被吃掉）。
- 上传：服务端读 xlsx（限大小与解压后大小，公式格不收），好行建任务、有实盘的写一笔记录，
  坏行跳过并逐条返回原因；同一 request_id 重放返回同一任务。
- 列表：本站点的公司资料下，可按名称 / 状态 / 行数 / 已盘 / 差异 / 创建时间排序，带总数。
- 删除：连同明细、记录、操作回执一起删（Cowork 不记库存账，没有要冲回的东西）。

任务仍是 `count_mode = scan`：导入的实盘落成逐笔记录，所以详情、结果对照与导出都复用原有实现。
"""
import hashlib
import io
import zipfile
from uuid import NAMESPACE_URL, uuid4, uuid5

import frappe

from pearnly_cowork.access import scope
from pearnly_cowork.stocktake import access, service, upload_sheet
from pearnly_cowork.stocktake.contracts import MAX_BYTES, MAX_UNCOMPRESSED_BYTES
from pearnly_cowork.stocktake.documents import persist
from pearnly_cowork.stocktake.native_contract import encode, identifier

TASK, ITEM, ENTRY, OPERATION = access.TYPES
SORTS = {"title": "t.title", "status": "t.status", "created_at": "t.created_at",
         "total": "total", "counted": "counted", "differences": "differences"}


def _labels():
    """四种语言的表头都认：模板是哪种语言下载的，上传回来都能对上。"""
    from pearnly_core.browser_text import all_languages

    labels = {}
    for field, key in upload_sheet.FIELD_LABELS:
        for text in all_languages(key).values():
            labels[text] = field
    return labels


def template():
    from openpyxl import Workbook
    from pearnly_core.browser_text import label

    access.require(scope.default_workspace())
    book = Workbook()
    sheet = book.active
    sheet.append([label(key) for _field, key in upload_sheet.FIELD_LABELS])
    for column, (field, _key) in enumerate(upload_sheet.FIELD_LABELS, 1):
        letter = sheet.cell(row=1, column=column).column_letter
        sheet.column_dimensions[letter].width = 18
        if field in ("product_code", "barcode"):
            for row in range(2, 1002):
                sheet.cell(row=row, column=column).number_format = "@"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _cells(content):
    if len(content) > MAX_BYTES:
        service.fail("stocktake.file_too_large")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                service.fail("stocktake.file_too_large")
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError):
        service.fail("stocktake.file_invalid")
    try:
        rows = book.worksheets[0].iter_rows()
        header = [cell.value for cell in next(rows, ())]
        cells = [[(cell.value, cell.data_type) for cell in row] for row in rows]
    finally:
        book.close()
    return header, cells


def create(title, request_id, content):
    workspace = scope.default_workspace()
    access.require(workspace)
    request_id = service.contract(identifier, request_id)
    title = str(title or "").strip()
    if not title or len(title) > 120:
        service.fail("stocktake.name_invalid")
    header, cells = _cells(content)
    items, problems = service.contract(upload_sheet.rows, header, cells, _labels())
    digest = hashlib.sha256(encode(items).encode()).hexdigest()
    if frappe.db.exists(TASK, request_id):
        return {**service._existing(request_id, workspace, title, digest), "replayed": True}
    counted = 0
    with service.atomic():
        timestamp, user = service.now(), frappe.session.user
        persist(service._new(TASK, request_id, workspace, title=title, source_digest=digest,
                             source_format="upload_xlsx_v1", count_mode="scan", status="active",
                             created_by=user, created_at=timestamp))
        task = frappe.get_doc(TASK, request_id)
        for row in items:
            actual = row.pop("actual_qty")
            item = service._new(ITEM, str(uuid4()), workspace, task=request_id, **row, actual_qty=None, version=0)
            persist(item)
            if actual is None:
                continue
            # 导入的实盘就是这一行的唯一一笔记录；记录落在账面仓位（表里只有一套仓位）。
            entry_id = str(uuid5(NAMESPACE_URL, f"cowork-stocktake-import:{request_id}:{item.name}"))
            persist(service._new(ENTRY, entry_id, workspace, task=request_id, item=item.name,
                warehouse=row["warehouse"] or "-", location=row["location"], quantity=actual, voided=0,
                version=0, counted_by=user, counted_at=timestamp, updated_by=user, updated_at=timestamp))
            item.actual_qty, item.version, item.counted_by, item.counted_at = actual, 1, user, timestamp
            persist(item)
            counted += 1
    return {"id": task.name, "replayed": False, "imported": len(items), "counted": counted,
            "problems": problems}


def records(start=0, limit=20, sort_by="created_at", sort_order="desc"):
    workspace = scope.default_workspace()
    access.require(workspace)
    start, limit = service.page(start, limit)
    column = SORTS.get(sort_by, SORTS["created_at"])
    order = "ASC" if str(sort_order).lower() == "asc" else "DESC"
    # 列名来自上面的固定映射，数值一律参数化。
    rows = frappe.db.sql(f"""SELECT t.name, t.title, t.status, t.created_at, t.created_by,
        COUNT(i.name) AS total, COUNT(NULLIF(i.actual_qty,'')) AS counted,
        COALESCE(SUM(CASE WHEN NULLIF(i.actual_qty,'') IS NOT NULL AND
            CAST(i.actual_qty AS DECIMAL(20,6)) <> CAST(i.book_qty AS DECIMAL(20,6)) THEN 1 ELSE 0 END),0) AS differences
        FROM `tabCowork Stocktake Task` t LEFT JOIN `tabCowork Stocktake Item` i
            ON i.task=t.name AND i.workspace=t.workspace
        WHERE t.workspace=%s GROUP BY t.name ORDER BY {column} {order}, t.name DESC
        LIMIT %s OFFSET %s""", (workspace, limit, start), as_dict=True)
    total = frappe.db.count(TASK, {"workspace": workspace})
    return {"tasks": [{"id": row.name, "name": row.title, "status": row.status, "created_at": row.created_at,
                       "created_by": row.created_by, "total": int(row.total), "counted": int(row.counted),
                       "differences": int(row.differences)} for row in rows], "total": total}


def delete(task_ids):
    task_ids = service.parse(task_ids)
    if not isinstance(task_ids, list) or not task_ids or len(task_ids) > 200:
        service.fail("stocktake.request_invalid")
    deleted = []
    with service.atomic():
        for task_id in task_ids:
            task = service._task(task_id)
            for doctype in (OPERATION, ENTRY, ITEM):
                frappe.db.delete(doctype, {"task": task.name, "workspace": task.workspace})
            frappe.db.delete(TASK, {"name": task.name, "workspace": task.workspace})
            deleted.append(task.name)
    return {"deleted": len(deleted), "ids": deleted}


def detail(task_id):
    """详情 + 每行结果（一致 / 盘盈 / 盘亏 / 未盘）与仓位对照，页面一次拿齐。"""
    task = service._task(task_id)
    if task.count_mode != "scan":
        return service.detail(task_id)
    return service.report(task_id)
