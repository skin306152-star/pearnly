"""对账结果页 = 导出的复核表，逐格照画（用户 2026-09-26）：取表、改后重算、保存。

“前端显示表格按导出表格 1:1 画，可以不用官方 UI 组件；使用说明不用出现，不要过多文案备注。”
随后（同日）：三张表统一外观、去掉 KPI 与说明文案——改在导出工作簿本身（`reconciliation/export/theme.py`），页面照画即同步。

页面上的每一页都由**导出同一个函数**在内存里生成的工作簿读出来：单元格文字（按单元格的数字格式排好）、
底色、字色、粗斜体、字号、对齐、换行、边框、合并、列宽、行高、冻结窗格一并交给页面照画，
所以页面与下载的 Excel 逐格一致（工作簿里已不再有使用说明页）。
2026-09-27 起页面外观改为原生风格（用户：“改为图二那种，可以点击排序”）：明细表按原生表格画、可点列头排序，
汇总拆成卡片；每格的文字与数据仍来自这里，计算与导出不变（见 `public/desk/recon.js`）。

可改的格子就是两边的输入落在表里的那几格（银行：对账单明细、总账明细；收入：差异明细的税表一侧与
总账金额；销项税：发票明细、VAT 报告明细），每格带着它对应的输入位置。改了就用建任务时同一套比对重算，
重新出表。保存整份替换结果，并用任务的修改时间防止两个人同时改。
"""
import io
import re
from datetime import date, datetime

import frappe
from frappe import _

from pearnly_cowork.reconciliation import normalization, presets, revision, vat_check_service

MAX_CELLS = 400000
_SUM = re.compile(r"^=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$")


def _lang(lang):
    lang = str(lang or frappe.local.lang or "th")[:2]
    return lang if lang in ("th", "en", "zh", "ja") else "th"


# ── 单元格 → 文字与样式 ─────────────────────────────────────────────────────


def _rgb(color):
    value = getattr(color, "rgb", None) if color is not None else None
    if not isinstance(value, str) or len(value) < 6 or getattr(color, "type", "rgb") != "rgb":
        return None
    return "#" + value[-6:]


def _section_text(number, section):
    body = re.sub(r"\[[^\]]*\]|\"[^\"]*\"|[_*\\].", "", section).strip()
    percent = "%" in body
    if percent:
        number *= 100
    decimals = sum(1 for ch in body.split(".", 1)[1] if ch in "0#") if "." in body else 0
    text = f"{number:,.{decimals}f}" if "," in body else f"{number:.{decimals}f}"
    if "(" in body and ")" in body:
        text = f"({text})"
    elif body.startswith("-"):
        text = "-" + text
    return text + ("%" if percent else "")


def _number(value, fmt):
    """按 Excel 数字格式排数字：`正;负;零` 三段，`[Red]` 红字，括号负数。返回 (文字, 是否红字)。"""
    fmt = fmt or "General"
    if fmt in ("General", "@"):
        if isinstance(value, float) and value.is_integer():
            return str(int(value)), False
        return (f"{value:.10g}" if isinstance(value, float) else str(value)), False
    sections = fmt.split(";")
    if value < 0 and len(sections) > 1:
        section, shown = sections[1], -value
    elif value == 0 and len(sections) > 2:
        section, shown = sections[2], 0.0
    else:
        section, shown = sections[0], value
    return _section_text(shown, section), "[Red]" in section


def _sum(sheet, formula):
    match = _SUM.match(formula.replace(" ", ""))
    if not match:
        return None
    from openpyxl.utils import column_index_from_string

    first, top, last, bottom = match.groups()
    total = 0.0
    for row in sheet.iter_rows(min_row=int(top), max_row=int(bottom), min_col=column_index_from_string(first),
                               max_col=column_index_from_string(last), values_only=True):
        total += sum(value for value in row if isinstance(value, (int, float)) and not isinstance(value, bool))
    return round(total, 2)


def _value(sheet, cell):
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        value = _sum(sheet, value)
    if value is None:
        return "", False, False
    if isinstance(value, bool):
        return str(value), False, False
    if isinstance(value, (int, float)):
        text, red = _number(value, cell.number_format)
        return text, red, True
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M") if (value.hour or value.minute) else value.strftime("%d/%m/%Y"), False, False
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y"), False, False
    return str(value), False, False


def _side(side):
    style = getattr(side, "style", None)
    if not style:
        return None
    width = {"medium": 2, "thick": 3, "double": 3}.get(style, 1)
    return [width, _rgb(side.color) or "#000000", "double" if style == "double" else
            "dashed" if "dash" in style else "dotted" if style in ("dotted", "hair") else "solid"]


def _style(cell, numeric, red):
    font, fill, align, border = cell.font, cell.fill, cell.alignment, cell.border
    style = {}
    if fill is not None and fill.fill_type == "solid":
        background = _rgb(fill.fgColor)
        if background:
            style["bg"] = background
    color = "#FF0000" if red else _rgb(font.color) if font else None
    if color:
        style["fg"] = color
    if font is not None:
        if font.b:
            style["b"] = 1
        if font.i:
            style["i"] = 1
        if font.u:
            style["u"] = 1
        if font.sz:
            style["sz"] = float(font.sz)
        if font.name:
            style["fn"] = font.name
    horizontal = align.horizontal if align is not None else None
    style["h"] = horizontal if horizontal in ("left", "center", "right", "justify") else (
        "center" if horizontal in ("centerContinuous", "distributed") else "right" if numeric else "left")
    style["v"] = {"center": "middle", "top": "top"}.get(align.vertical if align is not None else None, "bottom")
    if align is not None and align.wrap_text:
        style["w"] = 1
    if align is not None and align.indent:
        style["in"] = int(align.indent)
    sides = [_side(getattr(border, name, None)) for name in ("top", "right", "bottom", "left")] if border else []
    if any(sides):
        style["bd"] = sides
    return style


# ── 可改的格子：表里的一格 → 输入里的哪一行哪一栏 ───────────────────────────


def _num(value):
    try:
        return round(float(str(value).replace(",", "")), 2) if value not in (None, "") else 0.0
    except ValueError:
        return None


def _iso_to_sheet(value):
    try:
        return date.fromisoformat(str(value)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(value or "")


def _by_content(sheet, first_row, rows, key_of_row, key_of_cells, width):
    """导出时这两页按日期等重新排过序：按内容把表里每一行认回它的输入行（同内容的行可以互换）；
    认不出的行按剩下的输入行顺次补上。"""
    lines = [list(values) for values in sheet.iter_rows(min_row=first_row, max_col=width, values_only=True)]
    unused = {}
    for index, row in enumerate(rows):
        unused.setdefault(key_of_row(row), []).append(index)
    mapped, missing = [], []
    for offset, values in enumerate(lines):
        bucket = unused.get(key_of_cells(values))
        if bucket:
            mapped.append((first_row + offset, bucket.pop(0)))
        else:
            missing.append(first_row + offset)
    rest = sorted(index for bucket in unused.values() for index in bucket)
    mapped.extend(zip(missing, rest))
    return mapped


def _bank_edits(book, inputs):
    edits = {}
    statement, ledger = inputs.get("statement") or [], inputs.get("gl") or []
    stmt_key = lambda row: (_iso_to_sheet(row.get("date")), str(row.get("description") or ""),  # noqa: E731
        _num(row.get("withdrawal")), _num(row.get("deposit")), _num(row.get("balance")))
    stmt_cells = lambda v: (str(v[0] or ""), str(v[1] or ""), _num(v[2]), _num(v[3]), _num(v[4]))  # noqa: E731
    fields = ("date", "description", "withdrawal", "deposit", "balance")
    for sheet_row, index in _by_content(book.worksheets[2], 2, statement, stmt_key, stmt_cells, 5):
        for column, field in enumerate(fields, 1):
            edits[(2, sheet_row, column)] = ["statement", index, field]
    gl_key = lambda row: (_iso_to_sheet(row.get("date")), str(row.get("doc_no") or ""),  # noqa: E731
        str(row.get("account_code") or ""), str(row.get("description") or ""), _num(row.get("debit")), _num(row.get("credit")))
    gl_cells = lambda v: (str(v[0] or ""), str(v[1] or ""), str(v[2] or ""), str(v[3] or ""), _num(v[4]), _num(v[5]))  # noqa: E731
    fields = ("date", "doc_no", "account_code", "description", "debit", "credit", "balance")
    for sheet_row, index in _by_content(book.worksheets[3], 2, ledger, gl_key, gl_cells, 6):
        for column, field in enumerate(fields, 1):
            edits[(3, sheet_row, column)] = ["gl", index, field]
    return edits


def _income_edits(book, inputs):
    """明细页按日期、单号重排过：按内容认回输入行（原来按位置对应，改了日期 / 单号行一挪，
    下一次改的就落到别的行上 —— 2026-09-28 录制发现），与银行两页同一做法。"""
    from pearnly_cowork.reconciliation.export.income_workbook import show_date

    sheet = book.worksheets[0]
    first = sheet.freeze_panes and int(re.sub(r"\D", "", sheet.freeze_panes) or 0) or 5
    fields = {2: "doc_no", 3: "date", 4: "customer_name", 5: "vat_amount", 6: "gl_amount", 8: "account_code"}
    row_key = lambda row: (str(row.get("doc_no") or ""), show_date(row.get("date")),  # noqa: E731
        str(row.get("customer_name") or ""), _num(row.get("vat_amount")))
    cell_key = lambda v: (str(v[1] or ""), str(v[2] or ""), str(v[3] or ""), _num(v[4]))  # noqa: E731
    return {(0, sheet_row, column): ["vat", index, field]
            for sheet_row, index in _by_content(sheet, first, inputs.get("vat") or [], row_key, cell_key, 5)
            for column, field in fields.items()}


def _vat_edits(_book, inputs):
    edits = {}
    invoice_fields = {2: "buyer_tax_id", 3: "buyer_name", 4: "invoice_no", 5: "invoice_date",
                      7: "amount_pre_vat", 8: "vat_amount", 9: "total_amount"}
    report_fields = {2: "report_invoice_no", 3: "report_buyer_tax_id", 4: "report_buyer_branch", 5: "report_buyer_name",
                     6: "report_date", 8: "report_amount_pre_vat", 9: "report_vat_amount"}
    for sheet_index, key, fields in ((0, "invoices", invoice_fields), (1, "report_rows", report_fields)):
        for index in range(len(inputs.get(key) or [])):
            for column, field in fields.items():
                edits[(sheet_index, index + 2, column)] = [key, index, field]
    return edits


EDITS = {"bank": _bank_edits, "income": _income_edits, "vat": _vat_edits}


# ── 一页 → 页面照画的数据 ─────────────────────────────────────────────────


def _sheet(sheet, index, edits, inputs, styles, style_ids):
    from openpyxl.utils import get_column_letter

    rows, columns = sheet.max_row, sheet.max_column
    if rows * columns > MAX_CELLS:
        frappe.throw(_("reconciliation.rows_invalid"))
    covered, spans = set(), {}
    for merged in sheet.merged_cells.ranges:
        spans[(merged.min_row, merged.min_col)] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                if (row, column) != (merged.min_row, merged.min_col):
                    covered.add((row, column))
    widths = []
    for column in range(1, columns + 1):
        dimension = sheet.column_dimensions.get(get_column_letter(column))
        width = dimension.width if dimension is not None and dimension.width else 8.43
        widths.append(0 if dimension is not None and dimension.hidden else round(width * 7 + 5))
    out_rows = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=1, max_row=rows, max_col=columns), 1):
        dimension = sheet.row_dimensions.get(row_number)
        height = dimension.height if dimension is not None and dimension.height else None
        line = []
        for cell in cells:
            position = (row_number, cell.column)
            if position in covered:
                line.append(None)
                continue
            text, red, numeric = _value(sheet, cell)
            style = _style(cell, numeric, red)
            key = repr(sorted(style.items()))
            if key not in style_ids:
                style_ids[key] = len(styles)
                styles.append(style)
            item = [text, style_ids[key]]
            extra = {}
            if position in spans:
                extra["rs"], extra["cs"] = spans[position]
            edit = edits.get((index,) + position)
            if edit:
                raw = (inputs.get(edit[0]) or [])[edit[1]].get(edit[2])
                extra["e"] = edit + [None if raw is None else raw]
            if extra:
                item.append(extra)
            line.append(item)
        out_rows.append({"h": round(height * 4 / 3) if height else None, "c": line})
    freeze = sheet.freeze_panes
    frozen_rows = frozen_columns = 0
    if freeze:
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

        letters, number = coordinate_from_string(freeze)
        frozen_rows, frozen_columns = number - 1, column_index_from_string(letters) - 1
    return {"index": index, "name": sheet.title, "widths": widths, "rows": out_rows,
            "grid": bool(sheet.sheet_view.showGridLines is not False),
            "freeze": [frozen_rows, frozen_columns]}


def _render(kind, content, inputs):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content))
    edits = EDITS[kind](book, inputs)
    styles, style_ids = [], {}
    sheets = [_sheet(sheet, index, edits, inputs, styles, style_ids)
              for index, sheet in enumerate(book.worksheets)]
    return {"sheets": sheets, "styles": styles}


# ── 页面打进来的字 → 输入 ────────────────────────────────────────────────


AMOUNTS = {
    "bank": {"statement": ("withdrawal", "deposit", "balance"), "gl": ("debit", "credit", "balance")},
    "income": {"vat": ("vat_amount", "gl_amount")},
    "vat": {"invoices": ("amount_pre_vat", "vat_amount", "total_amount"),
            "report_rows": ("report_amount_pre_vat", "report_vat_amount", "report_amount")},
}
BLANK_IS_NONE = {("statement", "balance"), ("gl", "balance"), ("vat", "gl_amount")}


def _amount(value, blank):
    if value is None or isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(",", "").replace(" ", "")
    if not text:
        return blank
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        number = round(float(text), 2)
    except ValueError:
        frappe.throw(_("reconciliation.amount_invalid"))
    return -number if negative else number


def _bank_date(value):
    """复核表里日期显示成 日/月/年；打进来的日/月/年（佛历年自动减 543）换回任务里存的 年-月-日。"""
    text = str(value or "").strip()
    match = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})$", text)
    if not match:
        return text
    # 年份与识别同一规则（`normalization.parse_date`）：两位年是佛历缩写（69 → 2569 → 2026），
    # 不是 2069（2026-09-28 录制发现）；四位佛历减 543。
    day, month, year = match.groups()
    parsed = normalization.parse_date(f"{int(day):02d}/{int(month):02d}/{year}")
    return parsed.isoformat() if parsed else text


def _normalize(kind, inputs):
    for key, fields in AMOUNTS[kind].items():
        for row in inputs.get(key) or []:
            if not isinstance(row, dict):
                continue
            for field in fields:
                if field in row:
                    row[field] = _amount(row[field], None if (key, field) in BLANK_IS_NONE else 0.0)
            if kind in ("bank", "income") and "date" in row:
                row["date"] = _bank_date(row.get("date"))  # 两边同一规则：打的日/月/年换成存的年-月-日（佛历减 543）
            if kind == "income":
                row["has_gl"] = int(row.get("gl_amount") is not None)
            if key == "report_rows" and row.get("report_amount_pre_vat") not in (None, ""):
                # 报告明细里“含税”一栏就是不含税 + VAT（导出也是这样算的），改了两者之一跟着走。
                row["report_amount"] = round(float(row["report_amount_pre_vat"]) + float(row.get("report_vat_amount") or 0), 2)
    return inputs


# ── 接口 ────────────────────────────────────────────────────────────────


def _task(kind, task_id, for_update=False):
    if kind == "vat":
        return vat_check_service.task(task_id, for_update=for_update)
    return revision.task(kind, task_id, for_update=for_update)


# ── 页面上不显示的内容（用户 2026-09-27："有关模型的信息就不要留，和对账无关的备注文案取消掉"）──
# 只影响页面：导出的 Excel 与计算不变。去掉识别置信度列、各种"原文件"列、汇总里的 OCR 准确性检查段与文件信息段、
# 销项税备注里的"请核对原 PDF"；匹配层保留（它是对账依据：精确日期 / 日期容差 / 仅金额），只去掉 L1 / L2 / L3 代号。
HEADER_BG = "#1F2937"
_LAYER_CODE = re.compile(r"^L\d\s*-\s*")


def _hidden(lang):
    from pearnly_cowork.reconciliation import vat_check_core
    from pearnly_cowork.reconciliation.export import labels, vat_check_i18n

    t = lambda key: labels.t(key, lang)  # noqa: E731
    vat = vat_check_i18n._I18N.get(lang) or vat_check_i18n._I18N["th"]
    source = vat_check_i18n._I18N["zh"]["h_inv"].index("原文件")
    return {
        "columns": {t("col_confidence"), t("col_source_file"), t("col_source_stmt"), t("col_source_gl"),
                    vat["h_inv"][source]},
        "sections": {t("lbl_ocr_check"), t("sh_fileinfo")},
        "layers": tuple(t(f"layer_{n}") for n in (1, 2, 3, 4)),
        "notes": [vat_check_core.diff_labels(lang)["ocr_incomplete"]],
    }


def _for_page(view, lang):
    hidden, styles = _hidden(lang), view["styles"]
    header = lambda item: bool(item) and styles[item[1]].get("bg") == HEADER_BG  # noqa: E731
    for sheet in view["sheets"]:
        rows = sheet["rows"]
        # 1. 整列去掉：表头行里标题命中的列（明细表没有合并格）。
        head = next((row for row in rows if any(header(item) for item in row["c"])), None)
        drop = [index for index, item in enumerate(head["c"]) if header(item) and item[0] in hidden["columns"]] if head else []
        for index in reversed(drop):
            sheet["widths"].pop(index)
            for row in rows:
                if index < len(row["c"]):
                    row["c"].pop(index)
            if index < sheet["freeze"][1]:
                sheet["freeze"][1] -= 1
        # 2. 汇总里整段去掉：段标题行起，到下一个空行为止。
        kept, skipping = [], False
        for row in rows:
            first = row["c"][0] if row["c"] else None
            blank = all(not item or item[0] == "" for item in row["c"])
            if first and first[0] in hidden["sections"]:
                skipping = True
                continue
            if skipping and not blank:
                continue
            skipping = False
            kept.append(row)
        while kept and all(not item or item[0] == "" for item in kept[-1]["c"]):
            kept.pop()
        sheet["rows"] = kept
        # 3. 格子里的字：匹配层去代号，备注去掉"核对原 PDF"。
        for row in kept:
            for item in row["c"]:
                if not item or not item[0]:
                    continue
                if isinstance(item[0], str) and item[0].startswith(hidden["layers"]):  # 合计组后面带组号「 #3」
                    item[0] = _LAYER_CODE.sub("", item[0])
                for note in hidden["notes"]:
                    if note in item[0]:
                        item[0] = " · ".join(part for part in item[0].replace(note, "").split(" · ") if part.strip())
    return view


def _view(kind, doc, inputs, lang):
    if kind == "vat":
        content, _summary = vat_check_service.compute(vat_check_service._inputs(inputs), lang)
    elif kind == "bank":
        _rows, _header, results, summary = revision.bank_compute(doc, inputs)
        content = revision.bank_workbook_bytes(doc, results, summary, lang)
    else:
        _rows, results, summary = revision.income_compute(doc, inputs)
        content = revision.income_workbook_bytes(results, summary, lang)
    return {**_for_page(_render(kind, content, inputs), lang), "inputs": inputs}


def _parse(kind, inputs):
    value = frappe.parse_json(inputs) if isinstance(inputs, str) else inputs
    if not isinstance(value, dict):
        frappe.throw(_("reconciliation.request_invalid"))
    return _normalize(kind, value)


def purchase_invoice_notes(inputs):
    """发票的买方就是本公司（税报上的卖方税号）：多半是进项发票放进了销项核查（用户 2026-09-30 定只提示；V04 的
    12 份全是 GIZE 收到的发票）。点名是哪几份文件。"""
    from pearnly_cowork.reconciliation.normalization import normalize_tax_id

    from pearnly_cowork.reconciliation.normalization import company_name

    own = normalize_tax_id(inputs.get("seller_tax_id") or "")
    if not own:
        return []
    invoices = inputs.get("invoices") or []
    # 税号对上的那几张里本公司的写法；税号没读出 / 读错的发票，买方名称是同一写法也算（V04 的 RN3062065：税号读成
    # 别的号码，买方写「GIZE COMPANY LIMITED」）。
    names = {company_name(invoice.get("buyer_name")) for invoice in invoices
             if normalize_tax_id(invoice.get("buyer_tax_id") or "") == own} - {""}
    files = []
    for invoice in invoices:
        name = str(invoice.get("filename") or invoice.get("source_filename") or "")
        ours = (normalize_tax_id(invoice.get("buyer_tax_id") or "") == own
                or company_name(invoice.get("buyer_name")) in names)
        if ours and name and name not in files:
            files.append(name)
    return [_("vat_check.note.purchase_invoices").format(", ".join(files))] if files else []


@frappe.whitelist()
def load(kind, task_id, lang=None):
    doc = _task(kind, task_id)
    if kind == "vat":
        inputs, version = vat_check_service.stored_inputs(doc), vat_check_service.revision(doc)
    else:
        inputs = revision.bank_inputs(doc) if kind == "bank" else revision.income_inputs(doc)
        version = revision.revision(doc)
    notes = revision.bank_notes(doc) if kind == "bank" else revision.income_notes(doc) if kind == "income" else []
    questions = []
    if kind == "bank":
        # 多笔对一笔的疑似合计匹配：按当前（含复核改动的）结果算。这个账户还没定过就问一次（对账预设）；
        # 定了「是」的已按合计对上（标题旁说明），定了「不是」的不再提示。
        results = revision.bank_compute(doc, inputs)[2]
        questions = revision.bank_questions(doc, inputs, results)
        notes += revision.bank_preset_notes(results)
        if revision.bank_rules(doc, inputs.get("gl") or []).get(presets.GROUPED) != presets.NO:
            notes += revision.bank_group_notes(results)
    if kind == "vat" and inputs.get("failed_files"):
        # 读不出的发票文件原来存着却不显示（2026-09-28 录制读代码发现）：与银行同样写在标题下一句。
        notes.append(_("vat_check.note.unread").format(", ".join(str(name) for name in inputs["failed_files"])))
    if kind == "vat":
        notes += purchase_invoice_notes(inputs)
    return {"task": {"name": doc.name, "title": doc.title, "status": doc.status, "modified": str(doc.modified),
                     "revision": version, "notes": notes, "questions": questions}, **_view(kind, doc, inputs, _lang(lang))}


@frappe.whitelist(methods=["POST"])
def recalc(kind, task_id, inputs, lang=None):
    doc = _task(kind, task_id)
    return _view(kind, doc, _parse(kind, inputs), _lang(lang))


@frappe.whitelist(methods=["POST"])
def breaks(kind, task_id, inputs):
    """需复核的银行对账保存前：余额链还连不上的几处（逐个确认成「此处缺页 / 已知差异」才可保存）。"""
    doc = _task(kind, task_id)
    return revision.bank_breaks(doc, _parse(kind, inputs)) if kind == "bank" else []


@frappe.whitelist(methods=["POST"])
def save(kind, task_id, inputs, modified, lang=None, acknowledged=None):
    doc = _task(kind, task_id, for_update=True)
    if str(doc.modified) != str(modified):
        frappe.throw(_("reconciliation.edit_conflict"))
    inputs = _parse(kind, inputs)
    if kind == "vat":
        vat_check_service.save(doc, inputs, _lang(lang))
    elif kind == "bank":
        marks = frappe.parse_json(acknowledged) if isinstance(acknowledged, str) else acknowledged
        revision.bank_save(doc, inputs, marks if isinstance(marks, dict) else None)
    else:
        revision.income_save(doc, inputs)
    frappe.db.commit()
    return load(kind, task_id, lang)
