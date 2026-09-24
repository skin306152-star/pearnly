"""Bounded, lossless XLSX input for opening stock; no framework or storage IO."""
from io import BytesIO
from zipfile import ZipFile

MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 5000
FIELD_LABELS = (
    ("product", "stock.opening.productId"),
    ("product_code", "stock.col.code"),
    ("product_name", "stock.col.name"),
    ("barcode", "stock.col.barcode"),
    ("unit", "stock.col.unit"),
    ("location", "stock.count.location"),
    ("qty", "stock.opening.qty"),
    ("unit_cost", "stock.opening.cost"),
)


def _norm(value):
    return " ".join(str(value or "").strip().lower().split())


def header_aliases(labels=None):
    aliases = {_norm(field): field for field, _ in FIELD_LABELS}
    aliases.update({_norm(key): field for field, key in FIELD_LABELS})
    aliases.update({_norm(key): field for key, field in (labels or {}).items()})
    return aliases


def parse_workbook(data, aliases=None):
    """Reject malformed structure; preserve invalid row values for batch preview."""
    import openpyxl

    if not isinstance(data, bytes) or len(data) > MAX_BYTES:
        raise ValueError("file_too_large")
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("file_too_large")
            if any(entry.flag_bits & 1 for entry in entries):
                raise ValueError("file_invalid")
        book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("file_invalid") from error
    try:
        if len(book.worksheets) != 1:
            raise ValueError("sheet_invalid")
        sheet = book.worksheets[0]
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ValueError("too_many_rows")
        if sheet.max_column and sheet.max_column > len(FIELD_LABELS):
            raise ValueError("headers_invalid")
        values = iter(sheet.iter_rows())
        header = next(values, ())
        aliases = aliases or header_aliases()
        fields = []
        for cell in header:
            field = aliases.get(_norm(cell.value))
            if not field or field in fields or cell.data_type in ("f", "e"):
                raise ValueError("headers_invalid")
            fields.append(field)
        if not ({"product", "product_code", "barcode"} & set(fields)) or not {"qty", "unit_cost"} <= set(fields):
            raise ValueError("headers_invalid")
        rows = []
        for number, cells in enumerate(values, start=2):
            if number > MAX_ROWS + 1:
                raise ValueError("too_many_rows")
            if any(cell.data_type in ("f", "e") for cell in cells):
                raise ValueError("formula_invalid")
            if all(cell.value in (None, "") for cell in cells):
                continue
            row = {field: cell.value for field, cell in zip(fields, cells)}
            if all(row.get(field) is None or isinstance(row.get(field), str) and not row[field].strip() for field in ("qty", "unit_cost")):
                # Unfilled catalog template rows are not instructions to set zero.
                continue
            # Numeric codes destroy leading zeros in Excel. Require literal text;
            # silently converting them could select a different existing product.
            for field in ("product", "product_code", "barcode"):
                if row.get(field) is not None and not isinstance(row[field], str):
                    raise ValueError("identifier_text_required")
            for field in ("qty", "unit_cost"):
                if row.get(field) is not None:
                    row[field] = str(row[field])
            row["row_number"] = number
            rows.append(row)
        if not rows:
            raise ValueError("empty")
        return rows
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("file_invalid") from error
    finally:
        book.close()


def build_bytes(headers, rows):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Opening stock"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EEE7FC")
        cell.data_type = "s"
    for row in rows:
        sheet.append([row.get(field) for field, _ in FIELD_LABELS])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
    for number, (field, _) in enumerate(FIELD_LABELS, 1):
        sheet.column_dimensions[get_column_letter(number)].width = 32 if field == "product_name" else 22
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
