"""盘点表的生成与解析（纯表格层：不连库、不鉴权）。

为什么单独一层：表头认不出、数量不是数字、少了商品标识 —— 这些**都能在没有 Site 的
地方断言**。真实写入由 `inventory.count` 逐行处理，两边不各写一套规则。

字段照旧项目 `services/stocktake/excel.py`：商品编号 / 商品名称 / 条码 / 单位 /
账面数量 / 仓库 / 库位，**只多一列实盘数量** —— 用户 2026-09-22 定：
"上传导入……把这次数到的填进来"。

表头认三套（与商品导入同一套做法）：**字段名**、**旧项目那张表的写法**（泰/中/英）、
以及**界面上的那几句**（走语言包）—— 所以"下载得到的模板"与"能认的表头"永远是同一张表。
"""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation

# 字段 → 界面文案的键。模板表头就取这几句（按站点语言），顺序也是这个顺序。
FIELD_LABELS = (
    ("product_code", "stock.col.code"),
    ("product_name", "stock.col.name"),
    ("barcode", "stock.col.barcode"),
    ("unit", "stock.col.unit"),
    ("book_qty", "stock.count.book"),
    ("counted_qty", "stock.count.counted"),
    ("warehouse", "stock.col.warehouse"),
    ("location", "stock.count.location"),
)

# 旧项目那张表、导出件、客户手里的历史文件长这样，照认（比对前统一去空白、转小写）。
LEGACY_ALIASES = {
    "product_code": "product_code",
    "code": "product_code",
    "รหัสสินค้า": "product_code",
    "รหัส": "product_code",
    "商品编号": "product_code",
    "编码": "product_code",
    "货号": "product_code",
    "商品コード": "product_code",
    "product_name": "product_name",
    "name": "product_name",
    "ชื่อสินค้า": "product_name",
    "ชื่อ": "product_name",
    "商品名称": "product_name",
    "名称": "product_name",
    "商品名": "product_name",
    "barcode": "barcode",
    "บาร์โค้ด": "barcode",
    "บาร์โค้ด / รหัส qr": "barcode",
    "บาร์โค้ด หรือ รหัส qr": "barcode",
    # 商品那边的码合并成一格之后，盘点表这一列也跟着改叫"条码/二维码"（同一个含义：
    # 扫得出这条商品的码）。旧表头照旧认，新表头也要认。
    "barcode / qr code": "barcode",
    "条码/二维码": "barcode",
    "条码 / 二维码": "barcode",
    "条码": "barcode",
    "バーコード": "barcode",
    "unit": "unit",
    "หน่วย": "unit",
    "单位": "unit",
    "単位": "unit",
    "book_qty": "book_qty",
    "book quantity": "book_qty",
    "账面数量": "book_qty",
    "จำนวนตามบัญชี": "book_qty",
    "帳簿数量": "book_qty",
    "counted_qty": "counted_qty",
    "actual_qty": "counted_qty",
    "counted quantity": "counted_qty",
    "实盘数量": "counted_qty",
    "จำนวนที่นับได้": "counted_qty",
    "実地数量": "counted_qty",
    "warehouse": "warehouse",
    "仓库": "warehouse",
    "คลังสินค้า": "warehouse",
    "倉庫": "warehouse",
    "คลังสินค้าเริ่มต้น (ไม่บังคับ)": "warehouse",
    "warehouse (optional)": "warehouse",
    "location": "location",
    "库位": "location",
    "ตำแหน่ง": "location",
    "ช่องเก็บ": "location",
    "棚番": "location",
    "ตำแหน่งเริ่มต้น (ไม่บังคับ)": "location",
    "location (optional)": "location",
}

MAX_ROWS = 10000


def header_aliases(labels=None) -> dict:
    """表头 → 字段：字段名、旧写法、界面标签三套并集（与商品导入同一条规矩）。

    `labels` 是"文案 → 字段"（调用方从语言包取四语），本层不碰 Frappe 与语言包。
    """
    aliases = {_norm(field): field for field, _key in FIELD_LABELS}
    aliases.update({_norm(word): field for word, field in LEGACY_ALIASES.items()})
    for text, field in (labels or {}).items():
        if str(text or "").strip():
            aliases[_norm(text)] = field
    return aliases


def template_headers(label_of) -> list:
    return [label_of(key) for _field, key in FIELD_LABELS]


def _norm(value) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _quantity(value, field):
    """空 = 没填（不是 0）；账面可为负，实盘必须非负。"""
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field}:not_a_number") from None
    if not number.is_finite() or (field == "counted_qty" and number < 0):
        raise ValueError(f"{field}:invalid")
    return number


def parse_workbook(data: bytes, aliases: dict | None = None) -> tuple[list[dict], list[dict]]:
    """读表 → (行, 问题)。坏行只跳过自己，其余照收（与商品导入同一条规矩）。"""
    import openpyxl

    aliases = aliases or header_aliases()
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("file_invalid") from None
    try:
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("headers_invalid")
        columns = {}
        for index, cell in enumerate(header):
            field = aliases.get(_norm(cell))
            if field and field not in columns:
                columns[field] = index
        if not {"product_code", "barcode"} & set(columns):
            raise ValueError("headers_invalid")
        out, problems = [], []
        for number, cells in enumerate(rows, start=2):
            if number > MAX_ROWS + 1:
                raise ValueError("too_many_rows")
            if not cells or all(cell is None or str(cell).strip() == "" for cell in cells):
                continue
            value = lambda field: (cells[columns[field]]
                                   if field in columns and columns[field] < len(cells) else None)
            try:
                row = {
                    "product_code": str(value("product_code") or "").strip(),
                    "product_name": str(value("product_name") or "").strip(),
                    "barcode": str(value("barcode") or "").strip(),
                    "unit": str(value("unit") or "").strip(),
                    "warehouse": str(value("warehouse") or "").strip(),
                    "location": str(value("location") or "").strip(),
                    "book_qty": _quantity(value("book_qty"), "book_qty"),
                    "counted_qty": _quantity(value("counted_qty"), "counted_qty"),
                }
            except ValueError as error:
                problems.append({"row": number, "reason": str(error)})
                continue
            if not (row["product_code"] or row["barcode"]):
                problems.append({"row": number, "reason": "no_identifier"})
                continue
            out.append(row)
        if not out and not problems:
            raise ValueError("empty")
        return out, problems
    finally:
        book.close()


def build_bytes(headers: list, rows: list | None = None) -> bytes:
    """生成表格：只有表头（模版）或带数据（下载时已经填好的表）。"""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Stock count"
    sheet.append([str(header) for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(
            14, len(str(header)) + 4)
    # 编号与条码预置成文本：Excel 会把长数字的前导零吃掉，写回去就找不到商品了。
    text_columns = [index for index, (field, _key) in enumerate(FIELD_LABELS, start=1)
                    if field in ("product_code", "barcode")]
    for row in rows or []:
        line = []
        for field, _key in FIELD_LABELS:
            cell_value = row.get(field)
            if isinstance(cell_value, Decimal):
                cell_value = float(cell_value)
            line.append("" if cell_value is None else cell_value)
        sheet.append(line)
        for index in text_columns:
            sheet.cell(sheet.max_row, index).number_format = "@"
    if rows:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()
