"""Real-Site ERP flow acceptance, restricted to an isolated synthetic lab.

Creates uniquely marked fixtures and native private files; never deletes another
test's data. OCR input is a supplied extraction result, so no provider is called.
Service commits are intentional: run only on the disposable acceptance Site.
"""
import io
import json
import unittest
from decimal import Decimal
from uuid import uuid4

import frappe
import openpyxl

from pearnly_erp_portal.access import menu
from pearnly_erp_portal.access.setup import sync_roles
from pearnly_erp_portal.cashier import api as cashier_api
from pearnly_erp_portal.inventory import api as stock_api, stock, workbook
from pearnly_erp_portal.masters import api as master_api, company
from pearnly_erp_portal.products import contracts as product_contracts, service as products
from pearnly_erp_portal.purchasing import entry as purchase, recognition, suppliers
from pearnly_erp_portal.team import service as team


class EntryFlowAcceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        assert frappe.conf.pearnly_erp_pos_lab and frappe.conf.db_host == "db"
        frappe.set_user("Administrator")
        from pearnly_core.provisioning import ensure_feature_roles
        from pearnly_control.features import codes_of, roles_payload

        ensure_feature_roles(json.dumps(roles_payload("erp", codes_of("erp"))), admin_user="Administrator")
        frappe.conf.pearnly_features = codes_of("erp")
        sync_roles()
        cls.prefix = "flowaccept" + uuid4().hex[:10]
        cls.company = company.create_company(company_name=cls.prefix + " company", pos_enabled=True)
        company.set_active_company(cls.company["name"])
        cls.warehouse = cls.company["warehouse"]
        cls.terminal = cls.company["terminals"][0]
        cls.products = {}
        for key in ("zero", "blank", "counted_zero", "inventory", "card", "cashier", "cashier_mixed", "purchase"):
            cls.products[key] = products.create_product(product_name=cls.prefix + " " + key,
                code=cls.prefix + "-" + key, unit="pcs", default_cost=10, unit_price=20)["name"]
        for key, quantity in (("blank", 7), ("counted_zero", 5), ("inventory", 9), ("cashier", 8), ("cashier_mixed", 8)):
            stock.apply_movement(product=cls.products[key], warehouse=cls.warehouse,
                movement_type="stock_in", qty=quantity, unit_cost=10,
                reason=cls.prefix, client_uuid=cls.prefix + "-seed-" + key, internal=True)
        cls.tax_id = str(int(uuid4().hex[:10], 16)).zfill(13)
        cls.suppliers = {}
        for branch in ("head", "0001", "0002"):
            cls.suppliers[branch] = suppliers.create_supplier(
                supplier_name=cls.prefix + " supplier " + branch, tax_id=cls.tax_id,
                branch_type="head" if branch == "head" else "branch",
                branch_no=None if branch == "head" else branch)["name"]
        cls.users = {}
        for key in ("stock-count", "purchase", "inventory", "stock-card", "cashier", "company", "warehouses"):
            username = cls.prefix + key.replace("-", "")
            cls.users[key] = username + "@example.test"
            team.add_member(username=username, email=cls.users[key], password="lab-only",
                            roles=[menu.role(key)], companies=[cls.company["name"]])
        frappe.db.commit()

    def setUp(self):
        self.as_role("purchase")
        frappe.db.commit()

    def tearDown(self):
        frappe.db.rollback()
        frappe.clear_cache()
        frappe.set_user("Administrator")

    def as_role(self, key):
        frappe.set_user(self.users[key])
        company.set_active_company(self.company["name"])

    def denied(self, operation, *args, **kwargs):
        with self.assertRaises(frappe.PermissionError):
            operation(*args, **kwargs)

    def quantity(self, key):
        return stock.current_balance(self.products[key], self.warehouse).qty

    def upload(self, suffix, content):
        uploaded = frappe.get_doc({"doctype": "File", "file_name": self.prefix + "-" + suffix,
                                   "content": content, "is_private": 1})
        uploaded.insert(ignore_permissions=True)
        frappe.db.commit()
        return uploaded

    def payload(self, **changes):
        value = {"company": self.company["name"], "doc_date": "2026-09-24", "doc_no": uuid4().hex,
                 "supplier": self.suppliers["head"], "has_vat": 0, "currency": "THB", "fx_rate": 1,
                 "is_paid": 1, "payment_method": "cash", "auto_create_product": 0,
                 "lines": [{"product": self.products["purchase"], "description": self.prefix + " purchase",
                            "qty": 2, "unit": "pcs", "unit_price": 10, "warehouse": self.warehouse}]}
        value.update(changes)
        return value

    def test_xlsx_file_import_and_post_preserve_blank_and_dedupe(self):
        self.as_role("stock-count")
        previous_response = dict(frappe.local.response)
        try:
            stock_api.api_stock_count_template(warehouse=self.warehouse)
            template = frappe.local.response.filecontent
            self.assertTrue(frappe.local.response.filename.endswith(".xlsx"))
        finally:
            frappe.local.response.clear()
            frappe.local.response.update(previous_response)
        book = openpyxl.load_workbook(io.BytesIO(template))
        sheet = book.worksheets[0]
        aliases = workbook.header_aliases(stock_api._stock_label_texts())
        columns = {aliases[" ".join(str(cell.value).strip().lower().split())]: cell.column
                   for cell in sheet[1]}
        codes = {self.prefix + "-" + key: key for key in ("zero", "blank", "counted_zero")}
        retained = {}
        for index in range(sheet.max_row, 1, -1):
            code = str(sheet.cell(index, columns["product_code"]).value)
            if code not in codes:
                sheet.delete_rows(index)
        for index in range(2, sheet.max_row + 1):
            key = codes[str(sheet.cell(index, columns["product_code"]).value)]
            retained[key] = sheet.cell(index, columns["book_qty"]).value
            self.assertIsNone(sheet.cell(index, columns["counted_qty"]).value)
            sheet.cell(index, columns["counted_qty"], {"zero": 3, "counted_zero": 0}.get(key))
        self.assertEqual(retained, {"zero": 0, "blank": 7, "counted_zero": 5})
        data = io.BytesIO()
        book.save(data)
        book.close()
        uploaded = self.upload("filled-count.xlsx", data.getvalue())
        task = stock_api.api_create_stock_count(title=self.prefix + " XLSX count",
            warehouse=self.warehouse, file_url=uploaded.file_url)
        self.assertEqual((task["rows"], task["counted"], task["problems"]), (3, 2, []))
        repeated_import = stock_api.api_stock_count_import(name=task["name"], file_url=uploaded.file_url)
        self.assertEqual((repeated_import["rows"], repeated_import["counted"], repeated_import["problems"]),
                         (3, 2, []))
        items = stock_api.api_stock_count_items(count=task["name"], limit=0)
        by_product = {row["product"]: row for row in items}
        self.assertFalse(by_product[self.products["blank"]]["counted"])
        self.assertTrue(by_product[self.products["counted_zero"]]["counted"])
        self.assertEqual(len(stock_api.api_stock_count_entries(count=task["name"], limit=0)), 2)
        review = stock_api.api_stock_count_review(task["name"])
        posted = stock_api.api_stock_count_post(name=task["name"], review_token=review["review_token"])
        self.assertEqual(posted["status"], "Posted")
        self.assertEqual(len(posted["adjustments"]), 2)
        self.assertEqual(tuple(self.quantity(key) for key in ("zero", "blank", "counted_zero")),
                         (Decimal(3), Decimal(7), Decimal(0)))
        entries = [row["entry"] for row in posted["adjustments"]]
        for name in entries:
            doc = frappe.get_doc("ERP Stock Entry", name)
            self.assertEqual((doc.ref_type, doc.ref_id), ("count", task["name"]))
            self.assertTrue(doc.doc_no)
        movements = frappe.db.count("ERP Stock Movement", {"ref_id": ["in", entries]})
        self.assertEqual(movements, 2)
        replay = stock_api.api_stock_count_post(name=task["name"], review_token=review["review_token"])
        self.assertTrue(replay["deduped"])
        self.assertEqual(frappe.db.count("ERP Stock Movement", {"ref_id": ["in", entries]}), movements)
        self.assertEqual(tuple(self.quantity(key) for key in ("zero", "blank", "counted_zero")),
                         (Decimal(3), Decimal(7), Decimal(0)))

    def test_ocr_tax_and_branch_match_despite_wrong_name(self):
        for branch in ("head", "0001", "0002"):
            with self.subTest(branch=branch):
                raw = {"supplier": "1234", "supplier_name": self.prefix + " OCR misspelling",
                       "supplier_tax_id": self.tax_id, "supplier_branch_type": "head_office" if branch == "head" else "branch",
                       "supplier_branch_no": None if branch == "head" else branch,
                       "doc_no": uuid4().hex, "doc_date": "2026-09-24", "has_vat": 0,
                       "lines": [{"code": self.prefix + "-purchase", "description": "OCR item spelling",
                                  "qty": 1, "unit": "pcs", "unit_price": 10}]}
                clean = recognition.clean_result(raw)
                self.assertNotIn("supplier", clean)
                payload = self.payload(**clean)
                payload.pop("supplier")
                result = purchase.post(payload, request_id=uuid4().hex)
                doc = frappe.get_doc("Purchase Document", result["name"])
                self.assertEqual(doc.supplier, self.suppliers[branch])
                self.assertEqual(doc.lines[0].product, self.products["purchase"])
                self.assertEqual(doc.supplier_tax_id, self.tax_id)
                if branch != "head":
                    self.assertEqual(doc.supplier_branch_no, branch)
        self.assertEqual(frappe.db.count("ERP Supplier", {"tax_id": self.tax_id}), 3)

    def test_auto_created_supplier_and_product_roll_back_after_late_failure(self):
        label = self.prefix + " auto " + uuid4().hex[:8]
        request = uuid4().hex
        payload = self.payload(supplier=None, supplier_name=label, auto_create_product=1,
            lines=[{"description": label, "code": label.replace(" ", "-"), "qty": 2,
                    "unit": "pcs", "unit_price": 10, "warehouse": self.warehouse}])
        observed = {}

        def fail_after_all_writes(result):
            doc = frappe.get_doc("Purchase Document", result["name"])
            observed.update(document=doc.name, supplier=doc.supplier, product=doc.lines[0].product)
            self.assertEqual(stock.current_balance(observed["product"], self.warehouse).qty, Decimal(2))
            self.assertEqual(Decimal(str(doc.paid_amount)), Decimal(20))
            self.assertEqual(len(doc.payments), 1)
            raise RuntimeError("synthetic automatic masters finalization failure")

        with self.assertRaisesRegex(RuntimeError, "synthetic automatic masters"):
            purchase.post(payload, request_id=request, finalize=fail_after_all_writes)
        for doctype, key in (("Purchase Document", "document"), ("ERP Supplier", "supplier"), ("ERP Product", "product")):
            self.assertFalse(frappe.db.exists(doctype, observed[key]))
        self.assertFalse(frappe.db.exists("ERP Stock Movement", {"product": observed["product"]}))
        self.assertFalse(frappe.db.exists("ERP Stock Balance", {"product": observed["product"]}))
        self.assertFalse(frappe.db.exists("ERP Purchase Payment", {"parent": observed["document"]}))
        result = purchase.post(payload, request_id=request)
        doc = frappe.get_doc("Purchase Document", result["name"])
        self.assertEqual(frappe.db.count("ERP Supplier", {"supplier_name": label}), 1)
        # 票面货号只用来匹配，新建的商品按本系统续号（用户 2026-09-26）。
        self.assertEqual(frappe.db.count("ERP Product", {"product_name": label}), 1)
        self.assertFalse(frappe.db.exists("ERP Product", {"code": label.replace(" ", "-")}))
        self.assertTrue(frappe.db.get_value("ERP Product", doc.lines[0].product, "code").startswith(product_contracts.CODE_PREFIX))
        self.assertEqual(stock.current_balance(doc.lines[0].product, self.warehouse).qty, Decimal(2))
        self.assertTrue(purchase.post(payload, request_id=request)["duplicate"])
        self.assertEqual(stock.current_balance(doc.lines[0].product, self.warehouse).qty, Decimal(2))

    def test_amendment_copies_bill_and_persists_new_attachment_once(self):
        bill = self.upload("invoice.txt", self.prefix + " bill " + uuid4().hex)
        old_proof = self.upload("old-payment.txt", self.prefix + " old payment " + uuid4().hex)
        before = self.quantity("purchase")
        original = purchase.post(self.payload(attachments=[{"file": bill.name, "kind": "bill"},
            {"file": old_proof.name, "kind": "payment_proof"}]), request_id=uuid4().hex)
        proof = self.upload("new-payment.txt", self.prefix + " new payment " + uuid4().hex)
        payload = self.payload(attachments=[{"file": proof.name, "kind": "payment_proof"}])
        payload["lines"][0]["qty"] = 3
        request = uuid4().hex
        result = purchase.post(payload, request_id=request, amended_from=original["name"])
        corrected = frappe.get_doc("Purchase Document", result["name"])
        self.assertEqual(corrected.corrected_from, original["name"])
        self.assertEqual(frappe.db.get_value("Purchase Document", original["name"], "status"), "void")
        self.assertEqual(self.quantity("purchase"), before + Decimal(3))
        self.assertEqual((corrected.lines[0].qty, corrected.paid_amount), (3, 30))
        self.assertEqual(len(corrected.attachments), 2)
        copied = next(row for row in corrected.attachments if row.kind == "bill")
        self.assertNotEqual(copied.file, bill.name)
        self.assertEqual(frappe.get_doc("File", copied.file).get_content(), bill.get_content())
        self.assertEqual({row.file for row in corrected.attachments if row.kind == "payment_proof"}, {proof.name})
        for name, target in ((bill.name, original["name"]), (old_proof.name, original["name"]),
                             (copied.file, corrected.name), (proof.name, corrected.name)):
            saved = frappe.get_doc("File", name)
            self.assertEqual((saved.attached_to_doctype, saved.attached_to_name), ("Purchase Document", target))
            self.assertTrue(saved.is_private)
        self.assertEqual(corrected.image_sha256, frappe.db.get_value("Purchase Document", original["name"], "image_sha256"))
        attached = frappe.db.count("File", {"attached_to_doctype": "Purchase Document", "attached_to_name": corrected.name})
        replay = purchase.post(payload, request_id=request, amended_from=original["name"])
        self.assertTrue(replay["duplicate"])
        self.assertEqual(frappe.db.count("File", {"attached_to_doctype": "Purchase Document", "attached_to_name": corrected.name}), attached)
        self.assertEqual(self.quantity("purchase"), before + Decimal(3))

    def test_inventory_leaf_reads_business_stock_but_cannot_post_or_open_report(self):
        self.as_role("inventory")
        kwargs = {"warehouse": self.warehouse, "q": self.prefix + " inventory"}
        rows = stock_api.api_stock(**kwargs)
        self.assertEqual([(row["product"], row["qty_on_hand"]) for row in rows], [(self.products["inventory"], 9)])
        self.assertEqual(stock_api.api_stock_count(**kwargs), 1)
        self.assertEqual(stock_api.api_stock_totals(**kwargs)["sku_count"], 1)
        self.denied(stock_api.api_stock_card, product=self.products["inventory"], warehouse=self.warehouse)
        self.denied(stock_api.api_receive, warehouse=self.warehouse, lines=[{"product": self.products["inventory"], "qty": 1}])
        self.assertEqual(self.quantity("inventory"), Decimal(9))
        self.as_role("purchase")
        self.denied(stock_api.api_stock, **kwargs)

    def test_stock_card_leaf_writes_opening_and_reads_card_without_duplicate_stock(self):
        self.as_role("stock-card")
        args = {"product": self.products["card"], "warehouse": self.warehouse,
                "as_of_date": "2026-09-24", "qty": 4, "unit_cost": 10, "note": self.prefix}
        opening = stock_api.api_upsert_opening(**args)
        replay = stock_api.api_upsert_opening(**args)
        self.assertEqual(opening["name"], replay["name"])
        self.assertEqual(self.quantity("card"), Decimal(4))
        self.assertEqual(frappe.db.count("ERP Stock Movement", {"ref_type": "opening", "ref_id": opening["name"]}), 1)
        rows = stock_api.api_list_openings(product=self.products["card"], warehouse=self.warehouse)
        self.assertEqual([row["name"] for row in rows], [opening["name"]])
        card = stock_api.api_stock_card(product=self.products["card"], warehouse=self.warehouse,
                                       date_from="2026-09-23", date_to="2026-09-25")
        self.assertEqual(card["closing"]["qty"], 4)
        self.assertEqual(card["closing"]["value"], 40)
        self.denied(stock_api.api_stock, warehouse=self.warehouse)
        self.denied(stock_api.api_create_warehouse, company=self.company["name"], warehouse_name=self.prefix + " denied")
        self.as_role("inventory")
        self.denied(stock_api.api_upsert_opening, **args)

    def test_cashier_leaf_shift_sale_replay_and_settlement(self):
        self.as_role("cashier")
        self.assertEqual(cashier_api.api_current_sales(), [])
        found = cashier_api.api_counter_lookup(value=self.prefix + "-cashier", terminal=self.terminal)
        self.assertEqual((found["name"], found["warehouse"], found["qty_on_hand"]),
                         (self.products["cashier"], self.warehouse, 8))
        shift = cashier_api.api_open_shift(terminal=self.terminal, opening_float=100, note=self.prefix)
        active = cashier_api.api_current_shift(terminal=self.terminal)
        self.assertEqual(active["shift"]["name"], shift["name"])
        self.assertEqual(next(row["opening_amount"] for row in active["balances"] if row["method"] == "cash"), 100)
        args = {"terminal": self.terminal, "shift": shift["name"], "client_uuid": uuid4().hex,
                "lines": [{"product": self.products["cashier"], "description": self.prefix + " cashier",
                           "qty": 2, "unit": "pcs", "unit_price": 20}],
                "payments": [{"method": "cash", "amount": 40}]}
        sale = cashier_api.api_create_sale(**args)
        replay = cashier_api.api_create_sale(**args)
        self.assertEqual(sale["name"], replay["name"])
        self.assertTrue(replay["deduped"])
        self.assertEqual(self.quantity("cashier"), Decimal(6))
        self.assertEqual([row["name"] for row in cashier_api.api_current_sales()], [sale["name"]])
        self.denied(cashier_api.api_list_sales)
        self.denied(cashier_api.api_list_shifts)
        summary = cashier_api.api_shift_summary(name=shift["name"])
        self.assertEqual((summary["sales_count"], summary["gross"], summary["expected_cash"]), (1, 40, 140))
        self.assertEqual(next(row["expected_amount"] for row in summary["balances"] if row["method"] == "cash"), 140)
        closed = cashier_api.api_close_shift(name=shift["name"], counted_cash=140)
        self.assertEqual(closed["cash_diff"], 0)
        self.assertIsNone(cashier_api.api_current_shift(terminal=self.terminal))
        self.assertEqual(cashier_api.api_current_sales(), [])
        self.as_role("purchase")
        self.denied(cashier_api.api_open_shift, terminal=self.terminal)
        self.denied(cashier_api.api_create_sale, **args)
        self.assertEqual(self.quantity("cashier"), Decimal(6))

    def test_cashier_mixed_payments_and_legacy_cash_opening_reconcile(self):
        self.as_role("cashier")
        shift = cashier_api.api_open_shift(terminal=self.terminal,
            balances={"cash": 100, "card": 5, "transfer": 7}, note=self.prefix + " mixed")
        # Model a task-owned shift persisted by the old opening_float path: the
        # parent has 100 while the cash child was incorrectly initialized to 0.
        cash_row = frappe.db.get_value("ERP Shift Balance", {"parent": shift["name"], "method": "cash"}, "name")
        frappe.db.set_value("ERP Shift Balance", cash_row, "opening_amount", 0)
        sale = cashier_api.api_create_sale(terminal=self.terminal, shift=shift["name"], client_uuid=uuid4().hex,
            lines=[{"product": self.products["cashier_mixed"], "description": self.prefix + " mixed",
                    "qty": 2, "unit": "pcs", "unit_price": 50}],
            payments=[{"method": "cash", "amount": 60}, {"method": "card", "amount": 30},
                      {"method": "transfer", "amount": 20}])
        self.assertEqual((sale["grand_total"], sale["change_amount"]), (100, 10))
        summary = cashier_api.api_shift_summary(name=shift["name"])
        expected = {"cash": 150, "card": 35, "promptpay": 0, "transfer": 27}
        self.assertEqual(summary["expected_cash"], 150)
        self.assertEqual({row["method"]: row["expected_amount"] for row in summary["balances"]}, expected)
        closed = cashier_api.api_close_shift(name=shift["name"], counted=expected)
        self.assertEqual((closed["expected_cash"], closed["counted_cash"], closed["cash_diff"]), (150, 150, 0))
        saved = frappe.get_doc("ERP Shift", shift["name"])
        self.assertEqual({row.method: row.expected_amount for row in saved.balances}, expected)
        self.assertTrue(all(row.difference == 0 for row in saved.balances))
        self.assertEqual(next(row.opening_amount for row in saved.balances if row.method == "cash"), 100)
        self.assertEqual(self.quantity("cashier_mixed"), Decimal(6))

    def test_company_and_warehouse_leaf_management_boundaries(self):
        self.as_role("company")
        master_api.api_update_company(name=self.company["name"], phone="020000000", short_name=self.prefix)
        profile = master_api.api_get_profile(name=self.company["name"])
        self.assertEqual(profile["phone"], "020000000")
        self.assertTrue(profile["can_write"])
        self.assertTrue(any(row["name"] == self.terminal for row in master_api.api_company_terminals(company_name=self.company["name"])))
        self.denied(stock_api.api_create_warehouse, company=self.company["name"], warehouse_name=self.prefix + " denied")
        self.denied(cashier_api.api_open_shift, terminal=self.terminal)
        self.as_role("warehouses")
        created = stock_api.api_create_warehouse(company=self.company["name"], warehouse_name=self.prefix + " extra")
        stock_api.api_update_warehouse(name=created["name"], note=self.prefix + " updated", is_active=False)
        rows = stock_api.api_list_warehouses(company=self.company["name"], include_inactive=True, limit=0)
        saved = next(row for row in rows if row["name"] == created["name"])
        self.assertEqual((saved["company"], saved["note"], saved["is_active"]),
                         (self.company["name"], self.prefix + " updated", 0))
        self.denied(master_api.api_update_company, name=self.company["name"], phone="denied")
        self.assertEqual(frappe.db.get_value("ERP Company", self.company["name"], "phone"), "020000000")
        self.as_role("purchase")
        self.denied(stock_api.api_update_warehouse, name=created["name"], note="denied")
        self.assertEqual(frappe.db.get_value("ERP Warehouse", created["name"], "note"), self.prefix + " updated")
