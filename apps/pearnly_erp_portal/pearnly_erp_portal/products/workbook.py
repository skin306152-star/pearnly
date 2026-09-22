"""Excel 商品导入的解析层（只读表格，不连库、不鉴权）。

为什么单独一层：解析错了会把整批脏数据写进主数据，而"表头认不出""金额不是数字""少了商品名"
这些**都可以在没有 Site 的地方断言**。入库由 `service.import_products` 逐行调用同一个
`create_product`，所以导入与手工录入走的是同一条校验路径，不会有两套规则。

表头认三套：**字段名**、**历史写法**（客户手里那张表是旧的，改表头不现实）、
以及**界面上的那一句**（走语言包，中/泰/英/日都收）—— 下载的模板用的就是第三套，
所以"下载得到的模板"与"能认的表头"永远同一张表，不会各造一套。
字段名与旧线上 `services/sales/products.py` 的列一致 —— 一个不多一个不少。
"""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation

from pearnly_erp_portal.products import contracts

# 字段 → 界面文案的键。模板的表头就取这几句（按站点语言），顺序也是这个顺序。
FIELD_LABELS = (
    ("product_name", "product.field.name"),
    ("code", "product.field.code"),
    # 界面上一格放条码与二维码（用户 2026-09-22 口径），模板与导出跟着走一格；
    # 老表里的两列仍然认（见 LEGACY_ALIASES），只是新模板不再分两列。
    ("code_value", "product.field.codeValue"),
    ("unit", "product.field.unit"),
    ("base_unit", "product.field.baseUnit"),
    ("category_id", "product.field.category"),
    ("unit_price", "product.field.price"),
    ("default_cost", "product.field.cost"),
    ("min_stock", "product.field.minStock"),
    ("vat_applicable", "product.field.vat"),
    ("track_batch", "product.field.trackBatch"),
    ("track_expiry", "product.field.trackExpiry"),
    ("image_url", "product.field.image"),
    ("default_warehouse", "product.field.defaultWarehouse"),
    ("default_location", "product.field.defaultLocation"),
)

# 历史写法（去空白、转小写后比对）→ 我们的字段名。旧表格、导出件里长这样，照认。
LEGACY_ALIASES = {
    "product_name": "product_name",
    "name": "product_name",
    "name_th": "product_name",
    "ชื่อสินค้า": "product_name",
    "ชื่อ": "product_name",
    "商品名称": "product_name",
    "名称": "product_name",
    "code": "code",
    "รหัส": "code",
    "รหัสสินค้า": "code",
    "编码": "code",
    "货号": "code",
    "barcode": "barcode",
    "บาร์โค้ด": "barcode",
    "条码": "barcode",
    "qr payload": "qr_payload",
    "qr_payload": "qr_payload",
    "qr": "qr_payload",
    "二维码": "qr_payload",
    # 合并后的那一格（新模板的表头；四语标签由调用方从语言包补，这里兜住中英两套写法）。
    "barcode / qr code": "code_value",
    "条码/二维码": "code_value",
    "条码 / 二维码": "code_value",
    "unit": "unit",
    "หน่วย": "unit",
    "单位": "unit",
    "base_unit": "base_unit",
    "unit_price": "unit_price",
    "price": "unit_price",
    "ราคา": "unit_price",
    "单价": "unit_price",
    "价格": "unit_price",
    "default_cost": "default_cost",
    "cost": "default_cost",
    "ต้นทุน": "default_cost",
    "成本": "default_cost",
    "min_stock": "min_stock",
    "สต็อกต่ำ": "min_stock",
    "低库存": "min_stock",
    "vat_applicable": "vat_applicable",
    "vat": "vat_applicable",
    "含税": "vat_applicable",
    "category_id": "category_id",
    "分类": "category_id",
    "track_batch": "track_batch",
    "批次": "track_batch",
    "track_expiry": "track_expiry",
    "效期": "track_expiry",
    "image_url": "image_url",
    "图片": "image_url",
}


def header_aliases(labels=None) -> dict:
    """表头 → 字段：字段名、历史写法、界面标签三套并集。

    `labels` 是"文案 → 字段"（调用方从语言包取四语），本层不碰 Frappe 与语言包。
    """
    aliases = {_norm_header(field): field for field, _key in FIELD_LABELS}
    aliases.update({_norm_header(word): field for word, field in LEGACY_ALIASES.items()})
    for text, field in (labels or {}).items():
        if str(text or "").strip():
            aliases[_norm_header(text)] = field
    return aliases


def template_headers(label_of) -> list:
    """模板表头：界面上那几句，按站点语言给。"""
    return [label_of(key) for _field, key in FIELD_LABELS]

TEXT_FIELDS = ("product_name", "code", "code_value",
               "unit", "base_unit", "image_url", "default_warehouse", "default_location")
NUMBER_FIELDS = ("unit_price", "default_cost", "min_stock")
FLAG_FIELDS = ("vat_applicable", "track_batch", "track_expiry")

_TRUTHY = {"1", "true", "yes", "y", "ใช่", "是"}


def _norm_header(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _is_blank(row) -> bool:
    return row is None or all(cell is None or str(cell).strip() == "" for cell in row)


def _to_number(value) -> Decimal:
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise ValueError("not_a_number") from None
    if not number.is_finite() or number < 0:
        raise ValueError("negative_or_invalid")
    return number


def _to_flag(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in _TRUTHY


def _record(row, columns: dict) -> dict:
    record: dict = {}
    for index, field in columns.items():
        if index >= len(row):
            continue
        value = row[index]
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        if field in NUMBER_FIELDS:
            record[field] = _to_number(value)
        elif field in FLAG_FIELDS:
            record[field] = _to_flag(value)
        elif field == "category_id":
            try:
                record[field] = int(value)
            except (TypeError, ValueError):
                raise ValueError("not_a_number") from None
        else:
            record[field] = str(value).strip()
    # 合并的那一列按形状落回条码/二维码两列；表里另外还给了具体某列时，那一列说话
    # （老表既有合并列又有条码列时不至于互相覆盖）。
    merged = record.pop("code_value", None)
    if merged:
        record.setdefault(contracts.code_field(merged), merged)
    if not record.get("product_name"):
        raise ValueError("name_required")
    return record


def template_bytes(headers) -> bytes:
    """模板文件：只有一行表头，**不留示例行** —— 示例行会被当真数据导进来。"""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([str(header) for header in headers])
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = max(
            14, len(str(header)) + 4)
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def parse_workbook(data: bytes, aliases: dict | None = None) -> tuple[list[dict], list[dict]]:
    """xlsx 字节 → (可入库的行, 行级错误)。错误每条 `{row, error}`（行号从 2 起，表头是第 1 行）。

    认不出表头或整张表空时抛 `ValueError`，由调用方翻译成给用户看的提示。
    """
    if not data:
        raise ValueError("empty_file")
    lookup = header_aliases() if aliases is None else aliases
    # 延迟导入：本模块要能在**没有 openpyxl** 的宿主上被导入（解析逻辑的断言在
    # bench 环境里跑，见 tests 的跳过守卫）。
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ValueError("not_a_workbook") from None
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("empty_file") from None
        columns = {index: lookup[_norm_header(cell)]
                   for index, cell in enumerate(header)
                   if _norm_header(cell) in lookup}
        if "product_name" not in columns.values():
            raise ValueError("missing_name_column")

        valid: list[dict] = []
        errors: list[dict] = []
        for line_no, row in enumerate(rows, start=2):
            if _is_blank(row):
                continue
            try:
                valid.append(_record(row, columns))
            except ValueError as exc:
                errors.append({"row": line_no, "error": str(exc)})
        if not valid and not errors:
            raise ValueError("empty_file")
        return valid, errors
    finally:
        book.close()
