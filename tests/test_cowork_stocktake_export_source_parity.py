"""Run the real hash-bound legacy writer; no legacy application imports or I/O.

CLI: --source <fixed-source-root>, or --source-stdin with JSON mapping each
HASHES relative path to its source text. The real photo attach is executed with
empty photos, its source-supported early-return path; it is not a mock writer.
"""
import argparse
import ast
from copy import copy, deepcopy
from datetime import UTC, datetime, timedelta, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path
import random
import sys

import openpyxl

from pearnly_cowork.stocktake.export import workbook
from test_cowork_stocktake_export import entry, item, snapshot

SOURCE_SHA = "ec2e5c957e6b4be7dee93ffdd2dae9fa83f69188"
HASHES = {
    "services/stocktake/reports.py": "2051f30c8a073326724f5f2d506d6856d3d38f7058a035d8b15afef5da47db72",
    "services/stocktake/photo_report.py": "ea230dccf13abff1075e92e4a13615fea6834d2b926ebd3153b49366dc39e456",
}


class RemovePhotoModuleImport(ast.NodeTransformer):
    def visit_ImportFrom(self, node):
        if node.module == "services.stocktake.photo_report":
            assert [(name.name, name.asname) for name in node.names] == [("attach", None)]
            return None
        assert not (node.module or "").startswith(("services", "core", "frappe")), node.module
        return node


def load_legacy(root=None, source_text=None):
    trees = {}
    for path, digest in HASHES.items():
        data = source_text[path].encode("utf-8") if source_text is not None else (Path(root) / path).read_bytes()
        assert hashlib.sha256(data).hexdigest() == digest, f"Source hash mismatch: {path}"
        trees[path] = RemovePhotoModuleImport().visit(ast.parse(data.decode("utf-8")))
    photo = {"__name__": "fixed_stocktake_photo_source"}
    exec(compile(trees["services/stocktake/photo_report.py"], "fixed_photo_report.py", "exec"), photo)
    report = {"__name__": "fixed_stocktake_report_source", "attach": photo["attach"]}
    exec(compile(trees["services/stocktake/reports.py"], "fixed_reports.py", "exec"), report)
    return report["workbook"]


def compare(old_bytes, new_bytes):
    old = openpyxl.load_workbook(BytesIO(old_bytes), data_only=False)
    new = openpyxl.load_workbook(BytesIO(new_bytes), data_only=False)
    cells = 0
    try:
        assert old.sheetnames == new.sheetnames
        for left, right in zip(old, new, strict=True):
            for attr in ("max_row", "max_column", "freeze_panes", "sheet_state"):
                assert getattr(left, attr) == getattr(right, attr), (left.title, attr)
            assert left.auto_filter == right.auto_filter
            assert str(left.merged_cells) == str(right.merged_cells)
            assert not left._images and not right._images
            assert {k: (v.width, v.hidden) for k, v in left.column_dimensions.items()} == {
                k: (v.width, v.hidden) for k, v in right.column_dimensions.items()}
            assert {k: (v.height, v.hidden) for k, v in left.row_dimensions.items()} == {
                k: (v.height, v.hidden) for k, v in right.row_dimensions.items()}
            for arow, brow in zip(left, right, strict=True):
                for a, b in zip(arow, brow, strict=True):
                    assert (a.value, a.data_type) == (b.value, b.data_type), (left.title, a.coordinate, a.value, b.value)
                    assert a.hyperlink == b.hyperlink
                    for attr in ("font", "fill", "border", "alignment", "number_format", "protection"):
                        assert copy(getattr(a, attr)) == copy(getattr(b, attr)), (left.title, a.coordinate, attr)
                    cells += 1
    finally:
        old.close()
        new.close()
    return cells


def native_values(value):
    """Only adapt source NUMERICs to the native report's exact JSON strings."""
    result = deepcopy(value)
    for row in result["items"]:
        for key in ("book_qty", "actual_qty"):
            if row[key] is not None:
                row[key] = str(row[key])
    for row in result["entries"]:
        row["quantity"] = str(row["quantity"])
    return result


def fixtures():
    rng = random.Random(20260921)
    for batch in range(80):
        rows, events = [], []
        for number in range(rng.randrange(7)):
            row = item(f"I-{batch}-{number}", product_code=f"00{number}",
                       product_name=rng.choice(("สินค้า", "=SUM(A1)", "中文\nThai ไทย", "+123")),
                       warehouse=rng.choice(("A", "B", "")), location=rng.choice(("", "01", "02")),
                       book_qty=Decimal(rng.randrange(-10000000, 10000000)).scaleb(-6))
            active = []
            for index in range(rng.randrange(6)):
                qty = Decimal(rng.randrange(10000000)).scaleb(-6)
                voided = rng.choice((True, False, False))
                stamp = datetime(2026, 9, 21, 8, 2, index, tzinfo=timezone(timedelta(hours=7)))
                events.append(entry(row, id=f"E-{number}-{index}", quantity=qty, voided=voided,
                                    warehouse=rng.choice(("A", "B", "")), location=rng.choice(("", "01", "02")),
                                    counted_at=stamp, updated_at=stamp + timedelta(minutes=1)))
                if not voided:
                    active.append(qty)
            row["actual_qty"] = sum(active, Decimal("0.000000")) if active else None
            rows.append(row)
        yield snapshot(rows, events, photos=[])
    for book, actual in (("0.000000", "0.000000"), ("99999999999999.999999", "99999999999999.999998"),
                         ("-99999999999999.999999", "99999999999999.999999"), ("0.000001", "0.000000")):
        row = item(book_qty=Decimal(book), actual_qty=Decimal(actual), product_code="=1+1")
        yield snapshot([row], [entry(row, quantity=Decimal(actual))], photos=[])


def run(root=None, source_text=None):
    old_writer = load_legacy(root, source_text)
    books, cells = 0, 0
    for value in fixtures():
        cells += compare(old_writer(deepcopy(value)), workbook(native_values(value)))
        books += 1
    print(f"PASS source={SOURCE_SHA} workbooks={books} cells={cells} value_style_differences=0 photos=excluded")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--source")
    inputs.add_argument("--source-stdin", action="store_true")
    args = parser.parse_args()
    run(args.source, json.load(sys.stdin) if args.source_stdin else None)
