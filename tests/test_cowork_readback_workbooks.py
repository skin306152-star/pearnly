"""Express 回读出的工作簿交给对账现有的解析器原样能读（2026-10 CC 闭环 L1）：总账给银行对账（`gl_sheets` /
`bank_intake.read_ledger_file`）与收入对账（`income_parse`），销项 / 进项税报给 `vat_parse`（收入对账右口、销项核查左口），
收入对账整套配得上、作废行认得出。总账解析只改了一处（2026-10-13，对 Express 自己导出的全科目总账同样适用）：科目名里的
「ยกมา」不当期初标签、给了科目时别的科目段里的期初行不算。文字按文字存（「=」开头不变公式、控制字符去掉）。
要 openpyxl：在生产镜像跑（`scripts/ci-candidate-checks.sh` offline）。"""
import datetime as dt
import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps" / "pearnly_cowork"))
for name in [n for n in sys.modules if n.split(".")[0] == "pearnly_cowork"]:
    del sys.modules[name]

from openpyxl import load_workbook  # noqa: E402

from pearnly_cowork.readback import contracts, diff, layout, workbooks  # noqa: E402
from pearnly_cowork.reconciliation import (  # noqa: E402
    bank_intake, gl_sheets, income_parse, income_vat, vat_check_core, vat_parse)
from pearnly_cowork.tabular import readers  # noqa: E402

BANK, VAT_OUT, SALES, SERVICE, PURCHASES, FEES, OTHER = (
    "11-01-03-06", "21-04-00-00", "41-01-00-00", "41-02-00-00", "51-03-00-00", "53-00-00-00", "12-00-00-00")


def line(seq, account, day, trntyp, amount, descrp=""):
    return {"seq": seq, "account": account, "date": day, "descrp": descrp, "trntyp": trntyp, "amount": amount,
            "dept": ""}


def voucher(number, book, day, descrp, lines, docstat="N"):
    return {"voucher": number, "book": book, "date": day, "descrp": descrp, "docstat": docstat, "trnstat": "P",
            "creby": "PEARNLY", "lines": lines}


def sale(number, day, base, vat, account=SALES, ref=""):
    total = f"{float(base) + float(vat):.2f}"
    return voucher(number, "03", day, "ขายเชื่อ", [line("2", BANK, day, "0", total),
                                                   line("5", account, day, "1", base, "ค่าน้ำแข็ง"),
                                                   line("6", VAT_OUT, day, "1", vat)])


def vat_row(side, number, day, base, vat, ref="", name="ร้าน B", tax_id="0105556012345", docstat="N"):
    return {"vatrec": side, "rectyp": "3", "vatprd": day[:8] + "01", "vatdat": day, "docdat": day, "docnum": number,
            "refnum": ref, "descrp": name, "prenam": "", "taxid": tax_id, "branch": "00000", "docstat": docstat,
            "amt01": base, "vat01": vat, "extra": {}, "deleted": False}


def readback_raw():
    journal = [
        sale("IV690503-001", "2026-05-03", "500.00", "35.00"),
        sale("IV690503-002", "2026-05-03", "200.00", "14.00"),  # 同一天第二张：日期只印一次
        sale("IV690520-002", "2026-05-20", "300.00", "21.00"),
        sale("HS690525-001", "2026-05-25", "100.00", "7.00", account=SERVICE),
        voucher("HP690512-001", "01", "2026-05-12", "ซื้อสด", [line("0", PURCHASES, "2026-05-12", "0", "100.00"),
                                                              line("3", "11-05-00-00", "2026-05-12", "0", "7.00"),
                                                              line("6", BANK, "2026-05-12", "1", "107.00")]),
        voucher("JV690531-001", "00", "2026-05-31", "ค่าธรรมเนียมธนาคาร", [line("1", FEES, "2026-05-31", "0", "20.00"),
                                                                           line("2", BANK, "2026-05-31", "1", "20.00")]),
        voucher("IV690510-001", "03", "2026-05-10", "ยกเลิก", [line("2", BANK, "2026-05-10", "0", "99.00"),
                                                               line("5", SALES, "2026-05-10", "1", "99.00")],
                docstat="C"),
    ]
    raw = {
        "version": 1, "period": "2026-05", "account_set": "T9SINCER", "read_at": "2026-10-13T05:00:00Z",
        "company": {"name": "บริษัท ซินเซียร์ จำกัด", "tax_id": "0315551000297", "branch": "00000"},
        "calendar": [], "accounts": [{"code": code, "name": name, "name2": "", "group": code[0], "level": "",
                                      "parent": ""} for code, name in (
            (BANK, "SAVING-K Bank"), (VAT_OUT, "ภาษีขาย"), (SALES, "รายได้จากการขาย"), (SERVICE, "รายได้ค่าบริการ"),
            (PURCHASES, "ซื้อสุทธิ"), (FEES, "ค่าธรรมเนียม"), (OTHER, "ลูกหนี้"))],
        "balances": [{"code": BANK, "values": {"BEGLY": "1000.00", "DEBIT3LY": "150.00"}},
                     {"code": OTHER, "values": {"BEGLY": "40.00"}}],
        "movements": [{"code": BANK, "month": "2026-04", "trntyp": "1", "amount": "50.00"}],
        "journal": journal,
        "vat": [vat_row("S", "IV690503-001", "2026-05-03", "500.00", "35.00", ref="INV-0001"),
                vat_row("S", "IV690503-002", "2026-05-03", "200.00", "14.00"),
                vat_row("S", "IV690510-001", "2026-05-10", "99.00", "6.93", name="ร้าน C", docstat="C"),
                vat_row("S", "IV690520-002", "2026-05-20", "300.00", "21.00", tax_id="0994000123456"),
                vat_row("S", "HS690525-001", "2026-05-25", "100.00", "7.00"),
                vat_row("P", "HP690512-001", "2026-05-12", "100.00", "7.00", ref="A-77", name="ร้าน A"),
                vat_row("P", "RR690514-001", "2026-05-14", "1000.00", "70.00", ref="B-1", name="ร้าน D"),
                vat_row("P", "RR690515-001", "2026-05-15", "50.00", "3.50", ref="B-2", name="ร้าน D")],
        "documents": [], "diagnostics": {},
    }
    return raw


RAW = readback_raw()
DATA = contracts.parse(RAW, "2026-05")
GL = workbooks.general_ledger(DATA["company"]["name"], "2026-05", contracts.ledger(DATA), dt.date(2026, 10, 13))
SALES_VAT = workbooks.vat_report("S", DATA["company"], "2026-05", contracts.vat_rows(DATA, "S"))
PURCHASE_VAT = workbooks.vat_report("P", DATA["company"], "2026-05", contracts.vat_rows(DATA, "P"))


def sheets(content):
    return [rows for _name, rows in readers.load_workbook_sheets(content)]


class Layout(unittest.TestCase):
    def test_the_general_ledger_is_the_express_printout_with_an_account_column(self):
        sheet = load_workbook(io.BytesIO(GL)).active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(sheet.title, layout.GL_SHEET)
        self.assertEqual(rows[1][0], layout.GL_TITLE)
        self.assertEqual(rows[2][:2], [layout.GL_FROM, "1 พ.ค. 2569      ถึง  31 พ.ค. 2569"])
        self.assertEqual(tuple(rows[4]), layout.GL_HEADER)
        self.assertEqual(rows[5][0], f"{BANK} SAVING-K Bank")
        self.assertEqual(rows[5][7], 1100)  # 期初在余额列：1000 + 150 − 4 月贷 50
        self.assertEqual((rows[6][0].date(), rows[6][1], rows[6][2], rows[6][8]),
                         (dt.date(2026, 5, 3), "ขาย", "IV690503-001", BANK))
        self.assertIsNone(rows[7][0])  # 同一天的第二行不印日期
        self.assertEqual(sheet.cell(7, 1).number_format, layout.DATE_FORMAT)
        totals = [row for row in rows if row[3] == layout.GL_TOTAL]
        self.assertEqual(len(totals), len(contracts.ledger(DATA)))
        self.assertNotIn("IV690510-001", [row[2] for row in rows])  # 作废的凭证不列

    def test_the_vat_reports_carry_the_company_and_zero_the_cancelled_invoice(self):
        rows = sheets(SALES_VAT)[0]
        self.assertEqual(rows[0][0], layout.VAT_TITLE["S"])
        self.assertEqual(rows[3][:3], [layout.VAT_TAX_ID, "0315551000297", layout.HEAD_OFFICE])
        void = next(row for row in rows if len(row) > 3 and row[3] == "IV690510-001")
        self.assertEqual((void[4], void[7], void[8], void[9]), (f"ร้าน C {layout.VOID}", 0, 0, 0))
        self.assertEqual(rows[-1][0], layout.VAT_TOTAL)


class BankReconciliation(unittest.TestCase):
    def test_the_whole_ledger_reads_and_each_account_keeps_its_rows(self):
        parsed = gl_sheets.parse_gl_workbook(GL, "T9SINCER-GL-2026-05.xlsx")
        self.assertTrue(parsed["ok"], parsed)
        self.assertEqual(set(parsed["accounts"]), {BANK, VAT_OUT, SALES, SERVICE, PURCHASES, "11-05-00-00", FEES})
        self.assertEqual(parsed["period_start"], "2026-05-01")

    def test_the_bank_account_reads_with_its_opening_closing_and_lines(self):
        for parsed in (gl_sheets.parse_gl_workbook(GL, "gl.xlsx", BANK),
                       bank_intake.read_ledger_file(GL, "gl.xlsx", account_code=BANK)):
            self.assertTrue(parsed["ok"], parsed)
            self.assertEqual(parsed["opening"], 1100.0)
            # 1100 + 535 + 214 + 321 + 107 − 107 − 20
            self.assertEqual(parsed["closing"], 2150.0)
            self.assertTrue(parsed["closing_printed"])
            self.assertEqual(len(parsed["rows"]), 6)
            self.assertEqual([(row.date.isoformat(), row.doc_no, row.debit, row.credit) for row in parsed["rows"]][:3],
                             [("2026-05-03", "IV690503-001", 535.0, 0.0), ("2026-05-03", "IV690503-002", 214.0, 0.0),
                              ("2026-05-12", "HP690512-001", 0.0, 107.0)])
        self.assertTrue(gl_sheets.parse_gl_workbook(GL, "gl.xlsx", BANK)["opening_found"])

    def test_an_account_without_movement_this_month_still_reads_its_opening(self):
        parsed = gl_sheets.parse_gl_workbook(GL, "gl.xlsx", OTHER)
        self.assertEqual((parsed["ok"], parsed["rows"], parsed["opening"], parsed["opening_found"]),
                         (True, [], 40.0, True))


    def test_opening_words_in_other_accounts_do_not_replace_the_bank_opening(self):
        # 泰国科目表里常见「กำไร(ขาดทุน)สะสมยกมา」「สินค้าคงเหลือยกมา」：以前段头或说明带「ยกมา」的那一行被当成整本的期初，
        # 银行期初读成 −250,000 / 5,000（应为 1,100）
        raw = dict(RAW, accounts=RAW["accounts"] + [
            {"code": "32-01-00-00", "name": "กำไร(ขาดทุน)สะสมยกมา", "name2": "", "group": "3", "level": "",
             "parent": ""},
            {"code": "11-06-00-00", "name": "สินค้าคงเหลือ", "name2": "", "group": "1", "level": "", "parent": ""}],
            balances=RAW["balances"] + [{"code": "32-01-00-00", "values": {"BEGLY": "-250000.00"}},
                                        {"code": "11-06-00-00", "values": {"BEGLY": "3000.00"}}],
            journal=RAW["journal"] + [voucher("JV690501-001", "00", "2026-05-01", "", [
                line("1", "11-06-00-00", "2026-05-01", "0", "2000.00", "สินค้าคงเหลือยกมา"),
                line("2", "32-01-00-00", "2026-05-01", "1", "2000.00", "opening b/f")])])
        data = contracts.parse(raw, "2026-05")
        book = workbooks.general_ledger(data["company"]["name"], "2026-05", contracts.ledger(data),
                                        dt.date(2026, 10, 13))
        for parsed in (gl_sheets.parse_gl_workbook(book, "gl.xlsx", BANK),
                       bank_intake.read_ledger_file(book, "gl.xlsx", account_code=BANK)):
            self.assertTrue(parsed["ok"], parsed)
            self.assertEqual((parsed["opening"], parsed["closing"], len(parsed["rows"])), (1100.0, 2150.0, 6))
        whole = gl_sheets.parse_gl_workbook(book, "gl.xlsx")
        self.assertIn(BANK, whole["accounts"])
        self.assertNotIn("32-01-00-00", [row.doc_no for row in whole["rows"]])


class SafeText(unittest.TestCase):
    def test_control_characters_are_dropped_and_formulas_stay_text(self):
        # 说明、对方名来自 Express 与第三方发票的识别结果：Excel 不收的控制字符以前让出表报错，回读一直停在 reading
        raw = dict(RAW, journal=[voucher("JV690531-009", "00", "2026-05-31", "x", [
            line("1", FEES, "2026-05-31", "0", "20.00", "ค่า\x01ธรรม\x0bเนียม"),
            line("2", BANK, "2026-05-31", "1", "20.00", '=HYPERLINK("http://x")')])],
            vat=[vat_row("P", "RR690514-001", "2026-05-14", "1000.00", "70.00", ref="B\x1f-1",
                         name='=HYPERLINK("http://x")')])
        data = contracts.parse(raw, "2026-05")
        ledger = load_workbook(io.BytesIO(workbooks.general_ledger("=1+1", "2026-05", contracts.ledger(data),
                                                                   dt.date(2026, 10, 13)))).active
        cells = {cell.value: cell for row in ledger.iter_rows() for cell in row if isinstance(cell.value, str)}
        self.assertIn("ค่าธรรมเนียม", cells)
        for text in ('=HYPERLINK("http://x")', "=1+1"):
            self.assertEqual(cells[text].data_type, "s", text)
        report = load_workbook(io.BytesIO(workbooks.vat_report("P", data["company"], "2026-05",
                                                               contracts.vat_rows(data, "P")))).active
        row = [cell for cell in report[6]]
        self.assertEqual((row[2].value, row[4].value, row[4].data_type), ("B-1", '=HYPERLINK("http://x")', "s"))
        parsed = vat_parse.decode({"kind": "excel", "cells": sheets(workbooks.vat_report(
            "P", data["company"], "2026-05", contracts.vat_rows(data, "P")))[0]})
        self.assertEqual([item["report_invoice_no"] for item in parsed["rows"]], ["B-1"])
        found = {"changes": [{"docnum": "RR1", "direction": "purchase", "change": "changed", "field": "party",
                              "pushed": "=cmd|' /C calc'!A0", "now": "S1 ร้าน\x02 A"}], "direct": []}
        labels = {key: key for key in ("sheet_changes", "sheet_direct", "docnum", "direction", "change", "field",
                                       "pushed", "now", "date", "kind", "party", "before_vat", "vat", "total",
                                       "entered_by")}
        labels.update(directions={}, changes={}, fields={})
        sheet = load_workbook(io.BytesIO(workbooks.changes(found, labels)))["sheet_changes"]
        self.assertEqual([(cell.value, cell.data_type) for cell in sheet[2]][4:],
                         [("=cmd|' /C calc'!A0", "s"), ("S1 ร้าน A", "s")])


class IncomeReconciliation(unittest.TestCase):
    def test_only_revenue_rows_are_read_from_the_ledger(self):
        source = {"kind": "excel", "sheets": sheets(GL)}
        self.assertFalse(income_parse.is_trial_balance(source))
        parsed = income_parse.decode(source)
        self.assertTrue(parsed["ok"], parsed)
        self.assertEqual([(row.doc_no, row.date, row.account_code, row.credit) for row in parsed["rows"]], [
            ("IV690503-001", "2026-05-03", SALES, 500.0), ("IV690503-002", "2026-05-03", SALES, 200.0),
            ("IV690520-002", "2026-05-20", SALES, 300.0), ("HS690525-001", "2026-05-25", SERVICE, 100.0)])

    def test_the_sales_vat_report_reads_with_ref_invoice_tax_id_and_totals(self):
        cells = sheets(SALES_VAT)[0]
        self.assertFalse(vat_parse.is_sales_listing([" ".join(str(c or "") for c in row) for row in cells[:15]]))
        source = {"kind": "excel", "cells": cells}
        parsed = vat_parse.decode(source)
        self.assertTrue(parsed["ok"], parsed)
        rows = parsed["rows"]
        self.assertEqual([(row["report_invoice_no"], row["report_ref_no"], row["report_date"]) for row in rows], [
            ("INV-0001", "IV690503-001", "2026-05-03"), ("IV690503-002", "IV690503-002", "2026-05-03"),
            ("IV690510-001", "IV690510-001", "2026-05-10"), ("IV690520-002", "IV690520-002", "2026-05-20"),
            ("HS690525-001", "HS690525-001", "2026-05-25")])
        self.assertEqual(rows[3]["report_buyer_tax_id"], "0994000123456")
        self.assertEqual((rows[0]["report_buyer_branch"], rows[0]["report_amount_pre_vat"], rows[0]["report_vat_amount"],
                          rows[0]["report_amount"]), ("00000", 500.0, 35.0, 535.0))
        self.assertEqual((parsed["meta"]["total_amount_pre_vat"], parsed["meta"]["total_vat"]), (1100.0, 77.0))
        self.assertEqual(vat_parse.issuer_tax_id(source), "0315551000297")
        self.assertTrue(vat_check_core._void_report(rows[2]))
        self.assertFalse(vat_check_core._void_report(rows[0]))

    def test_ledger_and_sales_report_reconcile_with_no_difference(self):
        gl_rows = income_parse.decode({"kind": "excel", "sheets": sheets(GL)})["rows"]
        vat_rows = vat_parse.decode({"kind": "excel", "cells": sheets(SALES_VAT)[0]})["rows"]
        recon, summary = income_vat.reconcile_gl_vat(gl_rows, vat_rows)
        matched = [row for row in recon if row.gl_amount is not None]
        self.assertEqual(sorted(row.doc_no for row in matched),
                         ["HS690525-001", "IV690503-001", "IV690503-002", "IV690520-002"])
        self.assertTrue(all(abs(row.diff or 0) < 0.005 for row in matched))
        self.assertEqual((summary.gl_only_credit, summary.gl_only_debit, summary.vat_only_positive), (0, 0, 0))


class PurchaseReport(unittest.TestCase):
    def test_the_purchase_vat_report_reads_the_same_way(self):
        parsed = vat_parse.decode({"kind": "excel", "cells": sheets(PURCHASE_VAT)[0]})
        self.assertTrue(parsed["ok"], parsed)
        self.assertEqual([(row["report_invoice_no"], row["report_ref_no"], row["report_amount_pre_vat"])
                          for row in parsed["rows"]],
                         [("A-77", "HP690512-001", 100.0), ("B-1", "RR690514-001", 1000.0),
                          ("B-2", "RR690515-001", 50.0)])


class ChangesWorkbook(unittest.TestCase):
    def test_two_sheets_with_the_labels_given(self):
        found = {"changes": [{"docnum": "IV1", "direction": "sales", "change": "changed", "field": "total",
                              "pushed": "535.00", "now": "530.00"}],
                 "direct": [{"docnum": "JV1", "date": "2026-05-31", "kind": "", "party": "", "before_vat": "",
                             "vat": "", "total": "20.00", "entered_by": "BEE", "book": "00"}]}
        labels = {"sheet_changes": "Changes", "sheet_direct": "Entered in Express", "docnum": "Document No.",
                  "direction": "Direction", "change": "Change", "field": "Field", "pushed": "Pushed",
                  "now": "In Express", "date": "Date", "kind": "Type", "party": "Party", "before_vat": "Before VAT",
                  "vat": "VAT", "total": "Total", "entered_by": "Entered by",
                  "directions": {"sales": "Sale"}, "changes": {"changed": "Changed"}, "fields": {"total": "Total"}}
        book = load_workbook(io.BytesIO(workbooks.changes(found, labels)))
        self.assertEqual(book.sheetnames, ["Changes", "Entered in Express"])
        self.assertEqual([c.value for c in book["Changes"][2]], ["IV1", "Sale", "Changed", "Total", "535.00", "530.00"])
        self.assertEqual([c.value for c in book["Entered in Express"][2]][:3], ["JV1", "2026-05-31", layout.BOOKS["00"]])
        self.assertEqual(set(diff.CHANGES), {"changed", "deleted", "missing", "voided"})


if __name__ == "__main__":
    unittest.main()
