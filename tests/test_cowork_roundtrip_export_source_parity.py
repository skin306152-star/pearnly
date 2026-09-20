"""Hash-bound offline workbook differential. Never import the legacy application."""
import argparse
import ast
from copy import copy
from decimal import Decimal, ROUND_HALF_EVEN
import hashlib
import io
from pathlib import Path
import random
from types import SimpleNamespace

import openpyxl

from pearnly_cowork.roundtrip.export import build_review_workbook, validate_records
from pearnly_cowork.roundtrip.parser import parse_roundtrip_workbook
from test_cowork_roundtrip_export import record

SOURCE_SHA = "ec2e5c957e6b4be7dee93ffdd2dae9fa83f69188"
HASHES = {
    "erp_workbook.py": "a388fd3366a71b969a7618e877302505fe243de7d5ed75846b58fdcf8383f062",
    "excel_template_th.py": "ba63d2ab4270483e2d611e9955d6a1c01fda35b701f42f4540fbd677847b8af3",
    "xlsx_style.py": "8e9beb1bb4e176509291ab679ea0b72fa6ca0e80abd68c79483379faccf28615",
    "erp_roundtrip.py": "03aaf37f6396613811d3c5723fd3c685ec121415b88770d3632c38f7d76f33ad",
    "erp_roundtrip_reader.py": "26a367002ff778ae00e43ee309c3c90077accca1dfef2a4eb32420e09de171bc",
    "erp_money.py": "6bfe9ba4ed0e69a7417e2c1520e682615dbdd43a133a245db3fb25a6d7d7bdc8",
}


class NoApplicationImports(ast.NodeTransformer):
    def visit_ImportFrom(self, node):
        return None if (node.module or "").startswith("services") else node


def load_legacy(root):
    trees = {}
    for name, digest in HASHES.items():
        data = (Path(root) / "services/excel" / name).read_bytes()
        assert hashlib.sha256(data).hexdigest() == digest, name
        trees[name] = NoApplicationImports().visit(ast.parse(data.decode()))

    def execute(name, **bindings):
        namespace = {"__name__": "fixed_source_contract", **bindings}
        exec(compile(trees[name], name, "exec"), namespace)
        return namespace

    rt = execute("erp_roundtrip.py")
    spec = SimpleNamespace(**rt)
    sty = SimpleNamespace(**execute("xlsx_style.py"))
    cash = SimpleNamespace(**execute("erp_money.py", quantize_money=lambda value: value.quantize(Decimal(".01"), rounding=ROUND_HALF_EVEN)))
    sales = execute("excel_template_th.py", **{**rt, "rt": spec})
    writer = execute("erp_workbook.py", rt=spec, sty=sty, write_sales_sheet=sales["write_sales_sheet"],
                     sales_detail_filename=sales["sales_detail_filename"])
    reader = execute("erp_roundtrip_reader.py", rt=spec, money=cash)
    return writer["build_review_workbook"], reader["parse_roundtrip_workbook"]


def compare_workbooks(old_bytes, new_bytes):
    old = openpyxl.load_workbook(io.BytesIO(old_bytes), data_only=False)
    new = openpyxl.load_workbook(io.BytesIO(new_bytes), data_only=False)
    count = 0
    try:
        assert old.sheetnames == new.sheetnames
        for a, b in zip(old, new, strict=True):
            assert (a.max_row, a.max_column, a.freeze_panes) == (b.max_row, b.max_column, b.freeze_panes), a.title
            assert a.row_dimensions[1].height == b.row_dimensions[1].height
            assert {k: (v.width, v.hidden) for k, v in a.column_dimensions.items()} == {
                k: (v.width, v.hidden) for k, v in b.column_dimensions.items()}, a.title
            for row_a, row_b in zip(a, b, strict=True):
                for ca, cb in zip(row_a, row_b, strict=True):
                    assert (ca.value, ca.data_type) == (cb.value, cb.data_type), (a.title, ca.coordinate, ca.value, cb.value)
                    for attr in ("font", "fill", "border", "alignment", "number_format", "protection"):
                        assert copy(getattr(ca, attr)) == copy(getattr(cb, attr)), (a.title, ca.coordinate, attr)
                    count += 1
    finally:
        old.close()
        new.close()
    return count


def canonical(value):
    if isinstance(value, (float, Decimal)):
        return Decimal(str(value))
    if isinstance(value, dict):
        return {k: canonical(v) for k, v in value.items() if k not in ("source_row_keys",)}
    if isinstance(value, list):
        return [canonical(v) for v in value]
    return value


def run(root):
    writer, reader = load_legacy(root)
    rng, cells, documents = random.Random(20260922), 0, 0
    for batch in range(60):
        records, facts = [], {}
        for index in range(rng.randint(1, 6)):
            name = f"H-{batch}-{index}"
            direction = rng.choice(("sales", "purchase", "pending"))
            amounts = [rng.randint(1, 900) for _ in range(rng.randint(1, 4))]
            base = sum(amounts)
            items = [{"description": f"สินค้า {line}", "amount": str(amount),
                      **({"qty": "2", "unit_price": str(Decimal(amount) / 2)} if rng.choice((True, False)) else {})}
                     for line, amount in enumerate(amounts)]
            records.append(record(name, direction, subtotal=str(base), vat=str(Decimal(base) * Decimal(".07")),
                                  total_amount=str(Decimal(base) * Decimal("1.07")), items=items))
            if rng.choice((True, False)):
                facts[name] = {"actions": {"customer": rng.choice((True, False)),
                    "items": [rng.choice((True, False)) for _ in items],
                    "item_codes": [f"IT-{batch}-{index}-{line}" for line in range(len(items))],
                    "party_code": f"P-{index}", "created_party": True, "direction": direction,
                    "docnum": f"DOC-{index}", "stock_acc_group": "01 สินค้า"},
                    "state": {"status": rng.choice(("success", "failed", "pending", "manual")), "reason": "ตรวจสอบ"}}
        prepared = validate_records(records, push_facts=facts)
        old_bytes = writer(sales=list(prepared.sales), purchase=list(prepared.purchase),
                           pending=list(prepared.pending), created_masters=list(prepared.created_masters))
        new_bytes = build_review_workbook(records, push_facts=facts)
        cells += compare_workbooks(old_bytes, new_bytes)
        old_read, new_read = reader(old_bytes), parse_roundtrip_workbook(new_bytes)
        for key in ("documents", "pending", "sheets"):
            assert canonical(old_read[key]) == canonical(new_read[key]), (batch, key, old_read[key], new_read[key])
        documents += len(new_read["documents"]) + len(new_read["pending"])
    # Explicit float/Decimal half-cent difference: the source's literal VAT
    # round(1.5*.07,2)=.11; shared Decimal HALF_EVEN=.10. Fixed formulas stay.
    half = record(subtotal="1.5", vat="0.10", total_amount="1.60", items=[{"amount": "1.5"}])
    prepared = validate_records([half])
    old_bytes, new_bytes = writer(sales=list(prepared.sales)), build_review_workbook([half])
    old_wb, new_wb = (openpyxl.load_workbook(io.BytesIO(value)) for value in (old_bytes, new_bytes))
    try:
        assert old_wb["ขาย"]["J2"].value == 0.11
        assert new_wb["ขาย"]["J2"].value == 0.10
        assert reader(old_bytes)["documents"][0]["fields"]["total_amount"] == Decimal("1.60")
        assert parse_roundtrip_workbook(new_bytes)["documents"][0]["fields"]["total_amount"] == Decimal("1.60")
    finally:
        old_wb.close()
        new_wb.close()
    print(f"PASS source={SOURCE_SHA} workbooks=60 cells={cells} documents={documents} intentional_half_cent=1")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    run(parser.parse_args().source)
