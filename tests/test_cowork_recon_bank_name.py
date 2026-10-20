"""银行对账给人看的银行名取账单上印着的，读不出就空着（2026-10-20 教程录制问题清单第 3 条）。

ccdemo ดาวเหนือ 用虚构的 ธนาคารสาธิต / SATHIT DEMO BANK 账单（118-7-30554-2）对账：工具结果只有内部银行码 ``generic``，
CC 回答说成「ธนาคารกสิกรไทย」，导出的银行一栏写着「GENERIC」。现在：

- PDF 文字层第一页表头上印着的银行名记进结果（``bank_name``）与 ``parse_info.stmt_files``；识别读出的取 ``identity.bank``；
- 结果列表、结果要点（CC 工具）、导出写这个名字；读不出就空着，认得出银行的照旧写银行码（``KBANK``），
  内部代号 ``generic`` / ``GENERIC`` 不出现在任何给人看的地方。

``HeaderName`` / ``Labels`` 只调纯函数，本地可跑；``PdfIntake`` 把文字层换成给定文字；``Export`` 要 openpyxl；``Service`` 与
``RealStatement`` 要 frappe 与 pdfminer，在候选镜像里跑（scripts/ci-candidate-checks.sh）。
"""
import json
import os
import types
import unittest
from io import BytesIO
from unittest import mock

from pearnly_cowork.reconciliation import bank_intake, bank_text


def names():
	"""``bank_names``（修复前没有这个模块：其余各项照样跑，看修复前哪些断言不成立）。"""
	from pearnly_cowork.reconciliation import bank_names

	return bank_names

FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "bank-statements")
DEMO = "Statement_1187305542_202609.pdf"

# pdfminer 抽出的 D3 账单第一页开头（原样：声调符号抽成 \x00 并断行，英文名字母之间多了空格）。
MINER_HEADER = ("ธนาคารสาธิต\n\nSAT H I T   D E M O   BA N K\n\nรายการเดินบัญชี / Statement of Account\n"
	"บัญชีกระแสรายวัน (Current Account)\n\nชื\n\n\x00อบัญชี: บริษัท ดาวเหนือ ไอที ซัพพลาย จํากัด\n\n"
	"เลขที\n\n\x00บัญชี: 118-7-30554-2\n\nสาขาเจ้าของบัญชี: สาขาวิภาวดีรังสิต\n\n"
	"รอบรายการ: 01/09/2569 - 30/09/2569\n\nสกุลเงิน: THB\n\nหน้า: 1/1\n\nวันที\n\nเวลา\n\nรายการ\n\nถอนเงิน\n\nฝากเงิน\n\n"
	"ยอดคงเหลือ  ช่องทาง\n\n")
MINER_ROWS = ("                ยอดยกมา                       420,000.00\n\n"
	"01/09/2569  17:28    โอนเงินออก ไปยัง ธนาคารกสิกรไทย VIBHAVADI LAND   18,360.00      401,640.00  MOBILE\n\n"
	"02/09/2569  14:26    รับโอนเงิน จาก KBANK RUNGRUENG ACCOUNTING    50,439.80      452,079.80  ORFT\n\n"
	"03/09/2569  14:15    โอนเงินออก ไปยัง KRADAT DEE PANICH       32,667.10      419,412.70  MOBILE\n\n"
	"03/09/2569  14:55    หักบัญชีอัตโนมัติ SAIYAI NET BROADBAND    1,380.30      418,032.40  AUTO\n\n"
	"\nเอกสารนี\n\n\x00 จัดทําจากระบบอิเล็กทรอนิกส์ของธนาคารสาธิต\n\n")
# pypdf 抽同一页：声调符号成了空格，英文名完整。
PYPDF_HEADER = ("ธนาคารสาธิต\nSATHIT DEMO BANK\nรายการเดินบัญชี / Statement of Account\n"
	"บัญชีกระแสรายวัน (Current Account)\nชื อบัญชี: บริษัท ดาวเหนือ ไอที ซัพพลาย จํากัด เลขที บัญชี: 118-7-30554-2\n"
	"วันที  เวลา รายการ ถอนเงิน ฝากเงิน ยอดคงเหลือ  ช่องทาง\n")


def header(*lines, counterparties=True):
	"""一页账单：给定的表头行 + 两行流水（``counterparties``：流水里提到别的银行）。"""
	rows = (["01/09/2569 รับโอนจาก ธนาคารกสิกรไทย 590.00 180,590.00", "02/09/2569 โอนไป Bangkok Bank 100.00 180,490.00"]
		if counterparties else ["01/09/2569 รับโอน 590.00 180,590.00", "02/09/2569 โอนออก 100.00 180,490.00"])
	return "\n".join(["Statement of Account", *lines, "วันที่ รายการ ถอนเงิน ฝากเงิน ยอดคงเหลือ", *rows])


class HeaderName(unittest.TestCase):
	def test_the_demo_header_is_read_as_pdfminer_and_pypdf_give_it(self):
		self.assertEqual(bank_text.printed_bank_name(MINER_HEADER + MINER_ROWS), "ธนาคารสาธิต")
		self.assertEqual(bank_text.printed_bank_name(PYPDF_HEADER + MINER_ROWS), "ธนาคารสาธิต")

	def test_each_printed_form_counts(self):
		for line, name in (("ธนาคารกสิกรไทย จำกัด (มหาชน)", "ธนาคารกสิกรไทย จำกัด (มหาชน)"),
				("ธนาคารกรุงเทพ สาขาสีลม", "ธนาคารกรุงเทพ"), ("ธนาคารไทยพาณิชย์    เลขที่บัญชี 123-4-56789-0", "ธนาคารไทยพาณิชย์"),
				("BANGKOK BANK PUBLIC COMPANY LIMITED", "BANGKOK BANK PUBLIC COMPANY LIMITED"), ("KASIKORNBANK", "KASIKORNBANK"),
				("Bank of Ayudhya", "Bank of Ayudhya"), ("ธนาคาร: ยูโอบี", "ยูโอบี"), ("Bank Name : CIMB Thai", "CIMB Thai"),
				("ชื\x00อธนาคาร: ทิสโก้", "ทิสโก้")):
			with self.subTest(line=line):
				self.assertEqual(bank_text.printed_bank_name(header(line)), name)

	def test_nothing_printed_reads_nothing(self):
		# 流水行里的对方银行、页脚、第二页，表头的栏名与账单标题都不算；也不按账号格式猜。
		for text in (header("เลขที่บัญชี: 118-7-30554-2"), header("Bank Statement", "Account No. 118-7-30554-2"),
				header("ธนาคารผู้รับ  ธนาคารปลายทาง"), header("ธนาคาร"), header("To Bank"),
				"ยอดยกมา 420,000.00\nธนาคารสาธิต", header("สกุลเงิน: THB") + "\f" + MINER_HEADER, ""):
			with self.subTest(text=text[:40]):
				self.assertEqual(bank_text.printed_bank_name(text), "")


class Labels(unittest.TestCase):
	def test_the_printed_name_wins_and_internal_codes_never_show(self):
		self.assertEqual(names().task_label("generic", [{"ok": True, "bank_name": "ธนาคารสาธิต"}]), "ธนาคารสาธิต")
		# 流水里提到别的银行时内容识别会认成那家：照样写印着的
		self.assertEqual(names().task_label("kbank", [{"ok": True, "bank_name": "ธนาคารสาธิต"}]), "ธนาคารสาธิต")
		self.assertEqual(names().task_label("kbank", [{"ok": True}]), "KBANK")
		for code in ("generic", "GENERIC", "", None):
			with self.subTest(code=code):
				self.assertEqual(names().task_label(code, [{"ok": True, "bank_code": code}]), "")
				self.assertEqual(names().file_label({"bank_code": code}), "")
		# 读失败的那份不算；几份不同取多数
		self.assertEqual(names().task_label("generic", [{"ok": False, "bank_name": "X"}]), "")
		self.assertEqual(names().task_label("generic", [{"ok": True, "bank_name": "A"}, {"ok": True, "bank_name": "B"},
			{"ok": True, "bank_name": "B"}]), "B")

	def test_a_recognised_statement_gives_the_name_it_read(self):
		self.assertEqual(names().printed({"identity": {"bank": " ธนาคาร  กรุงไทย ", "last4": "1234"}}), "ธนาคาร กรุงไทย")
		self.assertEqual(names().printed({"bank_name": "ธนาคารสาธิต", "identity": {"bank": "X"}}), "ธนาคารสาธิต")
		self.assertEqual(names().printed({"identity": {"bank": ""}}), "")


class PdfIntake(unittest.TestCase):
	"""文字层换成给定文字（真接缝 ``pdf_text.extract_pdf_text_pages``），看读出的结果带不带 ``bank_name``。"""

	def read(self, text, filename="Statement_202609.pdf"):
		with mock.patch.object(bank_intake.pdf_text, "extract_pdf_text_pages", return_value=[text]):
			return bank_intake.read_statement_file(b"%PDF-1.4 fake", filename)

	def test_an_unknown_bank_keeps_its_printed_name(self):
		result = self.read(MINER_HEADER + MINER_ROWS)
		self.assertEqual((result["ok"], result["source"], result["bank_name"]), (True, "text_lines", "ธนาคารสาธิต"))

	def test_a_known_bank_is_still_recognised(self):
		result = self.read(header("BANGKOK BANK", "Account No. 123-4-56789-0", counterparties=False))
		self.assertEqual((result["ok"], result["bank_code"], result["bank_name"]), (True, "bbl", "BANGKOK BANK"))
		result = self.read(header("เลขที่บัญชี 123-4-56789-0", counterparties=False), "SCB_202609.pdf")
		self.assertEqual((result["ok"], result["bank_code"]), (True, "scb"))
		self.assertNotIn("bank_name", result)

	def test_no_printed_name_records_none(self):
		result = self.read(MINER_HEADER.replace("ธนาคารสาธิต\n\nSAT H I T   D E M O   BA N K\n\n", "") + MINER_ROWS)
		self.assertTrue(result["ok"])
		self.assertNotIn("bank_name", result)


class Export(unittest.TestCase):
	"""导出：银行一栏与标题写给人看的银行名；读不出就不写，内部代号 GENERIC 不出现。要 openpyxl。"""

	def book(self, bank_code, bank_label, parse_info=None):
		import openpyxl

		from pearnly_cowork.reconciliation.bank_matching import BankReconSummary
		from pearnly_cowork.reconciliation.export import labels
		from pearnly_cowork.reconciliation.export.workbook import workbook

		content = workbook([], BankReconSummary(bank_code=bank_code, gl_account_code="11-01-01-00"), lang="th",
			parse_info=parse_info, bank_label=bank_label)
		sheet = openpyxl.load_workbook(BytesIO(content))[labels.t("sh_summary", "th")]
		return sheet, labels

	def texts(self, sheet):
		return [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]

	def test_the_printed_name_is_written(self):
		sheet, labels = self.book("generic", "ธนาคารสาธิต", {"stmt_files": [{"file": DEMO, "rows": 26, "ok": True,
			"bank_code": "generic", "bank_name": "ธนาคารสาธิต"}], "gl_files": []})
		self.assertIn("ธนาคารสาธิต", sheet.cell(1, 1).value)
		bank_row = [row for row in sheet.iter_rows(values_only=True) if row[0] == labels.t("lbl_bank", "th")][0]
		self.assertEqual(bank_row[1], "ธนาคารสาธิต")
		self.assertFalse([text for text in self.texts(sheet) if "generic" in text.lower()])

	def test_an_unread_name_is_left_out(self):
		sheet, labels = self.book("generic", "", {"stmt_files": [{"file": DEMO, "rows": 26, "ok": True,
			"bank_code": "generic"}], "gl_files": []})
		self.assertNotIn("·", sheet.cell(1, 1).value)
		bank_row = [row for row in sheet.iter_rows(values_only=True) if row[0] == labels.t("lbl_bank", "th")][0]
		self.assertEqual(bank_row[1], "—")
		self.assertFalse([text for text in self.texts(sheet) if "generic" in text.lower()])

	def test_a_known_bank_without_a_printed_name_keeps_its_code(self):
		sheet, _labels = self.book("kbank", "KBANK", {"stmt_files": [{"file": "STM KBANK.pdf", "rows": 3, "ok": True,
			"bank_code": "kbank"}], "gl_files": []})
		self.assertTrue(sheet.cell(1, 1).value.endswith("· KBANK"))


class Service(unittest.TestCase):
	"""存下的任务：导出、结果要点（CC 工具读的）、列表都写给人看的银行名。要 frappe（不连站点，读库处换成假的）。"""

	PARSE_INFO = {"stmt_files": [{"file": DEMO, "rows": 26, "ok": True, "error": None, "error_code": None,
		"bank_code": "generic", "account": "5542", "account_read": True, "bank_name": "ธนาคารสาธิต"}],
		"gl_files": [{"file": "test-GL-2026-09.xlsx", "rows": 11, "ok": True, "error": None, "error_code": None,
			"accounts": ["11-01-01-00"]}]}

	def setUp(self):
		from pearnly_cowork.reconciliation import bank_recon_service

		self.service = bank_recon_service

	def doc(self, parse_info):
		summary = {"bank_code": "generic", "gl_account_code": "11-01-01-00", "matched_count": 8, "stmt_opening": 420000.0}
		values = {"name": "CBR-1", "workspace": "W", "status": "completed", "title": "กระทบยอดธนาคาร ก.ย. 2569",
			"bank_code": "generic", "gl_account_code": "11-01-01-00", "summary_json": json.dumps(summary),
			"parse_info_json": json.dumps(parse_info), "run_info_json": "{}"}
		doc = types.SimpleNamespace(**values)
		doc.get = values.get
		return doc

	def patched(self, doc):
		frappe = self.service.frappe
		return [mock.patch.object(self.service.bank_recon_access, "require", lambda workspace: None),
			mock.patch.object(frappe, "get_doc", lambda *args, **kwargs: doc),
			mock.patch.object(frappe, "get_all", lambda doctype, **kwargs: (
				[types.SimpleNamespace(name=doc.name, parse_info_json=doc.parse_info_json)]
				if "parse_info_json" in (kwargs.get("fields") or []) else
				[frappe._dict(name=doc.name, title=doc.title, bank_code=doc.bank_code)]
				if doctype == self.service.TASK else [])),
			mock.patch.object(frappe, "db", types.SimpleNamespace(count=lambda *args, **kwargs: 0))]

	def run_with(self, doc, call):
		patches = self.patched(doc)
		for patch in patches:
			patch.start()
		try:
			return call()
		finally:
			for patch in reversed(patches):
				patch.stop()

	def exported(self, doc):
		import openpyxl

		built = self.run_with(doc, lambda: self.service.export("CBR-1", "th"))
		book = openpyxl.load_workbook(BytesIO(built["content"]))
		return [str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row
			if cell.value is not None]

	def test_export_writes_the_printed_name_and_never_generic(self):
		texts = self.exported(self.doc(self.PARSE_INFO))
		self.assertFalse([text for text in texts if "generic" in text.lower()])
		self.assertTrue([text for text in texts if "ธนาคารสาธิต" in text])

	def test_the_result_page_workbook_writes_the_printed_name(self):
		"""结果页由同一本工作簿画出来（`recon_center.sheets._view` → `revision.bank_workbook_bytes`）：同样不写 GENERIC。"""
		import openpyxl

		from pearnly_cowork.reconciliation import revision

		doc = self.doc(self.PARSE_INFO)
		content = revision.bank_workbook_bytes(doc, [], json.loads(doc.summary_json), "th")
		texts = [str(cell.value) for sheet in openpyxl.load_workbook(BytesIO(content)).worksheets
			for row in sheet.iter_rows() for cell in row if cell.value is not None]
		self.assertFalse([text for text in texts if "generic" in text.lower()])
		self.assertTrue([text for text in texts if "ธนาคารสาธิต" in text])

	def test_export_without_a_name_writes_no_bank(self):
		parse_info = json.loads(json.dumps(self.PARSE_INFO))
		del parse_info["stmt_files"][0]["bank_name"]
		texts = self.exported(self.doc(parse_info))
		self.assertFalse([text for text in texts if "generic" in text.lower()])

	def test_the_result_and_the_list_give_the_name_not_the_code(self):
		doc = self.doc(self.PARSE_INFO)
		page = self.run_with(doc, lambda: self.service.TASKS.detail("CBR-1", 0, 200))
		self.assertEqual(page["bank"], "ธนาคารสาธิต")
		self.assertNotIn("generic", json.dumps(page["summary"]).lower())
		listed = self.run_with(doc, lambda: self.service.TASKS.records("W", 0, 20))
		self.assertEqual(listed["tasks"][0]["bank_code"], "ธนาคารสาธิต")
		parse_info = json.loads(json.dumps(self.PARSE_INFO))
		del parse_info["stmt_files"][0]["bank_name"]
		doc = self.doc(parse_info)
		page = self.run_with(doc, lambda: self.service.TASKS.detail("CBR-1", 0, 200))
		self.assertEqual(page["bank"], "")
		listed = self.run_with(doc, lambda: self.service.TASKS.records("W", 0, 20))
		self.assertEqual(listed["tasks"][0]["bank_code"], "")

	def test_the_name_is_stored_with_the_file_diagnostics(self):
		from pearnly_cowork.reconciliation import bank_recon_contracts

		stored = bank_recon_contracts.parse_info(self.PARSE_INFO)
		self.assertEqual(stored["stmt_files"][0]["bank_name"], "ธนาคารสาธิต")
		old = json.loads(json.dumps(self.PARSE_INFO))
		del old["stmt_files"][0]["bank_name"]
		self.assertNotIn("bank_name", bank_recon_contracts.parse_info(old)["stmt_files"][0])


class RealStatement(unittest.TestCase):
	"""教程 D3 的真 PDF（ธนาคารสาธิต / SATHIT DEMO BANK，118-7-30554-2）。要 pdfminer 与 frappe。"""

	def test_the_demo_statement_reads_its_printed_bank(self):
		from pearnly_cowork.reconciliation import bank_recon_service

		with open(os.path.join(FIXTURES, DEMO), "rb") as handle:
			result = bank_intake.read_statement_file(handle.read(), DEMO)
		self.assertEqual((result["ok"], result["bank_code"], result["row_count"], result["bank_name"]),
			(True, "generic", 26, "ธนาคารสาธิต"))
		self.assertEqual(names().printed(result), "ธนาคารสาธิต")
		self.assertEqual(names().task_label(result["bank_code"], [{"ok": True, "bank_name": names().printed(result),
			**bank_recon_service._account_of(result)}]), "ธนาคารสาธิต")


if __name__ == "__main__":
	unittest.main()
