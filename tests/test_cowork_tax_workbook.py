"""报税草稿工作簿（`tax.workbook`，openpyxl）：表名与栏名取自词条（泰文是税务局表样原词）、金额是数值且两位小数格式、
日期是佛历、合计与草稿一致、ภ.ง.ด.3 分 ชื่อ / ชื่อสกุล、ใบแนบ 每页小计。

需要 openpyxl：放在 `scripts/ci-candidate-checks.sh` 的 offline 清单（生产镜像里跑）。Frappe 与读回仍用
`test_cowork_tax_service.py` 的内存替身。
"""
import csv
import importlib.util
from io import BytesIO
from pathlib import Path
import sys
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("tax_service_fakes", ROOT / "tests/test_cowork_tax_service.py")
fakes = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fakes)

from openpyxl import load_workbook  # noqa: E402

GLOSSARY = ROOT / "apps/pearnly_cowork/pearnly_cowork/i18n/messages.tsv"
with open(GLOSSARY, encoding="utf-8", newline="") as handle:
    TERMS = {}
    for row in csv.DictReader(handle, delimiter="\t"):
        TERMS[row["key"] or row["en"]] = row
MISSING = []


def translate(text, lang=None):
    term = TERMS.get(text)
    if term is None:
        MISSING.append(text)
        return text
    return term[lang or "en"]


fakes.frappe._ = staticmethod(translate)
fakes.frappe.__dict__["_"] = translate


def values(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


class VatBook(fakes.Base):
    def test_form_reports_and_checks(self):
        records = [fakes.isvat(refnum="T1", amount="1000.00", vat="70.00"),
                   fakes.isvat(refnum="T2", amount="2000.50", vat="140.04", tax_date="2026-01-31"),
                   fakes.isvat(side="S", refnum="S1", amount="10000.00", vat="700.00")]
        with mock.patch.dict(sys.modules, fakes.readback_module(records)):
            summary = fakes.service.vat_draft("CONN-1", "ACC-2026", "2026-09")
        filename, content = fakes.service.workbook(summary, "th")
        self.assertTrue(filename.startswith("PP30-2026-09-" + fakes.COMPANY))
        book = load_workbook(BytesIO(content))
        self.assertEqual(book.sheetnames, ["ภ.พ.30", "รายงานภาษีขาย", "รายงานภาษีซื้อ", "ตรวจสอบ"])
        form = values(book["ภ.พ.30"])
        flat = [value for row in form for value in row]
        self.assertIn("ยอดขายในเดือนนี้", flat)
        self.assertIn("ภาษีที่ชำระเกินยกมา", flat)
        self.assertIn("09/2569", flat)
        self.assertIn("15/10/2569", flat)
        box7 = next(row for row in form if row[0] == 7)
        self.assertEqual(box7[2], 210.04)
        cell = next(cell for row in book["ภ.พ.30"].iter_rows() for cell in row if cell.value == 210.04)
        self.assertIn("#,##0.00", cell.number_format)
        box13 = next(row for row in form if row[0] == 13)
        self.assertIsNone(box13[2])
        purchases = values(book["รายงานภาษีซื้อ"])
        self.assertEqual(purchases[0][:3], ["ลำดับที่", "วัน เดือน ปี", "เลขที่ใบกำกับภาษี"])
        self.assertEqual(purchases[1][1], "31/01/2569")
        self.assertIn("เกินกำหนด 6 เดือน", purchases[1][9])
        self.assertEqual(purchases[-1][7:9], [3000.5, 210.04])
        check = values(book["ตรวจสอบ"])
        self.assertEqual(check[1][:2], ["ต้องแก้", "เกินกำหนด 6 เดือน"])
        self.assertEqual(MISSING, [])


class WhtBook(fakes.Base):
    def test_summary_attachments_and_review(self):
        fakes.frappe.lists["Cowork OCR History"] = []
        documents = []
        for number in range(8):
            payee = f"31005000000{number:02d}"
            documents.append({"ref": f"REF-{number}", "origin": "reference", "fields": {
                "document_flag": "withholding_certificate", "date": "2026-09-20", "seller_tax": fakes.COMPANY,
                "buyer_tax": payee, "buyer_name": f"นายสมชาย{number} ใจดี", "subtotal": "1000", "wht_amount": "50"}})
            fakes.service.answer([{"payee_tax_id": payee, "income_type": "rent"}])
        with mock.patch.object(fakes.certificates, "listing", return_value=documents):
            summary = fakes.service.wht_draft(fakes.COMPANY, "2026-09")
        filename, content = fakes.service.workbook(summary, "th")
        self.assertTrue(filename.startswith("PND3-PND53-2026-09-"))
        book = load_workbook(BytesIO(content))
        self.assertEqual(book.sheetnames, ["สรุป", "ใบแนบ ภ.ง.ด.53", "ใบแนบ ภ.ง.ด.3", "ต้องตรวจ"])
        overview = values(book["สรุป"])
        pnd3 = next(row for row in overview if row[0] == "ภ.ง.ด.3")
        self.assertEqual(pnd3[1:5], [8, 2, 8000.0, 400.0])
        attachment = values(book["ใบแนบ ภ.ง.ด.3"])
        self.assertEqual(attachment[0][4:7], ["คำนำหน้าชื่อ", "ชื่อ", "ชื่อสกุล"])
        first = attachment[1]
        self.assertEqual(first[:7], [1, 1, "3100500000000", "00000", "นาย", "สมชาย0", "ใจดี"])
        self.assertEqual(first[8:13], ["20/09/2569", "ค่าเช่า", "5", 1000.0, 50.0])
        subtotals = [row for row in attachment if row[1] == "รวมยอดเงินได้และภาษีที่นำส่ง"]
        self.assertEqual([row[0] for row in subtotals], [1, 2])
        self.assertEqual([row[12] for row in subtotals], [300.0, 100.0])
        self.assertEqual(attachment[-1][1], "รวม")
        self.assertEqual(MISSING, [])


if __name__ == "__main__":
    unittest.main()
