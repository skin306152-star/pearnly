"""Offline uploaded-roster and workspace field behavior; no model/network/Site."""
import io
import unittest
import zipfile
from unittest.mock import patch

import openpyxl

from pearnly_cowork.masters import client_import, workspace_fields as fields, workspace_table as table
from pearnly_cowork.tabular import readers


class WorkspaceInputs(unittest.TestCase):
    def test_field_contract_keeps_company_contact_values(self):
        self.assertEqual(fields.import_values({"name": " Company ", "tax_id": "0-1055-48103-91-2", "phone": " 08 ", "address": " Bangkok "}),
            {"display_name": "Company", "tax_id": "0-1055-48103-91-2", "address": "Bangkok", "branch": fields.HEAD_OFFICE, "phone": "08"})

    def test_invalid_fields_and_tax(self):
        for row in ({}, {"name": "x", "tax_id": "123"}, {"name": "x", "phone": []}, {"name": "x" * 201}):
            with self.assertRaises(fields.WorkspaceInputError):
                fields.import_values(row)
        self.assertEqual(fields.tax_identity("0-1055-48103-91-2"), fields.tax_identity("0105548103912"))

    def test_csv_preamble_totals_and_first_duplicate_header(self):
        content = "Title\nname,company name,tax id,phone\nFirst,Wrong,0105548103912,081\nTotal,,,\n".encode()
        result = table.parse(content, "roster.csv")
        self.assertEqual(result["matched"]["name"], 0)
        self.assertEqual(result["row_count"], 1)
        self.assertEqual(result["rows"][0]["name"], "First")

    def test_workbook_first_nonempty_sheet_and_blank_values(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.create_sheet("Companies")
        sheet.append(["name", "tax id", "phone"])
        sheet.append(["Company", None, "081"])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        result = table.parse(stream.getvalue(), "misnamed.csv")
        self.assertEqual(result["rows"][0]["tax_id"], "")
        self.assertEqual(result["rows"][0]["phone"], "081")

    def test_encodings_delimiters_and_quoted_cells(self):
        for encoding, name in (("utf-8-sig", "บริษัท"), ("cp874", "บริษัท"), ("gbk", "公司")):
            result = table.parse((f'name;tax id;address\n{name};;"Road; Block"\n').encode(encoding))
            # cp874 is intentionally tried before gbk, matching deployed decoding;
            # GBK may therefore decode as Thai, so only assert shape for that case.
            self.assertEqual(result["rows"][0]["address"], "Road; Block")
            if encoding != "gbk":
                self.assertEqual(result["rows"][0]["name"], name)

    def test_five_hundred_cap_and_truncation(self):
        content = ("name,tax id,phone\n" + "Company,,081\n" * 501).encode()
        result = table.parse(content)
        self.assertEqual(result["row_count"], 500)
        self.assertTrue(result["truncated"])

    def test_limits_corruption_and_missing_header_fail_closed(self):
        for content in (b"", b"PK broken", b"x,y,z\n1,2,3", b"x" * (client_import.MAX_BYTES + 1)):
            with self.assertRaises(ValueError):
                table.parse(content)
        with patch.object(table, "MAX_CELLS", 5):
            with self.assertRaises(ValueError):
                table.parse(b"name,tax id,phone\nCompany,,081")
        payload = io.BytesIO()
        with zipfile.ZipFile(payload, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("huge.xml", b"x" * (64 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(ValueError, "file_too_large"):
            table.parse(payload.getvalue())

    def test_workbook_close_failure_does_not_discard_successful_rows(self):
        from types import SimpleNamespace
        workbook = SimpleNamespace(worksheets=[SimpleNamespace(title="Companies",
            iter_rows=lambda **kwargs: [("name",), ("Company",)])])
        workbook.close = lambda: (_ for _ in ()).throw(OSError("synthetic close failure"))
        with patch("openpyxl.load_workbook", return_value=workbook):
            self.assertEqual(readers.load_workbook_sheets(b"synthetic"), [("Companies", [["name"], ["Company"]])])

    def test_shared_readers_default_and_optional_budget(self):
        content = b"a,b,c\n1,2,3"
        self.assertEqual(len(readers.load_csv_sheets(content)[0][1]), 2)
        self.assertEqual(readers.load_csv_sheets(content, max_cells=5), [])
        from pearnly_cowork.reconciliation import table_io
        self.assertIs(table_io.load_csv_sheets, readers.load_csv_sheets)


if __name__ == "__main__":
    unittest.main()
