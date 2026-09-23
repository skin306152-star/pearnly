"""Run the real hash-bound legacy bank-reconciliation workbook writer.

CLI: --source <fixed-source-root>. The legacy modules are executed from their
fixed source text with only the sibling-module imports removed (those objects are
injected), so no legacy application, database, file or network code runs.
"""
import argparse
import ast
from copy import copy, deepcopy
from datetime import date
import hashlib
from io import BytesIO
import json
from pathlib import Path
import random
import sys

import openpyxl

from pearnly_cowork.reconciliation.bank_matching import BankReconRow, BankReconSummary
from pearnly_cowork.reconciliation.export import labels, usage
from pearnly_cowork.reconciliation.export.workbook import workbook

SOURCE_SHA = "ec2e5c957e6b4be7dee93ffdd2dae9fa83f69188"
HASHES = {
    "services/recon/bank_recon_types.py":
        "df0508cccb9891a5ea628cbce104ee09001b976ade15afd4d599c5d766de149e",
    "services/recon/bank_recon_excel_styles.py":
        "038c2d4eb2590918f3988d61c943b76ab2b9f95f6732343713a2ff2fbe1d7011",
    "services/recon/bank_recon_excel_i18n.py":
        "d4a0071b860c201ec05c7ede48b24f403c96337a865e051dba1e621d381f4fde",
    "services/recon/bank_recon_excel_usage.py":
        "aeb4705687f5b96f6fbf51079f5fcbb423962e06db177b845e28ca0397664497",
    "services/recon/bank_recon_excel_summary.py":
        "cd818cd8ed2f573ab299ee6d0a55ade95fb655dc7ed08e0de6b3e409916495a7",
    "services/recon/bank_recon_excel.py":
        "2735766f0b1f59e8fd9b81854f9f0997f7d3987897b137ac64b33fed3a63a807",
}
ALLOWED_IMPORTS = {"io", "datetime", "typing", "hashlib", "dataclasses", "openpyxl"}
LEAF_NAMES = (
    "COLOR_MATCHED", "COLOR_L2", "COLOR_L3", "COLOR_GL_ONLY", "COLOR_ST_ONLY",
    "_hdr_style", "_border_range", "_fmt_date",
)


class RemoveSiblingImports(ast.NodeTransformer):
    """Drop the legacy sibling imports; every other import must be stdlib/openpyxl."""

    def visit_ImportFrom(self, node):
        if (node.module or "").startswith("services."):
            return None
        assert (node.module or "").split(".")[0] in ALLOWED_IMPORTS, node.module
        return node

    def visit_Import(self, node):
        for alias in node.names:
            assert alias.name.split(".")[0] in ALLOWED_IMPORTS, alias.name
        return node


def load_legacy(root=None, source_text=None):
    trees = {}
    for path, digest in HASHES.items():
        data = (source_text[path].encode("utf-8") if source_text is not None
                else (Path(root) / path).read_bytes())
        assert hashlib.sha256(data).hexdigest() == digest, f"Source hash mismatch: {path}"
        trees[path] = RemoveSiblingImports().visit(ast.parse(data.decode("utf-8")))

    def execute(path, namespace):
        exec(compile(trees[path], path, "exec"), namespace)
        return namespace

    types = execute("services/recon/bank_recon_types.py", {"__name__": "fixed_bank_recon_types"})
    styles = execute("services/recon/bank_recon_excel_styles.py", {"__name__": "fixed_styles"})
    i18n = execute("services/recon/bank_recon_excel_i18n.py", {"__name__": "fixed_i18n"})
    usage_text = execute("services/recon/bank_recon_excel_usage.py", {"__name__": "fixed_usage"})
    shared = {"BankReconRow": types["BankReconRow"], "BankReconSummary": types["BankReconSummary"],
              "_t": i18n["_t"]}
    summary = execute("services/recon/bank_recon_excel_summary.py", {
        "__name__": "fixed_summary", **shared, "_USAGE_BLOCKS": usage_text["_USAGE_BLOCKS"],
        "_border_range": styles["_border_range"]})
    excel = execute("services/recon/bank_recon_excel.py", {
        "__name__": "fixed_excel", **shared, "_build_summary_sheet": summary["_build_summary_sheet"],
        "_layer_label": i18n["_layer_label"], "_status_label": i18n["_status_label"],
        "_USAGE_BLOCKS": usage_text["_USAGE_BLOCKS"],
        **{name: styles[name] for name in LEAF_NAMES}})
    return excel["export_bank_recon_excel"], i18n, usage_text


def compare(old_bytes, new_bytes):
    old = openpyxl.load_workbook(BytesIO(old_bytes), data_only=False)
    new = openpyxl.load_workbook(BytesIO(new_bytes), data_only=False)
    cells = 0
    try:
        assert old.sheetnames == new.sheetnames, (old.sheetnames, new.sheetnames)
        for left, right in zip(old, new, strict=True):
            for attr in ("max_row", "max_column", "freeze_panes", "sheet_state"):
                assert getattr(left, attr) == getattr(right, attr), (left.title, attr)
            assert left.auto_filter == right.auto_filter
            assert str(left.merged_cells) == str(right.merged_cells)
            assert {k: (v.width, v.hidden) for k, v in left.column_dimensions.items()} == {
                k: (v.width, v.hidden) for k, v in right.column_dimensions.items()}, left.title
            assert {k: (v.height, v.hidden) for k, v in left.row_dimensions.items()} == {
                k: (v.height, v.hidden) for k, v in right.row_dimensions.items()}, left.title
            for arow, brow in zip(left, right, strict=True):
                for a, b in zip(arow, brow, strict=True):
                    assert (a.value, a.data_type) == (b.value, b.data_type), (
                        left.title, a.coordinate, a.value, b.value)
                    for attr in ("font", "fill", "border", "alignment", "number_format",
                                 "protection"):
                        assert copy(getattr(a, attr)) == copy(getattr(b, attr)), (
                            left.title, a.coordinate, attr)
                    cells += 1
    finally:
        old.close()
        new.close()
    return cells


def verify_language_tables(old_i18n, old_usage):
    """Every label, usage row and helper output must match the deployed tables."""
    assert tuple(old_i18n["_I18N_EXPORT"]) == tuple(labels.I18N_EXPORT)
    assert old_i18n["_I18N_EXPORT"] == labels.I18N_EXPORT
    assert old_usage["_USAGE_BLOCKS"] == usage.USAGE_BLOCKS
    cases = 0
    for key in list(labels.I18N_EXPORT) + ["missing_key", ""]:
        for lang in ("th", "en", "zh", "ja", "bogus", ""):
            assert old_i18n["_t"](key, lang) == labels.t(key, lang), (key, lang)
            cases += 1
    statuses = ["matched", "gl_debit_only", "gl_credit_only", "stmt_withdrawal_only",
                "stmt_deposit_only", "", "unknown"]
    for layer in [None, 0, 1, 2, 3, 4, "1"]:
        for lang in ("th", "en", "zh", "ja", "bogus"):
            assert old_i18n["_layer_label"](layer, lang) == labels.layer_label(layer, lang), layer
            cases += 1
        for status in statuses:
            assert old_i18n["_status_label"](status, lang) == labels.status_label(status, lang)
            cases += 1
    return cases


def row(status, layer, **values):
    return BankReconRow(
        match_status=status, match_layer=layer, stmt_date=values.get("stmt_date", date(2026, 6, 1)),
        stmt_desc=values.get("stmt_desc", "SWD ACME"), stmt_withdrawal=values.get("withdrawal", 400.0),
        stmt_deposit=values.get("deposit", 0.0), stmt_balance=values.get("balance", 600.0),
        gl_date=values.get("gl_date"), gl_doc_no=values.get("gl_doc_no", ""),
        gl_account_code=values.get("gl_account_code", ""), gl_desc=values.get("gl_desc", ""),
        gl_debit=values.get("debit", 0.0), gl_credit=values.get("credit", 0.0),
        gl_balance=values.get("gl_balance", 0.0), date_diff_days=values.get("date_diff_days"),
        source_stmt_file=values.get("source_stmt_file", "STM KBANK.pdf"),
        source_gl_file=values.get("source_gl_file", ""),
        stmt_confidence=values.get("confidence", "high"),
        stmt_balance_ok=values.get("balance_ok"), stmt_autocorrected=values.get("autocorrected", False))


def fixtures():
    rng = random.Random(20260930)
    langs = ["th", "en", "zh", "ja", "bogus"]
    descriptions = ["SWD ACME INV-77", "ค่าบริการ", "中文摘要", "=SUM(A1)", "", "x" * 60]
    for batch in range(40):
        rows = []
        for index in range(rng.randrange(0, 9)):
            status = rng.choice(["matched", "matched", "gl_debit_only", "gl_credit_only",
                                 "stmt_withdrawal_only", "stmt_deposit_only"])
            layer = rng.choice([1, 2, 3]) if status == "matched" else None
            amount = round(rng.choice([0.0, 12.5, 400.0, 12345.67, -75.25]), 2)
            rows.append(row(
                status, layer, stmt_date=rng.choice([date(2026, 6, index + 1), None]),
                stmt_desc=rng.choice(descriptions), withdrawal=amount if amount > 0 else 0.0,
                deposit=0.0 if amount > 0 else abs(amount), balance=round(rng.choice([0.0, 600.0]) , 2),
                gl_date=rng.choice([date(2026, 6, index + 1), None]),
                gl_doc_no=rng.choice(["GL-1", "", "GL-0002"]),
                gl_account_code=rng.choice(["1000", ""]),
                gl_desc=rng.choice(descriptions), debit=amount if amount > 0 else 0.0,
                credit=0.0 if amount > 0 else abs(amount),
                gl_balance=round(rng.choice([0.0, 12.5]), 2),
                date_diff_days=rng.choice([None, 0, 3]), source_gl_file=rng.choice(["", "GL.xlsx"]),
                confidence=rng.choice(["high", "medium", "low"]),
                balance_ok=rng.choice([True, False, None]),
                autocorrected=rng.choice([True, False])))
        summary = BankReconSummary(
            bank_code=rng.choice(["BBL", "kbank", ""]), gl_account_code=rng.choice(["1000", ""]),
            stmt_opening=round(rng.choice([0.0, 1000.0]), 2), stmt_closing=round(rng.choice([0.0, 510.0]), 2),
            gl_opening=round(rng.choice([0.0, 990.0]), 2), gl_closing=round(rng.choice([0.0, 500.0]), 2),
            stmt_total_deposit=400.0, stmt_total_withdrawal=490.0, gl_total_credit=400.0,
            gl_total_debit=75.0, matched_count=1, gl_debit_only_count=1, gl_credit_only_count=0,
            stmt_withdrawal_only_count=1, stmt_deposit_only_count=0, gl_debit_only_amount=75.0,
            gl_credit_only_amount=0.0, stmt_withdrawal_only_amount=90.0, stmt_deposit_only_amount=0.0,
            opening_diff=10.0, formula_stmt_closing=435.0, formula_diff=75.0)
        task_info = (None if rng.random() < 0.3 else {
            "stmt_files": "STM KBANK.pdf;STM SCB.pdf", "gl_files": "GL.xlsx",
            "stmt_row_count": rng.choice([0, 3]), "gl_row_count": rng.choice([0, 2]),
            "bank_code": "BBL", "gl_account": "1000", "created_at": "2026-09-21 12:00:00"})
        parse_info = (None if rng.random() < 0.3 else {
            "stmt_files": [{"file": "STM KBANK.pdf", "rows": rng.choice([0, 3]), "ok": rng.choice([True, False]),
                            "bank_code": "kbank", "error": rng.choice(["", "no rows"])}],
            "gl_files": [{"file": "GL.xlsx", "rows": 2, "ok": True, "accounts": ["1000", "1100"]}]})
        anchor_overrides = (None if rng.random() < 0.5 else {
            "stmt_opening": {"original": 1000.0, "override": 990.0},
            "gl_closing": {"original": 500.0, "override": 510.0}})
        warnings = (None if rng.random() < 0.5 else ["period mismatch", "", "科目不一致"])
        yield {"recon_rows": rows, "summary": summary, "lang": rng.choice(langs),
               "task_info": task_info, "parse_info": parse_info, "anchor_overrides": anchor_overrides,
               "anchor_ocr": rng.choice([None, {"stmt_opening": 1000.0}]),
               "warnings": warnings}
    # A workbook with no rows at all still has to match sheet by sheet.
    yield {"recon_rows": [], "summary": BankReconSummary(), "lang": "th"}


def run(root=None, source_text=None):
    old_writer, old_i18n, old_usage = load_legacy(root, source_text)
    label_cases = verify_language_tables(old_i18n, old_usage)
    books = cells = 0
    counts = {"rows": 0, "parse_info": 0, "task_info": 0, "overrides": 0, "warnings": 0,
              "autocorrected": 0, "balance_bad": 0, "low_confidence": 0}
    languages = {lang: 0 for lang in ("th", "en", "zh", "ja")}
    statuses, layers = {}, {}
    for value in fixtures():
        arguments = dict(value)
        rows = arguments.pop("recon_rows")
        summary = arguments.pop("summary")
        cells += compare(old_writer(deepcopy(rows), deepcopy(summary), **deepcopy(arguments)),
                         workbook(rows, summary, **arguments))
        books += 1
        counts["rows"] += len(rows)
        for key, argument in (("parse_info", "parse_info"), ("task_info", "task_info"),
                              ("overrides", "anchor_overrides"), ("warnings", "warnings")):
            if arguments.get(argument):
                counts[key] += 1
        if arguments["lang"] in languages:
            languages[arguments["lang"]] += 1
        for item in rows:
            statuses[item.match_status] = statuses.get(item.match_status, 0) + 1
            layers[item.match_layer] = layers.get(item.match_layer, 0) + 1
            counts["autocorrected"] += int(bool(item.stmt_autocorrected))
            counts["balance_bad"] += int(item.stmt_balance_ok is False)
            counts["low_confidence"] += int(item.stmt_confidence == "low")
    for name, minimum in (("rows", 100), ("parse_info", 8), ("task_info", 8), ("overrides", 8),
                          ("warnings", 8), ("autocorrected", 5), ("balance_bad", 5),
                          ("low_confidence", 5)):
        if counts[name] < minimum:
            raise AssertionError((f"{name} barely exercised", counts[name]))
    if books < 30:
        raise AssertionError(("workbooks barely exercised", books))
    for lang, minimum in (("th", 5), ("en", 5), ("zh", 5), ("ja", 5)):
        if languages[lang] < minimum:
            raise AssertionError((f"{lang} barely exercised", languages[lang]))
    for status, minimum in (("matched", 10), ("gl_debit_only", 5), ("gl_credit_only", 5),
                            ("stmt_withdrawal_only", 5), ("stmt_deposit_only", 5)):
        if statuses.get(status, 0) < minimum:
            raise AssertionError((f"{status} barely exercised", statuses.get(status, 0)))
    for layer, minimum in ((1, 5), (2, 5), (3, 5), (None, 5)):
        if layers.get(layer, 0) < minimum:
            raise AssertionError((f"layer {layer} barely exercised", layers.get(layer, 0)))
    print(f"PASS source={SOURCE_SHA} workbooks={books} cells={cells} label_cases={label_cases} "
          f"rows={counts['rows']} languages={languages} statuses={sorted(statuses)} "
          f"layers={sorted(layers, key=str)} value_style_differences=0")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--source", type=Path)
    group.add_argument("--source-stdin", action="store_true")
    arguments = parser.parse_args()
    if arguments.source_stdin:
        run(None, json.load(sys.stdin))
    else:
        run(arguments.source)
