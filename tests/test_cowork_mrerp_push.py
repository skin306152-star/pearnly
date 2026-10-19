"""MR.ERP 发票推送：分流 + 核对菜单 + 导入 5 步 + 逐行回执 + 回读 + 结果落定。假 MR.ERP 站点，不连真实系统、不碰数据库。

对照旧线上 `ec2e5c957e6b`：`mrerp_http/routing.py`（方向 / 付款 / 货费分流、匹配闸）、`mrerp_http/adapter.py`
（formupload → uploadexcel → formrdpc → importpc → report；"1" 全成功、纯数字读 report 逐行判、
非数字体整批失败）、`mrerp_http/autocreate.py`（缺客户 / 供应商 / 商品先自建）、
`confirmed_push_reservation` 的结果判定（导入确认发出后异常 = 结果不明）。

有意变化（2026-09-28，修 MR.ERP 假成功）：菜单按导入页选项 + 账套所选（不写死 118 / 125 / 67 / 453）；
受理后在所选菜单列表里回读到才算成功，「已存在」也要回读到；受理了却查不到 = 结果不明。
"""
import contextlib
import io
import re
import socket
import unittest
from datetime import date
from unittest.mock import patch

try:
    import requests
    from openpyxl import Workbook, load_workbook
    from pearnly_cowork.push.mrerp import routing, runner, transport
    from pearnly_cowork.push.mrerp.models import MRERPBusinessError
except ImportError as exc:  # 需要 Frappe 环境（bench 的 Python）
    raise unittest.SkipTest(f"needs the Frappe bench environment: {exc}")

BASE = "https://www.mrerp4sme.com"
PUBLIC = [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 0))]
LOGIN_PAGE = '<form><input name="txtusers"><input name="txtpasswords"></form>'
TODAY = date(2026, 9, 26)
OWN = "0105559000001"
SUPPLIER = "0105561012345"
CUSTOMER = "0105562055501"
PREVIEW = ('<form id="frmimport1"><input type="hidden" name="idus1" value="15">'
           '<input type="checkbox" name="cbimport[1]" value="1" checked></form>')


def response(url, text="", status=200, content=None):
    r = requests.Response()
    r.url, r.status_code = url, status
    r._content = content if content is not None else text.encode()
    r.encoding = "utf-8"
    return r


def report(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Worksheet"
    sheet.append(["เลขที่", "วันที่", "หมายเหตุ"])
    for invoice_no, note in rows:
        sheet.append([invoice_no, "2026-09-12", note])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# 默认的导入页菜单（照 test01 · TEST2019 的形态，号是合成的）：采购有一个已关联、一个未关联；销项各一个。
MENUS = {"impaptran": [("67", "ซื้อ-เดิม", False), ("68", "ซื้อ", True)],
         "impartran": [("119", "ขายเชื่อ", True)], "imparse": [("126", "ใบขายเงินสด", True)]}
LISTING = {"impaptran": "aptran", "impartran": "artran", "imparse": "arse"}


# 这家公司的主档（接口按 MR.ERP 源码的键与列位；每张小主档只有一个值 = 不用问，科目表有多个 = 要选过）。
CHART = [[1, "1111-01", "เงินสด"], [2, "1112-01", "เงินฝากธนาคาร"], [3, "1130-01", "ลูกหนี้การค้า"],
         [4, "2120-01", "เจ้าหนี้การค้า"], [5, "5300-01", "ค่าใช้จ่ายเบ็ดเตล็ด"], [6, "4110-01", "รายได้"],
         [7, "1150-30", "สินค้าคงเหลือ"], [8, "5100-01", "ต้นทุนขาย"]]
SMALL = {"txtothcombrh": [[1, "00000", "สำนักงานใหญ่"]], "txtothdept_h": [[1, "D001", "บัญชี"]],
         "txtothsuborder_h": [[1, "00001", "ทั่วไป"]], "txtemp": [[1, "E01", "นาย", "สมชาย ใจดี"]],
         "txtstore[]": [[1, "0001", "คลังหลัก"]], "txtararea": [[1, "000001", "กรุงเทพ"]],
         "txtardelivery": [[1, "000001", "ส่งเอง"]], "txtothcombrhcus": [[1, "00000", "สำนักงานใหญ่"]]}
LISTS = {
    "fabillrev": {k: SMALL[k] for k in ("txtothcombrh", "txtothdept_h", "txtothsuborder_h", "txtemp")},
    "artran": {**SMALL, "txtarmas": [[9, "C0001", "บริษัท", "ลูกค้าเก่า", 9, "C0001", "ลูกค้าเก่า", 2, "สมหญิง ขายเก่ง",
                                      2, "เชียงใหม่", 2, "ไปรษณีย์", 1, "00000", "สำนักงานใหญ่"]]},
    "armas": {**SMALL, "txtrectype": [[1, "C-01", "ลูกหนี้การค้า"]], "txtacfile": CHART, "inpdupdata": []},
    "apmas": {"txtsuptype": [[1, "2-11", "เจ้าหนี้การค้า"]], "txtacfile": CHART,
              "txtothcombrhcus": SMALL["txtothcombrhcus"], "inpdupdata": [[1, "เงินสด", "เงินสด"]]},
    "stktraniss": {"txtstore[]": SMALL["txtstore[]"]},
    "stkmas": {"txtstkgrp": [[1, "01", "สินค้า"]], "txtunit_usell": [[1, "PCS", "ชิ้น"]], "txtacfile_rev": CHART},
}
REFERENCE = "".join(f'<input type="text" name="txtacfile[{i}]" id="txtacfile{i}" value="{code}" readonly>'
                    for i, code in ((2, "1111-01"), (4, "1150-30"), (15, "4110-01"), (17, "4110-01"), (18, "4110-01"),
                                    (21, "5100-01"), (22, "5100-01"), (23, "5100-01"), (24, "5100-01")))
PREFIXES = '<select name="selprefix" id="selprefix"><option value="1">นาย</option><option value="2">บริษัท</option></select>'
# 账套选过的科目（科目表有多个，不会自动定）
CHOSEN = {"ar_account": "1130-01", "ap_account": "2120-01", "bank_account": "1112-01", "expense_account": "5300-01"}


def form_page(options):
    if options is None:
        return '<input type="hidden" name="idus" value="15"><h3>กรุณาป้อนข้อมูล "ตั้งค่าระบบสินค้า"</h3>'
    items = "".join(f'<option value="{v}" data-chklac="{1 if ok else 2}">{label}</option>' for v, label, ok in options)
    return f'<input type="hidden" name="idus" value="15"><select name="selmenu" id="selmenu">{items}</select>'


class FakeMrErp:
    """按路径回页面；记录每次导入落在哪个模块 / 菜单；发票导入的确认结果与落库可配置。

    `lands`：importpc 受理后单据是否真的出现在该菜单列表里；`prefix`：菜单单号前缀；
    `listed`：{菜单: [已有单号]}；`unlinked_listing`：列表页提示未关联日记账的菜单；`menus`：导入页选项。
    """

    def __init__(self, password="test01", importpc="1", notes=None, upload_error="", importpc_raises=False,
                 menus=None, lands=True, prefix=None, listed=None, unlinked_listing=(), change_menus_after=None,
                 lists=None, records=None, head="00000", company_tax=""):
        self.password, self.importpc, self.notes = password, importpc, notes
        self.upload_error, self.importpc_raises = upload_error, importpc_raises
        self.menus = dict(MENUS, **(menus or {}))
        self.lands, self.prefix = lands, {"68": "PG", "119": "SI", "126": "SE", **(prefix or {})}
        self.listed = {k: list(v) for k, v in (listed or {}).items()}
        self.unlinked_listing, self.change_menus_after = set(unlinked_listing), change_menus_after
        self.accepted, self.calls, self.uploads = False, [], []
        self.pending, self.session_menu, self.form_opens = {}, None, 0
        self.lists = {k: v for k, v in {**LISTS, **(lists or {})}.items() if v is not None}
        self.files = []
        self.records, self.head = dict(records or {}), head  # {(模块, id): dupdata 回应}；总店分店
        self.company_tax, self.company_reads = company_tax, []  # 本公司分店档案登记的税号；读过哪几个分店 id

    def __call__(self, session, method, url, **kwargs):
        path = url.replace(BASE, "").split("?")[0]
        query = url.split("?", 1)[1] if "?" in url else ""
        self.calls.append((method, path))
        if path == "/login/checklogin.php":
            session.cookies.set("PHPSESSID", "abc")
            self.accepted = kwargs["data"]["txtpasswords"] == self.password
            return response(url)
        if path == "/login/mainmenu.php":
            return response(url, "<p>menu</p>") if self.accepted else response(BASE + "/login/login.php", LOGIN_PAGE)
        module = path.strip("/").split("/")[0]
        if path.endswith("/formupload.php"):
            if module in LISTING:
                self.form_opens += 1
                if self.change_menus_after and self.form_opens > self.change_menus_after[0]:
                    self.menus.update(self.change_menus_after[1])
                return response(url, form_page(self.menus.get(module)))
            return response(url, '<input type="hidden" name="idus" value="15">')
        if path.endswith("/component/uploadexcel.php"):
            self.uploads.append((module, kwargs["data"]["selmenu"]))
            self.files.append((module, kwargs["files"]["uploadfile"][1]))
            if module in LISTING:
                book = load_workbook(io.BytesIO(kwargs["files"]["uploadfile"][1]))
                sheet = book.worksheets[0]
                self.pending = {"menu": kwargs["data"]["selmenu"],
                                "invoices": [str(r[0].value) for r in sheet.iter_rows(min_row=2) if r[0].value]}
            return response(url, self.upload_error if module in LISTING else "")
        if path.endswith("/formrdpc.php"):
            return response(url, PREVIEW)
        if path.endswith("/component/importpc.php"):
            if module in ("imparmas", "impapmas", "impstkmas"):
                return response(url, "1")  # 主档自建：全部受理
            if self.importpc_raises:
                raise requests.ConnectionError("connection reset after import confirmation")
            if self.lands and self.importpc == "1":
                menu = self.pending["menu"]
                self.listed.setdefault(menu, []).extend(self.prefix.get(menu, "") + i for i in self.pending["invoices"])
            return response(url, self.importpc)
        if path.endswith("/component/bshlistboxdata.php"):
            import json
            payload = self.lists.get(module)
            return response(url, json.dumps(payload) if payload is not None else
                            '<script>alert("กรุณา Login เข้าสู่ระบบ")</script>No database selected')
        if path == "/acreffile/acreffileform.php":
            return response(url, REFERENCE)
        if path == "/armas/allform.php":
            return response(url, PREFIXES)
        if path == "/artran/allform.php":
            return response(url, f'<input type="text" name="txtothcombrh" id="txtothcombrh" value="{self.head}" readonly>')
        if path == "/combrh/allform.php":
            self.company_reads.append(query)
            return response(url, '<label>เลขประจำตัวผู้เสียภาษี</label><input type="text" name="txttaxid" '
                                 f'id="txttaxid" size="25" maxlength="40" value="{self.company_tax}">')
        if path.endswith("/component/dupdata.php"):
            import json
            return response(url, json.dumps(self.records.get((module, kwargs["data"]["hiddupdata"]), [])))
        if path.endswith("/component/report.php"):
            return response(url, content=report(self.notes or []))
        if path.endswith("/allview.php"):
            self.session_menu = re.search(r"idmenu=(\d+)", query).group(1)
            if self.session_menu in self.unlinked_listing:
                return response(url, '<h3>กรุณากำหนดการเชื่อมโยง "สมุดรายวัน"</h3>')
            return response(url, '<div id="showdata"></div>')
        if path.endswith("/component/showdata.php"):
            want = kwargs["data"]["searchdataval"]
            rows = [d for d in self.listed.get(self.session_menu, []) if want in d]
            if not rows:
                return response(url, '<p id="nodata">No Data</p>')
            return response(url, "".join(f"<p><span>{d}</span><span>12/09/2569</span><span>ผู้ขาย</span>"
                                         f"<span>1,070.00</span><span><a></a></span></p>" for d in rows))
        return response(url, "<html></html>")


def flat(hid, direction, **fields):
    base = {"invoice_number": f"INV-{hid}", "date": "2026-09-12", "subtotal": "1000.00", "vat": "70.00",
            "total_amount": "1070.00", "direction": direction,
            "items": [{"name": "กระดาษ A4", "qty": "2", "unit_price": "500.00", "amount": "1000.00"}]}
    base.update(fields)
    return {"id": f"CWH-{hid:0>12}", "invoice_no": base["invoice_number"], "invoice_date": base["date"],
            "total_amount": base["total_amount"], "direction": direction, "fields": base}


def sales(hid="a1", **fields):
    return flat(hid, "sales", seller_name="บริษัท เรา จำกัด", seller_tax=OWN, buyer_name="บริษัท ลูกค้า จำกัด",
                buyer_tax=CUSTOMER, **fields)


def purchase(hid="b1", **fields):
    return flat(hid, "purchase", seller_name="บริษัท ผู้ขาย จำกัด", seller_tax=SUPPLIER, buyer_name="บริษัท เรา จำกัด",
                buyer_tax=OWN, **fields)


def push(site, histories, mappings=None):
    def send(session, method, url, **kwargs):
        r = site(session, method, url, **kwargs)
        for hook in session.hooks["response"]:
            hook(r)
        return r

    with patch.object(requests.Session, "request", autospec=True, side_effect=send), \
            patch.object(socket, "getaddrinfo", return_value=PUBLIC), \
            patch.object(transport, "account_lock", lambda *a, **k: contextlib.nullcontext(True)):
        importer = transport.MrErpImport(system_url=BASE, username="test01", password="test01",
                                         comidyear="15", seldb="1")
        with importer:
            return importer, importer.upload_routed_batch(
                histories, {"_own_tax_id": OWN, "_mrerp_defaults": dict(CHOSEN), **(mappings or {})}, today=TODAY)


def execute(site, histories):
    def send(session, method, url, **kwargs):
        r = site(session, method, url, **kwargs)
        for hook in session.hooks["response"]:
            hook(r)
        return r

    with patch.object(requests.Session, "request", autospec=True, side_effect=send), \
            patch.object(socket, "getaddrinfo", return_value=PUBLIC), \
            patch.object(transport, "account_lock", lambda *a, **k: contextlib.nullcontext(True)):
        return runner.execute({"histories": histories, "mappings": {"_own_tax_id": OWN, "_mrerp_defaults": dict(CHOSEN)}},
                              {"account_set": "15:1"}, {"username": "test01", "password": site.password_used},
                              BASE, TODAY)


class Routing(unittest.TestCase):
    """推什么按票面：用户选的方向与付款、复核改过的货 / 费。"""

    def test_sales_payment_follows_the_user_choice_then_the_invoice(self):
        self.assertEqual(routing.choose_doc_type(sales(posting_payment_manual="cash"), {},
                                                 own_tax_id=OWN, today=TODAY), "sales_cash")
        self.assertEqual(routing.choose_doc_type(sales(posting_payment_manual="credit", payment_method="cash"), {},
                                                 own_tax_id=OWN, today=TODAY), "sales_credit")

    def test_purchase_goods_or_expense_follows_the_review_choice(self):
        self.assertEqual(routing.choose_doc_type(purchase(posting_item_type_manual="goods"), {},
                                                 own_tax_id=OWN, today=TODAY), "purchase")
        self.assertEqual(routing.choose_doc_type(purchase(posting_item_type_manual="expense"), {},
                                                 own_tax_id=OWN, today=TODAY), "purchase_expense")

    def test_another_companys_invoice_is_confirmed_mismatch(self):
        other = flat("c1", None, seller_tax=SUPPLIER, buyer_tax=CUSTOMER)
        self.assertTrue(routing.confirmed_account_set_mismatch(other, other, own_tax_id=OWN,
                                                               expected_direction="sales"))
        self.assertFalse(routing.confirmed_account_set_mismatch(sales(), sales(), own_tax_id=OWN,
                                                                expected_direction="sales"))


class Customer(unittest.TestCase):
    """销项客户（用户 2026-09-26「按票面买方」）：关联客户映射 > 票面买方税号 > 现金客户。"""

    def test_customer_comes_from_the_invoice_buyer_without_a_linked_client(self):
        from pearnly_cowork.push.mrerp.xlsx.lookups import MRERP_CASH_CUSTOMER, resolve_customer_code

        self.assertEqual(resolve_customer_code(sales(), {}), CUSTOMER)
        walk_in = flat("w1", "sales", seller_tax=OWN)
        self.assertEqual(resolve_customer_code(walk_in, {}), MRERP_CASH_CUSTOMER)
        linked = dict(sales(), client_id="CWC-9")
        mapping = {"clients": [{"erp_type": "mrerp", "client_id": "CWC-9", "erp_code": "C0009"}]}
        self.assertEqual(resolve_customer_code(linked, mapping), "C0009")

    def test_invoice_buyer_is_created_in_mrerp_with_its_tax_id_as_code(self):
        from pearnly_cowork.push.mrerp import autocreate

        created = []

        class Adapter:
            def create_customers(self, records, mappings):
                created.extend(records)
                return {r["code"]: True for r in records}

        autocreate.provision_customers(Adapter(), [sales(buyer_addr="99 ถนนสุขุมวิท")], {})
        self.assertEqual([(r["code"], r["name"], r["tax_id"], r["address"]) for r in created],
                         [(CUSTOMER, "บริษัท ลูกค้า จำกัด", CUSTOMER, "99 ถนนสุขุมวิท")])


class Import(unittest.TestCase):
    def test_sales_credit_creates_missing_customer_and_product_then_imports(self):
        site = FakeMrErp()
        importer, result = push(site, [sales(posting_payment_manual="credit", client_id="7")])
        invoice_no = routing._gen.derive_mrerp_invoice_no(sales())
        self.assertEqual([s.mrerp_bill_no for s in result.success], ["SI" + invoice_no])  # 回读到的实际单号
        self.assertFalse(result.failed)
        self.assertTrue(importer.import_sent)
        # 客户（imparmas）、商品（impstkmas）先自建，再进销项赊账（impartran / 这家公司唯一已关联的 119）
        self.assertEqual(site.uploads, [("imparmas", ""), ("impstkmas", ""), ("impartran", "119")])

    def test_sales_cash_goes_to_the_cash_module(self):
        site = FakeMrErp()
        _, result = push(site, [sales(posting_payment_manual="cash")],
                         {"products": [{"item_name": "กระดาษ A4", "erp_code": "P-A4"}]})
        self.assertEqual(len(result.success), 1)
        self.assertEqual(site.uploads[-1], ("imparse", "126"))
        self.assertEqual((result.success[0].doc_type, result.success[0].menu), ("sales_cash", "ใบขายเงินสด"))

    def test_purchase_goods_and_expense_use_their_own_chosen_menus(self):
        menus = {"impaptran": [("67", "ซื้อ-เดิม", False), ("68", "ซื้อ", True), ("454", "ซื้อ-ค่าใช้จ่าย", True)]}
        chosen = {"_mrerp_menus": {"purchase": "68", "purchase_expense": "454"}}
        for kind, menu in (("goods", "68"), ("expense", "454")):
            with self.subTest(kind=kind):
                site = FakeMrErp(menus=menus, prefix={"454": "PN"})
                _, result = push(site, [purchase(posting_item_type_manual=kind)], chosen)
                self.assertEqual(len(result.success), 1)
                self.assertEqual(site.uploads[0], ("impapmas", ""))  # 缺供应商先建
                self.assertEqual(site.uploads[-1], ("impaptran", menu))

    def test_single_linked_menu_is_used_without_a_choice(self):
        site = FakeMrErp()  # 采购只有 68 已关联（67 未关联，不能用）
        _, result = push(site, [purchase(posting_item_type_manual="expense")])
        self.assertEqual(len(result.success), 1)
        self.assertEqual(site.uploads[-1], ("impaptran", "68"))

    def test_numeric_reply_reads_the_report_row_by_row(self):
        invoice_no = routing._gen.derive_mrerp_invoice_no(purchase())
        site = FakeMrErp(importpc="2", notes=[(invoice_no, "ไม่พบรหัสสินค้า")])
        _, result = push(site, [purchase()])
        self.assertFalse(result.success)
        self.assertEqual(result.failed[0].reasons, ["ไม่พบรหัสสินค้า"])

    def test_non_numeric_reply_fails_the_whole_batch_honestly(self):
        _, result = push(FakeMrErp(importpc="<script>alert('Error : ')</script>"), [purchase()])
        self.assertFalse(result.success)
        self.assertEqual(result.failed[0].reasons[0], "ERR_MRERP_IMPORT_ERROR")

    def test_upload_rejection_happens_before_the_import_is_sent(self):
        with self.assertRaises(MRERPBusinessError):
            push(FakeMrErp(upload_error="รูปแบบไฟล์ไม่ถูกต้อง"), [purchase()])

    def test_bookkeeping_for_a_client_goes_by_the_companys_own_tax_id(self):
        # T2（2026-10-19）：事务所的工作空间税号不是票上的买卖方；按这家公司在 MR.ERP 登记的税号（总店分店档案）判出进项，
        # 匹配闸也认它。以前判不出 → 默认赊销模块 → 匹配闸拦成 ERR_ACCOUNT_SET_MISMATCH。
        firm, client = "0107536000188", "0105551234567"  # 事务所工作空间 / 客户公司（校验位都对）
        bill = flat("t2", None, seller_name="บริษัท ผู้ขาย จำกัด", seller_tax=SUPPLIER, buyer_name="บริษัท ลูกค้า จำกัด",
                    buyer_tax=client)
        site = FakeMrErp(company_tax="0-1055-51234-56-7")
        _, result = push(site, [bill], {"_own_tax_id": firm})
        self.assertEqual(site.company_reads, ["id=1&status=vie"])  # 总店 00000 是分店列表里 id 1 那一行
        self.assertEqual((len(result.success), result.failed), (1, []))
        self.assertEqual([module for module, _ in site.uploads if module.startswith("imp")][-1:], ["impaptran"])
        # 这家公司没登记税号：照旧只按工作空间税号，拦下
        site = FakeMrErp()
        _, result = push(site, [bill], {"_own_tax_id": firm})
        self.assertEqual(result.failed[0].reasons, [routing.ACCOUNT_SET_MISMATCH_CODE])

    def test_the_company_tax_id_is_not_read_when_every_document_declares_its_direction(self):
        site = FakeMrErp(company_tax="0-1055-51234-56-7")
        push(site, [purchase()])
        self.assertEqual(site.company_reads, [])

    def test_another_companys_invoice_never_reaches_the_import(self):
        # 票面声明的方向优先于税号锚点（旧版同）：这道闸只在方向没声明时起作用。
        site = FakeMrErp()
        other = flat("c1", None, seller_tax=SUPPLIER, buyer_tax=CUSTOMER)
        _, result = push(site, [other])
        self.assertEqual(result.failed[0].reasons, [routing.ACCOUNT_SET_MISMATCH_CODE])
        self.assertFalse([u for u in site.uploads if u[0] in ("impartran", "imparse", "impaptran")])


def sheet_rows(blob, index=0):
    book = load_workbook(io.BytesIO(blob))
    return [[c.value for c in row] for row in book.worksheets[index].iter_rows(min_row=2)]


TEST2019 = ("BOI1", "กร ทดสอบ", "สุพรรณ", "TRUCK1", "1130-03", "2110-01", "5230-01", "04-SER", "03-FIG", "ผน",
            "CODE_NOAP", "1-11", "2-21")


class Defaults(unittest.TestCase):
    """导入值按这家公司的主档与账套所选（用户 2026-10-08）：不写 TEST2019 的值；定不下来就停下请人选。"""

    def test_header_and_lines_use_the_company_masters_and_the_customer_record(self):
        site = FakeMrErp()
        mapping = {"clients": [{"erp_type": "mrerp", "client_id": "7", "erp_code": "C0001"}],
                   "products": [{"item_name": "กระดาษ A4", "erp_code": "P-A4"}]}
        _, result = push(site, [dict(sales(posting_payment_manual="credit"), client_id="7")], mapping)
        self.assertEqual(len(result.success), 1)
        blob = dict(site.files)["impartran"]
        head = sheet_rows(blob)[0]
        # 分店 / 部门 / 工作取主档唯一值；员工 / 区域 / 运输取客户档 C0001 的
        self.assertEqual(head[3:7], ["00000", "D001", "00001", "สมหญิง ขายเก่ง"])
        self.assertEqual(head[12:14], ["เชียงใหม่", "ไปรษณีย์"])
        self.assertEqual(sheet_rows(blob, 1)[0][2:5], ["D001", "00001", "0001"])

    def test_nothing_from_test2019_reaches_any_import_file(self):
        for histories in ([sales(posting_payment_manual="credit", client_id="7")],
                          [purchase(posting_item_type_manual="goods")], [purchase(posting_item_type_manual="expense")]):
            site = FakeMrErp(menus={"impaptran": [("68", "ซื้อ", True)]})
            _, result = push(site, histories)
            self.assertEqual(len(result.success), 1, result.failed)
            for module, blob in site.files:
                book = load_workbook(io.BytesIO(blob))
                values = {str(c.value) for ws in book.worksheets for row in ws.iter_rows() for c in row if c.value}
                with self.subTest(module=module):
                    self.assertFalse(values & set(TEST2019), values & set(TEST2019))

    def test_new_customer_and_supplier_get_the_companys_codes(self):
        site = FakeMrErp()
        push(site, [sales(posting_payment_manual="credit", client_id="7")])
        customer = sheet_rows(dict(site.files)["imparmas"])[0]
        self.assertEqual(customer[1:3], ["C-01", "บริษัท"])          # 类型、称谓（按名称前缀）
        self.assertEqual(customer[12], "1130-01")                     # 应收科目（账套所选）
        self.assertEqual(customer[15:18], ["E01", "000001", "000001"])  # 销售员 / 区域 / 运输代码
        self.assertEqual(customer[24], "00000")
        product = sheet_rows(dict(site.files)["impstkmas"])[0]
        self.assertEqual(product[1], "01")
        self.assertEqual(product[10:13], ["PCS", "PCS", "PCS"])
        self.assertEqual(product[20:28], ["4110-01", "4110-01", "4110-01", "5100-01", "5100-01", "5100-01",
                                          "1150-30", "5100-01"])  # 参考科目
        self.assertEqual(product[28], "เงินสด")
        site = FakeMrErp()
        push(site, [purchase(posting_item_type_manual="expense")])
        supplier = sheet_rows(dict(site.files)["impapmas"])[0]
        self.assertEqual((supplier[1], supplier[12]), ("2-11", "2120-01"))
        expense_item = sheet_rows(dict(site.files)["impstkmas"])[0]
        self.assertEqual(set(expense_item[20:28]), {"5300-01"})       # 费用物料 8 科目 = 所选费用科目

    def test_cash_receipt_goes_in_the_cash_column_and_bank_receipt_to_the_chosen_bank_account(self):
        mapping = {"products": [{"item_name": "กระดาษ A4", "erp_code": "P-A4"}]}
        for method, cash, slot, account in (("cash", 1070.0, 0, "1111-01"), ("transfer", 0, 1070.0, "1112-01")):
            with self.subTest(method=method):
                site = FakeMrErp()
                push(site, [sales(posting_payment_manual="cash", payment_method=method, payment_status="paid")], mapping)
                head = sheet_rows(dict(site.files)["imparse"])[0]
                self.assertEqual((float(head[15] or 0), head[16], float(head[17] or 0)), (cash, account, slot))
                self.assertEqual((head[18], head[20]), (account, account))

    def test_several_departments_follow_the_mrerp_default_unless_one_was_chosen(self):
        # 不问人：没选过就跟 MR.ERP 网页新建单据的默认（按代码第一条），选过就用选过的
        two = {"fabillrev": {**LISTS["fabillrev"], "txtothdept_h": [[1, "D001", "บัญชี"], [2, "D002", "ขาย"]]}}
        site = FakeMrErp(lists=two)
        _, result = push(site, [purchase()])
        self.assertEqual(len(result.success), 1)
        self.assertEqual(sheet_rows(dict(site.files)["impaptran"])[0][4], "D001")
        site = FakeMrErp(lists=two)
        _, result = push(site, [purchase()], {"_mrerp_defaults": {**CHOSEN, "dept": "D002"}})
        self.assertEqual(len(result.success), 1)
        self.assertEqual(sheet_rows(dict(site.files)["impaptran"])[0][4], "D002")

    def test_a_chosen_value_that_is_gone_stops_and_is_not_replaced(self):
        site = FakeMrErp()
        _, result = push(site, [purchase()], {"_mrerp_defaults": {**CHOSEN, "dept": "D009"}})
        self.assertEqual(result.failed[0].reasons, ["ERR_MRERP_DEFAULT_CHANGED"])
        self.assertFalse(site.uploads)

    def test_unreadable_masters_stop(self):
        site = FakeMrErp(lists={"fabillrev": None})
        _, result = push(site, [purchase()])
        self.assertEqual(result.failed[0].reasons, ["ERR_MRERP_MASTER_UNREADABLE"])
        self.assertFalse(site.uploads)

    def test_new_customer_copies_the_companys_most_common_type_and_account(self):
        # 现有客户 3 个：2 个挂 1130-01，1 个挂 1111-01 → 新客户用 1130-01（MR.ERP「复制现有档案」同源）
        row = lambda account: [1, "C-01", "", 1, "x"] + [""] * 9 + [account]  # noqa: E731
        lists = {"armas": {**LISTS["armas"], "inpdupdata": [[1, "C1", "a"], [2, "C2", "b"], [3, "C3", "c"]]}}
        records = {("armas", "1"): row("1130-01"), ("armas", "2"): row("1130-01"), ("armas", "3"): row("1111-01")}
        site = FakeMrErp(lists=lists, records=records)
        _, result = push(site, [dict(sales(posting_payment_manual="credit"), client_id="7")],
                         {"_mrerp_defaults": {k: v for k, v in CHOSEN.items() if k != "ar_account"}})
        self.assertEqual(len(result.success), 1, result.failed)
        self.assertEqual(sheet_rows(dict(site.files)["imparmas"])[0][12], "1130-01")

    def test_head_office_is_the_branch_when_there_are_several(self):
        lists = {"artran": {**LISTS["artran"], "txtothcombrh": [[1, "00000", "สนญ"], [2, "00001", "สาขา"]]}}
        site = FakeMrErp(lists=lists, head="00001")
        push(site, [dict(sales(posting_payment_manual="credit"), client_id="7")],
             {"clients": [{"erp_type": "mrerp", "client_id": "7", "erp_code": "C0001"}],
              "products": [{"item_name": "กระดาษ A4", "erp_code": "P-A4"}]})
        self.assertEqual(sheet_rows(dict(site.files)["impartran"])[0][3], "00001")

    def test_account_not_chosen_stops_before_creating_the_customer(self):
        site = FakeMrErp()
        _, result = push(site, [sales(posting_payment_manual="credit", client_id="7")],
                         {"_mrerp_defaults": {k: v for k, v in CHOSEN.items() if k != "ar_account"}})
        self.assertEqual(result.failed[0].reasons, ["ERR_MRERP_DEFAULT_CHOICE_REQUIRED"])
        self.assertFalse([u for u in site.uploads if u[0] in ("imparmas", "impartran")])

    def test_master_mismatch_in_the_report_asks_to_choose_again(self):
        invoice_no = routing._gen.derive_mrerp_invoice_no(purchase())
        site = FakeMrErp(importpc="2", notes=[(invoice_no, "ไม่พบข้อมูลแผนก")])
        _, result = push(site, [purchase()])
        self.assertEqual(result.failed[0].reasons, ["ERR_MRERP_DEFAULT_CHANGED", "ไม่พบข้อมูลแผนก"])

    def test_only_an_existing_code_counts_as_created(self):
        self.assertTrue(transport._CODE_EXISTS.search("รหัสสินค้าดังกล่าวมีอยู่ในระบบแล้ว"))
        self.assertTrue(transport._CODE_EXISTS.search("รหัสลูกค้าดังกล่าวมีอยู่ในระบบแล้ว"))
        self.assertFalse(transport._CODE_EXISTS.search("ชื่อสินค้าดังกล่าวมีอยู่ในระบบแล้ว"))

    def test_masters_read_during_the_push_are_returned_for_the_choice_dialog(self):
        site = FakeMrErp()
        site.password_used = "test01"
        outcome = execute(site, [purchase()])
        self.assertEqual(outcome["status"], "success", outcome)
        self.assertEqual(outcome["masters"]["dept"], [{"code": "D001", "name": "บัญชี"}])


class Menu(unittest.TestCase):
    """菜单定不下来就在导入前停下（不建主档、不上传），原因码原样交给账本；不自动换菜单。"""

    def assert_stopped(self, site, histories, code, mappings=None):
        importer, result = push(site, histories, mappings)
        self.assertFalse(result.success)
        self.assertEqual(result.failed[0].reasons, [code])
        self.assertFalse(importer.import_sent)
        return importer

    def test_several_linked_menus_without_a_choice_ask_first(self):
        from pearnly_cowork.targets import mrerp_menus
        site = FakeMrErp(menus={"impaptran": [("68", "ซื้อ", True), ("454", "ซื้อ-ค่าใช้จ่าย", True)]})
        self.assert_stopped(site, [purchase()], mrerp_menus.CHOICE_REQUIRED)
        self.assertEqual(site.uploads, [])  # 供应商也没建

    def test_chosen_menu_that_is_gone_or_unlinked_stops_instead_of_switching(self):
        from pearnly_cowork.targets import mrerp_menus
        for menus in ([("68", "ซื้อ", True)], [("68", "ซื้อ", True), ("454", "ซื้อ-ค่าใช้จ่าย", False)]):
            with self.subTest(menus=menus):
                site = FakeMrErp(menus={"impaptran": menus})
                self.assert_stopped(site, [purchase()], mrerp_menus.CHANGED,
                                    {"_mrerp_menus": {"purchase": "454", "purchase_expense": "454"}})
                self.assertEqual(site.uploads, [])

    def test_no_linked_menu_or_unreadable_page(self):
        from pearnly_cowork.targets import mrerp_menus
        self.assert_stopped(FakeMrErp(menus={"impaptran": [("67", "ซื้อ", False)]}), [purchase()],
                            mrerp_menus.NOT_LINKED)
        self.assert_stopped(FakeMrErp(menus={"impaptran": None}), [purchase()], mrerp_menus.UNREADABLE)

    def test_menu_unlinked_while_masters_were_created_stops_before_upload(self):
        from pearnly_cowork.targets import mrerp_menus
        # 选账套抓 idus、核对菜单各打开一次导入页；第三次（主档自建之后、上传之前）时 68 已被取消关联。
        site = FakeMrErp(change_menus_after=(2, {"impaptran": [("68", "ซื้อ", False)]}))
        self.assert_stopped(site, [purchase()], mrerp_menus.CHANGED)
        self.assertNotIn("impaptran", [module for module, _ in site.uploads])

    def test_menu_failure_reaches_the_ledger_as_the_bare_code(self):
        from pearnly_cowork.targets import mrerp_menus
        site = FakeMrErp(menus={"impaptran": [("68", "ซื้อ", True), ("454", "ซื้อ-ค่าใช้จ่าย", True)]})
        site.password_used = "test01"
        out = execute(site, [purchase()])
        self.assertEqual((out["status"], out["reason"]), ("failed", mrerp_menus.CHOICE_REQUIRED))


class Outcome(unittest.TestCase):
    """落到账本的结果：成功 / 已推过 / 转人工 / 结果不明。"""

    def run_case(self, site, histories, password="test01"):
        site.password_used = password
        return execute(site, histories)

    def test_success_carries_the_mrerp_document_number(self):
        out = self.run_case(FakeMrErp(), [purchase()])
        self.assertEqual(out["status"], "success")
        self.assertEqual(out["docnum"], "PG" + routing._gen.derive_mrerp_invoice_no(purchase()))
        self.assertEqual([o["idmenu"] for o in out["menus"]["impaptran"]], ["67", "68"])  # 读到的选项交给调用方存
        # 进了哪类单据、哪个导入菜单记进回执（CC 说得出「赊购（费用），进了 ซื้อ」，2026-10-19）
        self.assertEqual((out["meta"]["doc_type"], out["meta"]["mrerp_menu"]), ("purchase_expense", "ซื้อ"))

    def test_already_existing_number_found_in_the_chosen_menu_is_already_pushed(self):
        invoice_no = routing._gen.derive_mrerp_invoice_no(purchase())
        site = FakeMrErp(importpc="2", notes=[(invoice_no, "เลขที่ดังกล่าวมีอยู่ในระบบแล้ว")],
                         listed={"68": ["PG" + invoice_no]})
        out = self.run_case(site, [purchase()])
        self.assertEqual((out["status"], out["docnum"], out["meta"].get("already_pushed")),
                         ("success", "PG" + invoice_no, True))

    def test_already_existing_number_not_found_in_the_chosen_menu_is_unknown(self):
        invoice_no = routing._gen.derive_mrerp_invoice_no(purchase())
        out = self.run_case(FakeMrErp(importpc="2", notes=[(invoice_no, "เลขที่ดังกล่าวมีอยู่ในระบบแล้ว")]),
                            [purchase()])
        self.assertEqual((out["status"], out["reason"]), ("unknown", transport.READBACK_NOT_FOUND))

    def test_accepted_but_not_listed_is_unknown_not_success(self):
        # MR.ERP 回 "1" 却在所选菜单里查不到（旧版就是在这里记了假成功）。
        out = self.run_case(FakeMrErp(lands=False), [purchase()])
        self.assertEqual((out["status"], out["reason"]), ("unknown", transport.READBACK_NOT_FOUND))

    def test_listing_that_says_the_menu_has_no_journal_is_unknown(self):
        out = self.run_case(FakeMrErp(unlinked_listing={"68"}), [purchase()])
        self.assertEqual((out["status"], out["reason"]), ("unknown", transport.READBACK_NOT_LINKED))

    def test_report_row_accepted_is_still_read_back(self):
        # "2" 路径：report 里空备注的行也要回读；混批另一张失败照旧失败。
        site = FakeMrErp(importpc="2", notes=[(routing._gen.derive_mrerp_invoice_no(purchase()), "")])
        self.assertEqual(self.run_case(site, [purchase()])["status"], "unknown")  # 没落进列表
        site = FakeMrErp(importpc="2", notes=[(routing._gen.derive_mrerp_invoice_no(purchase()), "")],
                         listed={"68": ["PG" + routing._gen.derive_mrerp_invoice_no(purchase())]})
        self.assertEqual(self.run_case(site, [purchase()])["status"], "success")

    def test_definite_failures_before_sending_are_not_unknown(self):
        for site, password, code in ((FakeMrErp(), "wrong", "ERR_AUTH"),
                                     (FakeMrErp(upload_error="ผิดพลาด"), "test01", "ERR_BUSINESS")):
            with self.subTest(code=code):
                out = self.run_case(site, [purchase()], password)
                self.assertEqual(out["status"], "failed")
                self.assertTrue(out["reason"].startswith(code))

    def test_an_error_after_the_import_was_sent_is_unknown(self):
        out = self.run_case(FakeMrErp(importpc_raises=True), [purchase()])
        self.assertEqual(out["status"], "unknown")


if __name__ == "__main__":
    unittest.main()
