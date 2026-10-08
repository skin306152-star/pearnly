"""Offline XLSX checks for the migrated income-reconciliation review workbook.

The workbook is pure presentation: the caller supplies the detail rows and the
summary, so these cases need no Site, database, model or file.
"""
import unittest
from io import BytesIO

import openpyxl

from pearnly_cowork.reconciliation.export import income_workbook
from pearnly_cowork.reconciliation.income_vat import GlVatSummary, ReconRow


def detail_row(doc_no="GL-1", **values):
	return ReconRow(
		doc_no=doc_no, date=values.get("date", "01/06/2569"),
		customer_name=values.get("customer_name", "ACME"),
		vat_amount=values.get("vat_amount", 1070.0),
		gl_amount=values.get("gl_amount", 1070.0),
		diff=values.get("diff", 0.0),
		account_codes=values.get("account_codes", "4110-01"))


def summary(**values):
	return GlVatSummary(
		gl_total=values.get("gl_total", 1493.0),
		gl_only_credit=values.get("gl_only_credit", 500.0),
		gl_only_debit=values.get("gl_only_debit", 0.0),
		vat_only_positive=values.get("vat_only_positive", 250.0),
		vat_only_negative=values.get("vat_only_negative", 0.0),
		vat_total=values.get("vat_total", 1243.0),
		gl_only_credit_items=values.get("gl_only_credit_items", [
			{"doc_no": "GL-3", "date": "03/06/2569", "name": "no vat row", "amount": 500.0}]),
		vat_only_positive_items=values.get("vat_only_positive_items", [
			{"doc_no": "INV-9", "date": "05/06/2569", "name": "BETA", "amount": 250.0}]))


def sheets(content):
	return openpyxl.load_workbook(BytesIO(content), data_only=False)


MONEY = "#,##0.00;[Red]-#,##0.00"
OK, WARN, MISS, HEADER, EMPHASIS = "00DCFCE7", "00FEF3C7", "00FEE2E2", "001F2937", "00F3F4F6"


class SheetLayout(unittest.TestCase):
	"""Cowork 统一复核表（用户 2026-09-26）：没有 KPI 与使用说明，外观走 export/theme.py。"""

	def test_the_two_sheets_are_named_in_the_requested_language(self):
		for lang, names in [("th", ("รายละเอียดผลต่าง", "สรุปการกระทบยอด")),
				("zh", ("差异明细", "对账汇总")),
				("en", ("Detail", "Reconciliation")),
				("ja", ("差異明細", "照合まとめ"))]:
			with self.subTest(lang=lang):
				book = sheets(income_workbook.export_gl_vat_excel([detail_row()], summary(),
					lang=lang))
				self.assertEqual(book.sheetnames, list(names))

	def test_an_unknown_language_falls_back_to_thai(self):
		book = sheets(income_workbook.export_gl_vat_excel([], summary(), lang="bogus"))
		self.assertEqual(book.sheetnames, ["รายละเอียดผลต่าง", "สรุปการกระทบยอด"])

	def test_the_detail_sheet_starts_with_its_header(self):
		content = income_workbook.export_gl_vat_excel([detail_row(), detail_row("GL-2")],
			summary(), lang="en")
		ws = sheets(content)["Detail"]
		self.assertEqual((ws.max_row, ws.max_column), (3, 8))
		self.assertEqual(ws.freeze_panes, "A2")
		self.assertEqual((ws.row_dimensions[1].height, ws.row_dimensions[2].height), (22, 18))
		self.assertEqual([ws.column_dimensions[letter].width for letter in "ABCDEFGH"],
			[14, 20, 14, 32, 18, 18, 14, 22])
		self.assertEqual([ws.cell(1, index).value for index in range(1, 9)],
			["Status", "Document No.", "Date", "Customer", "VAT Report Amount", "GL Amount",
				"Difference", "Revenue Account"])
		self.assertEqual((ws.cell(1, 1).fill.fgColor.rgb, ws.cell(1, 1).font.color.rgb), (HEADER, "00FFFFFF"))

	def test_a_row_carries_its_status_text_and_colour(self):
		rows = [detail_row("GL-1"), detail_row("GL-2", gl_amount=1100.0, diff=30.0),
			detail_row("INV-9", gl_amount=None, diff=None, account_codes="")]
		ws = sheets(income_workbook.export_gl_vat_excel(rows, summary(), lang="en"))["Detail"]
		self.assertEqual([ws.cell(row, 1).value for row in (2, 3, 4)],
			["✓ Matched", "⚠ Diff", "❗ GL Missing"])
		# 三张表同一口径：绿 = 一致，黄 = 有差异，红 = 缺一边。
		self.assertEqual([ws.cell(row, 1).fill.fgColor.rgb for row in (2, 3, 4)], [OK, WARN, MISS])
		self.assertEqual([ws.cell(2, 7).fill.fgColor.rgb, ws.cell(3, 7).fill.fgColor.rgb,
			ws.cell(4, 6).fill.fgColor.rgb, ws.cell(4, 7).fill.fgColor.rgb], [OK, WARN, MISS, MISS])
		self.assertEqual(ws.cell(2, 1).alignment.horizontal, "left")
		for column in (5, 6, 7):
			self.assertEqual(ws.cell(2, column).alignment.horizontal, "right")
			self.assertEqual(ws.cell(2, column).number_format, MONEY)

	def test_an_absent_ledger_side_reads_as_the_not_found_label(self):
		rows = [detail_row("INV-9", gl_amount=None, diff=None, account_codes="")]
		for lang, label in [("th", "ไม่พบข้อมูล"), ("zh", "未找到数据"), ("en", "Not found"),
				("ja", "データなし")]:
			with self.subTest(lang=lang):
				ws = sheets(income_workbook.export_gl_vat_excel(rows, summary(), lang=lang))[
					income_workbook._t(lang, "sheet1")]
				self.assertEqual((ws.cell(2, 6).value, ws.cell(2, 7).value), (label, label))

	def test_an_empty_task_still_renders_the_header(self):
		ws = sheets(income_workbook.export_gl_vat_excel([], summary(), lang="en"))["Detail"]
		self.assertEqual(ws.cell(1, 1).value, "Status")
		self.assertEqual(ws.max_row, 1)


class SummarySheet(unittest.TestCase):
	def test_the_ladder_reads_total_less_credits_plus_debits_then_the_report(self):
		ws = sheets(income_workbook.export_gl_vat_excel([], summary(), lang="en"))[
			"Reconciliation"]
		self.assertEqual(ws.max_column, 2)
		self.assertEqual(ws.cell(1, 1).value, "Description")
		self.assertEqual(ws.cell(1, 2).value, "Amount")
		# Only the two totals carry an amount; each adjustment heading leaves it to the
		# itemized rows the deployed report expands below it.
		self.assertEqual([(ws.cell(row, 1).value, ws.cell(row, 2).value)
			for row in (2, 5, 6, 8, 9)],
			[("Total per General Ledger", 1493.0),
				("Add: GL debits not in VAT Report", None),
				("Add: Sales in VAT Report not in GL", None),
				("Less: Credit notes in VAT Report not in GL", None),
				("Total per VAT Sales Report", 1243.0)])
		self.assertEqual([(ws.cell(row, 1).value, ws.cell(row, 2).value)
			for row in (3, 4, 7)],
			[("Less: GL credits not in VAT Report", None),
				("    · GL-3  ·  03/06/2026  ·  no vat row", -500.0),
				("    · INV-9  ·  05/06/2026  ·  BETA", 250.0)])

	def test_amount_differences_get_their_own_adjustment_so_the_ladder_adds_up(self):
		# 2026-09-28 录制：IV6909-003 税表 12,000 / GL 12,800，调节表直接跳到税表合计、看着算不平
		detail = [detail_row("IV6909-003", date="2026-09-18", customer_name="ACME", vat_amount=12000.0,
			gl_amount=12800.0, diff=-800.0), detail_row("IV6909-004")]
		ws = sheets(income_workbook.export_gl_vat_excel(detail, summary(gl_total=2293.0), lang="en"))["Reconciliation"]
		labels = [ws.cell(row, 1).value for row in range(1, ws.max_row + 1)]
		at = labels.index("Adjust: amount differences (VAT report − GL)")
		self.assertEqual((labels[at + 1], ws.cell(at + 2, 2).value), ("    · IV6909-003  ·  18/09/2026  ·  ACME", -800.0))
		self.assertEqual(labels[at + 2], "Total per VAT Sales Report")
		# 没有差额的任务版面不变
		plain = sheets(income_workbook.export_gl_vat_excel([detail_row()], summary(), lang="en"))["Reconciliation"]
		self.assertNotIn("Adjust: amount differences (VAT report − GL)",
			[plain.cell(row, 1).value for row in range(1, plain.max_row + 1)])

	def test_the_emphasized_rows_and_the_item_rows_share_the_common_look(self):
		ws = sheets(income_workbook.export_gl_vat_excel([], summary(), lang="en"))[
			"Reconciliation"]
		self.assertEqual(ws.cell(1, 1).fill.fgColor.rgb, HEADER)
		self.assertEqual((ws.cell(2, 2).fill.fgColor.rgb, ws.cell(2, 2).font.b), (EMPHASIS, True))
		self.assertEqual(ws.cell(2, 2).number_format, MONEY)
		self.assertEqual(ws.cell(3, 1).fill.fgColor.rgb, EMPHASIS)
		self.assertEqual(ws.cell(4, 1).value, "    · GL-3  ·  03/06/2026  ·  no vat row")
		# A credit row is shown negated.
		self.assertEqual(ws.cell(4, 2).value, -500.0)
		self.assertIsNone(ws.cell(4, 1).fill.fill_type)
		self.assertEqual((ws.cell(4, 1).font.size, ws.cell(4, 1).font.name), (10, "Arial"))
		self.assertEqual(ws.cell(4, 2).number_format, MONEY)
		# 列宽统一不超过 48 个字符（原来 90，金额列会被挤出屏幕）。
		self.assertEqual(ws.column_dimensions["A"].width, 48)
		self.assertEqual(ws.column_dimensions["B"].width, 22)

	def test_a_long_item_name_is_cut_with_an_ellipsis(self):
		value = summary(gl_only_credit_items=[{"doc_no": "GL-9", "date": "01/06/2569",
			"name": "x" * 80, "amount": 1.0}])
		ws = sheets(income_workbook.export_gl_vat_excel([], value, lang="en"))["Reconciliation"]
		self.assertEqual(ws.cell(4, 1).value, f"    · GL-9  ·  01/06/2026  ·  {'x' * 53}…")

	def test_a_positive_item_is_not_negated(self):
		ws = sheets(income_workbook.export_gl_vat_excel([], summary(), lang="en"))[
			"Reconciliation"]
		rows = {ws.cell(row, 1).value: ws.cell(row, 2).value for row in range(2, ws.max_row + 1)}
		self.assertEqual(rows["    · INV-9  ·  05/06/2026  ·  BETA"], 250.0)

	def test_every_row_of_the_ladder_is_ruled_alike(self):
		ws = sheets(income_workbook.export_gl_vat_excel([], summary(), lang="en"))[
			"Reconciliation"]
		for row in range(1, ws.max_row + 1):
			self.assertEqual(ws.cell(row, 1).border.top.style, "thin")
			self.assertEqual(ws.cell(row, 2).border.bottom.style, "thin")


class LabelTable(unittest.TestCase):
	def test_every_label_covers_the_four_languages_and_unknown_keys_fall_back(self):
		self.assertEqual(set(income_workbook._I18N), {"th", "zh", "en", "ja"})
		keys = set(income_workbook._I18N["th"])
		for lang in ("th", "zh", "en", "ja"):
			self.assertEqual(set(income_workbook._I18N[lang]), keys)
		for lang in ("th", "zh", "en", "ja"):
			for key in keys:
				self.assertTrue(income_workbook._t(lang, key))
		self.assertEqual(income_workbook._t("en", "missing_key"), "missing_key")
		self.assertEqual(income_workbook._t("bogus", "sheet1"),
			income_workbook._I18N["th"]["sheet1"])


if __name__ == "__main__":
	unittest.main()
