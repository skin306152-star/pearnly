"""对账提速（2026-10-01）：结果页 / 导出复用样式下标、不反复读 max_row；发往 Cloud Run 的对账识别不占本机识别名额。
在镜像里跑，不连站点。真实数据上的逐格比对见 docs/work/2026-10-01-recon-speed.md。
"""
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "apps" / "pearnly_cowork"))
sys.path.insert(0, str(ROOT / "apps" / "pearnly_core"))
for name in [name for name in sys.modules if name.startswith(("pearnly_cowork", "pearnly_core"))]:
    del sys.modules[name]

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

from pearnly_cowork.reconciliation.export import styles, theme  # noqa: E402

assert Path(theme.__file__).is_relative_to(ROOT), theme.__file__


class NeverCached(dict):
    """让 ``_style_sheet`` 每格都走逐格赋值（原来的做法），作对照。"""

    def __contains__(self, key):
        return False


def sample_book(rows=60):
    book = Workbook()
    sheet = book.active
    sheet.title = "detail"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    sheet.cell(1, 1, "标题")
    for column in range(1, 6):
        cell = sheet.cell(2, column, f"h{column}")
        cell.fill = PatternFill("solid", fgColor="1A3C5E")
    colors = ["D6F5E3", "FDE8E8", "FFF3CD", "EFF6FF", None]
    for row in range(3, rows):
        values = ["SV-%d" % row, "2026-04-%02d" % (row % 28 + 1), row * 1.5, None if row % 7 else "", "x"]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            color = colors[(row + column) % len(colors)]
            if color:
                cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=row % 3 == 0, color="DC2626" if column == 5 and row % 4 == 0 else None)
            if column == 3:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", indent=row % 2)
    sheet.cell(rows + 1, 1, "合计").fill = PatternFill("solid", fgColor="1A3C5E")
    second = book.create_sheet("summary")
    for row in range(1, 20):
        second.cell(row, 1, f"item {row}")
        second.cell(row, 2, row * 10.0).fill = PatternFill("solid", fgColor="DCFCE7" if row % 2 else "FEE2E2")
    return book


def snapshot(book):
    out = []
    for sheet in book.worksheets:
        out.append((sheet.title, sorted((k, v.height) for k, v in sheet.row_dimensions.items())))
        for row in sheet.iter_rows():
            for cell in row:
                out.append((cell.coordinate, cell.value, cell.number_format, repr(cell.font), repr(cell.fill),
                            repr(cell.border), repr(cell.alignment)))
    return out


class ThemeCacheTest(unittest.TestCase):
    def test_reused_styles_are_identical_to_styling_every_cell(self):
        fast, slow = sample_book(), sample_book()
        theme.apply(fast, "income")
        roles = theme.ROLES["income"]
        never = NeverCached()
        for index, sheet in enumerate(slow.worksheets):
            theme._style_sheet(sheet, roles, theme.STATUS_COLUMN.get("income", {}).get(index, 1), never)
        self.assertEqual(snapshot(fast), snapshot(slow))

    def test_border_range_matches_styling_every_cell(self):
        fast, slow = sample_book(), sample_book()
        styles.border_range(fast.active, 1, 40, 1, 6)
        thin = Side(style="thin", color="CCCCCC")
        for row in range(1, 41):
            for column in range(1, 7):
                slow.active.cell(row, column).border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.assertEqual(snapshot(fast), snapshot(slow))


class AppendedRowTest(unittest.TestCase):
    def test_appended_row_is_max_row(self):
        from pearnly_cowork.reconciliation.export import income_workbook

        sheet = Workbook().active
        sheet.append(["header"])
        sheet.cell(1, 5).font = Font(bold=True)
        for index in range(30):
            sheet.append([index, None, "x"])
            self.assertEqual(income_workbook._appended(sheet), sheet.max_row)
        self.assertNotIn("ws1.max_row\n\t\tfor", (ROOT / "apps/pearnly_cowork/pearnly_cowork/reconciliation/export/"
                                                    "income_workbook.py").read_text())


class ReadingSlotTest(unittest.TestCase):
    def test_cloud_run_reads_do_not_take_the_local_slot(self):
        from contextlib import nullcontext

        from pearnly_cowork.recon_center import document_ocr

        with patch.object(document_ocr.forward, "settings", return_value={"url": "x"}), \
                patch.object(document_ocr.core_pool, "local_slot") as local:
            self.assertIsInstance(document_ocr.reading_slot(), nullcontext)
        local.assert_not_called()
        with patch.object(document_ocr.forward, "settings", return_value=None), \
                patch.object(document_ocr.core_pool, "local_slot", return_value="slot") as local:
            self.assertEqual(document_ocr.reading_slot(), "slot")
        local.assert_called_once()

    def test_both_recon_readers_use_it(self):
        for path in ("recon_center/document_ocr.py", "recon_center/invoice_ocr.py"):
            source = (ROOT / "apps/pearnly_cowork/pearnly_cowork" / path).read_text()
            self.assertIn("reading_slot():", source, path)
            self.assertNotIn("with core_pool.local_slot()", source, path)


class PagesTogetherTest(unittest.TestCase):
    """对账文档各页第一遍同时抽取：发出的请求、再读给哪几页、最终结果与一页一页读完全相同。"""

    PAGES = 7

    def payload(self):
        texts = [f"page {number} text line" for number in range(1, self.PAGES + 1)]
        text, pages, start = "", [], 0
        for chunk in texts:
            text += chunk
            pages.append({"layout": {"textAnchor": {"textSegments": [{"startIndex": start, "endIndex": start + len(chunk)}]}}})
            start += len(chunk)
        return {"document": {"text": text, "pages": pages}}

    def read(self, category, together):
        import json as _json
        import threading
        import time

        from pearnly_cowork.ocr.engine.contracts import ModelResponse
        from pearnly_cowork.ocr.engine.errors import ContentError
        from pearnly_cowork.ocr.enterprise import pipeline, self_check
        from pearnly_cowork.ocr.remote.runner import page_mapper

        calls, lock = [], threading.Lock()
        # 第 2、3、5、6 页第一遍对不上（要再读）；第 4 页抽取结果不明（银行只丢这一页）；再读名额只有 2 页 —— 按页序给。
        bad, unread = {2, 3, 5, 6}, {4}

        def model_call(request):
            time.sleep(0.01 * (self.PAGES - request.page_number))  # 同时读时后面的页先回来
            reread = request.stage == pipeline.REREAD_STAGE
            with lock:
                calls.append((request.stage, request.page_number, request.prompt))
            if request.page_number in unread and not reread and category == "bank":
                raise ContentError(pipeline.UNREAD_PAGE, page_number=request.page_number)
            data = {"entries": [], "flag": int(request.page_number in bad and not reread), "page": request.page_number,
                    "sum": request.page_number}
            return ModelResponse(text=_json.dumps(data), raw={}, input_tokens=10, output_tokens=5, model="m",
                                 elapsed_ms=1)

        renders = []

        def render(number):
            renders.append(number)
            return b"img%d" % number, "image/png"

        state = lambda data, running: (self_check.MISMATCH if data.get("flag") or running.get("sum", 0) > 9 else
                                       self_check.PASSED, {})
        with patch.object(pipeline, "MAX_REREAD_PAGES", 2), \
                patch.object(pipeline, "_local", return_value=None), \
                patch.object(self_check, "bank_breaks", side_effect=lambda data: data.get("flag", 0)), \
                patch.object(self_check, "rows_short", return_value=0), \
                patch.object(self_check, "repair_single_break", return_value=None), \
                patch.object(self_check, "vat_state", side_effect=state), \
                patch.object(self_check, "vat_sums", side_effect=lambda data: {"sum": data.get("sum", 0)}), \
                patch.object(self_check, "printed_totals", return_value={}), \
                patch.object(pipeline, "check_document", return_value=[]):
            result = pipeline.run(category, page_count=self.PAGES, reads=[(b"pdf", "application/pdf", self.PAGES)],
                                  render=render, document_ai_call=lambda *a, **k: (self.payload(), 5, "k"),
                                  model_call=model_call, map_pages=page_mapper(4, lambda count: None) if together else None)
        dump = [(page.page_number, page.document.model_dump(), dict(page.extraction_audit), page.validation_warnings)
                for page in result.pages]
        return dump, calls

    def test_same_requests_rereads_and_result(self):
        for category in ("bank", "vat", "gl"):
            with self.subTest(category=category):
                serial, serial_calls = self.read(category, False)
                together, together_calls = self.read(category, True)
                self.assertEqual(together, serial)
                self.assertEqual(sorted(together_calls), sorted(serial_calls))  # 同一批请求，只是第一遍的先后不同
                from pearnly_cowork.ocr.enterprise import pipeline
                rereads = lambda calls: [(page, prompt) for stage, page, prompt in calls if stage == pipeline.REREAD_STAGE]
                self.assertEqual(rereads(together_calls), rereads(serial_calls))  # 再读的页与先后完全相同
                if category == "bank":
                    self.assertEqual([page for page, _prompt in rereads(serial_calls)], [2, 3])  # 名额按页序


if __name__ == "__main__":
    unittest.main()
