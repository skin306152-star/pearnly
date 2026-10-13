"""POS 导出的 Excel（.xlsx）读取：用 openpyxl 生成两种真实列结构（Ocha 式按单报表、Loyverse 式收据导出，数据合成），
经 `pos_import.tables.read` 读成表头与数据行，再认列、逐日汇总。要 openpyxl：在生产镜像里离线跑
（`scripts/ci-candidate-checks.sh` offline）。

另有测试机端到端用的真实版式导出（`tests/fixtures/pos/`，虚构的演示数据：标题三行、空行、第 5 行表头、一张票一行、
页脚「รวม (ไม่รวมบิลยกเลิก)」）从读表一路到推送载荷，逐分对上每天的数（2026-10-13：页脚曾被算成一张票，每天翻倍）。
"""
import io
import sys
import unittest
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "apps/pearnly_cowork"), str(ROOT / "apps/pearnly_core")]

from openpyxl import Workbook, load_workbook  # noqa: E402

from pearnly_cowork.contracts import pos as P  # noqa: E402
from pearnly_cowork.pos_import import contracts as C  # noqa: E402
from pearnly_cowork.pos_import import tables as T  # noqa: E402
from pearnly_cowork.push.express.sales_mapper import build_express_sales_payload  # noqa: E402

FIXTURES = ROOT / "tests/fixtures/pos"
QR, TRANSFER, CARD = "QR พร้อมเพย์", "โอนเงิน", "บัตรเครดิต/เดบิต"
RECEIPT_ACCOUNTS = {P.method_subject(QR): "11-02-01-01", P.method_subject(TRANSFER): "11-02-01-02",
                    P.method_subject(CARD): "11-02-01-03"}
CONFIG = {"account_set": "DATAT", "revenue_acc": "41-01-00-00", "vat_output_acc": "21-05-01-00",
          "ar_acc": "11-03-01-00", "choices": {"receipts": RECEIPT_ACCOUNTS}}
# 每天：含税合计、税额、税前额、销售票张数、作废张数、票号起止、各非现金付款方式转账（按出现先后）、现金
REAL_DAYS = {
    "SalesExport_20260928.xlsx": (date(2026, 9, 28), "13620.50", "891.06", "12729.44", 25, 0, "R6909-00837",
                                  "R6909-00861", [(QR, "3202.00"), (TRANSFER, "2484.30"), (CARD, "4088.20")],
                                  "3846.00"),
    "SalesExport_20260929.xlsx": (date(2026, 9, 29), "20132.30", "1317.06", "18815.24", 22, 0, "R6909-00862",
                                  "R6909-00883", [(TRANSFER, "7747.00"), (CARD, "5179.30"), (QR, "2819.00")],
                                  "4387.00"),
    "SalesExport_20260930.xlsx": (date(2026, 9, 30), "18939.00", "1239.00", "17700.00", 23, 0, "R6909-00884",
                                  "R6909-00906", [(TRANSFER, "1573.00"), (QR, "7244.00"), (CARD, "2956.00")],
                                  "7166.00"),
    # 有一张作废（R6909-00163，สถานะ ยกเลิก）：不算进当天，页脚也不含它
    "SalesExport_20260906.xlsx": (date(2026, 9, 6), "18577.00", "1215.32", "17361.68", 34, 1, "R6909-00147",
                                  "R6909-00181", [(CARD, "3779.00"), (QR, "8998.00"), (TRANSFER, "118.00")],
                                  "5682.00"),
}


def _book(sheets):
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


class WorkbookTests(unittest.TestCase):
    def test_ocha_style_report_with_title_rows_and_a_cover_sheet(self):
        content = _book([
            ("สรุป", [["ร้านทดสอบ"], ["ยอดขายรวม", 3959.0]]),
            ("รายบิล", [["รายงานยอดขายแยกตามบิล"], ["วันที่ 13/10/2569 - 14/10/2569"], [],
                        ["วันที่", "เวลา", "เลขที่บิล", "จำนวนลูกค้า", "ยอดก่อนส่วนลด", "ส่วนลด", "ภาษีมูลค่าเพิ่ม",
                         "ยอดสุทธิ", "ช่องทางชำระเงิน", "สถานะบิล"],
                        [datetime(2026, 10, 13), "10:01", "B0001", 2, 1070, 0, 70, 1070, "เงินสด", "ปกติ"],
                        [datetime(2026, 10, 13), "11:15", "B0002", 1, 535, 0, 35, 535, "พร้อมเพย์", "ปกติ"],
                        [datetime(2026, 10, 13), "12:00", "B0003", 1, 100, 0, 6.54, 100, "เงินสด", "ยกเลิก"],
                        [datetime(2026, 10, 14), "09:00", "B0004", 3, 2354, 0, 154, 2354, "บัตรเครดิต", "ปกติ"],
                        ["รวม", None, None, None, None, None, 265.54, 4059]]),
        ])
        headers, rows, first = T.read(content, "ocha.xlsx")
        self.assertEqual((headers[2], first), ("เลขที่บิล", 5))
        detection = C.detect(headers, rows)
        self.assertEqual(detection.questions, [])
        found, voids = C.parse(headers, rows, detection.layout, anchor_year=2026, first_row=first)
        self.assertEqual(voids, 1)
        self.assertEqual([(d.day, d.gross, d.vat, d.ref) for d in found],
                         [(date(2026, 10, 13), Decimal("1605.00"), Decimal("105.00"), "B0001-B0002"),
                          (date(2026, 10, 14), Decimal("2354.00"), Decimal("154.00"), "B0004")])

    def test_loyverse_style_receipts_with_floats_and_a_refund(self):
        content = _book([("Receipts", [
            ["Date", "Receipt number", "Receipt type", "Gross sales", "Discounts", "Net sales", "Taxes",
             "Total collected", "Payment type", "Status"],
            [datetime(2026, 10, 13, 9, 5), "1-1001", "Sale", 215.0, 5.0, 210.0, 13.74, 210.0, "Cash", "Closed"],
            [datetime(2026, 10, 13, 9, 40), "1-1002", "Sale", 65.1, 0.0, 65.1, 4.26, 65.1, "Card", "Closed"],
            [datetime(2026, 10, 13, 11, 0), "1-1003", "Refund", 30.0, 0.0, 30.0, 1.96, 30.0, "Cash", "Closed"],
        ])])
        headers, rows, first = T.read(content, "receipts.xlsx")
        detection = C.detect(headers, rows)
        self.assertEqual(detection.questions, [])
        found, _ = C.parse(headers, rows, detection.layout, anchor_year=2026, first_row=first)
        day = found[0]
        self.assertEqual((day.gross, day.vat, day.net), (Decimal("245.10"), Decimal("16.04"), Decimal("229.06")))
        self.assertEqual(day.payments, [("Cash", Decimal("180.00")), ("Card", Decimal("65.10"))])
        self.assertEqual((day.first, day.last, day.count, day.refunds), ("1-1001", "1-1002", 2, 1))

    def test_unreadable_bytes_are_reported(self):
        with self.assertRaises(C.PosError) as raised:
            T.read(b"not a workbook", "pos.xlsx")
        self.assertEqual(raised.exception.code, "pos.file_unreadable")
        with self.assertRaises(C.PosError) as raised:
            T.read(_book([("Sheet", [["a", "b"], [1, 2]])]), "pos.xlsx")
        self.assertEqual(raised.exception.code, "pos.no_header")


class RealExportTests(unittest.TestCase):
    """测试机端到端的真实版式导出：读表 → 认列 → 逐日 → 一天一张票 → 推送载荷（HS + 每种非现金付款方式一张转账凭证）。"""

    def _read(self, name, content=None):
        content = content if content is not None else (FIXTURES / name).read_bytes()
        headers, rows, first = T.read(content, name)
        detection = C.detect(headers, rows)
        self.assertEqual(detection.questions, [], name)
        return headers, rows, first, detection.layout

    def test_each_day_matches_the_pos_through_to_the_push_payload(self):
        for name, (day, gross, vat, net, count, voids, first_no, last_no, transfers, cash) in REAL_DAYS.items():
            with self.subTest(name):
                headers, rows, first, layout = self._read(name)
                self.assertEqual((first, layout["columns"]["amount"], layout["columns"]["vat"]),
                                 (6, "ยอดสุทธิ (รวม VAT)", "VAT 7%"))
                found, voided = C.parse(headers, rows, layout, anchor_year=2026, first_row=first)
                self.assertEqual(len(found), 1)
                summary = C.summary(found[0])
                self.assertEqual((summary["date"], summary["gross"], summary["vat"], summary["net"],
                                  summary["receipts"], voided), (day.isoformat(), gross, vat, net, count, voids))
                self.assertEqual(summary["range"], f"{first_no}-{last_no}")
                page = C.day_page(found[0], 1, layout_fingerprint=C.fingerprint(headers))
                fields = page["fields"]
                history = {"id": "H1", "invoice_date": fields["date"], "invoice_no": fields["invoice_number"],
                           "total_amount": fields["total_amount"], "fields": dict(fields), "pages": [page]}
                result = build_express_sales_payload(history, config=CONFIG, today=date(2026, 10, 13))
                self.assertTrue(result.ok, result.reason)
                payload = result.payload
                self.assertEqual((payload["base_amount"], payload["vat_amount"], payload["total_amount"]),
                                 (net, vat, gross))
                self.assertEqual((payload["pos"]["first"], payload["pos"]["last"], payload["pos"]["count"]),
                                 (first_no, last_no, count))
                self.assertEqual(payload["pos"]["transfers"],
                                 [{"acc": RECEIPT_ACCOUNTS[P.method_subject(label)], "amount": amount, "label": label}
                                  for label, amount in transfers])
                self.assertEqual(payload["pos"]["cash"], cash)
                moved = sum(Decimal(row["amount"]) for row in payload["pos"]["transfers"])
                self.assertEqual(moved + Decimal(payload["pos"]["cash"]), Decimal(gross))

    def test_a_real_export_whose_footer_does_not_add_up_stops(self):
        name = "SalesExport_20260928.xlsx"
        book = load_workbook(io.BytesIO((FIXTURES / name).read_bytes()))
        sheet = book["Receipts"]
        footer = sheet.max_row
        self.assertEqual(sheet.cell(footer, 5).value, "รวม (ไม่รวมบิลยกเลิก)")
        sheet.cell(footer, 9).value = 13620.51
        buffer = io.BytesIO()
        book.save(buffer)
        headers, rows, first, layout = self._read(name, buffer.getvalue())
        with self.assertRaises(C.PosError) as raised:
            C.parse(headers, rows, layout, anchor_year=2026, first_row=first)
        self.assertEqual((raised.exception.code, raised.exception.details), ("pos.total_mismatch", {"row": footer}))


if __name__ == "__main__":
    unittest.main()
