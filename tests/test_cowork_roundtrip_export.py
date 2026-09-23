"""Pure workbook preparation and real XLSX round-trip contracts; no Site or ERP IO."""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
import io
import unittest

import openpyxl

from pearnly_cowork.roundtrip import spec as rt
from pearnly_cowork.roundtrip.export import ExportError, build_review_workbook, validate_records
from pearnly_cowork.roundtrip.parser import parse_roundtrip_workbook


def record(name="H-1", direction="sales", **changes):
    fields = {"direction": direction, "date": "2026-09-21", "invoice_number": "001/69",
              "buyer_name": "ผู้ซื้อ", "buyer_tax": "0105551234567", "buyer_branch": "00000",
              "seller_name": "ผู้ขาย", "seller_tax": "0105557654321", "seller_branch": "00001",
              "subtotal": "100", "vat": "7", "total_amount": "107",
              "items": [{"description": "สินค้า", "qty": "2", "unit_price": "50", "amount": "100"}]}
    fields.update(changes)
    return {"history_id": name, "merged_fields": fields}


def workbook(records, **kwargs):
    return openpyxl.load_workbook(io.BytesIO(build_review_workbook(records, **kwargs)), data_only=False)


class ExportTests(unittest.TestCase):
    def assert_error(self, code, records, **kwargs):
        with self.assertRaises(ExportError) as caught:
            validate_records(records, **kwargs)
        self.assertEqual(caught.exception.code, code)

    def test_four_sheets_and_exact_known_formulas(self):
        wb = workbook([record(), record("H-2", "purchase"), record("H-3", "unknown")])
        self.assertEqual(wb.sheetnames, [rt.SHEET_SALES, rt.SHEET_PURCHASE, rt.SHEET_PENDING, rt.SHEET_SUMMARY])
        self.assertEqual([wb[s].max_column for s in wb.sheetnames], [23, 18, 14, 5])
        for sheet, col in ((rt.SHEET_SALES, "V"), (rt.SHEET_PURCHASE, "R"), (rt.SHEET_PENDING, "N")):
            self.assertTrue(wb[sheet].column_dimensions[col].hidden)
            self.assertEqual(wb[sheet].freeze_panes, "A2")
        self.assertEqual([wb[rt.SHEET_SALES].cell(2, c).value for c in (7, 8, 9, 10)],
                         ["=E2*F2", "=E2*F2", "=E2*F2", "=I2*0.07"])
        self.assertEqual(wb[rt.SHEET_SALES]["V2"].value, "PRN:H-1:0")
        wb.close()

    def test_all_empty_sheets_are_kept(self):
        wb = workbook([])
        self.assertEqual([s.max_row for s in wb], [1, 1, 1, 1])
        wb.close()

    def test_schema_subtotal_is_lossless_and_input_is_detached(self):
        value = record(items=[{"description": "สินค้า", "subtotal": "100"}])
        before = deepcopy(value)
        parsed = parse_roundtrip_workbook(build_review_workbook([value]))["documents"][0]["fields"]
        self.assertEqual((parsed["amount_before_vat"], parsed["vat_amount"], parsed["total_amount"]),
                         (Decimal("100.00"), Decimal("7.00"), Decimal("107.00")))
        self.assertEqual(parsed["items"][0]["amount"], Decimal("100"))
        self.assertEqual(value, before)

    def test_subtotal_zero_is_a_real_line_amount(self):
        value = record(subtotal="0", vat="0", total_amount="0", items=[{"subtotal": "0"}])
        fields = parse_roundtrip_workbook(build_review_workbook([value]))["documents"][0]["fields"]
        self.assertEqual(fields["items"][0]["amount"], Decimal("0"))
        self.assertEqual(fields["total_amount"], Decimal("0"))

    def test_subtotal_conflicts_and_missing_lines_are_rejected(self):
        self.assert_error("roundtrip.export_amount_not_roundtrippable",
                          [record(items=[{"amount": "100", "subtotal": "101"}])])
        self.assert_error("roundtrip.export_sales_items_required", [record(items=[])])
        self.assert_error("roundtrip.export_line_amount_missing", [record(items=[{"description": "unread"}])])

    def test_purchase_aliases_and_zero_require_explicit_vat(self):
        for fields in ({"subtotal": "100", "vat": "7"},
                       {"amount_before_vat": "100", "vat_amount": "7"}):
            value = record(direction="purchase", **fields)
            parsed = parse_roundtrip_workbook(build_review_workbook([value]))["documents"][0]["fields"]
            self.assertEqual(parsed["total_amount"], Decimal("107"))
        value = record(direction="purchase", subtotal="0", vat=None, total_amount="0")
        self.assert_error("roundtrip.export_amount_not_roundtrippable", [value])
        value["merged_fields"]["vat"] = "0"
        parsed = parse_roundtrip_workbook(build_review_workbook([value]))["documents"][0]["fields"]
        self.assertEqual(parsed["total_amount"], Decimal("0"))

    def test_tax_base_total_and_printed_tax_must_survive_source_reader(self):
        for changes in ({"total_amount": "100"}, {"vat": "6.97"}, {"subtotal": "99"}):
            self.assert_error("roundtrip.export_amount_not_roundtrippable", [record(**changes)])
        for tax in ("6.98", "7.02"):
            value = record(vat=tax, total_amount=str(Decimal(100) + Decimal(tax)))
            fields = parse_roundtrip_workbook(build_review_workbook([value]))["documents"][0]["fields"]
            self.assertEqual(fields["vat_amount"], Decimal(tax))

    def test_duplicate_items_are_preserved_but_double_merged_money_is_rejected(self):
        items = [{"subtotal": "100"}, {"subtotal": "200"}]
        first = record(subtotal="300", vat="21", total_amount="321", items=items)
        good = parse_roundtrip_workbook(build_review_workbook([first]))["documents"][0]["fields"]
        self.assertEqual(len(good["items"]), 2)
        # Real history shape: first.fields has all items; later fields retain
        # the original physical-page evidence. Appending it again is unsafe.
        bad = deepcopy(first)
        bad["merged_fields"]["items"].append({"subtotal": "200"})
        self.assert_error("roundtrip.export_amount_not_roundtrippable", [bad])
        same = record(subtotal="200", vat="14", total_amount="214", items=[{"subtotal": "100"}] * 2)
        self.assertEqual(len(parse_roundtrip_workbook(build_review_workbook([same]))["documents"][0]["fields"]["items"]), 2)

    def test_untrusted_push_metadata_never_becomes_success_evidence(self):
        value = record(erp_docnum="fake", erp_party_code="fake", customer_erp_action="new",
                       push_status="success", stock_acc_group="fake",
                       items=[{"amount": "100", "erp_action": "new", "erp_item_code": "fake"}])
        wb = workbook([value])
        sheet = wb[rt.SHEET_SALES]
        self.assertEqual([sheet.cell(2, c).value for c in (13, 14, 21)], ["-", "-", "-"])
        self.assertTrue(all(sheet.cell(2, c).value is None for c in (18, 19, 20, 23)))
        self.assertEqual(wb[rt.SHEET_SUMMARY].max_row, 1)
        wb.close()

    def test_verified_actions_and_latest_failure_remain_distinct(self):
        facts = {"H-1": {"actions": {"customer": False, "items": [True], "item_codes": ["IT001"],
                  "docnum": "IV69-001", "party_code": "C001", "created_party": True,
                  "direction": "sales", "stock_acc_group": "01 原材料"},
                  "state": {"status": "failed", "reason": "later attempt failed", "adapter": "express"}}}
        wb = workbook([record()], push_facts=facts)
        sheet = wb[rt.SHEET_SALES]
        self.assertEqual(sheet["M2"].value, "ใหม่")
        self.assertEqual(sheet["N2"].value, "เดิม")
        self.assertEqual(sheet["U2"].value, "ล้มเหลว · later attempt failed")
        self.assertEqual(sheet["R2"].value, "IV69-001")
        self.assertEqual(wb[rt.SHEET_SUMMARY].max_row, 3)
        self.assertEqual(wb[rt.SHEET_SUMMARY]["B2"].value, "IT001")
        wb.close()

    def test_user_text_is_explicit_string_and_leading_zeros_survive(self):
        value = record(invoice_number="=HYPERLINK(\"https://invalid.example\")", buyer_name="=1+1",
                       items=[{"description": "=WEBSERVICE(\"x\")", "qty": "2", "unit_price": "50"}])
        wb = workbook([value])
        for cell in ("B2", "C2", "D2"):
            self.assertEqual(wb[rt.SHEET_SALES][cell].data_type, "s")
        self.assertEqual(wb[rt.SHEET_SALES]["P2"].value, "00000")
        self.assertEqual(wb[rt.SHEET_SALES]["B2"].number_format, "@")
        wb.close()

    def test_dates_are_source_compatible_but_invalid_dates_do_not_vanish(self):
        for date in ("21/09/2569", "2026-09-21", "21.09.2026"):
            wb = workbook([record(date=date)])
            self.assertEqual(wb[rt.SHEET_SALES]["A2"].value, datetime(2026, 9, 21))
            wb.close()
        self.assert_error("roundtrip.export_date_invalid", [record(date="not a date")])

    def test_nonfinite_precision_and_xml_limits_are_explicit(self):
        for bad in ("NaN", "Infinity", True, []):
            self.assert_error("roundtrip.export_number_invalid", [record(total_amount=bad)])
        self.assert_error("roundtrip.export_number_precision", [record(total_amount="1234567890123456")])
        for bad in ("a\x00b", "a" * 32768):
            self.assert_error("roundtrip.export_text_invalid", [record(buyer_name=bad)])

    def test_history_keys_and_fact_scope_are_checked(self):
        for bad in ("", "bad:key", " key"):
            self.assert_error("roundtrip.export_history_invalid", [record(bad)])
        self.assert_error("roundtrip.export_history_invalid", [record(), record()])
        self.assert_error("roundtrip.export_facts_invalid", [record()], push_facts={"other": {}})
        self.assert_error("roundtrip.export_records_invalid", [record(str(i)) for i in range(101)])

    def test_pending_is_visible_review_only_and_not_promoted_to_financial_document(self):
        value = record(direction="unknown", pending_reason="review direction")
        parsed = parse_roundtrip_workbook(build_review_workbook([value]))
        self.assertEqual(parsed["documents"], [])
        self.assertEqual(parsed["pending"][0]["reason"], "review direction")
        self.assertNotIn("total_amount", parsed["pending"][0]["fields"])

    def test_decimal_half_cent_is_explicit(self):
        value = record(subtotal="0.5", vat="0.04", total_amount="0.54", items=[{"amount": "0.5"}])
        wb = workbook([value])
        self.assertEqual(wb[rt.SHEET_SALES]["J2"].value, 0.04)
        wb.close()


if __name__ == "__main__":
    unittest.main()
