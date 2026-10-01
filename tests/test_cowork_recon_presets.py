"""对账预设（用户 2026-09-30 批准）：规矩类的不确定问一次、按账户记住、以后照办。

第一类：银行「多笔对一笔」是不是汇总入账。只调纯函数或给假任务对象，不连数据库、网络或模型。
"""
import io
import json
import unittest
from datetime import date
from unittest import mock

from openpyxl import load_workbook

from pearnly_cowork.reconciliation import bank_groups, bank_matching, bank_recon_contracts, bank_recon_service
from pearnly_cowork.reconciliation import presets, recon_contracts, revision
from pearnly_cowork.reconciliation.bank_matching import BankReconRow
from pearnly_cowork.reconciliation.export import workbook


def stmt_row(day, deposit=0.0, withdrawal=0.0):
	return BankReconRow(match_status="stmt_deposit_only" if deposit else "stmt_withdrawal_only", match_layer=None,
		stmt_date=date(2026, 3, day), stmt_deposit=deposit, stmt_withdrawal=withdrawal)


def gl_row(day, debit=0.0, credit=0.0, doc="RE"):
	return BankReconRow(match_status="gl_debit_only" if debit else "gl_credit_only", match_layer=None,
		gl_date=date(2026, 3, day), gl_debit=debit, gl_credit=credit, gl_doc_no=doc)


class Scope(unittest.TestCase):
	def test_an_account_is_known_by_bank_tail_and_ledger_account(self):
		self.assertEqual(presets.bank_scope("kbank", "123-4-56789", "1112-01"), "bank|KBANK|6789|1112-01")
		self.assertEqual(presets.bank_scope("KBANK", "", "1112-01"), "bank|KBANK|-|1112-01")
		self.assertEqual(presets.bank_scope("KBANK", "4821", ""), "bank|KBANK|4821|-")

	def test_without_a_tail_or_a_ledger_account_nothing_is_remembered(self):
		# 一个工作区有多家客户公司：只凭银行记会把别家的答案套过来。
		self.assertEqual(presets.bank_scope("KBANK", "", ""), "")
		self.assertEqual(presets.bank_scope("generic", "4821", "1112-01"), "")
		self.assertEqual(presets.bank_scope("", "4821", "1112-01"), "")

	def test_the_most_common_tail_and_ledger_account_win(self):
		self.assertEqual(presets.account_tail(["xxx-x-x4821-x", "4821", "12", "", "9999"]), "4821")
		self.assertEqual(presets.ledger_account(["1112-01", "1112-01", "", "2100"]), "1112-01")

	def test_people_see_a_short_label(self):
		self.assertEqual(presets.scope_label("bank|KBANK|4821|1112-01"), "KBANK ···4821 · 1112-01")
		self.assertEqual(presets.scope_label("bank|KBANK|-|1112-01"), "KBANK · 1112-01")

	def test_only_known_answers_are_accepted(self):
		self.assertTrue(presets.valid("bank", presets.GROUPED, presets.YES))
		self.assertFalse(presets.valid("bank", presets.GROUPED, "maybe"))
		self.assertFalse(presets.valid("income", presets.GROUPED, presets.YES))


class GroupedAsMatched(unittest.TestCase):
	def test_a_unique_group_becomes_one_numbered_match(self):
		results = [stmt_row(20, deposit=600.0), gl_row(19, debit=100.0), gl_row(20, debit=200.0),
			gl_row(20, debit=300.0), gl_row(25, credit=50.0)]
		accepted, count = bank_groups.accept(results)
		self.assertEqual(count, 1)
		group = [row for row in accepted if row.group_no]
		self.assertEqual([row.group_no for row in group], [1, 1, 1, 1])
		self.assertEqual({(row.match_status, row.match_layer) for row in group}, {("matched", 4)})
		self.assertIsNotNone(group[0].stmt_date)  # 锚行（银行那一笔）排在组首
		self.assertEqual(accepted[-1].match_status, "gl_credit_only")

	def test_groups_that_do_not_add_up_uniquely_stay_one_sided(self):
		results = [stmt_row(10, deposit=300.0), gl_row(8, debit=100.0), gl_row(8, debit=200.0),
			gl_row(12, debit=150.0), gl_row(12, debit=150.0)]
		accepted, count = bank_groups.accept(results)
		self.assertEqual(count, 0)
		self.assertFalse(any(row.group_no for row in accepted))

	def test_the_formula_recounts_without_the_grouped_lines(self):
		statement = bank_recon_contracts.decode_rows(
			[{"date": "2026-03-20", "description": "BATCH", "withdrawal": "0", "deposit": "600", "balance": "1600"}],
			[{"date": "2026-03-19", "doc_no": "RE1", "account_code": "1112", "description": "", "debit": "100",
				"credit": "0", "balance": None},
			 {"date": "2026-03-20", "doc_no": "RE2", "account_code": "1112", "description": "", "debit": "500",
				"credit": "0", "balance": None}])
		header = {"stmt_opening": 1000, "stmt_closing": 1600, "gl_opening": 1000, "gl_closing": 1600,
			"bank_code": "KBANK", "gl_account_code": "1112"}
		_plain, before = bank_recon_service._match(statement, header)
		self.assertEqual((before["matched_count"], before["stmt_deposit_only_amount"]), (0, 600.0))
		results, after = bank_recon_service._match(statement, header, {presets.GROUPED: presets.YES})
		self.assertEqual(after["matched_count"], 3)
		self.assertEqual((after["stmt_deposit_only_amount"], after["gl_debit_only_amount"]), (0.0, 0.0))
		self.assertEqual(after["formula_diff"], 0.0)
		_no, declined = bank_recon_service._match(statement, header, {presets.GROUPED: presets.NO})
		self.assertEqual(declined, before)

	def test_the_group_number_is_stored_and_read_back(self):
		row = stmt_row(20, deposit=600.0)
		row.match_status, row.match_layer, row.group_no = "matched", 4, 2
		stored = bank_recon_service.row_values(row)
		self.assertEqual(stored["group_no"], 2)

		class Stored(dict):
			__getattr__ = dict.get
		back = bank_recon_service._export_row(Stored({**stored, "stmt_date": row.stmt_date}))
		self.assertEqual((back.group_no, back.match_layer), (2, 4))


class Workbook(unittest.TestCase):
	def sheets(self, lang="zh"):
		results = [stmt_row(5, deposit=100.0), gl_row(5, debit=100.0), stmt_row(20, deposit=600.0),
			gl_row(19, debit=100.0), gl_row(20, debit=500.0)]
		results = [BankReconRow(**{**vars(results[0]), "match_status": "matched", "match_layer": 1,
			"gl_date": date(2026, 3, 5), "gl_debit": 100.0, "gl_doc_no": "RV"})] + results[2:]
		accepted, _count = bank_groups.accept(results)
		content = workbook(accepted, bank_matching.BankReconSummary(), lang=lang)
		return load_workbook(io.BytesIO(content))

	def test_group_rows_follow_the_plain_matches_with_their_number(self):
		book = self.sheets()
		match = book.worksheets[1]
		layers = [match.cell(row, 2).value for row in range(2, 6)]
		self.assertEqual(layers[0], "L1-精确日期")
		self.assertEqual(layers[1:], ["L4-合计匹配（按预设） #1"] * 3)

	def test_the_ledger_sheet_skips_the_bank_side_of_a_group(self):
		book = self.sheets()
		ledger = book.worksheets[-1]
		dates = [ledger.cell(row, 1).value for row in range(2, ledger.max_row + 1) if ledger.cell(row, 1).value]
		self.assertEqual(len(dates), 3)  # 普通匹配一行 + 组里总账两行，不含银行那一笔


class FakeDoc:
	def __init__(self, **values):
		self.__dict__.update(values)

	def get(self, key, default=None):
		return self.__dict__.get(key, default)


class Questions(unittest.TestCase):
	STATEMENT = [{"date": "2026-03-20", "description": "BATCH", "withdrawal": "0", "deposit": "600", "balance": "1600"}]
	LEDGER = [{"date": "2026-03-19", "doc_no": "RE1", "account_code": "1112-01", "description": "", "debit": "100",
		"credit": "0", "balance": None},
		{"date": "2026-03-20", "doc_no": "RE2", "account_code": "1112-01", "description": "", "debit": "500",
		"credit": "0", "balance": None}]

	def doc(self):
		rows = bank_recon_contracts.decode_rows(self.STATEMENT, self.LEDGER)
		return FakeDoc(name="CBR-9", workspace="W", status="completed", review_json="{}",
			inputs_json=recon_contracts.inputs_json(statement=rows[0], gl=rows[1]),
			parse_info_json=json.dumps({"stmt_files": [{"file": "a.pdf", "ok": True, "account": "4821"}]}),
			stmt_opening=1000, stmt_closing=1600, gl_opening=1000, gl_closing=1600, bank_code="KBANK",
			gl_account_code="")

	def ask(self, stored):
		doc = self.doc()
		calls = []

		def rules(workspace, scope):
			calls.append((workspace, scope))
			return stored
		with mock.patch.object(bank_recon_service, "preset_rules", side_effect=rules), \
				mock.patch.object(revision, "_", side_effect=lambda text: text):
			given = revision.bank_inputs(doc)
			results = revision.bank_compute(doc, given)[2]
			return revision.bank_questions(doc, given, results), results, calls

	def test_an_undecided_account_is_asked_once_with_evidence(self):
		questions, results, calls = self.ask({})
		self.assertEqual(calls[0], ("W", "bank|KBANK|4821|1112-01"))
		self.assertEqual([item["rule"] for item in questions], [presets.GROUPED])
		self.assertEqual([option["value"] for option in questions[0]["options"]], [presets.YES, presets.NO])
		self.assertEqual(len(questions[0]["evidence"]), 1)
		self.assertFalse(any(row.group_no for row in results))

	def test_a_decided_account_is_not_asked_and_yes_matches_the_group(self):
		questions, results, _calls = self.ask({presets.GROUPED: presets.YES})
		self.assertEqual(questions, [])
		self.assertEqual({row.group_no for row in results}, {1})
		with mock.patch.object(revision, "_", side_effect=lambda text: text):
			self.assertEqual(revision.bank_preset_notes(results), ["reconciliation.preset.grouped_applied"])

	def test_no_is_remembered_too(self):
		questions, results, _calls = self.ask({presets.GROUPED: presets.NO})
		self.assertEqual(questions, [])
		self.assertFalse(any(row.group_no for row in results))


# ── 收入：税报里一笔都对不上的收入科目参不参与对账 ─────────────────────────────

TAX = "0105556012345"


def income_gl(number, account, credit, debit="0"):
	return {"doc_no": number, "norm_doc_no": "", "date": "2026-06-01", "account_code": account,
		"description": "", "debit": debit, "credit": credit}


def income_result(number, amount, accounts="4110"):
	import frappe

	return frappe._dict(doc_no=number, vat_date="2026-06-01", customer_name="ACME", vat_amount=amount,
		has_gl_match=1, gl_amount=amount, account_codes=accounts)


class IncomeScope(unittest.TestCase):
	def test_a_company_is_known_by_the_tax_id_on_the_report(self):
		self.assertEqual(presets.income_scope("0-1055-56012-34-5", "4201"), f"income|{TAX}|4201")
		self.assertEqual(presets.income_scope("12345", "4201"), "")
		self.assertEqual(presets.income_prefix(TAX), f"income|{TAX}|")
		self.assertEqual(presets.scope_label(f"income|{TAX}|4201"), f"{TAX} · 4201")

	def test_only_a_thirteen_digit_tax_id_is_stored(self):
		from pearnly_cowork.reconciliation import income_recon_contracts, income_recon_service

		self.assertEqual(income_recon_service._tax_field("0105556012345"), {"tax_id": TAX})
		self.assertEqual(income_recon_service._tax_field("01055560123456"), {})
		stored = income_recon_contracts.parse_info({"vat": [{"file": "r.pdf", "rows": 1, "tax_id": TAX}]})
		self.assertEqual(income_recon_service.report_tax_id(stored), TAX)


class IncomeQuestions(unittest.TestCase):
	def doc(self):
		from pearnly_cowork.reconciliation import income_recon_contracts

		gl = income_recon_contracts.decode_rows([income_gl("IV001", "4110", "1000"),
			income_gl("INT06", "4201", "35.50"), income_gl("INT06B", "4201", "12.00")], [])[0]
		return FakeDoc(name="CIR-9", workspace="W", gl_account_code="4", amount_tolerance=0.01, summary_json="{}",
			inputs_json=recon_contracts.inputs_json(gl=gl),
			parse_info_json=json.dumps({"vat": [{"file": "r.pdf", "rows": 1, "tax_id": TAX}]}))

	def run_page(self, decided):
		from pearnly_cowork.reconciliation import preset_store

		doc = self.doc()
		with mock.patch.object(revision.frappe, "get_all", return_value=[income_result("IV001", 1000)]), \
				mock.patch.object(preset_store, "under", return_value=decided), \
				mock.patch.object(revision, "_", side_effect=lambda text: text):
			inputs = revision.income_inputs(doc)
			rows, results, summary = revision.income_compute(doc, inputs)
			return (revision.income_questions(doc, rows, results), summary, revision.income_preset_notes(doc, rows),
				rows)

	def test_an_account_the_report_never_mentions_is_asked_about_once(self):
		questions, summary, notes, _rows = self.run_page({})
		self.assertEqual([(item["rule"], item["subject"]) for item in questions], [(presets.EXCLUDED, "4201")])
		self.assertEqual(len(questions[0]["evidence"]), 2)
		self.assertEqual(summary.gl_only_credit, 47.5)
		self.assertEqual(notes, [])

	def test_a_left_out_account_leaves_the_result_but_stays_in_the_inputs(self):
		questions, summary, notes, rows = self.run_page({f"income|{TAX}|4201": {presets.EXCLUDED: presets.YES}})
		self.assertEqual(questions, [])
		self.assertEqual(summary.gl_only_credit, 0.0)
		self.assertEqual(notes, ["reconciliation.preset.excluded_applied"])
		self.assertEqual(len(rows[0]), 3)  # 存回去的总账行不少：改了预设能拿回来

	def test_kept_accounts_are_not_asked_again(self):
		questions, summary, _notes, _rows = self.run_page({f"income|{TAX}|4201": {presets.EXCLUDED: presets.NO}})
		self.assertEqual(questions, [])
		self.assertEqual(summary.gl_only_credit, 47.5)

	def test_nothing_is_asked_without_a_tax_id(self):
		doc = self.doc()
		doc.parse_info_json = "{}"
		self.assertEqual(revision.income_questions(doc, ([income_gl("A", "4201", "1")], []), []), [])


# ── 银行：这个账户对总账哪个科目（答了用原文件重跑） ─────────────────────────────


class LedgerAccount(unittest.TestCase):
	def statement(self):
		from pearnly_cowork.reconciliation.bank_statement import StatementRow

		rows = [StatementRow(date=date(2026, 5, 3), description="", withdrawal=100.0, deposit=0.0, balance=900.0),
			StatementRow(date=date(2026, 5, 5), description="", withdrawal=0.0, deposit=300.0, balance=1200.0)]
		return {"ok": True, "rows": rows, "row_count": len(rows), "opening": 1000.0, "closing": 1200.0, "bank_code": "kbank",
			"source": "workbook", "parser_version": "x", "balance_warn_count": 0, "amount_fixed_count": 0,
			"low_conf_count": 0, "completeness": {"ok": True}, "needs_review": False,
			"identity": {"bank": "KBANK", "last4": "4821"}}

	def ledger(self, only=None):
		def gl(code, day, debit=0.0, credit=0.0):
			return bank_matching.GlRow(date=date(2026, 5, day), doc_no=f"{code}-{day}", account_code=code,
				description="", debit=debit, credit=credit)
		rows = [gl("1112-01", 3, credit=100.0), gl("1112-01", 20, debit=5.0),
			gl("1112-02", 5, debit=300.0), gl("1112-02", 21, credit=7.0)]
		rows = [row for row in rows if not only or row.account_code == only]
		return {"ok": True, "rows": rows, "row_count": len(rows), "accounts": sorted({row.account_code for row in rows}),
			"opening": 1000.0, "closing": 1200.0, "source": "workbook", "needs_review": False}

	def create(self, preset):
		reads = []

		def read_ledger(content, name, account_code="", model_rows=None):
			reads.append(account_code)
			return self.ledger(account_code or None)
		with mock.patch.object(bank_recon_service.bank_recon_access, "require"), \
				mock.patch.object(bank_recon_service.bank_intake, "read_statement_file", return_value=self.statement()), \
				mock.patch.object(bank_recon_service.bank_intake, "read_ledger_file", side_effect=read_ledger), \
				mock.patch.object(bank_recon_service, "preset_rules", return_value=preset), \
				mock.patch.object(bank_recon_service, "_", side_effect=lambda text, **_kw: text + " {0}"), \
				mock.patch.object(bank_recon_service, "create", return_value={"name": "T1"}) as create:
			result = bank_recon_service.create_from_documents("ws", "t", "r", b"", "", b"", "",
				statement_files=[(b"1", "kbank.xlsx")], gl_files=[(b"g", "gl.xlsx")])
		return result, create.call_args, reads

	def test_two_equally_likely_accounts_are_asked_about_with_the_files_kept(self):
		result, call, _reads = self.create({})
		self.assertEqual([item["rule"] for item in result["questions"]], [presets.GL_ACCOUNT])
		question = result["questions"][0]
		self.assertEqual(question["scope"], "bank|KBANK|4821|-")
		self.assertEqual({item["code"] for item in question["candidates"]}, {"1112-01", "1112-02"})
		self.assertEqual(call.args[5]["gl_account_code"], "")  # 照旧全部科目一起对

	def test_the_remembered_account_is_used_without_asking(self):
		result, call, reads = self.create({presets.GL_ACCOUNT: "1112-02"})
		self.assertEqual(result["questions"], [])
		self.assertEqual(call.args[5]["gl_account_code"], "1112-02")
		self.assertEqual(reads[-1], "1112-02")  # 按这个科目重读（期初按它的科目标题行）
		self.assertIn("reconciliation.preset.gl_account_applied 1112-02", call.kwargs["warnings"])

	def test_a_remembered_account_missing_from_this_ledger_is_not_forced(self):
		result, call, _reads = self.create({presets.GL_ACCOUNT: "9999"})
		self.assertEqual([item["rule"] for item in result["questions"]], [presets.GL_ACCOUNT])
		self.assertEqual(call.args[5]["gl_account_code"], "")

	def test_free_answers_are_account_codes(self):
		self.assertTrue(presets.valid("bank", presets.GL_ACCOUNT, "1112-01"))
		self.assertFalse(presets.valid("bank", presets.GL_ACCOUNT, ""))
		self.assertFalse(presets.valid("bank", presets.GL_ACCOUNT, "x" * 41))
		self.assertEqual(presets.account_scope("KBANK", ""), "")


class Rerun(unittest.TestCase):
	def test_a_rerun_reuses_what_the_first_run_recognised(self):
		from pearnly_cowork.recon_center import document_ocr

		db = mock.MagicMock()
		with mock.patch.object(document_ocr.frappe, "db", db), mock.patch.object(document_ocr, "submit") as submit, \
				mock.patch.object(document_ocr, "wait", return_value={"pages": []}) as wait:
			db.exists.return_value = True
			self.assertEqual(document_ocr.reader("job-1", "left0", "F1", "other@example.com")(
				"bank_statement", b"x", "a.jpg"), {"pages": []})
		submit.assert_not_called()
		# 按原对账号的参考号找到那一份（名字由类别 + 参考号算出，与共用层建任务时同一个算法）
		work = document_ocr.core_work.name_for(document_ocr.KIND, "job-1:left0:bank_statement")
		wait.assert_called_once_with(work)
		self.assertEqual(db.exists.call_args.args[1], work)

	def test_the_job_keeps_the_files_for_a_question_and_drops_them_after_the_rerun(self):
		from pearnly_cowork.recon_center import jobs

		def run(options, result):
			db = mock.MagicMock()
			db.exists.return_value = True
			with mock.patch.object(jobs, "_state"), mock.patch.object(jobs, "_read", return_value=[("a.xlsx", b"a")]), \
					mock.patch.object(jobs.frappe, "session", mock.Mock(user="u@example.com")), \
					mock.patch.object(jobs.frappe, "db", db), mock.patch.object(jobs.frappe, "delete_doc") as delete, \
					mock.patch.object(jobs.rerun, "keep") as keep, mock.patch.object(jobs.rerun, "drop") as drop, \
					mock.patch.object(jobs.bank_recon_service, "create_from_documents", return_value=result) as create:
				jobs.run("job-2", "bank", "ws", "t", {"left": ["L0"], "right": ["R0"]}, json.dumps(options))
			return keep, drop, delete, create

		keep, drop, delete, create = run({}, {"name": "T1", "questions": [{"rule": "gl_account"}]})
		keep.assert_called_once()
		self.assertEqual(keep.call_args.args[3], "job-2")  # 识别结果以后按这个对账号取
		delete.assert_not_called()
		drop.assert_not_called()

		keep, drop, delete, create = run({"ocr_job": "job-1", "replace_task": "T1"}, {"name": "T2", "questions": []})
		keep.assert_not_called()
		drop.assert_called_once_with("T1", files=False)
		self.assertEqual(create.call_args.kwargs["replace_task"], "T1")
		self.assertEqual(delete.call_count, 2)  # 重跑完原文件删掉


if __name__ == "__main__":
	unittest.main()
