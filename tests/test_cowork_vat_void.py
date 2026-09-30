"""销项税报告核查的「作废一致」（Cowork 有意改动，用户 2026-09-30 定）：税报作废行（金额全 0、标「*」或写着作废）与同一张
发票配成作废一致；单号只差开头的字母（识别掉了「I」）且两边候选都唯一也配；没有发票的作废行不算缺失；普通行照旧。"""
import io
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps" / "pearnly_cowork"))

try:
    import openpyxl  # noqa: F401
except ImportError:  # pragma: no cover - the bench environment has it
    openpyxl = None

from pearnly_cowork.reconciliation import vat_check_core  # noqa: E402

TEXTS = {"vat_check.status.void_match": "作废一致", "vat_check.status.void_no_invoice": "作废（无需发票）"}


def report_row(no, buyer="ACME CO., LTD.", pre=1000.0, vat=70.0, tax="0105551234567"):
    return {"report_invoice_no": no, "report_date": "2026-06-05", "report_buyer_name": buyer, "report_buyer_tax_id": tax,
            "report_buyer_branch": "00000", "report_amount_pre_vat": pre, "report_vat_amount": vat,
            "report_amount": round(pre + vat, 2) if pre is not None else None}


def void_row(no, buyer="*ยกเลิก*"):
    return report_row(no, buyer, 0.0, 0.0, "")


def invoice(no, total=1070.0, buyer="ACME CO., LTD."):
    return {"invoice_no": no, "invoice_date": "2026-06-05", "buyer_name": buyer, "buyer_tax_id": "0105551234567",
            "buyer_branch": "", "amount_pre_vat": round(total / 1.07, 2), "vat_amount": round(total - total / 1.07, 2),
            "total_amount": total, "filename": f"{no}.jpg"}


def kinds(result):
    return sorted(pair["kind"] for pair in result["pairs"])


class VoidPairTest(unittest.TestCase):
    def test_void_row_with_the_same_number_is_a_void_match(self):
        result = vat_check_core._build_recon_pairs([invoice("IV69/00581")], [void_row("IV69/00581")], "zh")
        self.assertEqual(kinds(result), ["void"])
        self.assertEqual((result["unmatched_inv"], result["unmatched_rep"], result["void_rep"]), ([], [], []))

    def test_markers_star_cancel_and_blank_amounts(self):
        for buyer in ("*", "ยกเลิก", "CANCELLED", "Void", "作废", "ACME *"):
            row = {**void_row("IV1", buyer), "report_amount_pre_vat": None, "report_vat_amount": "", "report_amount": None}
            result = vat_check_core._build_recon_pairs([invoice("IV1")], [row], "zh")
            self.assertEqual(kinds(result), ["void"], buyer)
        # 「*」标在单号上：规范化后单号不同，按「只差开头」照样配
        result = vat_check_core._build_recon_pairs([invoice("IV69/00581")], [void_row("*IV69/00581", "ACME")], "zh")
        self.assertEqual(kinds(result), ["void"])

    def test_number_that_lost_its_leading_letter_is_a_void_match_when_unique(self):
        for read in ("V69/00581", "V6900581", "V69.00581", "69/00581"):
            result = vat_check_core._build_recon_pairs(
                [invoice(read), invoice("IV69/00590")], [void_row("IV69/00581"), report_row("IV69/00590")], "zh")
            self.assertEqual(kinds(result), ["matched", "void"], read)
            void = next(pair for pair in result["pairs"] if pair["kind"] == "void")
            self.assertEqual((void["inv_idx"], void["rep_idx"]), (0, 0))

    def test_leading_letter_candidates_that_are_not_unique_are_not_paired(self):
        # 两张发票都像这一行作废行
        result = vat_check_core._build_recon_pairs([invoice("V69/00581"), invoice("69/00581")],
                                                   [void_row("IV69/00581")], "zh")
        self.assertNotIn("void", kinds(result))
        self.assertEqual(sorted(result["unmatched_inv"]), [0, 1])
        self.assertEqual(result["void_rep"], [0])  # 作废行本身不算缺失
        # 一张发票像两行作废行
        result = vat_check_core._build_recon_pairs([invoice("V69/00581")],
                                                   [void_row("IV69/00581"), void_row("AV69/00581")], "zh")
        self.assertEqual((result["pairs"], result["unmatched_inv"]), ([], [0]))
        # 换了开头字母（IV ↔ CN）不是「掉了字母」
        result = vat_check_core._build_recon_pairs([invoice("CN69/00581")], [void_row("IV69/00581")], "zh")
        self.assertEqual((result["pairs"], result["unmatched_inv"], result["void_rep"]), ([], [0], [0]))

    def test_zero_row_counts_as_void_when_only_the_invoice_is_marked(self):
        row = report_row("IV7", "ACME CO., LTD.", 0.0, 0.0)
        result = vat_check_core._build_recon_pairs([invoice("IV7", 0.0, "ยกเลิก")], [row], "zh")
        self.assertEqual(kinds(result), ["void"])
        # 两边都没有作废标记：照旧按普通配对
        result = vat_check_core._build_recon_pairs([invoice("IV7")], [row], "zh")
        self.assertEqual(kinds(result), ["matched"])
        # 没有标记的 0 金额行不会按「只差开头字母」去配
        result = vat_check_core._build_recon_pairs([invoice("V7")], [row], "zh")
        self.assertEqual((result["pairs"], result["unmatched_rep"], result["void_rep"]), ([], [0], []))

    def test_same_number_prefers_the_regular_row_over_the_void_one(self):
        result = vat_check_core._build_recon_pairs([invoice("IV9")], [void_row("IV9"), report_row("IV9")], "zh")
        self.assertEqual([(pair["kind"], pair["rep_idx"]) for pair in result["pairs"]], [("matched", 1)])
        self.assertEqual((result["unmatched_rep"], result["void_rep"]), ([], [0]))

    def test_marked_row_with_amounts_is_not_void(self):
        result = vat_check_core._build_recon_pairs([invoice("IV8", 2000.0)], [report_row("IV8", "*ACME")], "zh")
        self.assertEqual(kinds(result), ["matched"])


@unittest.skipIf(openpyxl is None, "openpyxl is not installed")
class VoidWorkbookTest(unittest.TestCase):
    def build(self, invoices, rows):
        from pearnly_cowork.reconciliation.export import vat_check_sheet3
        from pearnly_cowork.reconciliation.export.vat_check_workbook import build_excel

        with mock.patch.object(vat_check_sheet3, "_", side_effect=lambda key, lang=None: TEXTS[key]):
            return build_excel(invoices, rows, "Seller", "", 2026, 6, "zh")

    def test_void_rows_are_not_differences_or_missing(self):
        from pearnly_cowork.reconciliation import vat_check_service

        invoices = [invoice("V69/00581", 5350.0), invoice("IV69/00590"), invoice("IV69/00591", 2140.0)]
        rows = [void_row("IV69/00581"), report_row("IV69/00590"), report_row("IV69/00591"), void_row("IV69/00599")]
        content, summary = self.build(invoices, rows)
        by_status = {row["status_key"]: row for row in summary["rows"]}
        self.assertEqual(sorted(by_status), ["st_diff", "st_ok", "st_void", "st_void_no_inv"])
        self.assertEqual((summary["n_total"], summary["n_ok"], summary["n_void"], summary["n_diff"]), (4, 1, 2, 1))
        # 普通行的金额差照常是差异；作废行不进差异金额
        self.assertEqual(summary["diff_amount_total"], 1070.0)
        self.assertEqual(vat_check_service._counts(summary),
                         {"matched_count": 1, "diff_count": 1, "unmatched_count": 0, "diff_amount": 1070.0})
        sheet = openpyxl.load_workbook(io.BytesIO(content))["对账结果"]
        lines = {sheet.cell(r, 4).value: [sheet.cell(r, c).value for c in range(1, 16)] for r in range(2, sheet.max_row + 1)}
        self.assertEqual(lines["V69/00581"][1], "✓ 作废一致")
        # 金额（发 / 报 / 差），6 个维度差异列都空（空串写进工作簿读回来是 None）
        self.assertEqual(lines["V69/00581"][5:14], [5350.0, 0, 0] + [None] * 6)
        self.assertEqual(lines["IV69/00599"][1], "✓ 作废（无需发票）")
        self.assertEqual(lines["IV69/00591"][1], "⚠ 金额不一致")
        # 作废两行与「一致」同为绿底
        greens = {sheet.cell(r, 4).value for r in range(2, sheet.max_row + 1)
                  if (sheet.cell(r, 2).fill.fgColor.rgb or "")[-6:] == "DCFCE7"}
        self.assertEqual(greens, {"V69/00581", "IV69/00599", "IV69/00590"})

    def test_workbook_without_void_rows_needs_no_translation(self):
        from pearnly_cowork.reconciliation.export.vat_check_workbook import build_excel

        _content, summary = build_excel([invoice("IV1")], [report_row("IV1")], "Seller", "", 2026, 6, "zh")
        self.assertEqual((summary["n_ok"], summary["n_void"], summary["n_diff"]), (1, 0, 0))


if __name__ == "__main__":
    unittest.main()
