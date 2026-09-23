"""Memory-only synthetic table decoding; no Site, network, or external tools."""
import hashlib
import io
import json
import struct
import unittest
from unittest.mock import patch
import zipfile

import docx
import openpyxl

from pearnly_cowork.ocr.table_inputs import TableInputError, extract_table
from pearnly_cowork.ocr.table_inputs.archive import check_archive


def workbook_bytes(sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def document_bytes(paragraphs=(), tables=()):
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for rows in tables:
        table = document.add_table(rows=0, cols=max(len(row) for row in rows))
        for cells in rows:
            row = table.add_row()
            for cell, text in zip(row.cells, cells):
                cell.text = text
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def zipped(parts):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    return output.getvalue()


class TableInputTest(unittest.TestCase):
    def failure(self, code, content, filename):
        with self.assertRaises(TableInputError) as caught:
            extract_table(content, filename)
        self.assertEqual(code, caught.exception.code)

    def test_excel_logical_units_billing_and_original_rows(self):
        content = workbook_bytes([(" ขาย ", [[], [" Item ", "Amount"], ["A", 0], [], ["B", 2.675]]),
                                  ("Empty", [])])
        result = extract_table(content, "book.xlsx")
        self.assertEqual(2, result.page_count)
        self.assertEqual("table_path_excel", result.engine)
        self.assertEqual("logical_section", result.provenance["reference_kind"])
        self.assertEqual(["Item", "Amount"], result.pages[0].table_headers)
        self.assertEqual([{"Item": "A", "Amount": "0"}, {"Item": "B", "Amount": "2.675"}],
                         result.pages[0].table_rows)
        # Raw cells: len(' Item ')+len('Amount')+len('A')+len('0')+len('B')+len('2.675').
        self.assertEqual(20, result.billing_characters)
        self.assertEqual(2, result.provenance["units"][0]["header_source_row"])
        self.assertEqual([1, 2, 3, 4, 5], [row["source_row"] for row in result.provenance["units"][0]["rows"]])
        self.assertEqual([None, None, 1, None, 2],
                         [row["table_row_index"] for row in result.provenance["units"][0]["rows"]])
        self.assertEqual(hashlib.sha256(content).hexdigest(), result.source_sha256)
        self.assertTrue(result.coverage_complete)
        json.dumps(result.provenance)

    def test_excel_formula_row_one_is_separate_from_cached_model_headers(self):
        content = workbook_bytes([("ขาย", [['="header"', "Amount"], ["A", 10]])])
        result = extract_table(content, "formula.xlsm")
        self.assertEqual(['="header"', "Amount"], result.provenance["units"][0]["roundtrip_headers"])
        self.assertEqual(["", "Amount"], result.pages[0].table_headers)
        self.assertFalse(result.coverage_complete)
        self.assertEqual(9, result.billing_characters)

    def test_xls_retains_openpyxl_dispatch_does_not_introduce_biff_reader(self):
        content = workbook_bytes([("one", [["Name"], ["Item"]])])
        self.assertEqual(extract_table(content, "x.xlsx").pages, extract_table(content, "x.xls").pages)
        self.failure("ocr.table_invalid_archive", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 512, "binary.xls")

    def test_csv_physical_lines_include_multiline_and_blank_records(self):
        content = b'Name,Amount\n"first\nsecond",7\n\nthird,8\n'
        result = extract_table(content, "x.csv")
        rows = result.provenance["units"][0]["rows"]
        self.assertEqual([(1, 1), (2, 3), (4, 4), (5, 5)],
                         [(row["line_start"], row["line_end"]) for row in rows])
        self.assertEqual([None, 1, None, 2], [row["table_row_index"] for row in rows])
        self.assertEqual("first\nsecond", result.pages[0].table_rows[0]["Name"])
        self.assertTrue(result.coverage_complete)

    def test_encoding_and_character_billing_are_deliberately_different(self):
        content = "ชื่อ,ราคา\nกาแฟ,20".encode("cp874")
        result = extract_table(content, "thai.csv")
        self.assertEqual("cp874", result.provenance["units"][0]["encoding"])
        self.assertEqual("กาแฟ", result.pages[0].table_rows[0]["ชื่อ"])
        self.assertEqual(len(content.decode("utf-8", errors="ignore")), result.billing_characters)
        bom = extract_table(b"\xef\xbb\xbfhello", "bom.txt")
        self.assertEqual("hello", bom.pages[0].full_text)
        self.assertEqual(6, bom.billing_characters)
        self.assertEqual(0, extract_table(b"\xff", "zero.txt").billing_characters)

    def test_collisions_extra_columns_and_unnamed_values_keep_original_cells(self):
        result = extract_table(b"A,A,\none,two,three,four\n", "x.csv")
        self.assertEqual([{"A": "two", "": "three"}], result.pages[0].table_rows)
        self.assertEqual(["one", "two", "three", "four"], result.provenance["units"][0]["rows"][1]["cells"])
        self.assertEqual({"duplicate_headers", "columns_outside_headers", "unnamed_column_value"},
                         {issue["code"] for issue in result.provenance["coverage_issues"]})
        self.assertFalse(result.coverage_complete)

    def test_docx_primary_grid_and_uncovered_evidence(self):
        content = document_bytes(["前言", ""], [[["Item", "Amount"], ["Tea", "20"]], [["Second"], ["lost"]]])
        result = extract_table(content, "x.docx")
        self.assertEqual(1, result.page_count)
        self.assertEqual(2, result.billing_characters)
        self.assertEqual([{"Item": "Tea", "Amount": "20"}], result.pages[0].table_rows)
        self.assertIn("lost", result.pages[0].full_text)
        self.assertFalse(result.coverage_complete)
        self.assertEqual(2, len(result.provenance["units"][0]["tables"]))
        self.assertEqual({"docx_paragraphs_outside_primary_grid", "docx_tables_outside_primary_grid"},
                         {issue["code"] for issue in result.provenance["coverage_issues"]})

    def test_docx_no_grid_body_uses_all_text_and_table_only_billing_can_be_zero(self):
        result = extract_table(document_bytes(["text"], [[["Header"]], [["Other"], ["value"]]]), "x.docx")
        self.assertIsNone(result.pages[0].table_rows)
        self.assertTrue(result.coverage_complete)
        zero = extract_table(document_bytes([], [[["Name"], ["Tea"]]]), "zero.docx")
        self.assertEqual(0, zero.billing_characters)
        self.assertTrue(zero.coverage_complete)

    def test_legacy_doc_is_decoded_text_and_byte_billing_not_native_word_claim(self):
        content = b"Invoice\x00\r\nLine\t20\x01"
        result = extract_table(content, "x.doc")
        self.assertEqual("Invoice\nLine\t20", result.pages[0].full_text)
        self.assertEqual(.7, result.pages[0].avg_confidence)
        self.assertEqual(len(content) // 2, result.billing_characters)
        self.assertEqual("decoded_printable_text", result.provenance["units"][0]["line_basis"])
        self.assertFalse(result.coverage_complete)
        self.assertEqual("legacy_doc_best_effort", result.provenance["coverage_issues"][0]["code"])

    def test_docx_header_and_unread_textbox_cannot_claim_complete_coverage(self):
        from docx.oxml import OxmlElement
        document = docx.Document()
        document.add_paragraph("Body")
        document.sections[0].header.paragraphs[0].text = "Header-only seller tax ID"
        document.element.body.append(OxmlElement("w:txbxContent"))
        output = io.BytesIO()
        document.save(output)
        result = extract_table(output.getvalue(), "x.docx")
        self.assertEqual("Body", result.pages[0].full_text)
        self.assertFalse(result.coverage_complete)
        self.assertEqual({"docx_non_body_text", "docx_unread_body_structure"},
                         {issue["code"] for issue in result.provenance["coverage_issues"]})

    def test_formula_cache_missing_preserves_empty_value_and_original_coordinate(self):
        result = extract_table(workbook_bytes([("Sheet", [["Amount"], ["=1+2"]])]), "x.xlsx")
        self.assertEqual([], result.pages[0].table_rows)
        self.assertEqual([""], result.provenance["units"][0]["rows"][1]["cells"])
        self.assertFalse(result.coverage_complete)
        self.assertEqual({"code": "formula_cache_missing", "page_number": 1, "cell": "A2",
                          "source_row": 2, "source_column": 1}, result.provenance["coverage_issues"][0])

    def test_formula_object_does_not_put_memory_address_in_durable_manifest(self):
        from openpyxl.worksheet.formula import ArrayFormula
        content = workbook_bytes([("Sheet", [[ArrayFormula(ref="A1", text="=1+2")]])])
        first, second = extract_table(content, "x.xlsx"), extract_table(content, "x.xlsx")
        self.assertEqual(first.provenance, second.provenance)
        self.assertEqual(["[ArrayFormula]"], first.provenance["units"][0]["roundtrip_headers"])

    def test_errors_return_no_partial_document(self):
        self.failure("ocr.empty_file", b"", "x.csv")
        self.failure("ocr.table_invalid_bytes", "x", "x.txt")
        self.failure("ocr.table_format_not_supported", b"x", "x.png")
        self.failure("ocr.table_invalid_archive", b"not zip", "x.docx")
        self.failure("ocr.table_parse_failed", b"a\nb\rc", "x.csv")

    def test_archive_member_and_expanded_size_limits(self):
        with patch("pearnly_cowork.ocr.table_inputs.archive.MAX_ARCHIVE_MEMBERS", 1):
            self.failure("ocr.table_archive_member_limit", zipped({"a": "x", "b": "y"}), "x.xlsx")
        with patch("pearnly_cowork.ocr.table_inputs.archive.MAX_EXPANDED_BYTES", 100):
            self.failure("ocr.table_archive_size_limit", zipped({"padding": "x" * 101}), "x.xlsx")

    def test_xml_entity_declarations_are_rejected_in_utf8_and_utf16(self):
        text = '<!DOCTYPE x [<!ENTITY x "expansion">]><x>&x;</x>'
        for encoding in ("utf-8", "utf-16"):
            with self.subTest(encoding=encoding):
                self.failure("ocr.table_xml_declaration_not_allowed", zipped({"word/document.xml": text.encode(encoding)}), "x.docx")

    def test_sheet_dimension_and_actual_coordinate_limits_before_openpyxl(self):
        for xml, code in [
            ('<worksheet><dimension ref="A1:A100001"/></worksheet>', "ocr.table_row_limit"),
            ('<worksheet><dimension ref="A1:IW1"/></worksheet>', "ocr.table_column_limit"),
            ('<worksheet><dimension ref="A1:IV2000"/></worksheet>', "ocr.table_cell_limit"),
            ('<worksheet><dimension ref="A1:A1"/><sheetData><row r="1"><c r="B1"/></row></sheetData></worksheet>', "ocr.table_invalid_dimensions"),
            ('<worksheet><sheetData><row r="1"><c r="A1"/><c r="A1"/></row></sheetData></worksheet>', "ocr.table_invalid_dimensions"),
            ('<worksheet><sheetData><row r="1"/><row r="1"/></sheetData></worksheet>', "ocr.table_invalid_dimensions"),
            ('<worksheet><sheetData><row r="100001"/></sheetData></worksheet>', "ocr.table_row_limit")]:
            with self.subTest(code=code), patch("openpyxl.load_workbook", side_effect=AssertionError("must not load")):
                self.failure(code, zipped({"xl/worksheets/sheet1.xml": xml}), "x.xlsx")

    def test_delimited_dimensions_and_text_limits_reject_without_truncation(self):
        self.failure("ocr.table_column_limit", ("," * 256).encode(), "x.csv")
        with patch("pearnly_cowork.ocr.table_inputs.model.MAX_CELLS", 3):
            self.failure("ocr.table_cell_limit", b"A,B\none,two", "x.csv")
        with patch("pearnly_cowork.ocr.table_inputs.reader.MAX_TEXT_CHARACTERS", 3):
            self.failure("ocr.table_text_limit", b"1234", "x.csv")

    def test_custom_worksheet_relationship_cannot_bypass_preflight(self):
        base = workbook_bytes([("Sheet", [["Name"], ["Tea"]])])
        for target in ("custom.xml", "custom_without_extension"):
            with self.subTest(target=target), zipfile.ZipFile(io.BytesIO(base)) as original:
                parts = {name: original.read(name) for name in original.namelist()}
                parts["xl/" + target] = parts.pop("xl/worksheets/sheet1.xml")
                for name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                    parts[name] = parts[name].replace(b"/xl/worksheets/sheet1.xml", ("/xl/" + target).encode())
                valid = zipped(parts)
                self.assertEqual([{"Name": "Tea"}], extract_table(valid, "x.xlsx").pages[0].table_rows)
                parts["xl/" + target] = parts["xl/" + target].replace(b'A1:A2', b'A1:ZZZ1000000')
                with patch("openpyxl.load_workbook", side_effect=AssertionError("must not load")):
                    self.failure("ocr.table_row_limit", zipped(parts), "x.xlsx")

    def test_excel_comments_and_header_content_mark_unread_structure(self):
        from openpyxl.comments import Comment
        workbook = openpyxl.Workbook()
        workbook.active.append(["Name", "Amount"])
        workbook.active.append(["Tea", 20])
        workbook.active["A2"].comment = Comment("extra information", "synthetic")
        workbook.active.oddHeader.center.text = "Seller tax id in header"
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        result = extract_table(output.getvalue(), "x.xlsx")
        self.assertEqual([{"Name": "Tea", "Amount": "20"}], result.pages[0].table_rows)
        self.assertFalse(result.coverage_complete)
        features = {issue["feature"] for issue in result.provenance["coverage_issues"]
                    if issue["code"] == "worksheet_unread_structure"}
        self.assertEqual({"headerFooter", "legacyDrawing"}, features)

    def test_corrupt_archive_crc_is_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_STORED) as archive:
            archive.writestr("payload", b"abcdef")
        content = bytearray(output.getvalue())
        name_length, extra_length = struct.unpack_from("<HH", content, 26)
        content[30 + name_length + extra_length] ^= 1
        with self.assertRaises(TableInputError) as caught:
            check_archive(bytes(content))
        self.assertEqual("ocr.table_invalid_archive", caught.exception.code)


if __name__ == "__main__":
    unittest.main()
