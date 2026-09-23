"""Offline XLSX checks; authorized snapshots are supplied by the caller."""
from copy import deepcopy
from datetime import UTC, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
import unittest
from zipfile import ZipFile

import openpyxl

from pearnly_cowork.stocktake.contracts import StocktakeValidationError
from pearnly_cowork.stocktake.export import workbook


def item(name="I1", **values):
    return {"id": name, "product_code": "0012", "product_name": "สินค้า", "barcode": "0000012",
            "unit": "EA", "warehouse": "A", "location": "01", "book_qty": "2.000000",
            "actual_qty": "2.000000", **values}


def entry(row, **values):
    return {"id": "E1", "item_id": row["id"],
            **{key: row[key] for key in ("product_code", "product_name", "barcode", "unit", "warehouse", "location")},
            "quantity": "2.000000", "voided": False, "counted_by_name": "ผู้ตรวจ",
            "updated_by_name": "ผู้แก้ไข", "counted_at": datetime(2026, 9, 21, 1, 2, 3, tzinfo=UTC),
            "updated_at": datetime(2026, 9, 21, 2, 3, 4, tzinfo=UTC), **values}


def snapshot(items=None, entries=None, **values):
    items = [item()] if items is None else items
    entries = [entry(items[0])] if entries is None else entries
    return {"count_mode": "scan", "items": items, "entries": entries, "entry_total": len(entries), **values}


class StocktakeExport(unittest.TestCase):
    def load(self, value):
        result = openpyxl.load_workbook(BytesIO(workbook(value)), data_only=False)
        self.addCleanup(result.close)
        return result

    def test_thai_layout_filters_widths_and_blank_photo_columns(self):
        book = self.load(snapshot())
        self.assertEqual(book.sheetnames, ["สรุปผลตรวจนับ", "รายละเอียดการนับ"])
        summary, detail = book.worksheets
        self.assertEqual((summary.max_column, detail.max_column), (15, 17))
        self.assertEqual((summary["E1"].value, detail["G1"].value), ("จำนวนตามบัญชี", "จำนวนที่นับครั้งนี้"))
        for sheet, ref in ((summary, "A1:O2"), (detail, "A1:Q2")):
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, ref)
            self.assertEqual((sheet.column_dimensions["A"].width, sheet.column_dimensions["B"].width), (25, 36))
            self.assertEqual(sheet.row_dimensions[1].height, 26)
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00EAF0FF")
            self.assertEqual(sheet["A2"].border.left.style, "thin")
            self.assertTrue(sheet["A2"].alignment.wrap_text)
            self.assertIsNone(sheet.cell(2, sheet.max_column).value)

    def test_zero_is_counted_and_unseen_is_blank(self):
        zero, unseen = item("zero", actual_qty="0.000000"), item("unseen", actual_qty=None)
        sheet = self.load(snapshot([zero, unseen], [entry(zero, quantity="0.000000")])).worksheets[0]
        self.assertEqual((sheet["F2"].value, sheet["G2"].value, sheet["H2"].value), (0, -2, "ขาดบัญชี"))
        self.assertEqual((sheet["F3"].value, sheet["G3"].value, sheet["H3"].value), (None, None, "ยังไม่ได้นับ"))

    def test_persisted_total_is_used_and_voided_entries_remain_audit_rows(self):
        row = item(actual_qty="7.000000")
        active = entry(row, warehouse="B", location="02")
        void = entry(row, id="E2", quantity="999.000000", warehouse="Z", location="99", voided=True)
        book = self.load(snapshot([row], [void, active], summary=[{"actual_qty": "1000"}]))
        summary, detail = book.worksheets
        self.assertEqual((summary["F2"].value, summary["G2"].value), (7, 5))
        self.assertEqual((summary["J2"].value, summary["M2"].value), ("B", "02"))
        self.assertEqual(detail["J2"].value, "ยกเลิก (ไม่รวมยอด)")
        self.assertEqual(detail["G2"].value, 999)
        self.assertEqual(detail["D2"].fill.patternType, None)
        self.assertEqual(detail["D3"].fill.fgColor.rgb, "00FFF2CC")
        self.assertEqual(summary["E2"].fill.fgColor.rgb, "00FFF2CC")

    def test_locations_are_sorted_distinct_and_unspecified_keeps_source_label(self):
        row = item(warehouse="", location="")
        events = [entry(row, warehouse=w, location=p) for w, p in (("B", "02"), ("A", ""), ("B", "02"))]
        sheet = self.load(snapshot([row], events)).worksheets[0]
        self.assertEqual(sheet["J2"].value, "A\nB")
        self.assertEqual(sheet["M2"].value, "(ไม่ระบุ)\n02")
        self.assertEqual((sheet["K2"].value, sheet["N2"].value), ("ไม่ได้ระบุในบัญชี", "ไม่ได้ระบุในบัญชี"))

    def test_high_precision_remains_text_and_six_place_values_do_not_round(self):
        row = item(book_qty="99999999999999.999999", actual_qty="99999999999999.999998")
        book = self.load(snapshot([row], [entry(row, quantity="99999999999999.999998")]))
        self.assertEqual((book.worksheets[0]["E2"].value, book.worksheets[0]["E2"].data_type),
                         ("99999999999999.999999", "s"))
        self.assertEqual(book.worksheets[0]["G2"].value, -0.000001)
        self.assertEqual(book.worksheets[1]["G2"].value, "99999999999999.999998")

    def test_formula_like_user_values_are_literal_strings_and_no_external_links(self):
        row = item(product_code="=1+1", product_name="=HYPERLINK(\"https://example.invalid\",\"x\")",
                   barcode="0000012", warehouse="+SUM(A1)", location="@name")
        value = snapshot([row], [entry(row, counted_by_name="=2+2", updated_by_name="-3+4")])
        data = workbook(value)
        with ZipFile(BytesIO(data)) as archive:
            self.assertFalse(any("externalLink" in name for name in archive.namelist()))
        book = self.load(value)
        for sheet in book:
            self.assertFalse(any(cell.data_type == "f" for cells in sheet for cell in cells))
        self.assertEqual(book.worksheets[0]["A2"].value, "=1+1")
        self.assertEqual(book.worksheets[0]["C2"].value, "0000012")

    def test_naive_native_utc_and_aware_or_iso_timestamps(self):
        row = item()
        events = [entry(row, counted_at=datetime(2026, 9, 21, 1, 2, 3),
                        updated_at=datetime(2026, 9, 21, 8, 2, 3, tzinfo=timezone(timedelta(hours=7)))),
                  entry(row, counted_at="2026-09-21T08:02:03+07:00", updated_at="2026-09-21 01:02:03")]
        sheet = self.load(snapshot([row], events)).worksheets[1]
        for address in ("I2", "L2", "I3", "L3"):
            self.assertEqual(sheet[address].value, "2026-09-21T01:02:03+00:00")

    def test_snapshot_is_not_mutated_and_empty_workbook_keeps_headers(self):
        value = snapshot()
        before = deepcopy(value)
        self.load(value)
        self.assertEqual(value, before)
        book = self.load(snapshot([], []))
        self.assertEqual([sheet.max_row for sheet in book], [1, 1])

    def test_legacy_photos_and_incomplete_snapshot_are_explicitly_rejected(self):
        for value, code in ((snapshot(count_mode="legacy"), "stocktake.legacy_task"),
                            (snapshot(photos=[{"entry_id": "E1"}]), "stocktake.photos_not_supported"),
                            (snapshot(entry_total=101), "stocktake.export_incomplete")):
            with self.subTest(code=code), self.assertRaisesRegex(StocktakeValidationError, f"^{code}$"):
                workbook(value)
        value = snapshot()
        del value["entry_total"]
        with self.assertRaisesRegex(StocktakeValidationError, "stocktake.export_incomplete"):
            workbook(value)

    def test_invalid_quantity_date_or_orphan_entry_cannot_produce_partial_file(self):
        row = item()
        for value in (snapshot([row], [entry(row, quantity="NaN")]),
                      snapshot([row], [entry(row, counted_at="invalid")]),
                      snapshot([row], [entry(row, item_id="unrelated")])):
            with self.assertRaises(StocktakeValidationError):
                workbook(value)


if __name__ == "__main__":
    unittest.main()
