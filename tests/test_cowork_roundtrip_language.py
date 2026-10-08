"""复核表按导出人语言出、四种语言都导得回来（用户 2026-10-08）；纯本地，不碰站点。"""
from collections import defaultdict
import io
import unittest

import openpyxl

from pearnly_cowork.ocr.table_engine.roundtrip import has_roundtrip_sheet
from pearnly_cowork.roundtrip import spec as rt
from pearnly_cowork.roundtrip.export import build_review_workbook
from pearnly_cowork.roundtrip.parser import parse_roundtrip_workbook

LANGS = ("th", "zh", "en", "ja")


def record(name, direction):
    fields = {"direction": direction, "date": "2026-09-21", "invoice_number": f"{name}/69",
              "buyer_name": "ผู้ซื้อ", "buyer_tax": "0105551234567", "buyer_branch": "00000",
              "seller_name": "ผู้ขาย", "seller_tax": "0105557654321", "seller_branch": "00001",
              "subtotal": "100", "vat": "7", "total_amount": "107",
              "items": [{"description": "สินค้า", "qty": "2", "unit_price": "50", "amount": "100"}]}
    return {"history_id": name, "merged_fields": fields}


RECORDS = [record("H-1", "sales"), record("H-2", "purchase"), record("H-3", "unknown")]
FACTS = {"H-2": {"state": {"status": "manual", "reason": "x"}}}


def comparable(parsed):
    """去掉只与表格本身相关、不影响回导结果的部分（源文件摘要、表名证据）。"""
    keep = lambda entry: {k: v for k, v in entry.items() if k != "sheet"}  # noqa: E731
    return ([keep(e) for e in parsed["documents"]], [keep(e) for e in parsed["pending"]])


class Language(unittest.TestCase):
    def test_sheet_names_headers_and_status_follow_the_language(self):
        for lang, sales, purchase, status in (("th", "ขาย", "ซื้อ", "ตรวจสอบเอง"), ("zh", "销项", "进项", "待人工核对"),
                                              ("en", "Sales", "Purchases", "Manual check"), ("ja", "売上", "仕入", "手動確認")):
            with self.subTest(lang=lang):
                book = openpyxl.load_workbook(io.BytesIO(build_review_workbook(RECORDS, push_facts=FACTS, lang=lang)))
                self.assertEqual(book.sheetnames[:2], [sales, purchase])
                headers = [cell.value for cell in book[purchase][1]]
                self.assertEqual(headers[0], rt.localize("ลำดับ", lang))
                status = headers.index(rt.localize(rt.COL_PUSH_STATUS, lang)) + 1
                self.assertEqual(book[purchase].cell(2, status).value, f"{status_label(lang)} · x")
                # 系统引用列照旧隐藏
                key = headers.index(rt.localize(rt.COL_ROW_KEY, lang)) + 1
                self.assertTrue(book[purchase].column_dimensions[openpyxl.utils.get_column_letter(key)].hidden)
                book.close()

    def test_every_language_reads_back_the_same_as_thai(self):
        thai = comparable(parse_roundtrip_workbook(build_review_workbook(RECORDS, push_facts=FACTS)))
        for lang in LANGS[1:]:
            with self.subTest(lang=lang):
                data = build_review_workbook(RECORDS, push_facts=FACTS, lang=lang)
                parsed = parse_roundtrip_workbook(data)
                self.assertEqual(comparable(parsed), thai)
                self.assertEqual(parsed["sheets"], [rt.SHEET_SALES, rt.SHEET_PURCHASE, rt.SHEET_PENDING])

    def test_the_upload_gate_recognises_every_language(self):
        for lang in LANGS:
            book = openpyxl.load_workbook(io.BytesIO(build_review_workbook(RECORDS, lang=lang)))
            sheet = book.worksheets[1]
            unit = {"sheet_name": sheet.title, "roundtrip_headers": [cell.value for cell in sheet[1]]}
            with self.subTest(lang=lang):
                self.assertTrue(has_roundtrip_sheet({"units": [unit]}))
            book.close()
        self.assertFalse(has_roundtrip_sheet({"units": [{"sheet_name": "Sheet1", "roundtrip_headers": ["A"]}]}))

    def test_no_two_protocol_names_share_a_translation(self):
        # 读侧按译名反查规格名：同一个译名对应两个规格名就会读错列
        owners = defaultdict(set)
        for name in rt._READ_NAMES:
            for lang in LANGS:
                owners[rt.localize(name, lang).strip().casefold()].add(name)
        self.assertEqual({text: names for text, names in owners.items() if len(names) > 1}, {})

    def test_unknown_language_and_unknown_headers_pass_through(self):
        self.assertEqual(rt.localize(rt.SHEET_SALES, "fr"), rt.SHEET_SALES)
        self.assertEqual(rt.localize(rt.SHEET_SALES, "zh-TW"), "销项")
        self.assertEqual(rt.canonical("  Sales "), rt.SHEET_SALES)
        self.assertEqual(rt.canonical("备注"), "备注")


def status_label(lang):
    return rt.localize("ตรวจสอบเอง", lang)


if __name__ == "__main__":
    unittest.main()
