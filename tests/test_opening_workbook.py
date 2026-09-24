"""Opening workbook executes the real parser/exporter without a Site."""
from io import BytesIO
import importlib.util
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

SOURCE = Path(__file__).resolve().parents[1] / "apps/pearnly_erp_portal/pearnly_erp_portal/inventory/opening_workbook.py"
spec = importlib.util.spec_from_file_location("opening_workbook", SOURCE)
workbook = importlib.util.module_from_spec(spec)
spec.loader.exec_module(workbook)


def xlsx(rows):
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class OpeningWorkbook(unittest.TestCase):
    def test_template_has_text_identifiers_no_formulas_and_blank_quantity_cost(self):
        headers = [field for field, _ in workbook.FIELD_LABELS]
        data = workbook.build_bytes(headers, [{"product": "P1", "product_code": "0012", "barcode": "00003",
                                               "product_name": "=HYPERLINK(\"https://example.test\",\"x\")", "unit": "pcs"}])
        book = openpyxl.load_workbook(BytesIO(data), data_only=False)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet["B2"].value, "0012")
        self.assertEqual(sheet["D2"].value, "00003")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertIsNone(sheet["G2"].value)
        self.assertIsNone(sheet["H2"].value)
        self.assertEqual(sheet.freeze_panes, "A2")
        with self.assertRaisesRegex(ValueError, "empty"):
            workbook.parse_workbook(data)

    def test_zero_is_explicit_blank_catalog_rows_skip_and_partial_blanks_preserve(self):
        rows = workbook.parse_workbook(xlsx([
            ["product_code", "qty", "unit_cost"], ["001", None, None],
            ["002", " ", ""], ["003", 0, 0], ["004", 1, None], ["005", None, 2]]))
        self.assertEqual([row["row_number"] for row in rows], [4, 5, 6])
        self.assertEqual((rows[0]["product_code"], rows[0]["qty"], rows[0]["unit_cost"]), ("003", "0", "0"))
        self.assertIsNone(rows[1]["unit_cost"])
        self.assertIsNone(rows[2]["qty"])

    def test_bad_numbers_remain_rows_instead_of_silent_skipping(self):
        rows = workbook.parse_workbook(xlsx([["product", "qty", "unit_cost"], ["unknown", "not a number", -1]]))
        self.assertEqual(rows, [{"product": "unknown", "qty": "not a number", "unit_cost": "-1", "row_number": 2}])

    def test_translated_headers_and_leading_zero_codes(self):
        aliases = workbook.header_aliases({"物料编码": "product_code", "期初数量": "qty", "期初成本": "unit_cost"})
        rows = workbook.parse_workbook(xlsx([["物料编码", "期初数量", "期初成本"], ["0001", "1.25", "2.30"]]), aliases)
        self.assertEqual(rows[0]["product_code"], "0001")

    def test_duplicate_unknown_missing_headers_and_numeric_identifier_are_rejected(self):
        for values, message in (([["product", "qty", "qty", "unit_cost"]], "headers_invalid"),
                                ([["product", "qty"]], "headers_invalid"),
                                ([["product", "qty", "unit_cost", "extra"]], "headers_invalid"),
                                ([["product_code", "qty", "unit_cost"], [12, 1, 1]], "identifier_text_required")):
            with self.subTest(values=values), self.assertRaisesRegex(ValueError, message):
                workbook.parse_workbook(xlsx(values))

    def test_formulas_and_error_cells_are_rejected_even_without_quantities(self):
        for value in ("=1+1", "#DIV/0!"):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "formula_invalid"):
                workbook.parse_workbook(xlsx([["product", "qty", "unit_cost"], ["P1", value, 1]]))

    def test_zip_and_rows_are_bounded_before_expansion(self):
        with self.assertRaisesRegex(ValueError, "file_too_large"):
            workbook.parse_workbook(b"x" * (workbook.MAX_BYTES + 1))
        archive = BytesIO()
        with ZipFile(archive, "w", ZIP_DEFLATED) as output:
            output.writestr("sharedStrings.xml", b"x" * 2048)
        with patch.object(workbook, "MAX_EXPANDED_BYTES", 1024), self.assertRaisesRegex(ValueError, "file_too_large"):
            workbook.parse_workbook(archive.getvalue())
        with patch.object(workbook, "MAX_ROWS", 1), self.assertRaisesRegex(ValueError, "too_many_rows"):
            workbook.parse_workbook(xlsx([["product", "qty", "unit_cost"], ["P1", 1, 1], ["P2", 1, 1]]))

    def test_multiple_sheets_and_invalid_bytes_are_rejected(self):
        book = openpyxl.Workbook()
        book.create_sheet("extra")
        output = BytesIO()
        book.save(output)
        book.close()
        with self.assertRaisesRegex(ValueError, "sheet_invalid"):
            workbook.parse_workbook(output.getvalue())
        with self.assertRaisesRegex(ValueError, "file_invalid"):
            workbook.parse_workbook(b"not a spreadsheet")


if __name__ == "__main__":
    unittest.main()
