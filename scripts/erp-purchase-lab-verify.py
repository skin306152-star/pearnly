"""Purchase migration integration checks. Dedicated, synthetic lab writes only.

Requires pearnly-purchase-lab containers, local marker and two isolated Sites.
No OCR provider, LINE, Google or other external business call. Generated files and
records are intentionally retained for browser acceptance under a per-run prefix.
"""
import json
import os
from pathlib import Path
import sys
import subprocess
from uuid import uuid4
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from decimal import Decimal

import frappe

ROOT = Path('/home/frappe/frappe-bench/sites')


def connect(site='erp-pos-a.localhost'):
    if os.environ.get('PEARNLY_ERP_POS_LAB') != '1' or not (ROOT / '.erp-pos-lab').exists():
        raise RuntimeError('Dedicated lab required')
    if site not in ('erp-pos-a.localhost', 'erp-pos-b.localhost'):
        raise RuntimeError('Dedicated lab Site required')
    os.chdir(ROOT)
    frappe.init(site=site, sites_path=str(ROOT))
    frappe.connect()
    assert frappe.conf.get('pearnly_erp_pos_lab') and frappe.conf.db_host == 'db'
    frappe.flags.mute_emails = True
    frappe.set_user('Administrator')


def main():
    connect()
    from pearnly_core.provisioning import ensure_feature_roles
    sys.path.insert(0, '/home/frappe/frappe-bench/apps/pearnly_control')
    from pearnly_control import features
    from pearnly_erp_portal.purchasing import entry, access, records, files, output, recognition, api, service
    from pearnly_erp_portal.integrations import intake
    from pearnly_erp_portal.numbering import service as numbering
    from pearnly_erp_portal.inventory import stock
    ensure_feature_roles(json.dumps(features.roles_payload('erp', features.codes_of('erp'))), admin_user='Administrator')
    prefix = 'Purchase check ' + uuid4().hex[:8]
    payload = {'supplier_name': prefix, 'doc_no': prefix, 'doc_date': '2026-09-22',
               'supplier_branch_type': 'branch', 'supplier_branch_no': '00001',
               'supplier_branch_name': 'Bangkok', 'supplier_address': 'Test address',
               'currency': 'THB', 'fx_rate': 1,
               'lines': [{'description': prefix + ' goods', 'qty': '2.5', 'unit_price': '40',
                          'vat_rate': 7, 'wht_rate': 3, 'unit': 'pcs', 'warehouse': prefix}]}
    key = uuid4().hex
    result = entry.post(payload, request_id=key)
    doc = access.get_document(result['name'])
    assert doc.status == 'submitted' and doc.stock_posted_at
    assert doc.payment_status == 'paid' and float(doc.paid_amount) == 104
    assert float(doc.grand_total) == 107 and float(doc.wht_amount) == 3
    assert len(doc.payments) == 1 and doc.lines[0].product
    assert entry.post(payload, request_id=key)['duplicate']
    assert records.rows({'q': prefix})[0].name == doc.name
    checks = ['atomic posting', 'repeated submit', 'search']
    product, warehouse = doc.lines[0].product, doc.lines[0].warehouse
    assert stock.current_balance(product, warehouse, '').qty == Decimal('2.5')
    # Permission checks include generic native access, not just our route.
    buyer, other = 'purchase-buyer@example.test', 'purchase-other@example.test'
    for user in (buyer, other):
        if not frappe.db.exists('User', user):
            frappe.get_doc({'doctype': 'User', 'email': user, 'first_name': user.split('@')[0],
                'send_welcome_email': 0, 'roles': [{'role': 'Desk User'}, {'role': 'ERP Purchase'}]}).insert(ignore_permissions=True)
    frappe.db.commit()

    def rejects(fn, *args, **kwargs):
        try:
            fn(*args, **kwargs)
        except (frappe.ValidationError, frappe.PermissionError, frappe.AuthenticationError):
            frappe.db.rollback()
            return
        raise AssertionError('Expected rejection: ' + fn.__name__)

    bad = deepcopy(payload)
    bad.update(supplier_name=prefix + ' rollback', doc_no=prefix + ' rollback',
               attachments=[{'file': 'missing-file'}])
    bad['lines'][0]['description'] = prefix + ' rollback item'
    before = numbering.preview(direction='purchase', on='2026-09-22')
    try:
        entry.post(bad, request_id=uuid4().hex)
    except frappe.DoesNotExistError:
        pass
    else:
        raise AssertionError('invalid file must roll back')
    assert not frappe.db.exists('ERP Product', {'product_name': bad['lines'][0]['description']})
    assert not frappe.db.exists('ERP Supplier', {'supplier_name': bad['supplier_name']})
    assert numbering.preview(direction='purchase', on='2026-09-22') == before
    checks.append('rollback item supplier stock numbering')
    rejects(entry.post, payload, request_id=uuid4().hex)
    rejects(entry.post, {**payload, 'doc_no': 'different'}, request_id=key)
    checks.extend(['duplicate invoice refusal', 'request key payload binding'])
    # Full reversal, partial payment, overpayment refusal and settlement.
    entry.payment(doc.name, paid=False); frappe.db.commit()
    current = access.get_document(doc.name)
    assert current.paid_amount == 0 and len(current.payments) == 2 and current.payments[-1].amount == -104
    payment_key=uuid4().hex
    entry.payment(doc.name, amount=20, method='promptpay', date='2026-09-22', note='part', request_id=payment_key); frappe.db.commit()
    assert entry.payment(doc.name, amount=20, method='promptpay', date='2026-09-22', note='part', request_id=payment_key)['duplicate']
    rejects(entry.payment, doc.name, amount=21, request_id=payment_key)
    assert access.get_document(doc.name).payment_status == 'partial'
    rejects(entry.payment, doc.name, amount=1000)
    entry.payment(doc.name, paid=True); frappe.db.commit()
    assert access.get_document(doc.name).paid_amount == 104
    checks.append('payment reversal partial excess settlement')
    native = access.get_document(doc.name); native.supplier_name = 'tamper'
    rejects(native.save)
    native=access.get_document(doc.name);native.calculation_version=1
    rejects(native.save)
    rejects(frappe.delete_doc, 'Purchase Document', doc.name)
    # Private evidence and scoped workbook.
    file = frappe.get_doc({'doctype':'File', 'file_name':prefix+'.txt', 'is_private':1,
                           'content':('purchase test evidence '+prefix).encode()}).insert(ignore_permissions=True)
    files.add(doc.name,[{'file':file.name,'kind':'bill'}]);frappe.db.commit()
    changed = deepcopy(payload);changed['doc_no'] = prefix + ' amendment';changed['lines'][0]['qty']=3
    amended=entry.post(changed,request_id=uuid4().hex,amended_from=doc.name)
    original=access.get_document(doc.name);replacement=access.get_document(amended['name'])
    assert original.status=='void' and original.paid_amount==0
    assert replacement.status=='submitted' and replacement.paid_amount==124.8
    assert stock.current_balance(product,warehouse,'').qty==3
    assert len(replacement.attachments)==1
    entry.cancel(doc.name);frappe.db.commit()
    assert stock.current_balance(product,warehouse,'').qty==3
    checks.append('atomic amendment evidence copy and idempotent cancellation')
    frappe.set_user(buyer)
    personal=deepcopy(payload);personal.update(doc_no=prefix+' personal',supplier_name=prefix+' personal',is_paid=0,update_stock=0,auto_create_product=0)
    personal['lines'][0].update(item_type='service',description='=HYPERLINK("https://example.invalid")')
    own=entry.post(personal,request_id=uuid4().hex)
    assert len(records.rows({'q':prefix},limit=0))==1
    assert frappe.has_permission('Purchase Document','read',doc=own['name'])
    rejects(access.get_document, replacement.name)
    assert not frappe.has_permission('Purchase Document','read',doc=replacement.name)
    assert not frappe.has_permission('File','read',doc=replacement.attachments[0].file)
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(output.workbook({'q':prefix})))
    assert len(book.worksheets)==4 and book.worksheets[0].max_row==2
    line_sheet=book.worksheets[1]
    assert any(str(cell.value).startswith("'=HYPERLINK") and cell.data_type=='s' for row in line_sheet for cell in row)
    checks.append('owner scope native File access and scoped formula-safe Excel')
    frappe.set_user(other);rejects(access.get_document,own['name'])
    assert records.rows({'q':prefix},limit=0)==[]
    frappe.set_user('Guest');rejects(records.rows)
    frappe.set_user(buyer)
    # OCR is replaced at the provider boundary; no paid call.
    image=frappe.get_doc({'doctype':'File','file_name':prefix+'.png','is_private':1,
        'content':b'\x89PNG\r\n\x1a\nsynthetic-test'}).insert(ignore_permissions=True)
    frappe.db.commit()
    with patch('pearnly_core.ocr.is_configured',return_value=True), patch('pearnly_core.ocr.json_from_images',autospec=True,return_value=personal) as provider:
        result=recognition.recognise([{'file':image.name}])
        assert result['source']=='ocr' and len(result['image_sha256'])==64 and provider.call_count==1
        assert provider.call_args.kwargs['model']
    checks.append('OCR boundary mapping no external calls')
    received=intake.receive(line_user_id='lab-'+prefix,sent_by=buyer,media_type='image',
        content=b'line-file-'+prefix.encode(),filename=prefix+'-line.png')
    confirm_payload=deepcopy(personal);confirm_payload['doc_no']=prefix+' LINE'
    line_doc=intake.confirm(received['name'],**confirm_payload)
    assert access.get_document(line_doc['name']).status=='submitted'
    assert intake.get_intake(received['name'])['status']=='confirmed'
    assert len(access.get_document(line_doc['name']).attachments)==1
    assert intake.confirm(received['name'],**confirm_payload)['duplicate']
    assert intake.receive(line_user_id='lab-'+prefix,sent_by=buyer,media_type='image',content=b'line-file-'+prefix.encode(),filename='again.png')['duplicate']
    duplicate=deepcopy(confirm_payload);duplicate['doc_no']=prefix+' duplicate image';duplicate['image_sha256']=intake.get_intake(received['name'])['file_hash']
    rejects(entry.post,duplicate,request_id=uuid4().hex)
    frappe.set_user(other)
    assert not frappe.has_permission('ERP LINE Intake','read',doc=received['name'])
    source=frappe.db.get_value('File', {'attached_to_doctype':'ERP LINE Intake','attached_to_name':received['name']},'name')
    assert not frappe.has_permission('File','read',doc=source)
    frappe.set_user(buyer)
    checks.append('LINE intake native scope, repeat media and confirmed-file deduplication')
    frappe.set_user('Administrator')
    category=records.category(prefix+' category')['name']
    with_category=deepcopy(personal);with_category.update(doc_no=prefix+' category',doc_kind='expense')
    with_category['lines'][0]['category']=category
    classified=entry.post(with_category,request_id=uuid4().hex)
    assert [row.name for row in records.rows({'category':[category]},limit=0)]==[classified['name']]
    assert records.overview({'category':[category]})['expense_count']==1
    assert records.rows({'q':prefix,'month':'2026-08'},limit=0)==[]
    overview=records.overview({'q':prefix,'month':'2026-09'})
    assert overview['total']==len(records.rows({'q':prefix,'month':'2026-09'},limit=0))
    checks.append('category monthly filters count and summary agreement')
    # PDF rendering itself is exercised; not replaced by a mock HTTP 200.
    voucher=output.voucher(replacement.name,'wht_cert');frappe.db.commit()
    pdf=frappe.get_doc('File',voucher['file']).get_content()
    assert pdf.startswith(b'%PDF') and len(pdf)>1000
    assert output.voucher(replacement.name,'wht_cert')['file']==voucher['file']
    checks.append('native PDF private voucher and idempotent regeneration')
    service.set_payment_status(replacement.name,'unpaid')
    service.pay(replacement.name,amount=10,method='cash')
    assert access.get_document(replacement.name).paid_amount==10
    checks.append('legacy API callers route through new settlement contract')
    assert api.api_purchase_defaults()['capabilities']['ocr'] is False
    checks.append('unconfigured OCR entry unavailable')
    # A mixed invoice exercises every imported business field through native persistence and export.
    detailed=deepcopy(payload)
    detailed.update(supplier_name=prefix+' complete',supplier_tax_id='TEST-TAX',doc_no=prefix+' complete',
        doc_kind='purchase_order',currency='USD',fx_rate='34.25',
        price_mode='inclusive',due_date='2026-10-22',requester='Requester',requester_user=buyer,
        notes='Remarks preserved',external_ref=prefix+' ref',source='manual',is_paid=0,update_stock=0)
    detailed['lines']=[{'description':prefix+' complete item','item_type':'goods','qty':'2',
        'unit_price':'107','discount':'10.7','vat_rate':'7','wht_rate':'3','unit':'pcs',
        'code':prefix+'-sku','barcode':prefix+'-barcode','category':category,
        'batch_no':'LOT-1','expiry_date':'2027-01-01'},
        {'description':'Service','item_type':'service','qty':'1','unit_price':'53.5','vat_rate':'7','wht_rate':'3'}]
    complete=entry.post(detailed,request_id=uuid4().hex)
    reloaded=records.get(complete['name'])
    for field in ('supplier_name','supplier_tax_id','supplier_branch_type','supplier_branch_no',
                  'supplier_branch_name','supplier_address','doc_no','doc_kind','currency',
                  'price_mode','requester','requester_user','notes','external_ref','source'):
        assert reloaded[field]==detailed[field],field
    for field in ('code','barcode','batch_no','category','unit'):
        assert reloaded['lines'][0][field]==detailed['lines'][0][field],field
    assert str(reloaded['lines'][0]['expiry_date'])=='2027-01-01'
    assert reloaded['grand_total']==256.8 and reloaded['wht_amount']==7.2 and reloaded['net_payable']==249.6
    exported=load_workbook(BytesIO(output.workbook({'q':prefix+' complete'})))
    all_values=[cell.value for sheet in exported for row in sheet for cell in row]
    assert 'LOT-1' in all_values and detailed['external_ref'] in all_values and detailed['lines'][0]['barcode'] in all_values
    assert frappe._('Purchase Order') in all_values and 'purchase_order' not in all_values
    checks.append('mixed invoice all business fields preserved through reload and Excel')
    copied_file=frappe.get_doc({'doctype':'File','file_name':prefix+' duplicate.txt','is_private':1,
        'content':('purchase test evidence '+prefix).encode()}).insert(ignore_permissions=True)
    frappe.db.commit()
    duplicate_file=deepcopy(personal);duplicate_file.update(doc_no=prefix+' duplicate file',
        attachments=[{'file':copied_file.name,'kind':'bill'}])
    rejects(entry.post,duplicate_file,request_id=uuid4().hex)
    checks.append('manual uploaded evidence deduplicated by server content hash')

    # Version-one API remains compatible without requiring a pristine other tenant.
    legacy=service.create_purchase(supplier_name=prefix+' legacy',doc_date='2026-09-22',
        doc_kind='service',doc_no=prefix+' legacy',wht_rate=3,
        lines=[{'description':'Legacy service','qty':1,'unit_price':100,'vat_rate':7}])
    legacy_doc=access.get_document(legacy['name'])
    assert legacy_doc.calculation_version==1 and legacy_doc.net_payable==104
    service.set_status(legacy_doc.name,'submitted',receive_stock=False)
    service.pay(legacy_doc.name,amount=20,method='cash')
    assert access.get_document(legacy_doc.name).payment_status=='partial'
    historical=records.get(legacy_doc.name)
    assert historical['is_paid']==0 and historical['update_stock']==0 and historical['auto_create_product']==0
    checks.append('version-one amounts submit and payment remain compatible')
    frappe.db.commit()
    concurrent=deepcopy(payload);concurrent['doc_no']=prefix+' concurrent';concurrent['lines'][0]['description']=prefix+' concurrent item'
    job=json.dumps({'payload':concurrent,'request_id':uuid4().hex})
    workers=[subprocess.Popen([sys.executable,__file__,'--post-worker'],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True) for _ in range(2)]
    for worker in workers:
        worker.stdin.write(job);worker.stdin.close();worker.stdin=None
    replies=[]
    for worker in workers:
        out,err=worker.communicate(timeout=60)
        assert worker.returncode==0,err
        replies.append(json.loads(out))
    assert replies[0]['name']==replies[1]['name'] and sum(bool(row.get('duplicate')) for row in replies)==1
    frappe.db.rollback()
    assert frappe.db.count('ERP Product',{'product_name':prefix+' concurrent item'})==1
    parallel_doc=access.get_document(replies[0]['name'])
    assert stock.current_balance(parallel_doc.lines[0].product,warehouse,'').qty==Decimal('2.5')
    checks.append('two concurrent processes post once with one item and one stock receipt')
    frappe.db.commit();frappe.destroy();connect('erp-pos-b.localhost')
    assert not frappe.db.exists('Purchase Document',{'doc_no':prefix+' amendment'})
    assert not frappe.db.exists('Purchase Document',{'doc_no':prefix+' complete'})
    assert not frappe.db.exists('ERP Supplier',{'supplier_name':prefix+' complete'})
    checks.append('tenant isolation')
    print(json.dumps({'result':'PASS','checks':checks,'name':replacement.name,'prefix':prefix},ensure_ascii=False))
    frappe.destroy()


if __name__ == '__main__':
    if '--post-worker' in sys.argv:
        connect()
        from pearnly_erp_portal.purchasing import entry
        job=json.load(sys.stdin)
        print(json.dumps(entry.post(**job)))
        frappe.destroy()
    else:
        main()
