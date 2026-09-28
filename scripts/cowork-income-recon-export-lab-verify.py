"""Authorized income review-workbook downloads on the isolated Cowork lab; zero network.

Writes synthetic tasks only. Uses the real Frappe/WSGI/DB path; no model, ERP,
LINE, identity provider or external configuration is contacted.
"""
from io import BytesIO
import json
from pathlib import Path
import runpy
import unittest
from urllib.parse import unquote
from uuid import uuid4

import frappe
from openpyxl import load_workbook
from werkzeug.test import Client
from werkzeug.wrappers import Response

from pearnly_cowork.reconciliation.export import income_workbook

FIXTURE = runpy.run_path(str(Path(__file__).with_name("cowork-income-recon-lab-verify.py")))
BASE = FIXTURE["NativeIncomeReconciliation"]
connect = FIXTURE["connect"]
TASK = FIXTURE["TASK"]
t = income_workbook._t


class NativeIncomeReconExport(unittest.TestCase):
	"""Reuses the fixture class of the income store verification."""
	setUpClass = classmethod(BASE.setUpClass.__func__)
	setUp, tearDown = BASE.setUp, BASE.tearDown

	def parsed_task(self):
		return self.api.create_from_parsed(workspace=self.workspace, title="June 2569 output VAT",
			request_id=self.request_id, gl_source={"kind": "excel",
				"sheets": FIXTURE["PARSED_SHEETS"]}, vat_rows=None,
			header=FIXTURE["HEADER"], vat_source={"kind": "excel", "cells": FIXTURE["VAT_SHEETS"]},
			gl_filename="GL June.xlsx", vat_filename="VAT June.xlsx")

	def test_export_renders_the_stored_rows_and_summary(self):
		task = self.parsed_task()
		exported = self.service.export(task["name"], "th")
		self.assertEqual(exported["filename"], "June 2569 output VAT.xlsx")  # 任务名称（2026-09-28 起）
		book = load_workbook(BytesIO(exported["content"]))
		self.assertEqual(book.sheetnames,
			[t("th", "sheet1"), t("th", "sheet2"), t("th", "sheet3")])
		detail = book[t("th", "sheet1")]
		# Four ledger rows reached the store, three report rows the detail follows;
		# one report row found no counterpart, so it reads as GL missing.
		self.assertEqual([detail.cell(2, index).value for index in range(1, 6)], [3, 2, 0, 1, 0.0])
		self.assertEqual([detail.cell(row, 1).value for row in (5, 6, 7)],
			[t("th", "st_matched"), t("th", "st_matched"), t("th", "st_missing")])
		self.assertEqual([detail.cell(row, 2).value for row in (5, 6, 7)],
			["GL-1", "GL-2", "INV-9"])
		self.assertEqual((detail.cell(5, 6).value, detail.cell(5, 7).value), (1100.0, 0.0))
		self.assertEqual((detail.cell(6, 6).value, detail.cell(7, 6).value),
			(-107.0, t("th", "not_found")))
		summary = book[t("th", "sheet2")]
		labels = [summary.cell(row, 1).value for row in range(1, summary.max_row + 1)]
		self.assertEqual(summary.cell(2, 2).value, 1493.0)
		# The three adjustments carry one item each, and the ladder ends on the report.
		self.assertEqual(summary.cell(summary.max_row, 2).value, 1243.0)
		self.assertEqual(sum(labels.count(label) for label in (
			t("th", "s_minus_gl_cr"), t("th", "s_plus_vat_p"))), 2)
		self.assertIn("GL-3", " ".join(str(value) for value in labels))
		self.assertIn("INV-9", " ".join(str(value) for value in labels))
		# An unknown language falls back to Thai, exactly like the store contract.
		fallback = self.service.export(task["name"], "de")
		self.assertEqual(load_workbook(BytesIO(fallback["content"])).sheetnames, book.sheetnames)

	def test_export_is_scoped_and_entitled_like_the_task_api(self):
		task = self.parsed_task()
		frappe.set_user(self.outsider)
		with self.assertRaises(frappe.PermissionError):
			self.service.export(task["name"])
		frappe.set_user(self.norecon)
		frappe.conf.pearnly_features = ["intake", "master_data", "push_log", "stocktake"]
		with self.assertRaises(frappe.PermissionError):
			self.service.export(task["name"])
		frappe.conf.pearnly_features = ["intake", "master_data", "push_log", "recon", "stocktake"]
		frappe.set_user(self.actor)

	def test_export_download_is_a_private_attachment_and_writes_the_access_log(self):
		task = self.parsed_task()
		site = frappe.local.site
		frappe.set_user("Administrator")
		account = frappe.get_doc("User", self.actor)
		account.api_key, secret = uuid4().hex, uuid4().hex
		account.api_secret = secret
		account.save()
		headers = {"Authorization": f"token {account.api_key}:{secret}"}
		frappe.db.commit()
		frappe.destroy()
		from frappe.app import application
		client = Client(application, Response)
		url = ("/api/method/pearnly_cowork.reconciliation.income_recon_api.export?task_id="
			+ task["name"] + "&lang=th")

		def get(authenticated=True):
			return client.get(url, base_url="http://" + site,
				headers=headers if authenticated else {})

		try:
			response = get()
			self.assertEqual(response.status_code, 200,
				response.get_data(as_text=True) if response.status_code != 200 else "")
			self.assertIn("attachment", response.headers["Content-Disposition"])
			self.assertIn("June 2569 output VAT.xlsx", unquote(response.headers["Content-Disposition"]))
			self.assertIn("no-store", response.headers["Cache-Control"])
			self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
			self.assertIn("spreadsheetml.sheet", response.headers["Content-Type"])
			self.assertEqual(len(load_workbook(BytesIO(response.data)).sheetnames), 3)
			self.assertEqual(get(False).status_code, 403)
			self.assertEqual(client.post(url, base_url="http://" + site,
				headers=headers).status_code, 403)
		finally:
			connect()
		# The real request path defers Access Log writes to the per-Site Redis queue;
		# the isolated lab intentionally runs no scheduler.
		queued = [json.loads(value)
			for value in frappe.cache.lrange("insert_queue_for_Access Log", 0, -1)]
		rows = [row for batch in queued for row in (batch if isinstance(batch, list) else [batch])]
		self.assertTrue(any(row.get("user") == self.actor
			and row.get("file_type") == "income_recon_workbook"
			and row.get("reference_document") == task["name"] for row in rows))


if __name__ == "__main__":
	connect()
	try:
		result = unittest.TextTestRunner(verbosity=2).run(
			unittest.defaultTestLoader.loadTestsFromTestCase(NativeIncomeReconExport))
		print(json.dumps({"income_recon_export_native": "PASS" if result.wasSuccessful() else "FAIL",
			"tests": result.testsRun}))
		raise SystemExit(0 if result.wasSuccessful() else 1)
	finally:
		if getattr(frappe.local, "db", None):
			frappe.db.rollback()
		frappe.destroy()
