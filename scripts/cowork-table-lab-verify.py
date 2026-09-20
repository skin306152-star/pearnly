"""Guarded native table workflow tests: local MariaDB/Redis, fake model/billing.

Creates synthetic fixtures in the dedicated Cowork lab only. No worker, public
listener, provider call, ERP write or user-data deletion. Keeps labelled fixtures.
"""
from contextlib import ExitStack
import io
import json
from pathlib import Path
import runpy
import subprocess
import sys
import unittest
import uuid
from unittest.mock import patch

import frappe
from openpyxl import Workbook
from docx import Document

BASE = runpy.run_path(str(Path(__file__).with_name("cowork-execution-lab-verify.py")))
NativeExecution = BASE["NativeExecution"]
response, connect, grant = (BASE[k] for k in ("response", "connect", "grant"))


def workbook(sheets=2, duplicate=False):
    book = Workbook()
    for i in range(sheets):
        sheet = book.active if i == 0 else book.create_sheet()
        sheet.title = "Sheet " + str(i + 1)
        sheet.append(["invoice_number", "total_amount", "total_amount" if duplicate else "seller_name"])
        sheet.append(["IV69/100", "107", "Synthetic seller"])
    result = io.BytesIO()
    book.save(result)
    return result.getvalue()


def word(*, paragraph=False, extra=False):
    doc = Document()
    if paragraph:
        doc.add_paragraph("Second business document must remain private")
    table = doc.add_table(rows=2, cols=3)
    for row, values in zip(table.rows, (("invoice_number", "total_amount", "seller_name"),
                                      ("IV69/100", "107", "Synthetic seller"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    if extra:
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "A hidden second business document"
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


class NativeTableExecution(unittest.TestCase):
    setUpClass = classmethod(NativeExecution.setUpClass.__func__)
    setUp, tearDown = NativeExecution.setUp, NativeExecution.tearDown
    admit, run_work = NativeExecution.admit, NativeExecution.run_work

    def upload(self, filename="synthetic.xlsx", data=None):
        self.source = self.files.create(self.workspace, filename, data or workbook(), uuid.uuid4().hex)
        frappe.db.commit()
        return self.source

    def other_access(self):
        frappe.set_user("Administrator")
        if not frappe.db.exists("User Permission", {"user": self.other, "allow": "Cowork Workspace", "for_value": self.workspace}):
            grant(self.other, "Cowork Workspace", self.workspace)
            frappe.db.commit()
        frappe.set_user(self.other)

    def test_multisheet_logical_manifest_and_single_character_charge(self):
        from pearnly_cowork.ocr.table_inputs import extract_table
        data = workbook()
        parsed = extract_table(data, "synthetic.xlsx")
        self.upload(data=data)
        doc = self.admit()
        result = self.run_work(doc)
        self.assertEqual((result["state"], result["reference_kind"], result["page_count"]), ("saved", "logical_section", 2))
        self.assertEqual(self.http.call_count, 2)
        self.assertEqual(self.charge.call_args.args, ("cowork", "excel", parsed.billing_characters))
        self.assertEqual(self.status.call_args.args, ("excel", parsed.billing_characters))
        source = frappe.get_doc("Cowork Source File", self.source["name"])
        self.assertEqual((source.page_count, source.reference_count), (1, 2))
        versions = frappe.get_all("Version", filters={"ref_doctype": "Cowork Source File",
            "docname": source.name}, pluck="data", limit_page_length=0)
        self.assertTrue(all("provenance_json" not in value for value in versions))
        self.assertNotIn("provenance_json", self.files.project(source))
        history = frappe.get_doc("Cowork OCR History", result["history_records"][0])
        self.assertFalse(history.source_pdf_id)
        self.assertEqual(json.loads(history.source_page_indices), [1, 2])
        for call in self.http.call_args_list:
            self.assertNotIn("inlineData", json.dumps(call.kwargs["json"]))
        self.other_access()
        self.assertEqual(self.files.original(source.name)[1], data)
        with self.assertRaises(frappe.PermissionError):
            self.files.page(history.name, 1)

    def test_legacy_encoding_billing_uses_original_utf8_ignore(self):
        data = "invoice_number,total_amount,seller_name\nIV69/100,107,บริษัท ทดสอบ".encode("cp874")
        self.upload("synthetic.csv", data)
        result = self.run_work(self.admit())
        self.assertEqual(result["state"], "saved")
        self.assertEqual(self.charge.call_args.args, ("cowork", "excel", len(data.decode("utf-8", errors="ignore"))))

    def test_word_table_only_zero_units_not_replaced_by_one(self):
        self.upload("synthetic.docx", word())
        doc = self.admit()
        self.assertEqual((doc.billing_kind, doc.billing_units), ("excel", 0))
        self.assertEqual(self.run_work(doc)["state"], "saved")
        self.assertEqual(self.charge.call_args.args, ("cowork", "excel", 0))
        self.assertEqual(self.status.call_args.args, ("excel", 0))

    def test_preexisting_image_batch_without_new_fields_still_resumes(self):
        # Reproduce the schema defaults seen on an admitted fourth-batch row.
        # This exact synthetic row is owned by the test, not an operational task.
        doc = self.admit()
        frappe.db.set_value("Cowork OCR Batch", doc.name,
            {"reference_kind": None, "billing_kind": None, "billing_units": 0}, update_modified=False)
        frappe.db.commit()
        result = self.run_work(doc)
        self.assertEqual((result["state"], result["reference_kind"], result["billing_units"]),
                         ("saved", "physical_page", 1))
        self.assertEqual(self.charge.call_args.args, ("cowork", "pdf", 1))

    def test_duplicate_admission_after_another_connection_commits(self):
        self.upload()
        value = {"actor": self.actor, "source": self.source["name"], "key": self.key,
                 "conf": {k: v for k, v in frappe.conf.items() if k.startswith("pearnly_")}}
        child = subprocess.Popen([sys.executable, __file__, "--stale-admit-child"],
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        try:
            child.stdin.write(json.dumps(value) + "\n")
            child.stdin.flush()
            self.assertEqual(child.stdout.readline().strip(), "READY")
            first = self.admit()
            output, error = child.communicate("GO\n", timeout=30)
            self.assertEqual(child.returncode, 0, error)
            second = json.loads(output.strip())
            self.assertEqual((second["accepted"], second["name"]), (True, first.name))
            self.assertGreaterEqual(second["admission_attempts"], 1)
            self.assertLessEqual(second["admission_attempts"], 3)
            self.assertEqual(frappe.db.count("Cowork OCR Batch", {"source_file": self.source["name"]}), 1)
        finally:
            if child.poll() is None:
                child.kill()
                child.communicate()

    def test_word_unread_sections_do_not_unlock_original(self):
        data = word(paragraph=True, extra=True)
        self.upload("synthetic.docx", data)
        result = self.run_work(self.admit())
        self.assertEqual(result["state"], "saved")
        self.assertEqual(self.files.original(self.source["name"])[1], data)
        self.other_access()
        with self.assertRaises(frappe.PermissionError):
            self.files.original(self.source["name"])

    def test_duplicate_headers_do_not_unlock_original(self):
        self.upload(data=workbook(duplicate=True))
        self.assertEqual(self.run_work(self.admit())["state"], "saved")
        self.other_access()
        with self.assertRaises(frappe.PermissionError):
            self.files.original(self.source["name"])

    def test_placeholder_document_is_not_a_complete_manifest(self):
        self.upload("synthetic.txt", b"Synthetic document not parsed yet")
        frappe.db.set_value("Cowork Source File", self.source["name"], "binding_complete", 1)
        frappe.db.commit()
        self.other_access()
        with self.assertRaises(frappe.PermissionError):
            self.files.original(self.source["name"])
        self.http.assert_not_called()

    def test_table_quota_resumes_completed_sheet_without_repay(self):
        self.upload()
        self.http.side_effect = [response(), response(status=429), response()]
        doc = self.admit()
        self.assertEqual(self.run_work(doc)["state"], "deferred")
        self.assertEqual(frappe.db.count("Cowork OCR History", {"source_file": self.source["name"]}), 0)
        self.charge.assert_not_called()
        with patch("pearnly_core.execution.budget.permit", side_effect=lambda *a, **kw: ExitStack()):
            result = self.run_work(doc, retry=True)
        self.assertEqual(result["state"], "saved")
        self.assertEqual((self.http.call_count, self.charge.call_count), (3, 1))
        self.assertEqual(self.charge.call_args.args[1], "excel")

    def test_incomplete_second_sheet_saves_and_charges_nothing(self):
        from pearnly_cowork.ocr.results import OcrResultError
        self.upload()
        self.http.side_effect = [response(), response({})]
        doc = self.admit()
        with self.assertRaises(OcrResultError):
            self.work.run(doc.work_name)
        self.assertEqual(self.api.detail(doc.name)["state"], "failed")
        self.assertEqual(frappe.db.count("Cowork OCR History", {"source_file": self.source["name"]}), 0)
        self.assertFalse(frappe.db.get_value("Cowork Source File", self.source["name"], "binding_complete"))
        self.charge.assert_not_called()

    def test_oversized_model_input_rejected_before_balance(self):
        self.upload("synthetic.txt", b"A" * 30001)
        result = self.api.submit(self.source["name"], self.key)
        self.assertEqual(result, {"accepted": False, "code": "ocr.table_text_too_long"})
        self.status.assert_not_called()
        self.http.assert_not_called()
        self.charge.assert_not_called()
        self.assertFalse(frappe.db.exists("Cowork OCR Batch", {"source_file": self.source["name"]}))

    def test_broken_workbook_never_admitted(self):
        self.upload("synthetic.xlsx", b"not an OOXML archive")
        result = self.api.submit(self.source["name"], self.key)
        self.assertFalse(result["accepted"])
        self.assertTrue(result["code"].startswith("ocr.table_"), result)
        self.status.assert_not_called()
        self.http.assert_not_called()

    def test_roundtrip_workbook_is_not_sent_through_a_model(self):
        from pearnly_cowork.ocr.table_engine.roundtrip import REQUIRED_HEADERS
        book = Workbook()
        book.active.title = "ขาย"
        book.active.append(sorted(REQUIRED_HEADERS))
        book.active.append(["synthetic"] * len(REQUIRED_HEADERS))
        data = io.BytesIO()
        book.save(data)
        self.upload(data=data.getvalue())
        result = self.api.submit(self.source["name"], self.key)
        self.assertEqual(result, {"accepted": False, "code": "ocr.table_roundtrip_not_supported"})
        self.status.assert_not_called()
        self.http.assert_not_called()

    def test_partial_sheet_assignment_does_not_unlock_whole_workbook(self):
        from pearnly_cowork.history.recognition import persist_batch
        self.upload()
        self.admit()
        pages = [{"page_index": 1, "fields": BASE["GOOD"].copy()}]
        persist_batch(self.source["name"], [{"workspace": self.workspace, "source": "manual",
            "source_page_indices": [1], "pages": pages, "ai_raw": pages, "confidence": "high"}], uuid.uuid4().hex)
        frappe.db.commit()
        self.other_access()
        with self.assertRaises(frappe.PermissionError):
            self.files.original(self.source["name"])

    def test_manifest_cannot_be_mutated_with_native_ignore_validate(self):
        from pearnly_cowork.files.guard import WRITE_CAPABILITY
        self.upload()
        self.admit()
        source = frappe.get_doc("Cowork Source File", self.source["name"])
        source.reference_count = 1
        source.flags.cowork_original_write = WRITE_CAPABILITY
        source.flags.ignore_validate = True
        with self.assertRaises(frappe.PermissionError):
            source.save(ignore_permissions=True)

    def test_table_lost_receipt_retries_same_characters_without_model(self):
        self.upload()
        self.charge.return_value = {"ok": False}
        doc = self.admit()
        first = self.run_work(doc)
        self.assertEqual(first["billing_state"], "pending")
        expected = self.charge.call_args
        self.charge.return_value = {"ok": True, "duplicate": True}
        self.http.side_effect = AssertionError("receipt recovery must not call a model")
        self.status.side_effect = AssertionError("receipt recovery must not recheck balance")
        result = self.run_work(doc, retry=True)
        self.assertEqual(result["billing_state"], "charged")
        self.assertEqual(self.charge.call_args, expected)
        self.assertEqual(self.http.call_count, 2)

    def test_revocation_between_sheets_blocks_remaining_calls(self):
        self.upload()
        def revoke(*args, **kwargs):
            frappe.db.delete("User Permission", {"name": self.grant})
            frappe.db.commit()
            return response()
        self.http.side_effect = revoke
        doc = self.admit()
        with self.assertRaises(frappe.PermissionError):
            self.work.run(doc.work_name)
        self.assertEqual(frappe.db.get_value("Pearnly OCR Work", doc.work_name, "status"), "failed")
        with self.assertRaises(frappe.PermissionError):
            self.api.detail(doc.name)
        self.assertEqual(self.http.call_count, 1)
        self.charge.assert_not_called()


if __name__ == "__main__":
    connect()
    try:
        if "--stale-admit-child" in sys.argv:
            value = json.loads(sys.stdin.readline())
            frappe.conf.update(value["conf"])
            frappe.set_user(value["actor"])
            # Establish MariaDB RR's consistent read before the other admission
            # commits. The source row is deliberately not locked in this phase.
            frappe.db.sql("SELECT name FROM `tabCowork OCR Batch` WHERE source_file=%s", value["source"])
            print("READY", flush=True)
            if sys.stdin.readline().strip() != "GO":
                raise RuntimeError("Missing local test synchronization")
            assert frappe.db.sql("SELECT @@session.tx_isolation")[0][0] == "REPEATABLE-READ"
            assert not frappe.db.sql("SELECT name FROM `tabCowork OCR Batch` WHERE source_file=%s", value["source"])
            from pearnly_cowork.ocr import api
            with patch("pearnly_core.billing_client.status", return_value={"allowed": True}), patch("pearnly_core.execution.work.kick"), patch.object(api.service, "submit", wraps=api.service.submit) as attempts:
                result = api.submit(value["source"], value["key"], direction="purchase")
                print(json.dumps({**result, "admission_attempts": attempts.call_count}))
        else:
            suite = (unittest.TestSuite([NativeTableExecution(sys.argv[1])]) if len(sys.argv) > 1
                else unittest.defaultTestLoader.loadTestsFromTestCase(NativeTableExecution))
            result = unittest.TextTestRunner(verbosity=2).run(suite)
            raise SystemExit(0 if result.wasSuccessful() else 1)
    finally:
        frappe.destroy()
