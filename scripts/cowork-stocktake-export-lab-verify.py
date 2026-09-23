"""Authorized complete XLSX downloads on the isolated Cowork lab; zero network.

Writes synthetic stocktakes and access logs only. Uses real Frappe/WSGI/DB.
"""
from io import BytesIO
import json
from pathlib import Path
import runpy
import unittest
from uuid import uuid4

import frappe
from openpyxl import load_workbook

FIXTURE = runpy.run_path(str(Path(__file__).with_name("cowork-stocktake-lab-verify.py")))
BASE = FIXTURE["NativeStocktake"]
connect = FIXTURE["connect"]


class NativeStocktakeExport(unittest.TestCase):
    setUpClass = classmethod(BASE.setUpClass.__func__)
    setUp, tearDown, create, add, error = BASE.setUp, BASE.tearDown, BASE.create, BASE.add, BASE.error

    def book(self, name):
        from pearnly_cowork.stocktake.service import export
        return load_workbook(BytesIO(export(name)))

    def test_complete_closed_export_includes_all_entries_and_void_history(self):
        task = self.create()
        first = self.add(task, "1")
        self.api.edit_entry(task["id"], first["entry_id"], str(uuid4()), "1", "A", 0, "01", True)
        for _ in range(100):
            self.add(task, "0.000001")
        preview = self.api.detail(task["id"])
        self.assertEqual((len(preview["entries"]), preview["entry_total"]), (100, 101))
        self.api.close(task["id"])
        book = self.book(task["id"])
        self.assertEqual(book.sheetnames, ["สรุปผลตรวจนับ", "รายละเอียดการนับ"])
        summary, detail = book.worksheets
        self.assertEqual((summary.max_row, detail.max_row), (3, 102))
        self.assertEqual(summary["A2"].value, "0012")
        self.assertEqual(summary["F2"].value, 0.0001)
        self.assertEqual(summary["G2"].value, -2.499901)
        self.assertEqual(sum(detail.cell(i, 10).value == "ยกเลิก (ไม่รวมยอด)" for i in range(2, 103)), 1)
        self.assertEqual(summary.freeze_panes, "A2")
        self.assertEqual(detail.auto_filter.ref, "A1:Q102")

    def test_current_scope_user_feature_and_mode_rechecked(self):
        task = self.create()
        with self.assertRaises(frappe.PermissionError):
            self.book(self.foreign_reference)
        frappe.conf.pearnly_features = [f for f in self.original_features if f != "stocktake"]
        with self.assertRaises(frappe.PermissionError):
            self.book(task["id"])
        frappe.conf.pearnly_features = self.original_features
        frappe.db.set_value("User", self.actor, "enabled", 0)
        with self.assertRaises(frappe.PermissionError):
            self.book(task["id"])
        frappe.db.set_value("User", self.actor, "enabled", 1)
        frappe.db.set_value("Cowork Stocktake Task", task["id"], "count_mode", "legacy")
        with self.error("stocktake.legacy_task"):
            self.book(task["id"])

    def test_download_http_headers_auth_and_formula_text(self):
        from werkzeug.test import Client
        from werkzeug.wrappers import Response
        from frappe.app import application
        data = FIXTURE["rows"]()
        data[0][1]["value"] = "=HYPERLINK(\"https://invalid.example\",\"x\")"
        task = self.create(data=data)
        self.add(task)
        frappe.set_user("Administrator")
        account = frappe.get_doc("User", self.actor)
        account.api_key, secret = uuid4().hex, uuid4().hex
        account.api_secret = secret
        account.save()
        headers = {"Authorization": f"token {account.api_key}:{secret}"}
        frappe.db.commit()
        frappe.destroy()
        client = Client(application, Response)
        url = "/api/method/pearnly_cowork.stocktake.api.export?task_id="
        def get(name, authenticated=True):
            return client.get(url + name, base_url="http://" + FIXTURE["SITE"],
                              headers=headers if authenticated else {})
        try:
            response = get(task["id"])
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True) if response.status_code != 200 else "")
            self.assertIn("attachment", response.headers["Content-Disposition"])
            self.assertIn("cowork-stocktake.xlsx", response.headers["Content-Disposition"])
            self.assertIn("no-store", response.headers["Cache-Control"])
            self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
            self.assertIn("spreadsheetml.sheet", response.headers["Content-Type"])
            book = load_workbook(BytesIO(response.data))
            self.assertEqual(book.worksheets[0]["B2"].value, data[0][1]["value"])
            self.assertEqual(book.worksheets[0]["B2"].data_type, "s")
            self.assertEqual(get(task["id"], False).status_code, 403)
            self.assertEqual(get(self.foreign_reference).status_code, 403)
            rejected = client.post(url + task["id"], base_url="http://" + FIXTURE["SITE"], headers=headers)
            self.assertEqual(rejected.status_code, 403)
        finally:
            connect()
        # Frappe's real request path defers Access Log writes to its per-Site
        # Redis queue; the isolated lab intentionally has no scheduler.
        queued = [json.loads(value) for value in frappe.cache.lrange("insert_queue_for_Access Log", 0, -1)]
        rows = [row for batch in queued for row in (batch if isinstance(batch, list) else [batch])]
        self.assertTrue(any(row.get("user") == self.actor and row.get("export_from") == "Cowork Stocktake Task"
            and row.get("reference_document") == task["id"] and row.get("file_type") == "stocktake_workbook"
            for row in rows))


if __name__ == "__main__":
    connect()
    try:
        result = unittest.TextTestRunner(verbosity=2).run(
            unittest.defaultTestLoader.loadTestsFromTestCase(NativeStocktakeExport))
        print(json.dumps({"stocktake_export_native": "PASS" if result.wasSuccessful() else "FAIL",
                          "tests": result.testsRun}))
        raise SystemExit(0 if result.wasSuccessful() else 1)
    finally:
        if getattr(frappe.local, "db", None):
            frappe.db.rollback()
        frappe.destroy()
