"""Bounded review-export checks on the isolated native Cowork lab.

Synthetic local records and machine ACKs only. Reuses fixture helpers, not prior
suites. No ERP call, LINE message, worker, scheduler or external credentials.
"""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import runpy
import sys
import unittest
from uuid import uuid4

import frappe
from openpyxl import load_workbook

BASE = runpy.run_path(str(Path(__file__).with_name('cowork-ack-effects-lab-verify.py')))
A = BASE['AckEffects']
P = BASE['PREP']
connect = BASE['connect']


class ExportFacts(unittest.TestCase):
    setUpClass = classmethod(A.setUpClass.__func__)
    tearDown, install, prepare, report, confirm, row = A.tearDown, A.install, A.prepare, A.report, A.confirm, A.row
    claimed, ack, count = A.claimed, A.ack, A.count

    def setUp(self):
        A.setUp(self)
        from pearnly_cowork.roundtrip import service, spec
        from pearnly_cowork.push import export_facts
        self.exporter, self.spec, self.facts = service, spec, export_facts
        pages = deepcopy(self.history['pages'])
        pages[0]['fields'].update(erp_docnum='FORGED', push_status='FORGED',
            items=[{'name': 'Synthetic first', 'qty': '1', 'unit_price': '50', 'subtotal': '50', 'erp_item_code': 'FORGED'},
                   {'name': 'Synthetic second', 'qty': '1', 'unit_price': '50', 'subtotal': '50'}])
        self.history = self.storage.update_pages(self.history['name'], pages, self.history['record_version'])

    def success(self, number='SYNTHETIC-ERP-001'):
        row = self.prepare()
        self.assertEqual(row['status'], 'prepared', row['reason'])
        self.confirm(row)
        handle = self.transport.lease(self.machine, {})['jobs'][0]['log_id']
        self.ack(handle, express_docnum=number,
            line_modes=[{'seq': 2, 'mode': 'stock_item', 'created': False, 'stkcod': 'ITEM-OLD'},
                        {'seq': 1, 'mode': 'stock_item', 'created': True, 'stkcod': 'ITEM-NEW'}],
            meta={'created_supplier': True, 'party_code': 'SUP-NEW', 'doc_type': 'purchase',
                  'unknown_secret': 'DO-NOT-EXPORT'})
        return row

    def actions(self):
        return {name: facts['actions'] for name, facts in self.facts.for_histories(
            [self.storage.read(self.history['name'])]).items() if 'actions' in facts}

    def workbook(self):
        return load_workbook(BytesIO(self.exporter.export([self.history['name']])), data_only=False)

    def new_connection(self):
        frappe.set_user(self.manager)
        self.connection = self.connections.configure(self.workspace, 'express', uuid4().hex,
            'Synthetic second connection', P['BASE']['CONFIG'], users=[self.actor, self.manager])
        self.install()
        self.issued = self.credentials.issue(self.connection['name'], uuid4().hex)
        self.machine = self.credentials.authenticate(self.issued['token'])
        self.report()
        self.key = uuid4().hex
        frappe.set_user(self.actor)

    def test_success_download_uses_receipt_and_created_master_summary(self):
        self.success()
        before = self.storage.read(self.history['name']).pages
        actions = self.actions()[self.history['name']]
        self.assertEqual((actions['items'], actions['item_codes']), ([True, False], ['ITEM-NEW', 'ITEM-OLD']))
        from pearnly_cowork.roundtrip import api
        api.export([self.history['name']])
        self.assertTrue(frappe.local.cowork_private_response)
        wb = load_workbook(BytesIO(frappe.local.response.filecontent))
        values = [str(v) for sheet in wb for row in sheet.values for v in row if v is not None]
        for expected in ('SYNTHETIC-ERP-001', 'SUP-NEW', 'ITEM-NEW', 'Synthetic first', 'สำเร็จ'):
            self.assertIn(expected, values)
        for secret in ('FORGED', 'DO-NOT-EXPORT', self.issued['token']):
            self.assertNotIn(secret, str(values))
        self.assertEqual(wb[self.spec.SHEET_SUMMARY].max_row, 3)
        self.assertEqual(self.storage.read(self.history['name']).pages, before)
        wb.close()

    def test_latest_failure_keeps_prior_success_reference(self):
        first = self.success()
        self.new_connection()
        second, handle = self.claimed(line=False)
        self.ack(handle, result='failed', outcome='needs_review', express_docnum='FAILED-NOT-A-REFERENCE')
        self.assertNotEqual(first['name'], second['name'])
        self.assertEqual(self.actions()[self.history['name']]['docnum'], 'SYNTHETIC-ERP-001')
        facts = self.facts.for_histories([self.storage.read(self.history['name'])])
        self.assertEqual(facts[self.history['name']]['state'],
                         {'status': 'manual', 'reason': self.row(second).reason})
        self.assertTrue(self.row(second).reason)
        wb = self.workbook()
        values = str([list(sheet.values) for sheet in wb])
        self.assertIn('SYNTHETIC-ERP-001', values)
        self.assertNotIn('FAILED-NOT-A-REFERENCE', values)
        self.assertNotIn('สำเร็จ', values)
        wb.close()

    def test_latest_success_wins_and_prepared_does_not_hide_it(self):
        self.assertEqual(self.prepare()['status'], 'prepared')
        original = self.connection, self.issued, self.machine, self.key
        self.new_connection()
        self.success('CREATED-LATER-SUCCESS')
        self.connection, self.issued, self.machine, self.key = original
        self.success('COMPLETED-LATER-SUCCESS')
        self.new_connection()
        self.assertEqual(self.prepare()['status'], 'prepared')
        # Match source order by record creation, not completion or selection.
        self.assertEqual(self.actions()[self.history['name']]['docnum'], 'CREATED-LATER-SUCCESS')

    def test_disabled_connection_can_read_but_revoked_grant_cannot(self):
        self.success()
        frappe.set_user(self.manager)
        self.connection = self.connections.configure(self.workspace, 'express', uuid4().hex,
            'Disabled', P['BASE']['CONFIG'], name=self.connection['name'],
            expected_revision=self.connection['revision'], users=[self.actor, self.manager], enabled=False)
        frappe.set_user(self.actor)
        self.assertIn(self.history['name'], self.actions())
        frappe.set_user(self.manager)
        self.connections.configure(self.workspace, 'express', uuid4().hex,
            'Revoked', P['BASE']['CONFIG'], name=self.connection['name'],
            expected_revision=self.connection['revision'], users=[self.manager], enabled=False)
        frappe.set_user(self.actor)
        self.assertEqual(self.actions(), {})
        wb = self.workbook()
        self.assertNotIn('SYNTHETIC-ERP-001', str([list(s.values) for s in wb]))
        wb.close()

    def test_log_role_is_required_without_breaking_invoice_export(self):
        self.success()
        frappe.db.sql("DELETE FROM `tabHas Role` WHERE parent=%s AND role='Cowork Push Log'", (self.actor,))
        frappe.clear_cache(user=self.actor)
        self.assertEqual(self.actions(), {})
        wb = self.workbook()
        self.assertEqual(wb[self.spec.SHEET_SUMMARY].max_row, 1)
        wb.close()

    def test_actor_and_history_scope_are_independent(self):
        self.success()
        # Admin has both workspace access and an explicit connection grant, but
        # does not inherit another operator's ERP actions in their own export.
        frappe.set_user(self.manager)
        self.assertEqual(self.actions(), {})
        frappe.set_user(self.other)
        with self.assertRaises(frappe.PermissionError): self.workbook()

    def test_edit_keeps_reference_without_mislabeling_new_lines(self):
        self.success()
        pages = deepcopy(self.history['pages']); pages[0]['fields']['items'].reverse()
        self.history = self.storage.update_pages(self.history['name'], pages, self.history['record_version'])
        actions = self.actions()[self.history['name']]
        self.assertEqual(actions['docnum'], 'SYNTHETIC-ERP-001')
        self.assertEqual(actions['items'], [])
        wb = self.workbook()
        self.assertEqual(wb[self.spec.SHEET_SUMMARY].max_row, 1)
        self.assertNotIn('ITEM-NEW', str([list(s.values) for s in wb]))
        wb.close()

    def test_stock_group_labels_require_exact_report_origin(self):
        group = {'acccod': 'DM', 'stock_acc': '1100', 'stock_acc_name': 'Synthetic inventory', 'fit': True}
        self.report(catalog={'products': [], 'customers': [], 'isacc_candidates': [group]})
        row = self.success()
        target = json.loads(self.row(row).target_json)
        groups = self.connections.stock_groups_for_log(self.connection['name'], target)
        self.assertEqual(groups[0]['stock_acc_name'], group['stock_acc_name'])
        target['account_set'] = 'OTHER-ACCOUNT'
        self.assertEqual(self.connections.stock_groups_for_log(self.connection['name'], target), [])
        target = json.loads(self.row(row).target_json)
        target['preparation_evidence']['credential']['credential_id'] = 'OTHER-CREDENTIAL'
        self.assertEqual(self.connections.stock_groups_for_log(self.connection['name'], target), [])


if __name__ == '__main__':
    connect()
    try:
        suite = (unittest.TestSuite(ExportFacts(name) for name in sys.argv[1:]) if len(sys.argv)>1
                 else unittest.defaultTestLoader.loadTestsFromTestCase(ExportFacts))
        result = unittest.TextTestRunner(verbosity=2).run(suite)
        if not result.wasSuccessful(): raise SystemExit(1)
    finally:
        frappe.db.rollback(); frappe.destroy()
