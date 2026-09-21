"""Verify the ERP product master on the dedicated laboratory.

Covers create / list / update / soft-delete, duplicate code and barcode rejection,
per-job access (a shop role may manage products but must not see cost), and tenant
isolation. Synthetic data and lab volume only.
"""
import json
import os
from pathlib import Path
import re
import sys

import frappe

CONTROL_APP = "/home/frappe/frappe-bench/apps/pearnly_control"
SITE = os.environ.get("PEARNLY_ERP_POS_LAB_SITE", "erp-pos-a.localhost")
OTHER_SITE = "erp-pos-b.localhost" if SITE != "erp-pos-b.localhost" else "erp-pos-a.localhost"
ALLOWED_SITES = {"erp-pos-a.localhost", "erp-pos-b.localhost"}
SITE_ROOT = Path("/home/frappe/frappe-bench/sites")

OWNER = "erp-pos-owner@example.com"
SHOP = "erp-pos-products@example.com"
STOCK = "erp-pos-stock@example.com"
OUTSIDER = "erp-pos-outsider@example.com"


def connect(site=SITE):
    if os.environ.get("PEARNLY_ERP_POS_LAB") != "1":
        raise SystemExit("Dedicated ERP/POS laboratory required")
    if site not in ALLOWED_SITES or not (SITE_ROOT / ".erp-pos-lab").is_file():
        raise SystemExit("Dedicated lab Site marker required")
    os.chdir(SITE_ROOT)
    frappe.init(site=site, sites_path=str(SITE_ROOT))
    frappe.connect()
    if not frappe.conf.get("pearnly_erp_pos_lab") or frappe.conf.get("db_host") != "db":
        raise SystemExit("Refusing a non-lab Site")
    frappe.flags.mute_emails = True
    return site


def as_user(user, func, *args, **kwargs):
    frappe.set_user(user)
    try:
        return True, func(*args, **kwargs)
    except Exception as exc:  # noqa: BLE001 - the exception type is the evidence
        return False, type(exc).__name__


def ensure_member(email, roles):
    if frappe.db.exists("User", email):
        user = frappe.get_doc("User", email)
        # 只重设 ERP 岗位，保留框架角色；child table 必须用 append 加行。
        user.roles = [r for r in user.roles if not r.role.startswith("ERP ")]
        have = {r.role for r in user.roles}
        if "Desk User" not in have:
            user.append("roles", {"role": "Desk User"})
        for role in roles:
            user.append("roles", {"role": role})
        user.save(ignore_permissions=True)
    else:
        frappe.get_doc({"doctype": "User", "email": email, "first_name": email.split("@")[0],
                        "send_welcome_email": 0,
                        "roles": [{"role": "Desk User"}, *({"role": r} for r in roles)]
                        }).insert(ignore_permissions=True)
    frappe.db.commit()
    frappe.clear_cache(user=email)


def main():
    site = connect()
    frappe.set_user("Administrator")
    from pearnly_core.provisioning import ensure_feature_roles
    from pearnly_erp_portal.products import service

    if CONTROL_APP not in sys.path:
        sys.path.insert(0, CONTROL_APP)
    from pearnly_control import features as control_features

    spec = control_features.roles_payload("erp", control_features.codes_of("erp"))
    ensure_feature_roles(json.dumps(spec), admin_user=OWNER)
    ensure_member(SHOP, ["ERP Products"])
    ensure_member(STOCK, ["ERP Stock"])
    ensure_member(OUTSIDER, [])

    results = {}
    results["spec_roles"] = sorted(r["role"] for r in spec["roles"])

    existing = frappe.db.get_value("ERP Product", {"code": "LAB-P1"}, "name")
    if existing:
        frappe.delete_doc("ERP Product", existing, force=True, ignore_permissions=True)
        frappe.db.commit()

    # 老板建品（数量用 Decimal 语义，价格留空不等于 0）。
    ok, created = as_user(OWNER, service.create_product, product_name="ผงซักฟอก",
                          code="LAB-P1", barcode="8850000000001", unit="ถุง",
                          unit_price="25", min_stock="5", default_cost="18")
    results["owner_create"] = ok
    results["created"] = created
    product = created["name"] if ok else None

    # 撞码 / 撞条码分别报错，且两处都不能绕过。
    ok, err = as_user(OWNER, service.create_product, product_name="ซ้ำรหัส", code="LAB-P1")
    results["duplicate_code_allowed"] = ok
    results["duplicate_code_error"] = None if ok else err
    ok, err = as_user(OWNER, service.create_product, product_name="ซ้ำบาร์โค้ด",
                      barcode="8850000000001")
    results["duplicate_barcode_allowed"] = ok
    results["duplicate_barcode_error"] = None if ok else err

    # 编码是系统发的（旧线上 ERP 门户里这一格只读）：留空建品拿到 P-000001 起的号。
    ok, auto = as_user(OWNER, service.create_product, product_name="Lab auto code")
    results["auto_code_ok"] = ok
    auto_name = auto["name"] if ok else None
    auto_code = frappe.db.get_value("ERP Product", auto_name, "code") if auto_name else None
    results["auto_code_allocated"] = bool(re.fullmatch(r"P-\d{6}", auto_code or ""))
    # 编码锁死：改成别的号要被挡下，否则历史单据行会指向另一个货。
    ok, err = as_user(OWNER, service.update_product, product, code="LAB-P2")
    results["code_change_allowed"] = ok
    results["code_change_error"] = None if ok else err
    results["code_after_change_attempt"] = frappe.db.get_value("ERP Product", product, "code")

    # 商品岗位能管商品，但成本对它遮蔽；库存岗位能看到成本。
    ok, shop_list = as_user(SHOP, service.list_products, limit=50)
    results["shop_can_list"] = ok
    shop_row = next((r for r in shop_list if r["name"] == product), None) if ok else None
    results["shop_sees_cost"] = None if shop_row is None else shop_row["default_cost"]
    ok, stock_list = as_user(STOCK, service.list_products, limit=50)
    results["stock_can_list"] = ok
    stock_row = next((r for r in stock_list if r["name"] == product), None) if ok else None
    results["stock_sees_cost"] = None if stock_row is None else stock_row["default_cost"]
    # 库存岗位能看商品，但不能改商品主数据。
    ok, err = as_user(STOCK, service.update_product, product, unit_price="30")
    results["stock_can_edit_products"] = ok
    results["stock_edit_error"] = None if ok else err
    ok, owner_list = as_user(OWNER, service.list_products, limit=50)
    owner_row = next((r for r in owner_list if r["name"] == product), None) if ok else None
    results["owner_sees_cost"] = None if owner_row is None else owner_row["default_cost"]

    # 没有商品岗位的账号读不到列表。
    ok, err = as_user(OUTSIDER, service.list_products, limit=50)
    results["outsider_can_list"] = ok
    results["outsider_list_error"] = None if ok else err

    # 商品岗位不能管理团队（能力分开）。
    from pearnly_erp_portal.team import service as team_service
    ok, err = as_user(SHOP, team_service.list_members)
    results["shop_manages_team"] = ok

    # 单位：换算系数必须为正；默认销售单位只留一个。
    ok, units = as_user(SHOP, service.add_unit, product, unit_name="โหล",
                        factor_to_base="12", barcode="8850000000018", price="280",
                        is_default_sell=True)
    results["add_unit"] = ok
    ok, err = as_user(SHOP, service.add_unit, product, unit_name="ผิด", factor_to_base="0")
    results["zero_factor_allowed"] = ok
    ok, second = as_user(SHOP, service.add_unit, product, unit_name="ลัง",
                         factor_to_base="144", is_default_sell=True)
    doc = frappe.get_doc("ERP Product", product)
    results["default_sell_units"] = len([r for r in doc.units if r.is_default_sell])

    # 停用是软删除：列表默认看不到，但仍能按名字取回并重新上架。
    ok, off = as_user(SHOP, service.set_active, product, 0)
    results["deactivate"] = off
    ok, listed = as_user(SHOP, service.list_products, limit=50)
    results["hidden_after_deactivate"] = all(r["name"] != product for r in listed)
    ok, back = as_user(SHOP, service.set_active, product, 1)
    results["reactivate"] = back

    # 改价用空值 = 清空，不是写 0。
    ok, _ = as_user(SHOP, service.update_product, product, unit_price="")
    results["price_after_blank"] = frappe.db.get_value("ERP Product", product, "unit_price")
    results["price_set_after_blank"] = bool(
        frappe.db.get_value("ERP Product", product, "price_set"))
    ok, _ = as_user(SHOP, service.update_product, product, unit_price="0")
    results["price_set_after_real_zero"] = bool(
        frappe.db.get_value("ERP Product", product, "price_set"))
    # 分类：Frappe 的 Int 列是 NOT NULL DEFAULT 0，"没填"在库里就是 0（存不下 NULL，
    # 与 Currency 的"没设价"同一类事实）；填了号就存号。
    ok, _ = as_user(SHOP, service.update_product, product, category_id="0")
    results["category_zero_stored"] = frappe.db.get_value("ERP Product", product, "category_id")
    ok, _ = as_user(SHOP, service.update_product, product, category_id="7")
    results["category_seven_stored"] = frappe.db.get_value("ERP Product", product, "category_id")

    frappe.set_user("Administrator")

    # 干净起点：上一轮导入留下的行（编码规则是本脚本自己的），否则第二次跑会撞唯一键。
    for stale in frappe.get_all("ERP Product", filters={"code": ["like", "LAB-IMP-%"]},
                                pluck="name"):
        frappe.delete_doc("ERP Product", stale, force=True, ignore_permissions=True)
    for stale in frappe.get_all("ERP Product", filters={"product_name": "Lab auto code"},
                                pluck="name"):
        frappe.delete_doc("ERP Product", stale, force=True, ignore_permissions=True)
    for stale in frappe.get_all("ERP Product", filters={"product_name": "Lab template row"},
                                pluck="name"):
        frappe.delete_doc("ERP Product", stale, force=True, ignore_permissions=True)
    for stale in frappe.get_all("ERP Product", filters={"product_name": "Lab image row"},
                                pluck="name"):
        frappe.delete_doc("ERP Product", stale, force=True, ignore_permissions=True)
    for stale in frappe.get_all("File", filters={"file_name": "lab-product-image.png"},
                                pluck="name"):
        frappe.delete_doc("File", stale, force=True, ignore_permissions=True)
    for stale in frappe.get_all("ERP Product", filters={"product_name": "Lab image outside"},
                                pluck="name"):
        frappe.delete_doc("ERP Product", stale, force=True, ignore_permissions=True)
    frappe.db.commit()

    # Excel 导入：解析在纯层（bench 里有 openpyxl），入库走同一条 create_product。
    import io

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["商品名称", "รหัส", "单价", "vat", "批次"])
    sheet.append(["Lab import A", "LAB-IMP-1", "12.5", "ใช่", "1"])
    sheet.append(["Lab import C", "", "7", "", ""])
    sheet.append(["Lab import B", "", "ไม่ใช่ตัวเลข", "", ""])
    sheet.append([None, None, None, None, None])
    buffer = io.BytesIO()
    book.save(buffer)
    ok, imported = as_user(OWNER, service.import_products, buffer.getvalue())
    results["import_ok"] = ok
    results["import_error"] = None if ok else imported
    results["import_created"] = imported.get("created") if ok else None
    results["import_failed"] = imported.get("failed") if ok else None
    results["import_error_rows"] = [row["row"] for row in imported.get("errors", [])] if ok else None
    # 行级错误给用户看的是文案，不是解析层的 token。
    results["import_raw_tokens_leaked"] = [row["error"] for row in imported.get("errors", [])
                                           if row["error"] in ("not_a_number",
                                                               "negative_or_invalid",
                                                               "name_required")] if ok else None
    imported_name = frappe.db.get_value("ERP Product", {"code": "LAB-IMP-1"}, "name")
    results["import_flags"] = frappe.db.get_value(
        "ERP Product", imported_name, ["vat_applicable", "track_batch", "base_unit"],
        as_dict=True) if imported_name else None
    # 表里没带码的行同样由系统发号（旧线上 ERP 导入也走这条 `prepare()`）。
    results["import_auto_code"] = bool(frappe.db.exists(
        "ERP Product", {"product_name": "Lab import C", "code": ["like", "P-%"]}))

    # 模板下载：表头按站点语言现生成。下载 → 填一行 → 传回来必须能认（闭环）。
    ok, _ = as_user(OWNER, service.api_import_template)
    downloaded = frappe.local.response.get("filecontent") or b""
    results["template_download_ok"] = bool(ok and downloaded)
    book = openpyxl.load_workbook(io.BytesIO(downloaded), data_only=True)
    sheet = book.active
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    results["template_headers"] = len(headers)
    # 语言包里缺哪条，模板上就会露出键名 —— 露出来就是漏翻。
    results["template_keys_leaked"] = [header for header in headers
                                       if str(header).startswith("product.")]
    sheet.append(["Lab template row"])
    buffer = io.BytesIO()
    book.save(buffer)
    ok, templated = as_user(OWNER, service.import_products, buffer.getvalue())
    results["template_import_ok"] = ok
    results["template_created"] = templated.get("created") if ok else None
    results["template_failed"] = templated.get("failed") if ok else None
    # 空模板（一行数据都没有）要明确报"没有数据行"，不能默默什么都不做。
    ok, empty = as_user(OWNER, service.import_products, downloaded)
    results["template_empty_allowed"] = ok
    results["template_empty_error"] = None if ok else empty

    # 商品图：私有文件挂在商品上，读权限才跟着商品走。
    from frappe.utils.file_manager import save_file

    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000a49444154789c6360000002000154a24f5f0000000049454e44ae426082")
    lab_file = save_file("lab-product-image.png", png, None, None, is_private=1)
    ok, with_image = as_user(OWNER, service.create_product, product_name="Lab image row",
                             image_url=lab_file.file_url)
    results["image_product_ok"] = ok
    attached = frappe.db.get_value("File", lab_file.name,
                                   ["attached_to_doctype", "attached_to_name"], as_dict=True)
    results["image_attached_doctype"] = (attached or {}).get("attached_to_doctype")
    results["image_attached_same_row"] = bool(
        attached and with_image and attached.attached_to_name == with_image.get("name"))
    # 外链不碰：不是本站 File 的地址照原样存。
    ok, outside = as_user(OWNER, service.create_product, product_name="Lab image outside",
                          image_url="https://example.com/x.png")
    results["image_outside_ok"] = ok
    results["image_outside_value"] = frappe.db.get_value(
        "ERP Product", (outside or {}).get("name"), "image_url") if ok else None

    # 扫码当场查重：编码 / 条码 / 二维码内容 / 售卖单位条码都指向同一条，自己不算撞自己。
    ok, _ = as_user(SHOP, service.update_product, product, qr_payload="PEARNLY-LAB-QR")
    ok, by_barcode = as_user(SHOP, service.find_product, "8850000000001")
    results["find_by_barcode"] = bool(ok and (by_barcode or {}).get("name") == product)
    ok, by_code = as_user(SHOP, service.find_product, "LAB-P1")
    results["find_by_code"] = bool(ok and (by_code or {}).get("matched_by") == "code")
    ok, by_qr = as_user(SHOP, service.find_product, "PEARNLY-LAB-QR")
    results["find_by_qr"] = bool(ok and (by_qr or {}).get("matched_by") == "qr_payload")
    ok, by_unit = as_user(SHOP, service.find_product, "8850000000018")
    results["find_by_unit_barcode"] = bool(ok and (by_unit or {}).get("matched_by") == "unit_barcode")
    ok, self_hit = as_user(SHOP, service.find_product, "LAB-P1", product)
    results["find_excludes_self"] = bool(ok and self_hit == {})
    ok, missing = as_user(SHOP, service.find_product, "NO-SUCH-CODE")
    results["find_missing_empty"] = bool(ok and missing == {})
    # 查重与列表同一道权限：没有商品岗位的账号查不了。
    ok, err = as_user(OUTSIDER, service.find_product, "LAB-P1")
    results["outsider_can_find"] = ok

    # 字段对齐：与旧线上 products 的列一一对应（不多不少）。
    meta = frappe.get_meta("ERP Product")
    ours = {field.fieldname for field in meta.fields
            if field.fieldtype not in ("Section Break", "Column Break", "Tab Break")}
    # `name` 是框架隐式主键，不在 meta.fields 里，所以不列进来。
    expected = {"product_name", "code", "barcode",
                "qr_payload", "unit", "base_unit", "category_id", "unit_price",
                "price_set", "cost_set", "vat_applicable", "is_active", "min_stock",
                "default_cost", "image_url", "track_batch", "track_expiry",
                "units"}
    results["fields_missing"] = sorted(expected - ours)
    results["fields_extra"] = sorted(ours - expected)

    results["row_count_site_a"] = frappe.db.count("ERP Product")

    frappe.destroy()
    connect(OTHER_SITE)
    results["other_site_product_rows"] = frappe.db.count("ERP Product")
    results["other_site_has_code"] = bool(
        frappe.db.exists("ERP Product", {"code": "LAB-P1"}))

    expected = {
        "owner_create": True,
        "duplicate_code_allowed": False,
        "duplicate_code_error": "ValidationError",
        "duplicate_barcode_allowed": False,
        "duplicate_barcode_error": "ValidationError",
        "auto_code_ok": True,
        "auto_code_allocated": True,
        "code_change_allowed": False,
        "code_change_error": "ValidationError",
        "code_after_change_attempt": "LAB-P1",
        "shop_can_list": True,
        "shop_sees_cost": None,
        "stock_sees_cost": 18.0,
        "owner_sees_cost": 18.0,
        "outsider_can_list": False,
        "outsider_list_error": "PermissionError",
        "shop_manages_team": False,
        "add_unit": True,
        "zero_factor_allowed": False,
        "default_sell_units": 1,
        "hidden_after_deactivate": True,
        # Currency 列存不了 NULL，"没设价"由 price_set 标记承担（下一步断言）。
        "price_after_blank": 0.0,
        "price_set_after_blank": False,
        "price_set_after_real_zero": True,
        "category_zero_stored": 0,
        "category_seven_stored": 7,
        "import_ok": True, "import_created": 2, "import_failed": 1, "import_error_rows": [4],
        "import_raw_tokens_leaked": [],
        "import_auto_code": True,
        "template_download_ok": True,
        "template_headers": 14,
        "template_keys_leaked": [],
        "template_import_ok": True,
        "template_created": 1,
        "template_failed": 0,
        "template_empty_allowed": False,
        "template_empty_error": "ValidationError",
        "image_product_ok": True,
        "image_attached_doctype": "ERP Product",
        "image_attached_same_row": True,
        "image_outside_ok": True,
        "image_outside_value": "https://example.com/x.png",
        "find_by_barcode": True,
        "find_by_code": True,
        "find_by_qr": True,
        "find_by_unit_barcode": True,
        "find_excludes_self": True,
        "find_missing_empty": True,
        "outsider_can_find": False,
        "fields_missing": [], "fields_extra": [],
        "other_site_product_rows": 0,
        "other_site_has_code": False,
        "stock_can_list": True,
        "stock_can_edit_products": False,
        "stock_edit_error": "PermissionError",
    }
    problems = {k: {"expected": v, "actual": results.get(k)} for k, v in expected.items()
                if results.get(k) != v}
    for key in ("deactivate", "reactivate"):
        if not (results.get(key) or {}).get("is_active") == (key == "reactivate"):
            problems[key] = results.get(key)
    if "ERP Products" not in (results.get("spec_roles") or []):
        problems["spec_missing_products_role"] = results.get("spec_roles")

    print(json.dumps({"site": site, "results": results, "problems": problems},
                     ensure_ascii=False, indent=2, default=str))
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
