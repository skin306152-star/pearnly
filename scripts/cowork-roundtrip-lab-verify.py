"""Synthetic review-workbook integration on the guarded, offline Cowork Site.

Real native queue/DB/permissions/history; fake billing and forbidden provider IO.
Preserves labelled fixtures, never deletes history or invokes an ERP adapter.
"""
import io
import json
from pathlib import Path
import runpy
import sys
import unittest
import uuid
from unittest.mock import patch

import frappe
from openpyxl import Workbook, load_workbook

BASE = runpy.run_path(str(Path(__file__).with_name("cowork-execution-lab-verify.py")))
NativeExecution = BASE["NativeExecution"]
connect, grant = BASE["connect"], BASE["grant"]


def book(rows):
    from pearnly_cowork.roundtrip import spec as rt
    wb = Workbook()
    wb.remove(wb.active)
    for title, values in rows:
        ws = wb.create_sheet(title)
        headers = list(dict.fromkeys([*rt.ROUNDTRIP_HEADERS,
            rt.SALES_COL_DATE, rt.SALES_COL_INVOICE, rt.SALES_COL_PARTY, rt.SALES_COL_ITEM,
            rt.SALES_COL_QTY, rt.SALES_COL_PRICE, rt.SALES_COL_AMOUNT, rt.SALES_COL_PRE_VAT,
            *rt.PURCHASE_HEADERS, *rt.PENDING_HEADERS]))
        ws.append(headers)
        for value in values:
            ws.append([value.get(h) for h in headers])
    wb.create_sheet(rt.SHEET_SUMMARY).append(["Unassigned synthetic original content"])
    result = io.BytesIO()
    wb.save(result)
    return result.getvalue()


class NativeRoundtrip(unittest.TestCase):
    setUpClass = classmethod(NativeExecution.setUpClass.__func__)
    setUp, tearDown = NativeExecution.setUp, NativeExecution.tearDown
    admit, run_work = NativeExecution.admit, NativeExecution.run_work

    def prior(self, direction="sales", workspace=None):
        from pearnly_cowork.history import storage
        fields = {**BASE["GOOD"], "direction": direction, "buyer_name": "Synthetic buyer",
            "items": [{"name": "Synthetic item", "qty": "1", "price": "100", "subtotal": "100"}]}
        return storage.create_manual(workspace or self.workspace, [{"page_index": 1, "fields": fields}], uuid.uuid4().hex)

    def row(self, prior=None, **overrides):
        from pearnly_cowork.roundtrip import spec as rt
        row = {rt.SALES_COL_INVOICE: "SAME-001", rt.SALES_COL_DATE: "2026-09-21",
            rt.SALES_COL_PARTY: "Synthetic buyer", rt.SALES_COL_ITEM: "Synthetic item",
            rt.SALES_COL_QTY: 1, rt.SALES_COL_PRICE: 100, rt.COL_DOC_VAT: 7,
            rt.PURCHASE_COL_PRE_VAT: 100, rt.PURCHASE_COL_VAT: 7,
            rt.PURCHASE_COL_PARTY: "Synthetic seller"}
        if prior:
            row[rt.COL_ROW_KEY] = rt.encode_row_key(prior["name"], 0)
        row.update(overrides)
        return row

    def upload(self, data):
        self.source = self.files.create(self.workspace, "synthetic-review.xlsx", data, uuid.uuid4().hex)
        frappe.db.commit()

    def test_export_edit_reimport_preserves_money_and_creates_new_history(self):
        from pearnly_cowork.roundtrip import service, spec as rt
        from pearnly_cowork.ocr.table_inputs import extract_table
        old = self.prior()
        original = frappe.get_doc("Cowork OCR History", old["name"]).as_dict()
        output = service.export([old["name"]])
        wb = load_workbook(io.BytesIO(output))
        self.assertEqual(wb.sheetnames, [rt.SHEET_SALES, rt.SHEET_PURCHASE, rt.SHEET_PENDING, rt.SHEET_SUMMARY])
        wb[rt.SHEET_SALES]["B2"] = "EDITED-001"
        edited = io.BytesIO()
        wb.save(edited)
        data = edited.getvalue()
        self.upload(data)
        frappe.conf.pearnly_ocr_credentials = None
        with patch("pearnly_cowork.ocr.execution.workflow.build_client", side_effect=AssertionError("No model client")):
            result = self.run_work(self.admit())
        self.assertEqual((result["state"], result["page_count"], result["reference_kind"]), ("saved", 1, "logical_section"))
        saved = frappe.get_doc("Cowork OCR History", result["history_records"][0])
        self.assertNotEqual(saved.name, old["name"])
        fields = json.loads(saved.pages)[0]["fields"]
        self.assertEqual((fields["invoice_number"], fields["history_id"], fields["direction"]),
            ("EDITED-001", old["name"], "sales"))
        self.assertEqual((fields["subtotal"], fields["vat"], fields["total_amount"]), ("100.00", "7.00", "107.00"))
        self.assertEqual(frappe.get_doc("Cowork OCR History", old["name"]).as_dict(), original)
        self.assertEqual(self.charge.call_args.args, ("cowork", "excel", extract_table(data, "x.xlsx").billing_characters))
        self.http.assert_not_called()
        self.assertEqual(json.loads(frappe.get_doc("Cowork OCR Batch", result["name"]).provider_evidence), [])

    def test_same_number_different_keys_directions_and_pending_all_survive(self):
        from pearnly_cowork.roundtrip import spec as rt
        first, second, third = self.prior(), self.prior(), self.prior("purchase")
        self.upload(book([(rt.SHEET_SALES, [self.row(first), self.row(second)]),
            (rt.SHEET_PURCHASE, [self.row(third)]), (rt.SHEET_PENDING, [{rt.PENDING_COL_REASON: "Needs human direction"}])]))
        doc = self.admit()  # requests purchase, must not override sheet direction/pending
        result = self.run_work(doc)
        self.assertEqual((result["state"], result["page_count"], len(result["history_records"])), ("saved", 4, 4))
        pages = [json.loads(frappe.get_doc("Cowork OCR History", name).pages)[0] for name in result["history_records"]]
        self.assertEqual([p["fields"].get("direction") for p in pages], ["sales", "sales", "purchase", None])
        self.assertFalse(any(p["is_duplicate"] for p in pages))
        self.assertTrue(pages[-1]["_needs_manual_review"])
        self.assertTrue(result["requires_review"])
        self.assertEqual([p["page_index"] for p in pages], [1, 2, 3, 4])
        source = frappe.get_doc("Cowork Source File", self.source["name"])
        self.assertEqual(json.loads(source.provenance_json)["provenance"]["unit_kind"], "roundtrip_document")
        self.assertFalse(source.coverage_complete)
        self.http.assert_not_called()

    def test_unknown_legacy_link_rejected_before_billing(self):
        from pearnly_cowork.roundtrip import spec as rt
        self.upload(book([(rt.SHEET_SALES, [self.row({"name": "legacy-unmapped-uuid"})])]))
        with self.assertRaises(frappe.PermissionError):
            self.admit()
        self.status.assert_not_called()
        self.charge.assert_not_called()
        self.http.assert_not_called()
        self.assertFalse(frappe.db.exists("Cowork OCR Batch", {"source_file": self.source["name"]}))

    def test_lost_charge_receipt_retries_accounting_only(self):
        from pearnly_cowork.roundtrip import spec as rt
        self.upload(book([(rt.SHEET_SALES, [self.row()])]))
        self.charge.side_effect = [{"ok": False}, {"ok": True, "duplicate": True}]
        doc = self.admit()
        first = self.run_work(doc)
        self.assertEqual((first["state"], first["billing_state"]), ("saved", "pending"))
        with patch("pearnly_cowork.roundtrip.intake.decode", side_effect=AssertionError("No reparse")):
            final = self.run_work(doc, retry=True)
        self.assertEqual(final["billing_state"], "charged")
        self.assertEqual(first["history_records"], final["history_records"])
        self.assertEqual({c.kwargs["reference"] for c in self.charge.call_args_list}, {doc.billing_reference})
        self.http.assert_not_called()

    def test_links_rechecked_before_saving_frozen_result(self):
        from pearnly_cowork.roundtrip import spec as rt
        from pearnly_cowork.history import storage
        from pearnly_core.execution.errors import Deferred
        old = self.prior()
        self.upload(book([(rt.SHEET_SALES, [self.row(old)])]))
        doc = self.admit()
        with patch("pearnly_cowork.ocr.execution.workflow.persist_batch", side_effect=Deferred("local_test", 1)):
            self.assertEqual(self.run_work(doc)["state"], "deferred")
        self.assertTrue(frappe.db.get_value("Cowork OCR Batch", doc.name, "result_json"))
        actual_read = storage.read
        def revoked(name, **kwargs):
            if name == old["name"]:
                raise frappe.PermissionError
            return actual_read(name, **kwargs)
        with patch.object(storage, "read", side_effect=revoked), self.assertRaises(frappe.PermissionError):
            self.run_work(doc, retry=True)
        self.assertEqual(frappe.db.count("Cowork OCR History", {"source_file": self.source["name"]}), 0)
        self.charge.assert_not_called()

    def test_original_summary_does_not_unlock_for_another_reader(self):
        from pearnly_cowork.roundtrip import spec as rt
        self.upload(book([(rt.SHEET_SALES, [self.row()])]))
        self.run_work(self.admit())
        self.files.original(self.source["name"])
        frappe.set_user("Administrator")
        grant(self.other, "Cowork Workspace", self.workspace)
        frappe.db.commit()
        frappe.set_user(self.other)
        with self.assertRaises(frappe.PermissionError):
            self.files.original(self.source["name"])

    def test_export_requires_every_record_permission_and_scope(self):
        from pearnly_cowork.roundtrip import service
        old = self.prior()
        with self.assertRaises(frappe.DoesNotExistError):
            service.export([old["name"], "missing-history"])
        with self.assertRaises(frappe.ValidationError):
            service.export([old["name"], old["name"]])
        frappe.set_user(self.other)
        with self.assertRaises(frappe.PermissionError):
            service.export([old["name"]])

    def test_native_export_returns_private_attachment(self):
        from pearnly_cowork.roundtrip import api
        old = self.prior()
        api.export([old["name"]])
        self.assertEqual(frappe.local.response.type, "download")
        self.assertEqual(frappe.local.response.filename, "cowork-review.xlsx")
        self.assertTrue(frappe.local.response.filecontent.startswith(b"PK"))
        self.assertTrue(frappe.local.cowork_private_response)
        frappe.local.response.clear()

    def test_http_export_is_an_authenticated_private_post_download(self):
        from werkzeug.test import Client
        from werkzeug.wrappers import Response
        from frappe.app import application
        old = self.prior()
        frappe.set_user("Administrator")
        account = frappe.get_doc("User", self.actor)
        account.api_key, secret = uuid.uuid4().hex, uuid.uuid4().hex
        account.api_secret = secret
        account.save()
        headers = {"Authorization": f"token {account.api_key}:{secret}"}
        site = frappe.local.site
        frappe.db.commit()
        frappe.destroy()
        client = Client(application, Response)
        endpoint = "/api/method/pearnly_cowork.roundtrip.api.export"
        try:
            result = client.post(endpoint, base_url="http://" + site, headers=headers, json={"names": [old["name"]]})
            self.assertEqual(result.status_code, 200)
            self.assertTrue(result.data.startswith(b"PK"))
            self.assertIn("no-store", result.headers.get("Cache-Control", ""))
            self.assertIn("attachment", result.headers.get("Content-Disposition", ""))
            denied = client.post(endpoint, base_url="http://" + site, json={"names": [old["name"]]})
            self.assertEqual(denied.status_code, 403)
            self.assertEqual(client.get(endpoint, base_url="http://" + site, headers=headers).status_code, 403)
        finally:
            connect(site)

    def test_empty_recognized_template_is_rejected_without_generic_model(self):
        from pearnly_cowork.roundtrip import spec as rt
        self.upload(book([(rt.SHEET_SALES, [])]))
        self.assertEqual(self.api.submit(self.source["name"], self.key),
            {"accepted": False, "code": "ocr.roundtrip_empty"})
        self.status.assert_not_called()
        self.http.assert_not_called()

    def test_multi_page_saved_invoice_exports_each_item_once(self):
        from pearnly_cowork.history.recognition import persist_batch
        from pearnly_cowork.ocr.results import prepare_history_batch
        from pearnly_cowork.ocr.schema import PipelinePageResult, PipelineResult, validate_invoice
        from pearnly_cowork.roundtrip import service
        from pearnly_cowork.roundtrip.parser import parse_roundtrip_workbook
        data = BASE["pdf"]()
        self.source = self.files.create(self.workspace, "synthetic-two-pages.pdf", data, uuid.uuid4().hex)
        pipeline = PipelineResult(page_count=2, elapsed_ms=0, pages=[PipelinePageResult(page_number=i,
            layer_chain=["L2"], invoice=validate_invoice({**BASE["GOOD"], "subtotal": "300", "vat": "21",
                "total_amount": "321", "direction": "sales", "buyer_name": "Synthetic buyer",
                "items": [{"name": "Item " + str(i), "qty": "1", "price": str(i * 100), "subtotal": str(i * 100)}]},
                anchor_year=2026)) for i in (1, 2)])
        plan = prepare_history_batch(pipeline, expected_page_count=2, workspace=self.workspace,
                                     source="manual", direction="sales")
        rows = persist_batch(self.source["name"], plan["groups"], uuid.uuid4().hex, anchor_year=2026)
        self.assertEqual(len(rows), 1)
        self.assertEqual(len(rows[0]["pages"][0]["fields"]["items"]), 2)
        self.assertEqual(len(rows[0]["pages"][1]["fields"]["items"]), 1)
        frappe.db.commit()
        returned = parse_roundtrip_workbook(service.export([rows[0]["name"]]))["documents"]
        self.assertEqual(len(returned[0]["fields"]["items"]), 2)
        self.assertEqual(str(returned[0]["fields"]["total_amount"]), "321.00")

    def test_unbound_multi_page_invoice_is_not_guessed(self):
        from pearnly_cowork.history import storage
        from pearnly_cowork.roundtrip import service
        from pearnly_cowork.roundtrip.export import ExportError
        row = storage.create_manual(self.workspace, [{"page_index": i, "fields": BASE["GOOD"]} for i in (1, 2)], uuid.uuid4().hex)
        with self.assertRaises(ExportError) as error:
            service.export([row["name"]])
        self.assertEqual(error.exception.code, "roundtrip.export_multipage_unverified")

    def test_authorized_other_company_still_cannot_be_mixed_or_relinked(self):
        from pearnly_cowork.masters import api as masters
        from pearnly_cowork.roundtrip import service, spec as rt
        own = self.prior()
        frappe.set_user("Administrator")
        other_workspace = masters.create_workspace("Synthetic roundtrip other " + uuid.uuid4().hex)["name"]
        grant(self.actor, "Cowork Workspace", other_workspace)
        frappe.db.commit()
        frappe.set_user(self.actor)
        other = self.prior(workspace=other_workspace)
        with self.assertRaises(frappe.PermissionError):
            service.export([own["name"], other["name"]])
        self.upload(book([(rt.SHEET_SALES, [self.row(other)])]))
        with self.assertRaises(frappe.PermissionError):
            self.admit()
        self.status.assert_not_called()

    def test_authorized_other_client_cannot_supply_an_old_link(self):
        from pearnly_cowork.masters import api as masters
        from pearnly_cowork.history import storage
        from pearnly_cowork.roundtrip import spec as rt
        frappe.set_user("Administrator")
        client = masters.create_client(self.workspace, "Synthetic roundtrip client " + uuid.uuid4().hex)["name"]
        grant(self.actor, "Cowork Client", client)
        frappe.db.commit()
        frappe.set_user(self.actor)
        old = storage.create_manual(self.workspace, [{"page_index": 1, "fields": BASE["GOOD"]}],
                                    uuid.uuid4().hex, client)
        self.upload(book([(rt.SHEET_SALES, [self.row(old)])]))
        with self.assertRaises(frappe.PermissionError):
            self.admit()
        self.status.assert_not_called()

    def test_identical_admission_and_worker_replay_have_one_result_and_receipt(self):
        from pearnly_cowork.roundtrip import spec as rt
        self.upload(book([(rt.SHEET_SALES, [self.row()])]))
        first = self.admit()
        self.assertEqual(self.admit().name, first.name)
        completed = self.run_work(first)
        self.assertEqual(self.run_work(first)["history_records"], completed["history_records"])
        self.api.retry_billing(first.name)
        self.assertEqual(self.charge.call_count, 1)
        self.assertEqual(frappe.db.count("Cowork OCR History", {"source_file": self.source["name"]}), 1)
        self.assertFalse(frappe.get_all("Version", filters={"ref_doctype": "Cowork Source File",
            "docname": self.source["name"], "data": ["like", "%provenance_json%"]}, pluck="name"))


if __name__ == "__main__":
    connect()
    try:
        suite = (unittest.TestSuite([NativeRoundtrip(sys.argv[1])]) if len(sys.argv) > 1
                 else unittest.defaultTestLoader.loadTestsFromTestCase(NativeRoundtrip))
        result = unittest.TextTestRunner(verbosity=2).run(suite)
        raise SystemExit(0 if result.wasSuccessful() else 1)
    finally:
        frappe.destroy()
